"""Project-wide frequency-domain QC and exclusion metadata helpers."""

from __future__ import annotations

import hashlib
import json
import logging
import sys
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from fractions import Fraction
from pathlib import Path
from time import perf_counter
from typing import Any

import numpy as np
import pandas as pd

from Main_App.processing.frequency_qc_identity import (
    resolve_frequency_qc_recording_decisions,
)
from Main_App.processing.provisional_harmonic_cache import ProvisionalHarmonicCache
from Main_App.projects import (
    ProjectDatasetIndex,
    load_project_dataset_index,
    normalize_experimental_qc_settings,
    normalize_frequency_protocol,
)
from Main_App.projects.experimental_qc_settings import (
    SUMMED_BCA_SCREENING_BRIEF_TEXT,
    SummedBcaScreeningSettings,
)
from Main_App.projects.preprocessing_settings import (
    normalize_manual_excluded_participants,
    normalize_manual_excluded_recordings,
    normalize_preprocessing_settings,
)

logger = logging.getLogger(__name__)

QUALITY_CHECK_FOLDER = "Quality Check"
FREQUENCY_DOMAIN_QC_REPORT_NAME = "Frequency_Domain_QC_Review.txt"
FREQUENCY_DOMAIN_QC_METADATA_PATH = ("tools", "frequency_domain_qc")
FREQUENCY_DOMAIN_QC_SCHEMA_VERSION = 4
FREQUENCY_DOMAIN_QC_METHOD_VERSION = "experimental_summed_bca_electrode_review_v5"
REPEATED_FREQUENCY_DOMAIN_QC_METHOD_VERSION = (
    "experimental_summed_bca_recording_electrode_review_v5"
)
FREQUENCY_DOMAIN_QC_INTEGRITY_METHOD_VERSION = "selected_bca_finite_v1"
FREQUENCY_DOMAIN_QC_DECISION_VERSION = "frequency_qc_review_decision_v2"
FREQUENCY_DOMAIN_QC_LOOP_VERSION = "frequency_qc_bounded_recompute_v1"
FREQUENCY_DOMAIN_QC_REVIEW_EVIDENCE_VERSION = "frequency_qc_review_evidence_v1"
FREQUENCY_DOMAIN_QC_INDEPENDENT_EVIDENCE_VERSION = (
    "frequency_qc_independent_evidence_v1"
)
FREQUENCY_DOMAIN_QC_MAX_REVIEW_ITERATIONS = 64
SPECTRAL_METRIC_QC_SHEET_NAME = "Spectral Metric QC"
_HARMONIC_CACHE_ANNOTATIONS = frozenset({
    "selection_cache_source", "selection_cache_saved_at", "selection_cache_key",
})

DECISION_RETAIN = "retain"
# Read old receipts for audit/reconfirmation, but never apply this retired action.
DECISION_EXCLUDE_CONDITION_ELECTRODE = "exclude_condition_electrode"
DECISION_INTERPOLATE_CONDITION_ELECTRODE = "interpolate_condition_electrode"
DECISION_EXCLUDE_CONDITION = "exclude_condition"
DECISION_EXCLUDE_RECORDING = "exclude_recording"
DECISION_EXCLUDE_PARTICIPANT = "exclude_participant"

REVIEW_DECISIONS = (
    DECISION_RETAIN,
    DECISION_INTERPOLATE_CONDITION_ELECTRODE,
    DECISION_EXCLUDE_CONDITION,
    DECISION_EXCLUDE_RECORDING,
    DECISION_EXCLUDE_PARTICIPANT,
)
_BROAD_EXCLUSION_DECISIONS = frozenset({
    DECISION_EXCLUDE_CONDITION, DECISION_EXCLUDE_RECORDING, DECISION_EXCLUDE_PARTICIPANT,
})

_BCA_AUDIT_REQUIRED_COLUMNS = (
    "Electrode",
    "Target Frequency Exact (Hz)",
    "BCA Status",
    "Reason Codes",
)
_MAX_BCA_AUDIT_FREQUENCY_TOKENS = 256
_MAX_BCA_AUDIT_FREQUENCY_TOKEN_LENGTH = 128

WARNING_REASON_UNUSUAL_VALUES = "Unusual frequency-domain values"
WARNING_REASON_NOISY_SPECTRUM = "Noisy spectrum"
WARNING_REASON_KNOWN_ACQUISITION = "Known acquisition issue"
WARNING_REASON_OTHER = "Other reviewed concern"
MANUAL_EXCLUSION_REASONS = (
    WARNING_REASON_UNUSUAL_VALUES,
    WARNING_REASON_NOISY_SPECTRUM,
    WARNING_REASON_KNOWN_ACQUISITION,
    WARNING_REASON_OTHER,
)


@dataclass(frozen=True)
class FrequencyDomainQcThresholds:
    warning_summed_bca_uv: float = 10.0
    strong_warning_summed_bca_uv: float = 50.0
    extreme_review_summed_bca_uv: float = 250.0
    concentrated_review_flagged_cells: int = 5
    broad_extreme_review_unique_electrodes: int = 11

    @classmethod
    def from_settings(
        cls,
        settings: SummedBcaScreeningSettings,
    ) -> "FrequencyDomainQcThresholds":
        return cls(
            warning_summed_bca_uv=settings.warning_summed_bca_uv,
            strong_warning_summed_bca_uv=settings.strong_warning_summed_bca_uv,
            extreme_review_summed_bca_uv=settings.extreme_review_summed_bca_uv,
            concentrated_review_flagged_cells=(
                settings.concentrated_review_flagged_cells
            ),
            broad_extreme_review_unique_electrodes=(
                settings.broad_extreme_review_unique_electrodes
            ),
        )

    @property
    def hard_electrode_summed_bca_uv(self) -> float:
        """Compatibility alias for historical report readers."""

        return self.extreme_review_summed_bca_uv

    @property
    def repeated_warning_cells(self) -> int:
        """Compatibility alias for historical report readers."""

        return self.concentrated_review_flagged_cells

    @property
    def hard_participant_unique_electrodes(self) -> int:
        """Compatibility alias for historical report readers."""

        return self.broad_extreme_review_unique_electrodes - 1

    def to_manifest(self) -> dict[str, object]:
        return {
            "warning_summed_bca_uv": float(self.warning_summed_bca_uv),
            "strong_warning_summed_bca_uv": float(self.strong_warning_summed_bca_uv),
            "extreme_review_summed_bca_uv": float(
                self.extreme_review_summed_bca_uv
            ),
            "concentrated_review_flagged_cells": int(
                self.concentrated_review_flagged_cells
            ),
            "broad_extreme_review_unique_electrodes": int(
                self.broad_extreme_review_unique_electrodes
            ),
        }


DEFAULT_FREQUENCY_DOMAIN_QC_THRESHOLDS = FrequencyDomainQcThresholds()


@dataclass(frozen=True)
class FrequencyDomainExclusions:
    excluded_participants: frozenset[str]
    auto_excluded_participants: frozenset[str]
    manual_excluded_participants: frozenset[str]
    auto_excluded_electrodes_by_participant: dict[str, frozenset[str]]
    downstream_outputs_stale: bool
    excluded_recordings: frozenset[str] = frozenset()
    auto_excluded_recordings: frozenset[str] = frozenset()
    manual_excluded_recordings: frozenset[str] = frozenset()
    auto_excluded_electrodes_by_recording: dict[str, frozenset[str]] = field(
        default_factory=dict
    )
    excluded_participant_conditions: frozenset[tuple[str, str]] = frozenset()
    excluded_recording_conditions: frozenset[tuple[str, str]] = frozenset()
    excluded_electrodes_by_participant_condition: dict[
        tuple[str, str], frozenset[str]
    ] = field(default_factory=dict)
    excluded_electrodes_by_recording_condition: dict[
        tuple[str, str], frozenset[str]
    ] = field(default_factory=dict)


@dataclass(frozen=True)
class FrequencyDomainCoverageDecisions:
    """Authoritative reviewed QC-17 exclusions for QC-20/QC-21 consumers."""

    decision_fingerprint: str
    review_complete: bool
    excluded_participants: frozenset[str]
    excluded_recordings: frozenset[str]
    excluded_participant_conditions: frozenset[tuple[str, str]]
    excluded_recording_conditions: frozenset[tuple[str, str]]
    excluded_electrodes_by_participant_condition: dict[
        tuple[str, str], frozenset[str]
    ]
    excluded_electrodes_by_recording_condition: dict[
        tuple[str, str], frozenset[str]
    ]
    reviewed_decisions: tuple[dict[str, object], ...]

    def to_payload(self) -> dict[str, object]:
        return {
            "decision_fingerprint": self.decision_fingerprint,
            "review_complete": self.review_complete,
            "excluded_participants": sorted(self.excluded_participants),
            "excluded_recordings": sorted(self.excluded_recordings),
            "excluded_participant_conditions": [
                {"participant_id": participant_id, "condition": condition}
                for participant_id, condition in sorted(
                    self.excluded_participant_conditions
                )
            ],
            "excluded_recording_conditions": [
                {"recording_id": recording_id, "condition": condition}
                for recording_id, condition in sorted(
                    self.excluded_recording_conditions
                )
            ],
            "excluded_electrodes_by_participant_condition": [
                {
                    "participant_id": key[0],
                    "condition": key[1],
                    "electrodes": sorted(electrodes),
                }
                for key, electrodes in sorted(
                    self.excluded_electrodes_by_participant_condition.items()
                )
            ],
            "excluded_electrodes_by_recording_condition": [
                {
                    "recording_id": key[0],
                    "condition": key[1],
                    "electrodes": sorted(electrodes),
                }
                for key, electrodes in sorted(
                    self.excluded_electrodes_by_recording_condition.items()
                )
            ],
            "reviewed_decisions": [dict(item) for item in self.reviewed_decisions],
        }


class FrequencyDomainQcIntegrityError(RuntimeError):
    """Raised when technical workbook defects make frequency QC incomplete."""


@dataclass(frozen=True)
class _SummedBcaInspection:
    flags: tuple[dict[str, object], ...]
    technical_integrity_failures: tuple[dict[str, object], ...]
    unavailable_by_method: tuple[dict[str, object], ...]


@dataclass(frozen=True)
class _IndependentQcContext:
    source_identity: dict[str, object]
    cells: dict[tuple[str, str], dict[str, object]]
    processing_entries: dict[str, dict[str, object]]


def run_frequency_domain_qc_review(
    project: Any,
    *,
    log_func: Callable[[str], None] | None = None,
    dataset_index: ProjectDatasetIndex | None = None,
    provisional_cache: ProvisionalHarmonicCache | None = None,
) -> dict[str, object]:
    """Build a provisional summed-BCA QC report for the active project."""

    def _log(message: str) -> None:
        if log_func is not None:
            log_func(str(message))

    stage_name = ""
    stage_started = perf_counter()

    def _finish_stage() -> None:
        if stage_name:
            elapsed = perf_counter() - stage_started
            logger.info(
                "frequency_domain_qc_stage_complete stage=%s elapsed_s=%.3f",
                stage_name, elapsed,
                extra={"stage": stage_name, "elapsed_s": elapsed},
            )

    def _start_stage(name: str, description: str) -> None:
        nonlocal stage_name, stage_started
        _finish_stage()
        stage_name, stage_started = name, perf_counter()
        logger.info("frequency_domain_qc_stage_started stage=%s", name, extra={"stage": name})
        _log(f"Frequency-domain QC: {description}")

    _start_stage("canonical_inputs", "Checking project inputs…")
    project_root = Path(project.project_root).resolve()
    from Main_App.processing.condition_interpolation_state import (
        require_no_pending_condition_interpolation,
    )

    require_no_pending_condition_interpolation(project_root)
    screening_settings = _experimental_summed_bca_settings(project, project_root)
    interpolation_enabled = normalize_experimental_qc_settings(
        _read_manifest(project_root / "project.json").get("experimental_qc")
    ).condition_specific_interpolation_enabled
    thresholds = FrequencyDomainQcThresholds.from_settings(screening_settings)
    from Main_App.processing.harmonic_selection_qc import (
        resolve_processing_harmonic_selection_inputs,
    )

    if dataset_index is None:
        dataset_index = load_project_dataset_index(project_root)
    elif dataset_index.project_root.resolve() != project_root:
        raise ValueError(
            "The supplied dataset index belongs to a different project root."
        )
    canonical_inputs = resolve_processing_harmonic_selection_inputs(
        project,
        log_func=_log,
        dataset_index=dataset_index,
    )
    subjects = list(canonical_inputs.subjects)
    subject_data = canonical_inputs.subject_data
    ordered_conditions = list(canonical_inputs.conditions)
    repeated_session = canonical_inputs.is_repeated_session
    recording_assignments = canonical_inputs.recording_assignments
    rois = canonical_inputs.rois
    settings = canonical_inputs.settings
    active_exclusions = active_frequency_domain_exclusions(project_root)
    condition_exclusions = (
        active_exclusions.excluded_recording_conditions
        if repeated_session
        else active_exclusions.excluded_participant_conditions
    )
    condition_electrode_exclusions = (
        active_exclusions.excluded_electrodes_by_recording_condition
        if repeated_session
        else active_exclusions.excluded_electrodes_by_participant_condition
    )
    _start_stage("independent_evidence", "Checking preprocessing evidence…")
    independent_qc_context = _load_independent_qc_context(project_root)
    _start_stage("provisional_harmonics", "Preparing candidate harmonics…")
    provisional_inputs = dict(
        project_root=project_root,
        subjects=subjects,
        conditions=ordered_conditions,
        subject_data=subject_data,
        rois=rois,
        settings=settings,
        log_func=_log,
        recording_assignments=(recording_assignments if repeated_session else None),
        declared_session_ids=(
            canonical_inputs.declared_session_ids
            if repeated_session
            else None
        ),
        participant_group_ids=(
            canonical_inputs.participant_group_ids
            if repeated_session
            else None
        ),
        declared_group_ids=(
            canonical_inputs.declared_group_ids
            if repeated_session
            else None
        ),
        base_frequency_hz=canonical_inputs.base_frequency_hz,
        oddball_frequency_hz=canonical_inputs.oddball_frequency_hz,
        eligible_harmonic_orders=canonical_inputs.eligible_harmonic_orders,
        spectral_eligibility_fingerprint=(
            canonical_inputs.spectral_eligibility_fingerprint
        ),
        electrode_exclusions_by_subject_condition=(
            condition_electrode_exclusions
        ),
        expected_scalp_channels_by_subject_condition=(
            _expected_scalp_channels_by_subject_condition(
                independent_qc_context,
                excluded_conditions=condition_exclusions,
            )
        ),
    )
    selected_harmonics, provisional_metadata = (
        _provisional_harmonics(**provisional_inputs)
        if provisional_cache is None
        else provisional_cache.resolve(_provisional_harmonics, **provisional_inputs)
    )
    _log(
        "Frequency-domain QC is reviewing provisional summed BCA values "
        f"across {len(selected_harmonics)} harmonic(s)."
    )
    _start_stage("absolute_screening", "Checking electrode amplitudes…")
    inspection = _collect_summed_bca_flags(
        subjects=subjects,
        conditions=ordered_conditions,
        subject_data=subject_data,
        selected_harmonics=selected_harmonics,
        thresholds=thresholds,
        log_func=_log,
        recording_assignments=(recording_assignments if repeated_session else None),
        screening_enabled=screening_settings.enabled,
        selected_harmonics_metadata=provisional_metadata,
        excluded_electrodes_by_subject_condition=(
            condition_electrode_exclusions
        ),
    )
    harmonic_selection_fingerprint = _harmonic_selection_fingerprint(
        provisional_metadata
    )
    flags = list(inspection.flags)
    protocol_metadata = _project_protocol_review_metadata(project, project_root)
    for flag in flags:
        flag.update(protocol_metadata)
        flag["harmonic_selection_fingerprint"] = harmonic_selection_fingerprint
        _attach_independent_qc_evidence(flag, independent_qc_context)
        flag["finding_fingerprint"] = _frequency_qc_finding_fingerprint(flag)
    _start_stage("report_integrity", "Preparing review findings…")
    machine_findings = list(flags)
    technical_integrity_failures = list(
        inspection.technical_integrity_failures
    )
    unavailable_by_method = list(inspection.unavailable_by_method)
    technical_integrity_failed = bool(technical_integrity_failures)
    finite_input_status = {
        "method_version": FREQUENCY_DOMAIN_QC_INTEGRITY_METHOD_VERSION,
        "status": (
            "technical_output_integrity_failed"
            if technical_integrity_failed
            else "complete"
        ),
        "technical_integrity_failure_count": len(
            technical_integrity_failures
        ),
        "unavailable_by_method_count": len(unavailable_by_method),
    }
    finite_input_fingerprint = _hash_payload(
        {
            **finite_input_status,
            "technical_integrity_failures": [
                _finite_input_fingerprint_diagnostic(project_root, item)
                for item in technical_integrity_failures
            ],
            "unavailable_by_method": [
                _finite_input_fingerprint_diagnostic(project_root, item)
                for item in unavailable_by_method
            ],
        }
    )
    finite_input_status["fingerprint"] = finite_input_fingerprint
    source_workbooks = _source_workbook_rows(
        project_root=project_root,
        subjects=subjects,
        conditions=ordered_conditions,
        subject_data=subject_data,
        recording_assignments=(recording_assignments if repeated_session else None),
    )
    source_fingerprint = _hash_payload(
        {
            "source_workbooks": source_workbooks,
            "finite_input_fingerprint": finite_input_fingerprint,
            "harmonic_selection_fingerprint": harmonic_selection_fingerprint,
            "frequency_protocol_fingerprint": protocol_metadata.get(
                "frequency_protocol_fingerprint"
            ),
            "independent_qc_source_fingerprint": independent_qc_context.source_identity.get(
                "fingerprint"
            ),
        }
    )
    state = load_frequency_domain_qc_state(project_root)
    previous_review_evidence = (
        _validated_review_evidence_from_state(project_root, state)
        if state.get("review_complete") else None
    )
    evidence_context_fingerprint = _analysis_fingerprint(
        project_root=project_root,
        subjects=subjects,
        conditions=ordered_conditions,
        subject_data=subject_data,
        selected_harmonics=selected_harmonics,
        thresholds=thresholds,
        flags=machine_findings,
        finite_input_fingerprint=finite_input_fingerprint,
        recording_assignments=(recording_assignments if repeated_session else None),
        screening_settings=screening_settings,
        provisional_metadata=provisional_metadata,
        source_workbooks=source_workbooks,
        source_fingerprint=source_fingerprint,
        previous_review_evidence=previous_review_evidence,
    )
    reconfirmation_findings, stable_exclusion_decisions = (
        _review_exclusion_reconfirmation_findings(
            state,
            evidence_context_fingerprint=evidence_context_fingerprint,
            selected_harmonics=selected_harmonics,
            harmonic_selection_fingerprint=harmonic_selection_fingerprint,
            protocol_metadata=protocol_metadata,
            independent_qc_context=independent_qc_context,
        )
    )
    review_findings = [*machine_findings, *reconfirmation_findings]
    if repeated_session:
        summaries, machine_electrodes, machine_subjects = _summarize_recording_flags(
            review_findings,
            thresholds,
        )
    else:
        summaries, machine_electrodes, machine_subjects = _summarize_flags(
            review_findings,
            thresholds,
        )
    if technical_integrity_failed:
        summaries = []
        machine_electrodes = []
        machine_subjects = []
        review_findings = machine_findings
        reconfirmation_findings = []
        stable_exclusion_decisions = []
    analysis_fingerprint = evidence_context_fingerprint
    review_loop = _frequency_qc_review_loop_status(
        state,
        analysis_fingerprint=analysis_fingerprint,
    )
    current_manual = _manual_entries_from_state(state)
    current_manual_recordings = _manual_recording_entries_from_state(state)
    current_review_decisions = _current_review_decisions(
        report_flags=review_findings,
        state=state,
    )
    active_review_decisions = _merge_review_decision_rows(
        stable_exclusion_decisions,
        current_review_decisions,
        _resolved_reconfirmation_retains(
            state=state,
            analysis_fingerprint=analysis_fingerprint,
            review_evidence=previous_review_evidence,
        ),
    )
    current_decision_fingerprint = _decision_fingerprint(
        analysis_fingerprint=analysis_fingerprint,
        auto_electrodes=(),
        auto_participants=(),
        manual_participants=current_manual,
        auto_recording_electrodes=(() if repeated_session else None),
        auto_recordings=(() if repeated_session else None),
        manual_recordings=(current_manual_recordings if repeated_session else None),
        review_decisions=active_review_decisions,
        identity_scope=("recording" if repeated_session else "participant"),
    )
    last_review = state.get("last_review")
    reviewed_decision_fingerprint = ""
    if isinstance(last_review, Mapping):
        reviewed_decision_fingerprint = str(last_review.get("decision_fingerprint") or "")

    pause_subjects = [
        summary
        for summary in summaries
        if bool(summary.get("pause_review"))
    ]
    review_reused = bool(
        screening_settings.enabled
        and
        not technical_integrity_failed
        and pause_subjects
        and len(current_review_decisions) == len(review_findings)
        and reviewed_decision_fingerprint
        and reviewed_decision_fingerprint == current_decision_fingerprint
    )
    review_required = bool(
        screening_settings.enabled
        and not technical_integrity_failed
        and pause_subjects
        and not review_reused
    )
    if not screening_settings.enabled:
        review_reused = False
        review_required = False
    legacy_machine_suggestions = _legacy_machine_suggestions_from_state(state)
    report: dict[str, object] = {
        "schema_version": FREQUENCY_DOMAIN_QC_SCHEMA_VERSION,
        "method_version": (
            REPEATED_FREQUENCY_DOMAIN_QC_METHOD_VERSION
            if repeated_session
            else FREQUENCY_DOMAIN_QC_METHOD_VERSION
        ),
        "project_root": str(project_root),
        "screening_status": (
            "performed" if screening_settings.enabled else "not_performed"
        ),
        "screening_enabled": screening_settings.enabled,
        "condition_specific_interpolation_enabled": interpolation_enabled,
        "screening_policy_version": screening_settings.policy_version,
        "screening_explanation": SUMMED_BCA_SCREENING_BRIEF_TEXT,
        "screening_settings": screening_settings.to_manifest(),
        "thresholds": thresholds.to_manifest(),
        "subjects": list(subjects),
        "conditions": list(ordered_conditions),
        "selected_harmonics_hz": list(selected_harmonics),
        "harmonic_policy": settings.name,
        "frequency_protocol": protocol_metadata,
        "provisional_harmonic_metadata": provisional_metadata,
        "flags": flags,
        "cohort_relative_rows": [],
        "cohort_relative_flags": [],
        "reconfirmation_findings": reconfirmation_findings,
        "review_findings": review_findings,
        "harmonic_selection_fingerprint": harmonic_selection_fingerprint,
        "source_workbooks": source_workbooks,
        "source_fingerprint": source_fingerprint,
        "independent_qc_source": independent_qc_context.source_identity,
        "qc_complete": not technical_integrity_failed,
        "result_status": finite_input_status["status"],
        "technical_integrity_failed": technical_integrity_failed,
        "technical_integrity_failures": technical_integrity_failures,
        "unavailable_by_method": unavailable_by_method,
        "finite_input_status": finite_input_status,
        "threshold_flags_diagnostic_only": technical_integrity_failed,
        "participant_summaries": summaries,
        "machine_participant_electrode_suggestions": machine_electrodes,
        "machine_participant_suggestions": machine_subjects,
        "legacy_machine_suggestions": legacy_machine_suggestions,
        "auto_participant_electrode_exclusions": [],
        "auto_participant_exclusions": [],
        "manual_participant_exclusions": current_manual,
        "review_decisions": current_review_decisions,
        "active_review_decisions": active_review_decisions,
        "review_prefill_decisions": [
            {
                "finding_fingerprint": str(
                    finding.get("finding_fingerprint") or ""
                ),
                "decision": str(finding.get("prior_decision") or ""),
                "reason": str(finding.get("prior_reason") or ""),
            }
            for finding in reconfirmation_findings
        ],
        "analysis_fingerprint": analysis_fingerprint,
        "current_decision_fingerprint": current_decision_fingerprint,
        "review_required": review_required,
        "review_reused": review_reused,
        "review_subject_count": len(pause_subjects),
        "review_loop": review_loop,
        "generated_at": _now_utc_iso(),
    }
    if repeated_session:
        report.update(
            {
                "identity_scope": "recording",
                "subjects": sorted(
                    {
                        str(recording_assignments[recording_id]["participant_id"])
                        for recording_id in subjects
                    },
                    key=str.casefold,
                ),
                "recordings": list(subjects),
                "recording_assignments": [
                    dict(recording_assignments[recording_id])
                    for recording_id in subjects
                ],
                "recording_summaries": summaries,
                "participant_summaries": [],
                "machine_recording_electrode_suggestions": machine_electrodes,
                "machine_recording_suggestions": machine_subjects,
                "auto_recording_electrode_exclusions": [],
                "auto_recording_exclusions": [],
                "manual_recording_exclusions": current_manual_recordings,
                "review_recording_count": len(pause_subjects),
            }
        )
    _finish_stage()
    return report


def require_frequency_domain_qc_complete(
    report: Mapping[str, object],
) -> None:
    """Reject finalization when selected computable BCA inputs are invalid."""

    failures = [
        dict(item)
        for item in _iter_mapping_entries(
            report.get("technical_integrity_failures")
        )
    ]
    if not bool(report.get("technical_integrity_failed")) and not failures:
        return
    first = failures[0] if failures else {}
    location = "/".join(
        value
        for value in (
            str(first.get("recording_id") or first.get("participant_id") or ""),
            str(first.get("condition") or ""),
            str(first.get("electrode") or ""),
            str(first.get("harmonic_column") or ""),
        )
        if value
    )
    suffix = f" First affected cell: {location}." if location else ""
    raise FrequencyDomainQcIntegrityError(
        "Frequency-domain QC is incomplete because one or more selected, "
        "method-computable BCA cells failed technical output-integrity "
        f"validation.{suffix} Reprocess the affected condition workbook(s) "
        "before continuing."
    )


def is_roi_frequency_qc_entry(row: Mapping[str, object]) -> bool:
    """Identify retired ROI review evidence, including historical decision rows."""

    if str(row.get("roi") or "").strip():
        return True
    if any(
        str(row.get(key) or "").strip().casefold().replace("-", "_")
        in {"roi", "condition_roi", "participant_condition_roi", "recording_condition_roi"}
        for key in ("scope", "decision_scope", "target_type", "target_scope")
    ):
        return True
    metric = str(row.get("metric") or "").strip().casefold()
    finding_type = str(row.get("finding_type") or "").strip().casefold()
    if "roi" in metric.split("_") or "roi" in finding_type.split("_"):
        return True
    if (
        finding_type == "cohort_relative_summed_bca_context"
        and not str(row.get("electrode") or "").strip()
    ):
        return True
    evidence = row.get("evidence")
    return isinstance(evidence, Mapping) and is_roi_frequency_qc_entry(
        {"electrode": row.get("electrode"), **evidence}
    )


def validate_frequency_domain_qc_review_decisions(
    report: Mapping[str, object],
    decisions: Mapping[str, object] | Sequence[Mapping[str, object]] | None,
) -> tuple[dict[str, object], ...]:
    """Bind explicit GUI choices to every current QC-17 machine finding."""

    raw_findings = _iter_mapping_entries(
        report.get("review_findings")
        if report.get("review_findings") is not None
        else report.get("flags")
    )
    retired_fingerprints = {
        str(item.get("finding_fingerprint") or "")
        for item in raw_findings if is_roi_frequency_qc_entry(item)
    }
    raw_findings = [item for item in raw_findings if not is_roi_frequency_qc_entry(item)]
    flags = {
        str(item.get("finding_fingerprint") or ""): dict(item)
        for item in raw_findings
        if str(item.get("finding_fingerprint") or "")
    }
    if len(flags) != len(raw_findings):
        raise ValueError(
            "Every summed-BCA review finding requires a unique evidence fingerprint."
        )
    raw_by_fingerprint: dict[str, Mapping[str, object]] = {}
    if isinstance(decisions, Mapping):
        for raw_key, raw_value in decisions.items():
            key = str(raw_key or "").strip()
            if isinstance(raw_value, Mapping):
                row = dict(raw_value)
            else:
                row = {"decision": raw_value}
            row.setdefault("finding_fingerprint", key)
            raw_by_fingerprint[key] = row
    elif isinstance(decisions, Sequence) and not isinstance(
        decisions,
        (str, bytes),
    ):
        for raw_value in decisions:
            if not isinstance(raw_value, Mapping):
                raise ValueError("Summed-BCA review decisions must be mappings.")
            key = str(raw_value.get("finding_fingerprint") or "").strip()
            if not key or key in raw_by_fingerprint:
                raise ValueError(
                    "Summed-BCA review decisions contain a blank or duplicate finding fingerprint."
                )
            raw_by_fingerprint[key] = raw_value
    elif decisions is not None:
        raise ValueError("Summed-BCA review decisions must be a mapping or list.")

    if any(
        key in retired_fingerprints or is_roi_frequency_qc_entry(row)
        for key, row in raw_by_fingerprint.items()
    ):
        raise ValueError(
            "ROI-level summed-BCA review is no longer supported; regenerate the electrode review."
        )
    if not bool(report.get("screening_enabled", True)):
        if raw_by_fingerprint:
            raise ValueError(
                "Summed-BCA review decisions cannot be submitted while screening is disabled."
            )
        return ()
    unknown = sorted(set(raw_by_fingerprint).difference(flags))
    missing = sorted(set(flags).difference(raw_by_fingerprint))
    if unknown:
        raise ValueError(
            "Summed-BCA review contains stale or unknown findings: "
            + ", ".join(unknown[:5])
        )
    if missing:
        raise ValueError(
            "Choose an explicit decision for every summed-BCA finding before continuing."
        )

    scope = str(report.get("identity_scope") or "participant").strip().casefold()
    normalized: list[dict[str, object]] = []
    for finding_fingerprint, finding in flags.items():
        submitted = raw_by_fingerprint[finding_fingerprint]
        decision = str(submitted.get("decision") or "").strip().casefold()
        if decision not in REVIEW_DECISIONS:
            raise ValueError(
                "Choose Retain, an enabled electrode repair, or a whole condition, "
                "recording or participant exclusion. Electrode and ROI exclusions "
                "are no longer supported."
            )
        if decision == DECISION_EXCLUDE_RECORDING and scope != "recording":
            raise ValueError(
                "A recording exclusion requires a recording-scoped QC report."
            )
        reason = (
            "No reason provided"
            if decision == DECISION_RETAIN
            else str(submitted.get("reason") or "").strip() or "No reason provided"
        )
        participant_id = _normalize_participant_id(finding.get("participant_id"))
        recording_id = _normalize_recording_id(finding.get("recording_id"))
        condition = str(finding.get("condition") or "").strip()
        electrode = _normalize_electrode(finding.get("electrode"))
        roi = str(finding.get("roi") or "").strip()
        if not participant_id or not condition or not electrode:
            raise ValueError(
                "Summed-BCA finding identity is incomplete; regenerate the review."
            )
        if decision == DECISION_INTERPOLATE_CONDITION_ELECTRODE:
            if report.get("condition_specific_interpolation_enabled") is not True:
                raise ValueError("Enable experimental condition-specific interpolation in Settings first.")
            if not electrode or roi:
                raise ValueError("Condition-specific interpolation requires an electrode-level finding, not an ROI.")
            if scope == "recording" and not recording_id:
                raise ValueError("A recording-scoped electrode repair requires the exact recording ID.")
            if submitted.get("artifact_confirmed") is not True:
                raise ValueError("Confirm an electrode artifact before requesting interpolation; a large BCA alone is insufficient.")
        row: dict[str, object] = {
            "version": FREQUENCY_DOMAIN_QC_DECISION_VERSION,
            "finding_fingerprint": finding_fingerprint,
            "analysis_fingerprint": str(report.get("analysis_fingerprint") or ""),
            "decision": decision,
            "decision_scope": (
                "recording_condition_electrode"
                if scope == "recording"
                else "participant_condition_electrode"
            ),
            "participant_id": participant_id,
            "recording_id": recording_id,
            "session_id": str(finding.get("session_id") or ""),
            "visit_index": finding.get("visit_index"),
            "condition": condition,
            "electrode": electrode,
            "roi": roi,
            "reason": reason,
            "source": "explicit_gui_review",
            "outcome_informed": decision != DECISION_RETAIN,
            "replaces_decision_fingerprint": str(
                finding.get("replaces_decision_fingerprint") or ""
            ),
            "evidence": {
                key: _json_safe(finding.get(key))
                for key in (
                    "finding_type",
                    "summed_bca_uv",
                    "abs_summed_bca_uv",
                    "severity",
                    "band_crossed",
                    "selected_harmonics_hz",
                    "selected_harmonic_count",
                    "selection_fingerprint",
                    "harmonic_selection_fingerprint",
                    "frequency_protocol_fingerprint",
                    "expected_analyzed_oddball_cycles",
                    "analyzed_duration_seconds",
                    "independent_qc",
                    "independent_qc_status",
                    "independent_qc_authority",
                    "independent_qc_fingerprint",
                    "metric",
                    "value_uv",
                    "robust_center_uv",
                    "robust_spread_uv",
                    "robust_spread_method",
                    "robust_score",
                    "threshold_used",
                    "absolute_floor_used_uv",
                    "peak_harmonic_hz",
                    "peak_signed_roi_mean_uv",
                    "roi_definition_fingerprint",
                    "cohort_fingerprint",
                )
            },
        }
        if decision == DECISION_INTERPOLATE_CONDITION_ELECTRODE:
            row["artifact_confirmed"] = True
        row["decision_fingerprint"] = _hash_payload(row)
        normalized.append(row)

    _require_consistent_broad_decisions(normalized)
    return tuple(
        sorted(
            normalized,
            key=lambda item: str(item["finding_fingerprint"]),
        )
    )


def _require_consistent_broad_decisions(
    decisions: Sequence[Mapping[str, object]],
) -> None:
    by_recording: dict[str, set[str]] = defaultdict(set)
    by_participant: dict[str, set[str]] = defaultdict(set)
    by_condition: dict[tuple[str, str], set[str]] = defaultdict(set)
    by_electrode: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for item in decisions:
        decision = str(item.get("decision") or "")
        participant_id = _normalize_participant_id(item.get("participant_id"))
        recording_id = _normalize_recording_id(item.get("recording_id"))
        if participant_id:
            by_participant[participant_id].add(decision)
        if recording_id:
            by_recording[recording_id].add(decision)
        identity = recording_id or participant_id
        condition = str(item.get("condition") or "")
        if identity and condition:
            by_condition[(identity, condition)].add(decision)
            electrode = _normalize_electrode(item.get("electrode"))
            if electrode:
                by_electrode[(identity, condition, electrode)].add(decision)
    for (identity, condition, electrode), values in by_electrode.items():
        if DECISION_INTERPOLATE_CONDITION_ELECTRODE in values and len(values) != 1:
            raise ValueError(
                f"Conflicting repair decisions for {identity}/{condition}/{electrode}. "
                "Apply the same choice to each finding for this electrode in this condition."
            )
    for label, grouped, broad_decision in (
        ("recording", by_recording, DECISION_EXCLUDE_RECORDING),
        ("participant", by_participant, DECISION_EXCLUDE_PARTICIPANT),
    ):
        for identity, values in grouped.items():
            if broad_decision in values and values != {broad_decision}:
                raise ValueError(
                    f"A whole-{label} exclusion conflicts with another decision "
                    f"for {identity}. Apply the same broader choice to all of its findings."
                )
    for (identity, condition), values in by_condition.items():
        if DECISION_EXCLUDE_CONDITION in values and values != {
            DECISION_EXCLUDE_CONDITION
        }:
            raise ValueError(
                "A whole-condition exclusion conflicts with another decision "
                f"for {identity}/{condition}. Apply the same condition choice "
                "to all of its findings."
            )


def apply_frequency_domain_qc_decision(
    project_root: str | Path,
    report: Mapping[str, object],
    *,
    review_decisions: Mapping[str, object]
    | Sequence[Mapping[str, object]]
    | None = None,
    manual_participant_reasons: Mapping[str, str] | None = None,
    manual_recording_reasons: Mapping[str, str] | None = None,
) -> dict[str, object]:
    """Persist a reviewed QC decision and write the human-readable report."""

    require_frequency_domain_qc_complete(report)
    normalized_review_decisions = [
        dict(item)
        for item in validate_frequency_domain_qc_review_decisions(
            report,
            review_decisions,
        )
    ]
    resolved_recording_decisions = resolve_frequency_qc_recording_decisions(
        report,
        manual_recording_reasons,
    )
    root = Path(project_root).resolve()
    manifest_path = root / "project.json"
    manifest = _read_manifest(manifest_path)
    repair_decisions = [
        item for item in normalized_review_decisions
        if item.get("decision") == DECISION_INTERPOLATE_CONDITION_ELECTRODE
    ]
    if repair_decisions:
        if Path(str(report.get("project_root") or "")).resolve() != root:
            raise ValueError("The electrode-repair review belongs to a different project. Reopen QC.")
        if not _source_workbooks_are_current(root, report.get("source_workbooks")):
            raise ValueError("The reviewed condition outputs changed. Regenerate QC before requesting a repair.")
        if not normalize_experimental_qc_settings(
            manifest.get("experimental_qc")
        ).condition_specific_interpolation_enabled:
            raise ValueError("Condition-specific interpolation was disabled. Reopen the QC review.")
        from Main_App.processing.condition_interpolation_state import (
            queue_condition_interpolation_decisions,
        )

        queue_condition_interpolation_decisions(manifest, repair_decisions, report)
    state = _metadata_from_manifest(manifest)
    now = _now_utc_iso()
    repeated_session = str(report.get("identity_scope") or "") == "recording"
    replaced_decision_fingerprints = {
        str(item.get("replaces_decision_fingerprint") or "")
        for item in normalized_review_decisions
        if str(item.get("replaces_decision_fingerprint") or "")
    }
    current_finding_fingerprints = {
        str(item.get("finding_fingerprint") or "")
        for item in normalized_review_decisions
    }
    preserved_review_decisions = [
        item
        for item in _review_decisions_from_state(state)
        if str(item.get("decision") or "") in _BROAD_EXCLUSION_DECISIONS
        and str(item.get("decision_fingerprint") or "")
        not in replaced_decision_fingerprints
        and str(item.get("finding_fingerprint") or "")
        not in current_finding_fingerprints
    ]
    combined_review_decisions = _merge_review_decision_rows(
        preserved_review_decisions,
        normalized_review_decisions,
    )
    legacy_machine_suggestions = _merge_legacy_machine_suggestions(
        state,
        report.get("legacy_machine_suggestions"),
    )
    existing_manual = _manual_entries_from_state(state)
    manual_by_pid = {entry["participant_id"]: dict(entry) for entry in existing_manual}
    reviewed_participant_exclusions = {
        _normalize_participant_id(item.get("participant_id"))
        for item in normalized_review_decisions
        if item.get("decision") == DECISION_EXCLUDE_PARTICIPANT
    }
    for raw_pid, raw_reason in (manual_participant_reasons or {}).items():
        pid = _normalize_participant_id(raw_pid)
        if not pid or pid in reviewed_participant_exclusions:
            continue
        reason = str(raw_reason or "").strip() or "No reason provided"
        previous = manual_by_pid.get(pid, {})
        manual_by_pid[pid] = {
            "participant_id": pid,
            "reason": reason,
            "source": "manual_qc_review",
            "added_at": str(previous.get("added_at") or now),
            "updated_at": now,
        }
    manual_entries = _normalize_manual_entries(
        sorted(manual_by_pid.values(), key=lambda item: item["participant_id"])
    )
    existing_manual_recordings = _manual_recording_entries_from_state(state)
    manual_by_recording = {
        str(entry["recording_id"]).casefold(): dict(entry)
        for entry in existing_manual_recordings
    }
    reviewed_recording_exclusions = {
        _normalize_recording_id(item.get("recording_id"))
        for item in normalized_review_decisions
        if item.get("decision") == DECISION_EXCLUDE_RECORDING
    }
    for decision in resolved_recording_decisions:
        recording_id = str(decision.identity.recording_id)
        if _normalize_recording_id(recording_id) in reviewed_recording_exclusions:
            continue
        reason = str(decision.reason or "").strip() or "No reason provided"
        recording_key = recording_id.casefold()
        previous = manual_by_recording.get(recording_key, {})
        manual_by_recording[recording_key] = {
            "recording_id": recording_id,
            "participant_id": decision.identity.participant_id,
            "session_id": str(decision.identity.session_id or ""),
            "reason": reason,
            "source": "manual_qc_review",
            "added_at": str(previous.get("added_at") or now),
            "updated_at": now,
        }
    manual_recording_entries = _normalize_manual_recording_entries(
        sorted(
            manual_by_recording.values(),
            key=lambda item: str(item["recording_id"]).casefold(),
        )
    )
    analysis_fingerprint = str(report.get("analysis_fingerprint") or "")
    identity_scope = "recording" if repeated_session else "participant"
    review_evidence = _build_review_evidence_payload(report)
    decision_fingerprint = _decision_fingerprint(
        analysis_fingerprint=analysis_fingerprint,
        auto_electrodes=(),
        auto_participants=(),
        manual_participants=manual_entries,
        auto_recording_electrodes=(() if repeated_session else None),
        auto_recordings=(() if repeated_session else None),
        manual_recordings=(manual_recording_entries if repeated_session else None),
        review_decisions=combined_review_decisions,
        identity_scope=identity_scope,
    )
    _require_non_oscillating_review_decision(
        state,
        analysis_fingerprint=analysis_fingerprint,
        decision_fingerprint=decision_fingerprint,
    )
    for decision in normalized_review_decisions:
        decision["reviewed_at"] = now
    combined_review_decisions = _merge_review_decision_rows(
        preserved_review_decisions,
        normalized_review_decisions,
    )
    report_path = _write_frequency_domain_qc_text_report(
        root,
        report=report,
        manual_participants=manual_entries,
        manual_recordings=manual_recording_entries,
        review_decisions=combined_review_decisions,
        decision_fingerprint=decision_fingerprint,
        reviewed_at=now,
    )
    update: dict[str, object] = {
            "schema_version": FREQUENCY_DOMAIN_QC_SCHEMA_VERSION,
            "method_version": str(
                report.get("method_version") or FREQUENCY_DOMAIN_QC_METHOD_VERSION
            ),
            "identity_scope": identity_scope,
            "screening_settings": dict(
                report.get("screening_settings")
                if isinstance(report.get("screening_settings"), Mapping)
                else {}
            ),
            "thresholds": dict(
                report.get("thresholds")
                if isinstance(report.get("thresholds"), Mapping)
                else {}
            ),
            "auto_participant_electrode_exclusions": [],
            "auto_participant_exclusions": [],
            "auto_recording_electrode_exclusions": [],
            "auto_recording_exclusions": [],
            "legacy_machine_suggestions": legacy_machine_suggestions,
            "manual_participant_exclusions": manual_entries,
            "manual_recording_exclusions": (
                manual_recording_entries if repeated_session else []
            ),
            "review_decisions": combined_review_decisions,
            "review_evidence": review_evidence,
            "review_complete": True,
            "downstream_outputs_stale": True,
            "last_review": {
                "reviewed_at": now,
                "analysis_fingerprint": analysis_fingerprint,
                "decision_fingerprint": decision_fingerprint,
                "evidence_fingerprint": str(
                    review_evidence["evidence_fingerprint"]
                ),
                "identity_scope": identity_scope,
                "report_path": _manifest_safe_path(root, report_path),
                "review_subject_count": int(report.get("review_subject_count") or 0),
                "screening_status": str(report.get("screening_status") or ""),
                "screening_policy_version": str(
                    report.get("screening_policy_version") or ""
                ),
            },
        }
    if repeated_session:
        update.update(
            {
                "manual_recording_exclusions": manual_recording_entries,
            }
        )
        last_review = update.get("last_review")
        if isinstance(last_review, dict):
            last_review["review_recording_count"] = int(
                report.get("review_recording_count") or 0
            )
    update["review_history"] = _updated_review_history(
        state,
        report=report,
        decision_fingerprint=decision_fingerprint,
        reviewed_at=now,
    )
    update["retired_review_decisions"] = _superseded_narrow_decisions(
        state, combined_review_decisions,
    )
    state.update(update)
    _set_metadata_in_manifest(manifest, state)
    _write_manifest_if_changed(manifest_path, manifest)
    return state


def sync_frequency_domain_qc_automatic_state(
    project_root: str | Path,
    report: Mapping[str, object],
) -> dict[str, object]:
    """Persist a no-prompt/reused review without granting BCA automatic authority."""

    require_frequency_domain_qc_complete(report)
    root = Path(project_root).resolve()
    manifest_path = root / "project.json"
    manifest = _read_manifest(manifest_path)
    state = _metadata_from_manifest(manifest)
    previous_auto_electrodes = _auto_electrode_entries_from_state(state)
    previous_auto_participants = _auto_participant_entries_from_state(state)
    previous_auto_recording_electrodes = _auto_recording_electrode_entries_from_state(state)
    previous_auto_recordings = _auto_recording_entries_from_state(state)
    retired_roi_authority_removed = any(
        is_roi_frequency_qc_entry(row)
        and row.get("decision") in _BROAD_EXCLUSION_DECISIONS
        for row in _review_decisions_from_state(state, include_retired_roi=True)
    )
    legacy_authority_removed = bool(
        previous_auto_electrodes
        or previous_auto_participants
        or previous_auto_recording_electrodes
        or previous_auto_recordings
    )
    legacy_machine_suggestions = _merge_legacy_machine_suggestions(
        state,
        report.get("legacy_machine_suggestions"),
    )
    now = _now_utc_iso()
    current_decisions = _current_review_decisions(
        report_flags=_iter_mapping_entries(
            report.get("review_findings")
            if report.get("review_findings") is not None
            else report.get("flags")
        ),
        state=state,
    )
    current_decisions = _merge_review_decision_rows(
        [
            dict(item)
            for item in _iter_mapping_entries(
                report.get("active_review_decisions")
            )
            if not is_roi_frequency_qc_entry(item)
        ],
        current_decisions,
    )
    repeated_session = str(report.get("identity_scope") or "") == "recording"
    identity_scope = "recording" if repeated_session else "participant"
    manual_entries = _manual_entries_from_state(state)
    manual_recording_entries = _manual_recording_entries_from_state(state)
    analysis_fingerprint = str(report.get("analysis_fingerprint") or "")
    decision_fingerprint = _decision_fingerprint(
        analysis_fingerprint=analysis_fingerprint,
        auto_electrodes=(),
        auto_participants=(),
        manual_participants=manual_entries,
        auto_recording_electrodes=(() if repeated_session else None),
        auto_recordings=(() if repeated_session else None),
        manual_recordings=(manual_recording_entries if repeated_session else None),
        review_decisions=current_decisions,
        identity_scope=identity_scope,
    )
    review_evidence = _build_review_evidence_payload(report)
    previous_last_review = (
        dict(state.get("last_review"))
        if isinstance(state.get("last_review"), Mapping)
        else {}
    )
    if report.get("review_reused"):
        previous_evidence = _validated_review_evidence_from_state(root, state)
        if previous_evidence is not None and (
            previous_evidence.get("analysis_fingerprint") == analysis_fingerprint
        ):
            # Preserve what the user actually reviewed. In particular, an older
            # receipt may bind its identity to the original cache annotations.
            review_evidence = previous_evidence
    update: dict[str, object] = {
        "schema_version": FREQUENCY_DOMAIN_QC_SCHEMA_VERSION,
        "method_version": str(
            report.get("method_version") or FREQUENCY_DOMAIN_QC_METHOD_VERSION
        ),
        "identity_scope": identity_scope,
        "screening_settings": dict(
            report.get("screening_settings")
            if isinstance(report.get("screening_settings"), Mapping)
            else {}
        ),
        "thresholds": dict(
            report.get("thresholds")
            if isinstance(report.get("thresholds"), Mapping)
            else {}
        ),
        "auto_participant_electrode_exclusions": [],
        "auto_participant_exclusions": [],
        "auto_recording_electrode_exclusions": [],
        "auto_recording_exclusions": [],
        "legacy_machine_suggestions": legacy_machine_suggestions,
        "manual_participant_exclusions": manual_entries,
        "manual_recording_exclusions": (
            manual_recording_entries if repeated_session else []
        ),
        "review_decisions": current_decisions,
        "review_evidence": review_evidence,
        "review_complete": not bool(report.get("review_required")),
        "last_review": {
            "reviewed_at": str(previous_last_review.get("reviewed_at") or now),
            "synced_at": now,
            "analysis_fingerprint": analysis_fingerprint,
            "decision_fingerprint": decision_fingerprint,
            "evidence_fingerprint": str(
                review_evidence["evidence_fingerprint"]
            ),
            "identity_scope": identity_scope,
            "report_path": str(previous_last_review.get("report_path") or ""),
            "review_subject_count": int(
                report.get("review_subject_count") or 0
            ),
            "review_recording_count": int(
                report.get("review_recording_count") or 0
            ),
            "screening_status": str(report.get("screening_status") or ""),
            "screening_policy_version": str(
                report.get("screening_policy_version") or ""
            ),
        },
        "last_automatic_qc": {
            "reviewed_at": now,
            "analysis_fingerprint": analysis_fingerprint,
            "review_required": bool(report.get("review_required")),
            "review_reused": bool(report.get("review_reused")),
            "screening_status": str(report.get("screening_status") or ""),
            "authority": "review_only",
        }
    }
    update["retired_review_decisions"] = _superseded_narrow_decisions(state, current_decisions)
    state.update(update)
    if legacy_authority_removed or retired_roi_authority_removed:
        state["downstream_outputs_stale"] = True
        state["stale_reason"] = (
            "Retired ROI-level summed-BCA review exclusions are no longer applied; "
            "regenerate downstream outputs."
            if retired_roi_authority_removed else
            "Legacy summed-BCA automatic exclusions were converted to review-only suggestions."
        )
        state["stale_at"] = now
    _set_metadata_in_manifest(manifest, state)
    _write_manifest_if_changed(manifest_path, manifest)
    return state


def mark_frequency_domain_outputs_stale(
    project_root: str | Path,
    *,
    reason: str,
) -> None:
    root = Path(project_root).resolve()
    manifest_path = root / "project.json"
    manifest = _read_manifest(manifest_path)
    state = _metadata_from_manifest(manifest)
    state["downstream_outputs_stale"] = True
    state["stale_reason"] = str(reason)
    state["stale_at"] = _now_utc_iso()
    _set_metadata_in_manifest(manifest, state)
    _write_manifest_if_changed(manifest_path, manifest)


def mark_frequency_domain_outputs_current(project_root: str | Path) -> None:
    from Main_App.processing.condition_interpolation_state import (
        require_no_pending_condition_interpolation,
    )

    require_no_pending_condition_interpolation(project_root)
    root = Path(project_root).resolve()
    manifest_path = root / "project.json"
    manifest = _read_manifest(manifest_path)
    state = _metadata_from_manifest(manifest)
    if not state:
        return
    state["downstream_outputs_stale"] = False
    state.pop("stale_reason", None)
    state["last_outputs_refreshed_at"] = _now_utc_iso()
    _set_metadata_in_manifest(manifest, state)
    _write_manifest_if_changed(manifest_path, manifest)


def is_frequency_domain_output_stale(project_root: str | Path | None) -> bool:
    if project_root in (None, ""):
        return False
    state = load_frequency_domain_qc_state(project_root)
    return bool(state.get("downstream_outputs_stale", False))


def load_frequency_domain_qc_state(project_root: str | Path | None) -> dict[str, object]:
    if project_root in (None, ""):
        return {}
    manifest = _read_manifest(Path(project_root).resolve() / "project.json")
    state = _metadata_from_manifest(manifest)
    if any(
        is_roi_frequency_qc_entry(row)
        and row.get("decision") != DECISION_RETAIN
        for row in _review_decisions_from_state(state, include_retired_roi=True)
    ):
        state["downstream_outputs_stale"] = True
        state["stale_reason"] = (
            "Retired ROI-level summed-BCA review exclusions are no longer applied; "
            "regenerate downstream outputs."
        )
    if any(row.get("decision") not in REVIEW_DECISIONS
           for row in _review_decisions_from_state(state)):
        # Surface the retired policy through the common tool-readiness gate,
        # without rewriting project history merely when a project is opened.
        state["downstream_outputs_stale"] = True
        state["stale_reason"] = (
            "Saved electrode exclusions need a new QC decision. "
            "These exclusions are no longer applied; resume post-processing review."
        )
    return state


def active_frequency_domain_exclusions(
    project_root: str | Path | None,
) -> FrequencyDomainExclusions:
    state = load_frequency_domain_qc_state(project_root)
    return frequency_domain_exclusions_from_state(state)


def frequency_domain_exclusions_from_state(
    state: Mapping[str, object],
) -> FrequencyDomainExclusions:
    """Resolve an already-read QC state using the same active-decision guards."""

    return _frequency_domain_exclusions_from_rows(
        state=state,
        decisions=_review_decisions_from_state(state),
        manual_entries=_manual_entries_from_state(state),
        manual_recording_entries=_manual_recording_entries_from_state(state),
    )


def _frequency_domain_exclusions_from_rows(
    *,
    state: Mapping[str, object],
    decisions: Sequence[Mapping[str, object]],
    manual_entries: Sequence[Mapping[str, object]],
    manual_recording_entries: Sequence[Mapping[str, object]],
) -> FrequencyDomainExclusions:
    manual_participants = {
        _normalize_participant_id(entry.get("participant_id"))
        for entry in manual_entries
        if _normalize_participant_id(entry.get("participant_id"))
    }
    manual_recordings = {
        _normalize_recording_id(entry.get("recording_id"))
        for entry in manual_recording_entries
        if _normalize_recording_id(entry.get("recording_id"))
    }
    participant_conditions: set[tuple[str, str]] = set()
    recording_conditions: set[tuple[str, str]] = set()
    participant_condition_electrodes: dict[tuple[str, str], set[str]] = defaultdict(set)
    recording_condition_electrodes: dict[tuple[str, str], set[str]] = defaultdict(set)
    for decision in decisions:
        if is_roi_frequency_qc_entry(decision):
            continue
        action = str(decision.get("decision") or "")
        participant_id = _normalize_participant_id(decision.get("participant_id"))
        recording_id = _normalize_recording_id(decision.get("recording_id"))
        condition = str(decision.get("condition") or "").strip()
        if action == DECISION_EXCLUDE_PARTICIPANT and participant_id:
            manual_participants.add(participant_id)
        elif action == DECISION_EXCLUDE_RECORDING and recording_id:
            manual_recordings.add(recording_id)
        elif action == DECISION_EXCLUDE_CONDITION:
            if recording_id and condition:
                recording_conditions.add((recording_id, condition))
            elif participant_id and condition:
                participant_conditions.add((participant_id, condition))
    return FrequencyDomainExclusions(
        excluded_participants=frozenset(manual_participants),
        auto_excluded_participants=frozenset(),
        manual_excluded_participants=frozenset(manual_participants),
        auto_excluded_electrodes_by_participant={},
        downstream_outputs_stale=bool(state.get("downstream_outputs_stale", False)),
        excluded_recordings=frozenset(manual_recordings),
        auto_excluded_recordings=frozenset(),
        manual_excluded_recordings=frozenset(manual_recordings),
        auto_excluded_electrodes_by_recording={},
        excluded_participant_conditions=frozenset(participant_conditions),
        excluded_recording_conditions=frozenset(recording_conditions),
        excluded_electrodes_by_participant_condition={
            key: frozenset(sorted(electrodes))
            for key, electrodes in participant_condition_electrodes.items()
        },
        excluded_electrodes_by_recording_condition={
            key: frozenset(sorted(electrodes))
            for key, electrodes in recording_condition_electrodes.items()
        },
    )


def _raw_state_rows(
    state: Mapping[str, object],
    key: str,
) -> list[Mapping[str, object]] | None:
    raw = state.get(key, [])
    if not isinstance(raw, list):
        return None
    if any(not isinstance(item, Mapping) for item in raw):
        return None
    return [item for item in raw if isinstance(item, Mapping)]


def _manual_state_rows_are_canonical(
    state: Mapping[str, object],
    *,
    key: str,
    normalizer: Callable[[object], list[dict[str, object]]],
) -> tuple[bool, list[dict[str, object]]]:
    raw = _raw_state_rows(state, key)
    if raw is None:
        return False, []
    normalized = normalizer(raw)
    return _json_safe(raw) == _json_safe(normalized), normalized


def _review_decision_state_rows_are_hash_valid(
    state: Mapping[str, object],
) -> tuple[bool, list[dict[str, object]]]:
    raw = _raw_state_rows(state, "review_decisions")
    if raw is None:
        return False, []
    raw = [row for row in raw if not is_roi_frequency_qc_entry(row)]
    normalized = _review_decisions_from_state(state)
    decision_fingerprints = [
        str(item.get("decision_fingerprint") or "") for item in normalized
    ]
    finding_fingerprints = [
        str(item.get("finding_fingerprint") or "") for item in normalized
    ]
    valid = (
        len(normalized) == len(raw)
        and len(decision_fingerprints) == len(set(decision_fingerprints))
        and len(finding_fingerprints) == len(set(finding_fingerprints))
        and all(decision_fingerprints)
        and all(finding_fingerprints)
        and all(item.get("decision") in REVIEW_DECISIONS for item in normalized)
    )
    if not raw:
        valid = True
    return valid, normalized


def resolve_frequency_qc_coverage_decisions(
    project_root: str | Path | None,
) -> FrequencyDomainCoverageDecisions:
    """Return current explicit QC-17 decisions for final ROI/release gates."""

    root = Path(project_root).resolve() if project_root not in (None, "") else None
    state = load_frequency_domain_qc_state(project_root)
    decisions_valid, decision_rows = _review_decision_state_rows_are_hash_valid(
        state
    )
    participant_rows_valid, manual_entries = _manual_state_rows_are_canonical(
        state,
        key="manual_participant_exclusions",
        normalizer=_normalize_manual_entries,
    )
    recording_rows_valid, manual_recording_entries = (
        _manual_state_rows_are_canonical(
            state,
            key="manual_recording_exclusions",
            normalizer=_normalize_manual_recording_entries,
        )
    )
    last_review = state.get("last_review")
    saved_decision_fingerprint = (
        str(last_review.get("decision_fingerprint") or "").strip()
        if isinstance(last_review, Mapping)
        else ""
    )
    saved_analysis_fingerprint = (
        str(last_review.get("analysis_fingerprint") or "").strip()
        if isinstance(last_review, Mapping)
        else ""
    )
    identity_scope = str(state.get("identity_scope") or "").strip().casefold()
    last_identity_scope = (
        str(last_review.get("identity_scope") or "").strip().casefold()
        if isinstance(last_review, Mapping)
        else ""
    )
    identity_valid = (
        identity_scope in {"participant", "recording"}
        and identity_scope == last_identity_scope
    )
    if identity_scope == "participant" and manual_recording_entries:
        recording_rows_valid = False
    computed_decision_fingerprint = ""
    if saved_analysis_fingerprint and identity_valid:
        computed_decision_fingerprint = _decision_fingerprint(
            analysis_fingerprint=saved_analysis_fingerprint,
            auto_electrodes=(),
            auto_participants=(),
            manual_participants=manual_entries,
            auto_recording_electrodes=(
                () if identity_scope == "recording" else None
            ),
            auto_recordings=(() if identity_scope == "recording" else None),
            manual_recordings=(
                manual_recording_entries
                if identity_scope == "recording"
                else None
            ),
            review_decisions=decision_rows,
            identity_scope=identity_scope,
        )
    evidence = (
        _validated_review_evidence_from_state(root, state)
        if root is not None
        else None
    )
    review_complete = bool(
        state.get("review_complete")
        and decisions_valid
        and participant_rows_valid
        and recording_rows_valid
        and identity_valid
        and evidence is not None
        and saved_decision_fingerprint
        and computed_decision_fingerprint == saved_decision_fingerprint
    )
    exclusions = (
        _frequency_domain_exclusions_from_rows(
            state=state,
            decisions=decision_rows,
            manual_entries=manual_entries,
            manual_recording_entries=manual_recording_entries,
        )
        if review_complete
        else _frequency_domain_exclusions_from_rows(
            state=state,
            decisions=(),
            manual_entries=(),
            manual_recording_entries=(),
        )
    )
    decisions = tuple(dict(item) for item in decision_rows) if review_complete else ()
    return FrequencyDomainCoverageDecisions(
        decision_fingerprint=(
            computed_decision_fingerprint if review_complete else ""
        ),
        review_complete=review_complete,
        excluded_participants=exclusions.excluded_participants,
        excluded_recordings=exclusions.excluded_recordings,
        excluded_participant_conditions=exclusions.excluded_participant_conditions,
        excluded_recording_conditions=exclusions.excluded_recording_conditions,
        excluded_electrodes_by_participant_condition=(
            exclusions.excluded_electrodes_by_participant_condition
        ),
        excluded_electrodes_by_recording_condition=(
            exclusions.excluded_electrodes_by_recording_condition
        ),
        reviewed_decisions=decisions,
    )


def filter_frequency_domain_subjects(
    project_root: str | Path | None,
    subjects: Sequence[str],
    subject_data: Mapping[str, Mapping[str, str]],
) -> tuple[list[str], dict[str, dict[str, str]], list[str]]:
    exclusions = active_frequency_domain_exclusions(project_root)
    excluded = {pid.upper() for pid in exclusions.excluded_participants}
    filtered_subjects = [str(pid) for pid in subjects if str(pid).upper() not in excluded]
    filtered_data = {
        pid: {
            condition: path
            for condition, path in dict(subject_data.get(pid, {})).items()
            if (_normalize_participant_id(pid), str(condition))
            not in exclusions.excluded_participant_conditions
        }
        for pid in filtered_subjects
        if subject_data.get(pid)
    }
    filtered_subjects = [pid for pid in filtered_subjects if filtered_data.get(pid)]
    removed = sorted(str(pid) for pid in subjects if str(pid).upper() in excluded)
    return filtered_subjects, filtered_data, removed


def filter_frequency_domain_recordings(
    project_root: str | Path | None,
    recording_ids: Sequence[str],
    recording_data: Mapping[str, Mapping[str, str]],
    *,
    recording_participant_ids: Mapping[str, str],
) -> tuple[list[str], dict[str, dict[str, str]], list[str]]:
    """Apply recording and participant exclusions without collapsing visits."""

    exclusions = active_frequency_domain_exclusions(project_root)
    excluded_recordings = {
        recording_id.casefold() for recording_id in exclusions.excluded_recordings
    }
    excluded_participants = {
        participant_id.casefold()
        for participant_id in exclusions.excluded_participants
    }
    participant_lookup = {
        str(recording_id).casefold(): str(participant_id)
        for recording_id, participant_id in recording_participant_ids.items()
    }
    filtered_recordings = [
        str(recording_id)
        for recording_id in recording_ids
        if str(recording_id).casefold() not in excluded_recordings
        and participant_lookup.get(str(recording_id).casefold(), "").casefold()
        not in excluded_participants
    ]
    filtered_data = {
        recording_id: {
            condition: path
            for condition, path in dict(recording_data.get(recording_id, {})).items()
            if (_normalize_recording_id(recording_id), str(condition))
            not in exclusions.excluded_recording_conditions
        }
        for recording_id in filtered_recordings
        if recording_data.get(recording_id)
    }
    filtered_recordings = [
        recording_id
        for recording_id in filtered_recordings
        if filtered_data.get(recording_id)
    ]
    removed = sorted(
        str(recording_id)
        for recording_id in recording_ids
        if str(recording_id) not in filtered_recordings
    )
    return filtered_recordings, filtered_data, removed


def frequency_domain_excluded_electrodes_for_subject(
    project_root: str | Path | None,
    participant_id: object,
    condition: object | None = None,
) -> frozenset[str]:
    exclusions = active_frequency_domain_exclusions(project_root)
    pid = _normalize_participant_id(participant_id)
    if condition in (None, ""):
        return frozenset()
    return exclusions.excluded_electrodes_by_participant_condition.get(
        (pid, str(condition).strip()),
        frozenset(),
    )


def frequency_domain_excluded_electrodes_for_recording(
    project_root: str | Path | None,
    recording_id: object,
    condition: object | None = None,
) -> frozenset[str]:
    exclusions = active_frequency_domain_exclusions(project_root)
    normalized = _normalize_recording_id(recording_id)
    if condition in (None, ""):
        return frozenset()
    return exclusions.excluded_electrodes_by_recording_condition.get(
        (normalized, str(condition).strip()),
        frozenset(),
    )


def clear_manual_frequency_domain_participant_exclusions(
    project_root: str | Path,
    participant_ids: Iterable[object],
) -> list[str]:
    root = Path(project_root).resolve()
    manifest_path = root / "project.json"
    manifest = _read_manifest(manifest_path)
    state = _metadata_from_manifest(manifest)
    to_clear = {
        _normalize_participant_id(pid)
        for pid in participant_ids
        if _normalize_participant_id(pid)
    }
    if not to_clear:
        return []
    existing = _manual_entries_from_state(state)
    retained = [
        entry for entry in existing if entry.get("participant_id") not in to_clear
    ]
    cleared = sorted(
        entry["participant_id"]
        for entry in existing
        if entry.get("participant_id") in to_clear
    )
    if not cleared:
        return []
    state["manual_participant_exclusions"] = retained
    state["downstream_outputs_stale"] = True
    state["stale_reason"] = "Manual frequency-domain exclusions changed."
    state["stale_at"] = _now_utc_iso()
    state.pop("last_review", None)
    _set_metadata_in_manifest(manifest, state)
    _write_manifest_if_changed(manifest_path, manifest)
    return cleared


def clear_manual_frequency_domain_recording_exclusions(
    project_root: str | Path,
    recording_ids: Iterable[object],
) -> list[str]:
    root = Path(project_root).resolve()
    manifest_path = root / "project.json"
    manifest = _read_manifest(manifest_path)
    state = _metadata_from_manifest(manifest)
    to_clear = {
        _normalize_recording_id(recording_id)
        for recording_id in recording_ids
        if _normalize_recording_id(recording_id)
    }
    if not to_clear:
        return []
    existing = _manual_recording_entries_from_state(state)
    retained = [
        entry
        for entry in existing
        if str(entry.get("recording_id") or "").upper() not in to_clear
    ]
    cleared = sorted(
        str(entry["recording_id"])
        for entry in existing
        if str(entry.get("recording_id") or "").upper() in to_clear
    )
    if not cleared:
        return []
    state["manual_recording_exclusions"] = retained
    state["downstream_outputs_stale"] = True
    state["stale_reason"] = "Manual frequency-domain recording exclusions changed."
    state["stale_at"] = _now_utc_iso()
    state.pop("last_review", None)
    _set_metadata_in_manifest(manifest, state)
    _write_manifest_if_changed(manifest_path, manifest)
    return cleared


def thresholds_summary_lines() -> list[str]:
    thresholds = DEFAULT_FREQUENCY_DOMAIN_QC_THRESHOLDS
    return [
        "Experimental summed-BCA screening is review-only; no threshold "
        "automatically excludes data.",
        f"Warning review band: abs(summed BCA) > {thresholds.warning_summed_bca_uv:g} uV",
        (
            "Concentrated review: "
            f"{thresholds.concentrated_review_flagged_cells} or more flagged cells "
            "per recording or participant"
        ),
        (
            "Strong review band: "
            f"abs(summed BCA) > {thresholds.strong_warning_summed_bca_uv:g} uV"
        ),
        (
            "Extreme review band: "
            f"abs(summed BCA) > {thresholds.extreme_review_summed_bca_uv:g} uV"
        ),
        (
            "Broad review: at least "
            f"{thresholds.broad_extreme_review_unique_electrodes:g} unique "
            "extreme-band electrodes"
        ),
    ]


def _provisional_harmonics(
    *,
    project_root: Path,
    subjects: list[str],
    conditions: list[str],
    subject_data: dict[str, dict[str, str]],
    rois: dict[str, list[str]],
    settings: Any,
    log_func: Callable[[str], None],
    recording_assignments: Mapping[str, Mapping[str, object]] | None = None,
    declared_session_ids: Sequence[str] | None = None,
    participant_group_ids: Mapping[str, str] | None = None,
    declared_group_ids: Sequence[str] | None = None,
    base_frequency_hz: float,
    oddball_frequency_hz: float,
    eligible_harmonic_orders: Sequence[int],
    spectral_eligibility_fingerprint: str,
    electrode_exclusions_by_subject_condition: Mapping[
        tuple[str, str], frozenset[str]
    ] | None = None,
    expected_scalp_channels_by_subject_condition: Mapping[
        tuple[str, str], Sequence[str]
    ] | None = None,
) -> tuple[tuple[float, ...], dict[str, object]]:
    from Tools.Stats.analysis.dv_policy_fixed_predefined import (
        build_fixed_harmonic_selection,
    )
    from Tools.Stats.analysis.dv_policy_group_significant import (
        build_group_significant_harmonic_selection,
    )
    from Tools.Stats.analysis.dv_policy_settings import GROUP_SIGNIFICANT_POLICY_NAME

    if settings.name == GROUP_SIGNIFICANT_POLICY_NAME:
        selection = build_group_significant_harmonic_selection(
            subjects=subjects,
            conditions=conditions,
            subject_data=subject_data,
            base_frequency_hz=base_frequency_hz,
            rois=rois,
            log_func=log_func,
            settings=settings,
            max_freq=None,
            project_root=project_root,
            recording_assignments=recording_assignments,
            declared_session_ids=declared_session_ids,
            participant_group_ids=participant_group_ids,
            declared_group_ids=declared_group_ids,
            oddball_frequency_hz=oddball_frequency_hz,
            eligible_harmonic_orders=eligible_harmonic_orders,
            spectral_eligibility_fingerprint=spectral_eligibility_fingerprint,
            electrode_exclusions_by_subject_condition=(
                electrode_exclusions_by_subject_condition
            ),
            expected_scalp_channels_by_subject_condition=(
                expected_scalp_channels_by_subject_condition
            ),
        )
        return (
            tuple(round(float(freq), 4) for freq in selection.selected_harmonics_hz),
            selection.to_metadata(),
        )

    columns = _find_first_bca_columns(subjects, conditions, subject_data)
    if not columns:
        raise RuntimeError("Frequency-domain QC could not read BCA harmonic columns.")
    selection = build_fixed_harmonic_selection(
        requested_values=settings.fixed_harmonic_frequencies_hz,
        bca_columns=columns,
        base_frequency_hz=base_frequency_hz,
        auto_exclude_base_overlaps=settings.fixed_harmonic_auto_exclude_base,
        base_overlap_tolerance_hz=settings.fixed_harmonic_base_tolerance_hz,
        matching_tolerance_hz=settings.fixed_harmonic_matching_tolerance_hz,
        input_mode=settings.fixed_harmonic_input_mode,
        upper_harmonic_index=settings.fixed_harmonic_upper_harmonic_index,
        upper_frequency_hz=settings.fixed_harmonic_upper_frequency_hz,
        oddball_frequency_hz=oddball_frequency_hz,
        eligible_harmonic_orders=eligible_harmonic_orders,
    )
    metadata = selection.to_metadata()
    metadata.update(
        {
            "frequency_protocol_base_rate_hz": float(base_frequency_hz),
            "frequency_protocol_oddball_rate_hz": float(oddball_frequency_hz),
            "eligible_harmonic_orders": [
                int(order) for order in eligible_harmonic_orders
            ],
            "spectral_eligibility_fingerprint": str(
                spectral_eligibility_fingerprint
            ),
        }
    )
    return (
        tuple(round(float(freq), 4) for freq in selection.included_frequencies_hz),
        metadata,
    )


def _collect_summed_bca_flags(
    *,
    subjects: list[str],
    conditions: list[str],
    subject_data: dict[str, dict[str, str]],
    selected_harmonics: Sequence[float],
    thresholds: FrequencyDomainQcThresholds,
    log_func: Callable[[str], None],
    recording_assignments: Mapping[str, Mapping[str, object]] | None = None,
    screening_enabled: bool = True,
    selected_harmonics_metadata: Mapping[str, object] | None = None,
    excluded_electrodes_by_subject_condition: Mapping[
        tuple[str, str], frozenset[str]
    ] | None = None,
) -> _SummedBcaInspection:
    from Main_App.io import (
        MissingXlsxColumnsError,
        read_xlsx_sheet_selected_columns,
    )

    columns = [f"{float(freq):.4f}_Hz" for freq in selected_harmonics]
    flags: list[dict[str, object]] = []
    integrity_failures: list[dict[str, object]] = []
    unavailable_by_method: list[dict[str, object]] = []
    for subject in subjects:
        for condition in conditions:
            file_path = subject_data.get(subject, {}).get(condition)
            if not file_path or not Path(file_path).exists():
                continue
            try:
                frame = read_xlsx_sheet_selected_columns(
                    file_path,
                    sheet_name="BCA (uV)",
                    required_columns=["Electrode", *columns],
                )
            except MissingXlsxColumnsError as exc:
                missing = [column for column in columns if column in exc.missing_columns]
                if missing:
                    raise RuntimeError(
                        "Frequency-domain QC requires exact selected BCA harmonic "
                        f"columns in every included workbook. Missing columns in {file_path}: "
                        f"{missing[:8]}"
                    ) from exc
                log_func(f"Frequency-domain QC could not read BCA sheet for {file_path}: {exc}")
                continue
            if "Electrode" not in frame.columns:
                log_func(f"Frequency-domain QC skipped {file_path}: missing Electrode column.")
                continue
            audit_rows = _read_bca_method_audit_rows(
                file_path=file_path,
                reader=read_xlsx_sheet_selected_columns,
                log_func=log_func,
            )
            identity = _frequency_qc_cell_identity(
                subject=subject,
                condition=condition,
                file_path=file_path,
                recording_assignments=recording_assignments,
            )
            for row_offset, row in frame.iterrows():
                electrode = _normalize_electrode(row.get("Electrode"))
                if not electrode:
                    continue
                cell_exclusions = (
                    excluded_electrodes_by_subject_condition or {}
                ).get(
                    (_normalize_recording_id(subject), str(condition)),
                    frozenset(),
                )
                if electrode in cell_exclusions:
                    continue
                finite_values: list[float] = []
                score_unavailable = False
                for harmonic_hz, column in zip(selected_harmonics, columns):
                    method_evidence = audit_rows.get((electrode, column), ())
                    method_state, reason_codes = _bca_method_state(
                        method_evidence
                    )
                    diagnostic = {
                        **identity,
                        "electrode": electrode,
                        "worksheet_row": int(row_offset) + 2,
                        "harmonic_hz": round(float(harmonic_hz), 4),
                        "harmonic_column": column,
                    }
                    if method_state == "technical_integrity_failed":
                        _append_unique_diagnostic(
                            integrity_failures,
                            {
                                **diagnostic,
                                "failure_type": "spectral_metric_input_nonfinite",
                                "value_category": _source_bca_value_category(
                                    file_path=file_path,
                                    worksheet_row=int(row_offset) + 2,
                                    harmonic_column=column,
                                    fallback_value=row.get(column),
                                ),
                                "reason_codes": list(reason_codes),
                            },
                        )
                        score_unavailable = True
                        continue
                    if method_state == "unavailable_by_method":
                        _append_unique_diagnostic(
                            unavailable_by_method,
                            {
                                **diagnostic,
                                "status": "unavailable_by_method",
                                "reason_codes": list(reason_codes),
                            },
                        )
                        score_unavailable = True
                        continue

                    value, _ = _finite_bca_value(row.get(column))
                    if value is None:
                        _append_unique_diagnostic(
                            integrity_failures,
                            {
                                **diagnostic,
                                "failure_type": "invalid_selected_bca_cell",
                                "value_category": _source_bca_value_category(
                                    file_path=file_path,
                                    worksheet_row=int(row_offset) + 2,
                                    harmonic_column=column,
                                    fallback_value=row.get(column),
                                ),
                                "reason_codes": [],
                            },
                        )
                        score_unavailable = True
                        continue
                    finite_values.append(value)

                if score_unavailable:
                    continue
                value = float(pd.Series(finite_values, dtype=float).sum())
                if not np.isfinite(value):
                    _append_unique_diagnostic(
                        integrity_failures,
                        {
                            **identity,
                            "electrode": electrode,
                            "worksheet_row": int(row_offset) + 2,
                            "harmonic_hz": None,
                            "harmonic_column": "<selected harmonic sum>",
                            "failure_type": "nonfinite_selected_bca_sum",
                            "value_category": _bca_value_category(value),
                            "reason_codes": [],
                        },
                    )
                    continue
                abs_value = abs(value)
                if not screening_enabled or abs_value <= thresholds.warning_summed_bca_uv:
                    continue
                severity = "warning"
                if abs_value > thresholds.extreme_review_summed_bca_uv:
                    severity = "extreme"
                elif abs_value > thresholds.strong_warning_summed_bca_uv:
                    severity = "strong"
                flag: dict[str, object] = {
                    **identity,
                    "finding_type": "absolute_electrode_summed_bca",
                    "electrode": electrode,
                    "summed_bca_uv": value,
                    "abs_summed_bca_uv": float(abs_value),
                    "severity": severity,
                    "band_crossed": severity,
                    "selected_harmonics_hz": [
                        round(float(freq), 4) for freq in selected_harmonics
                    ],
                    "selected_harmonic_count": len(selected_harmonics),
                    "selection_fingerprint": str(
                        (selected_harmonics_metadata or {}).get(
                            "selection_fingerprint"
                        )
                        or (selected_harmonics_metadata or {}).get(
                            "selection_input_fingerprint"
                        )
                        or ""
                    ),
                    "independent_qc": [],
                    "independent_qc_status": "not_supplied_to_summed_bca_review",
                }
                if recording_assignments is not None:
                    assignment = recording_assignments.get(subject, {})
                    flag.update(
                        {
                            "recording_id": _normalize_recording_id(subject),
                            "participant_id": _normalize_participant_id(
                                assignment.get("participant_id")
                            ),
                            "session_id": str(
                                assignment.get("session_id") or ""
                            ),
                            "visit_index": assignment.get("visit_index"),
                        }
                    )
                flag["finding_fingerprint"] = _frequency_qc_finding_fingerprint(
                    flag
                )
                flags.append(flag)
    return _SummedBcaInspection(
        flags=tuple(
            sorted(
                flags,
                key=lambda item: (
                    str(
                        item.get("recording_id")
                        or item.get("participant_id")
                        or ""
                    ),
                    -float(item.get("abs_summed_bca_uv") or 0.0),
                    str(item.get("condition") or ""),
                    str(item.get("electrode") or ""),
                ),
            )
        ),
        technical_integrity_failures=tuple(
            sorted(integrity_failures, key=_bca_diagnostic_sort_key)
        ),
        unavailable_by_method=tuple(
            sorted(unavailable_by_method, key=_bca_diagnostic_sort_key)
        ),
    )


def _read_bca_method_audit_rows(
    *,
    file_path: str,
    reader: Callable[..., pd.DataFrame],
    log_func: Callable[[str], None],
) -> dict[tuple[str, str], tuple[dict[str, object], ...]]:
    """Read optional QC-12/QC-14 per-cell availability from a workbook."""

    from Main_App.io import MissingXlsxColumnsError

    try:
        frame = reader(
            file_path,
            sheet_name=SPECTRAL_METRIC_QC_SHEET_NAME,
            required_columns=list(_BCA_AUDIT_REQUIRED_COLUMNS),
        )
    except MissingXlsxColumnsError as exc:
        log_func(
            "Frequency-domain QC ignored malformed optional spectral audit "
            f"metadata in {file_path}: {exc}"
        )
        return {}
    except (OSError, ValueError):
        # Historical and externally supplied workbooks can predate the audit
        # sheet. Their literal selected BCA cells still receive the finite gate.
        return {}

    rows_by_cell: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    if frame.columns.is_unique:
        positions = frame.columns.get_indexer(_BCA_AUDIT_REQUIRED_COLUMNS)
        # iterrows uses this same common-dtype array, then allocates one Series
        # per row. Preserve its scalar values without those temporary objects.
        row_values = (
            tuple(values[position] if position >= 0 else None for position in positions)
            for values in frame.to_numpy(copy=False)
        )
    else:
        # Retain historical behavior for malformed duplicate-column inputs.
        row_values = (
            tuple(row.get(column) for column in _BCA_AUDIT_REQUIRED_COLUMNS)
            for _, row in frame.iterrows()
        )
    # Each target repeats across electrodes. Reuse only immutable built-in text
    # within this table; other scalar types retain the original conversion.
    try:
        frequency_columns: dict[str, str] | None = dict()
    except MemoryError:
        frequency_columns = None
    for raw_electrode, raw_frequency, raw_status, raw_reasons in row_values:
        electrode = _normalize_electrode(
            _optional_cell_text(raw_electrode)
        )
        cacheable = (
            type(raw_frequency) is str
            and len(raw_frequency) <= _MAX_BCA_AUDIT_FREQUENCY_TOKEN_LENGTH
        )
        column = None
        if cacheable and frequency_columns is not None:
            try:
                column = frequency_columns.get(raw_frequency)
            except MemoryError:
                frequency_columns = None
        if column is None:
            # Keep converter errors outside cache-only allocation recovery.
            column = _exact_frequency_column(raw_frequency)
            if cacheable and frequency_columns is not None:
                try:
                    if len(frequency_columns) < _MAX_BCA_AUDIT_FREQUENCY_TOKENS:
                        frequency_columns[raw_frequency] = column
                except MemoryError:
                    frequency_columns = None
        if not electrode or not column:
            continue
        reason_codes = tuple(
            reason.strip()
            for reason in _optional_cell_text(raw_reasons).split(";")
            if reason.strip()
        )
        rows_by_cell[(electrode, column)].append(
            {
                "bca_status": _optional_cell_text(
                    raw_status
                ).casefold(),
                "reason_codes": reason_codes,
            }
        )
    return {
        key: tuple(value)
        for key, value in rows_by_cell.items()
    }


def _exact_frequency_column(value: object) -> str:
    text = _optional_cell_text(value)
    if not text:
        return ""
    try:
        frequency = Fraction(text)
    except (ValueError, ZeroDivisionError):
        return ""
    return f"{float(frequency):.4f}_Hz"


def _optional_cell_text(value: object) -> str:
    if value is None:
        return ""
    try:
        if bool(pd.isna(value)):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _bca_method_state(
    evidence: Sequence[Mapping[str, object]],
) -> tuple[str, tuple[str, ...]]:
    if not evidence:
        return "method_computable", ()
    statuses = {
        str(item.get("bca_status") or "").strip().casefold()
        for item in evidence
    }
    reasons = tuple(
        dict.fromkeys(
            str(reason)
            for item in evidence
            for reason in item.get("reason_codes", ()) or ()
            if str(reason)
        )
    )
    if statuses - {"available", "unavailable"}:
        return "technical_integrity_failed", (
            *reasons,
            "invalid_bca_availability_status",
        )
    if "unavailable" in statuses:
        if any(reason.startswith("nonfinite_") for reason in reasons):
            return "technical_integrity_failed", reasons
        if not reasons:
            return "technical_integrity_failed", (
                "missing_bca_unavailability_reason",
            )
        return "unavailable_by_method", reasons
    return "method_computable", reasons


def _finite_bca_value(value: object) -> tuple[float | None, str]:
    if value is None:
        return None, "blank"
    if isinstance(value, str) and not value.strip():
        return None, "blank"
    try:
        numeric = float(value)
    except (TypeError, ValueError, OverflowError):
        return None, "text"
    if np.isnan(numeric):
        return None, "nan"
    if np.isposinf(numeric):
        return None, "positive_infinity"
    if np.isneginf(numeric):
        return None, "negative_infinity"
    return numeric, "finite"


def _bca_value_category(value: object) -> str:
    _, category = _finite_bca_value(value)
    return category


def _source_bca_value_category(
    *,
    file_path: str,
    worksheet_row: int,
    harmonic_column: str,
    fallback_value: object,
) -> str:
    """Recover blank versus literal NaN text only on an invalid-cell path."""

    from Main_App.io.condition_data import (
        declared_condition_companion,
        read_condition_sheet,
    )

    if declared_condition_companion(file_path) is not None:
        frame = read_condition_sheet(file_path, sheet_name="BCA (uV)")
        row_index = int(worksheet_row) - 2
        if harmonic_column in frame.columns and 0 <= row_index < len(frame):
            return _bca_value_category(frame.iloc[row_index][harmonic_column])
        return _bca_value_category(fallback_value)

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook["BCA (uV)"]
            headers = {
                str(cell.value): index
                for index, cell in enumerate(sheet[1], start=1)
                if cell.value not in (None, "")
            }
            column_index = headers.get(harmonic_column)
            if column_index is None:
                return _bca_value_category(fallback_value)
            source_value = sheet.cell(
                row=int(worksheet_row),
                column=column_index,
            ).value
            return _bca_value_category(source_value)
        finally:
            workbook.close()
    except (KeyError, OSError, ValueError):
        return _bca_value_category(fallback_value)


def _frequency_qc_cell_identity(
    *,
    subject: str,
    condition: str,
    file_path: str,
    recording_assignments: Mapping[str, Mapping[str, object]] | None,
) -> dict[str, object]:
    identity: dict[str, object] = {
        "participant_id": _normalize_participant_id(subject),
        "condition": str(condition),
        "workbook_path": str(file_path),
    }
    if recording_assignments is not None:
        assignment = recording_assignments.get(subject, {})
        identity.update(
            {
                "recording_id": _normalize_recording_id(subject),
                "participant_id": _normalize_participant_id(
                    assignment.get("participant_id")
                ),
                "session_id": str(assignment.get("session_id") or ""),
                "visit_index": assignment.get("visit_index"),
            }
        )
    return identity


def _append_unique_diagnostic(
    collection: list[dict[str, object]],
    diagnostic: Mapping[str, object],
) -> None:
    normalized = dict(diagnostic)
    if normalized not in collection:
        collection.append(normalized)


def _bca_diagnostic_sort_key(item: Mapping[str, object]) -> tuple[object, ...]:
    return (
        str(item.get("recording_id") or item.get("participant_id") or "").casefold(),
        str(item.get("condition") or "").casefold(),
        str(item.get("workbook_path") or "").casefold(),
        str(item.get("electrode") or "").casefold(),
        str(item.get("harmonic_column") or "").casefold(),
        int(item.get("worksheet_row") or 0),
    )


def _finite_input_fingerprint_diagnostic(
    project_root: Path,
    diagnostic: Mapping[str, object],
) -> dict[str, object]:
    payload = dict(diagnostic)
    workbook_path = payload.get("workbook_path")
    if workbook_path not in (None, ""):
        payload["workbook_path"] = _manifest_safe_path(
            project_root,
            Path(str(workbook_path)),
        )
    return payload


def _finding_abs_value(finding: Mapping[str, object]) -> float:
    value = finding.get("abs_summed_bca_uv")
    if value is None:
        value = finding.get("value_uv")
    if value is None:
        value = finding.get("summed_bca_uv")
    try:
        parsed = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return 0.0
    return abs(parsed) if np.isfinite(parsed) else 0.0


def _summarize_flags(
    flags: Sequence[Mapping[str, object]],
    thresholds: FrequencyDomainQcThresholds,
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    by_pid: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    extreme_by_pid_electrode: dict[
        tuple[str, str], list[Mapping[str, object]]
    ] = defaultdict(list)
    for flag in flags:
        pid = _normalize_participant_id(flag.get("participant_id"))
        electrode = _normalize_electrode(flag.get("electrode"))
        if not pid:
            continue
        by_pid[pid].append(flag)
        if str(flag.get("severity") or "") == "extreme" and electrode:
            extreme_by_pid_electrode[(pid, electrode)].append(flag)

    machine_electrodes: list[dict[str, object]] = []
    extreme_electrodes_by_pid: dict[str, set[str]] = defaultdict(set)
    for (pid, electrode), entries in sorted(extreme_by_pid_electrode.items()):
        extreme_electrodes_by_pid[pid].add(electrode)
        max_entry = max(entries, key=_finding_abs_value)
        machine_electrodes.append(
            {
                "participant_id": pid,
                "electrode": electrode,
                "reason": "abs summed BCA exceeded the experimental extreme-review threshold",
                "threshold_uv": float(thresholds.extreme_review_summed_bca_uv),
                "max_abs_summed_bca_uv": _finding_abs_value(max_entry),
                "triggering_conditions": sorted(
                    {str(entry.get("condition") or "") for entry in entries if entry.get("condition")}
                ),
                "source": "experimental_summed_bca_machine_suggestion",
                "authority": "review_only",
            }
        )

    machine_subjects: list[dict[str, object]] = []
    summaries: list[dict[str, object]] = []
    for pid, entries in sorted(by_pid.items()):
        warning_count = len(entries)
        strong_count = sum(1 for item in entries if str(item.get("severity") or "") in {"strong", "extreme"})
        extreme_electrode_count = len(extreme_electrodes_by_pid.get(pid, set()))
        max_entry = max(entries, key=_finding_abs_value)
        broad_review = extreme_electrode_count >= int(
            thresholds.broad_extreme_review_unique_electrodes
        )
        if broad_review:
            machine_subjects.append(
                {
                    "participant_id": pid,
                    "reason": "broad experimental review threshold reached",
                    "extreme_electrode_count": int(extreme_electrode_count),
                    "source": "experimental_summed_bca_machine_suggestion",
                    "authority": "review_only",
                }
            )
        pause_reasons: list[str] = []
        if warning_count:
            pause_reasons.append("experimental review flag")
        if broad_review:
            pause_reasons.append("broad high-priority review")
        if extreme_electrode_count:
            pause_reasons.append("extreme value review")
        if strong_count:
            pause_reasons.append("strong warning")
        if warning_count >= int(thresholds.concentrated_review_flagged_cells):
            pause_reasons.append("concentrated warning pattern")
        summaries.append(
            {
                "participant_id": pid,
                "max_abs_summed_bca_uv": _finding_abs_value(max_entry),
                "max_condition": str(max_entry.get("condition") or ""),
                "max_electrode": str(max_entry.get("electrode") or ""),
                "warning_cell_count": int(warning_count),
                "strong_or_extreme_cell_count": int(strong_count),
                "extreme_electrode_count": int(extreme_electrode_count),
                "broad_review": bool(broad_review),
                "automatic_action": "none",
                "pause_review": bool(pause_reasons),
                "pause_reasons": pause_reasons,
            }
        )
    return summaries, machine_electrodes, machine_subjects


def _summarize_recording_flags(
    flags: Sequence[Mapping[str, object]],
    thresholds: FrequencyDomainQcThresholds,
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    """Summarize repeated-session QC without promoting a visit to a person."""

    by_recording: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    extreme_by_recording_electrode: dict[
        tuple[str, str], list[Mapping[str, object]]
    ] = defaultdict(list)
    for flag in flags:
        recording_id = _normalize_recording_id(flag.get("recording_id"))
        electrode = _normalize_electrode(flag.get("electrode"))
        if not recording_id:
            continue
        by_recording[recording_id].append(flag)
        if str(flag.get("severity") or "") == "extreme" and electrode:
            extreme_by_recording_electrode[(recording_id, electrode)].append(flag)

    machine_electrodes: list[dict[str, object]] = []
    extreme_electrodes_by_recording: dict[str, set[str]] = defaultdict(set)
    for (recording_id, electrode), entries in sorted(
        extreme_by_recording_electrode.items()
    ):
        extreme_electrodes_by_recording[recording_id].add(electrode)
        max_entry = max(entries, key=_finding_abs_value)
        machine_electrodes.append(
            {
                "recording_id": recording_id,
                "participant_id": _normalize_participant_id(
                    max_entry.get("participant_id")
                ),
                "session_id": str(max_entry.get("session_id") or ""),
                "visit_index": max_entry.get("visit_index"),
                "electrode": electrode,
                "reason": "abs summed BCA exceeded the experimental extreme-review threshold",
                "threshold_uv": float(thresholds.extreme_review_summed_bca_uv),
                "max_abs_summed_bca_uv": _finding_abs_value(max_entry),
                "triggering_conditions": sorted(
                    {
                        str(entry.get("condition") or "")
                        for entry in entries
                        if entry.get("condition")
                    }
                ),
                "source": "experimental_summed_bca_machine_suggestion",
                "authority": "review_only",
            }
        )

    machine_recordings: list[dict[str, object]] = []
    summaries: list[dict[str, object]] = []
    for recording_id, entries in sorted(by_recording.items()):
        warning_count = len(entries)
        strong_count = sum(
            1
            for item in entries
            if str(item.get("severity") or "") in {"strong", "extreme"}
        )
        extreme_electrode_count = len(
            extreme_electrodes_by_recording.get(recording_id, set())
        )
        max_entry = max(entries, key=_finding_abs_value)
        broad_review = extreme_electrode_count >= int(
            thresholds.broad_extreme_review_unique_electrodes
        )
        identity = {
            "recording_id": recording_id,
            "participant_id": _normalize_participant_id(
                max_entry.get("participant_id")
            ),
            "session_id": str(max_entry.get("session_id") or ""),
            "visit_index": max_entry.get("visit_index"),
        }
        if broad_review:
            machine_recordings.append(
                {
                    **identity,
                    "reason": "broad experimental review threshold reached",
                    "extreme_electrode_count": int(extreme_electrode_count),
                    "source": "experimental_summed_bca_machine_suggestion",
                    "authority": "review_only",
                }
            )
        pause_reasons: list[str] = []
        if warning_count:
            pause_reasons.append("experimental review flag")
        if broad_review:
            pause_reasons.append("broad high-priority review")
        if extreme_electrode_count:
            pause_reasons.append("extreme value review")
        if strong_count:
            pause_reasons.append("strong warning")
        if warning_count >= int(thresholds.concentrated_review_flagged_cells):
            pause_reasons.append("concentrated warning pattern")
        summaries.append(
            {
                **identity,
                "max_abs_summed_bca_uv": _finding_abs_value(max_entry),
                "max_condition": str(max_entry.get("condition") or ""),
                "max_electrode": str(max_entry.get("electrode") or ""),
                "warning_cell_count": int(warning_count),
                "strong_or_extreme_cell_count": int(strong_count),
                "extreme_electrode_count": int(extreme_electrode_count),
                "broad_review": bool(broad_review),
                "automatic_action": "none",
                "pause_review": bool(pause_reasons),
                "pause_reasons": pause_reasons,
            }
        )
    return summaries, machine_electrodes, machine_recordings


def _analysis_fingerprint(
    *,
    project_root: Path,
    subjects: Sequence[str],
    conditions: Sequence[str],
    subject_data: Mapping[str, Mapping[str, str]],
    selected_harmonics: Sequence[float],
    thresholds: FrequencyDomainQcThresholds,
    flags: Sequence[Mapping[str, object]],
    finite_input_fingerprint: str,
    recording_assignments: Mapping[str, Mapping[str, object]] | None = None,
    screening_settings: SummedBcaScreeningSettings | None = None,
    provisional_metadata: Mapping[str, object] | None = None,
    source_workbooks: Sequence[Mapping[str, object]] | None = None,
    source_fingerprint: str = "",
    previous_review_evidence: Mapping[str, object] | None = None,
) -> str:
    workbooks = list(source_workbooks or ())
    if not workbooks:
        workbooks = _source_workbook_rows(
            project_root=project_root,
            subjects=subjects,
            conditions=conditions,
            subject_data=subject_data,
            recording_assignments=recording_assignments,
        )
    payload = {
        "method_version": (
            REPEATED_FREQUENCY_DOMAIN_QC_METHOD_VERSION
            if recording_assignments is not None
            else FREQUENCY_DOMAIN_QC_METHOD_VERSION
        ),
        "thresholds": thresholds.to_manifest(),
        "screening_settings": (
            screening_settings.to_manifest()
            if screening_settings is not None
            else {}
        ),
        "subjects": list(map(str, subjects)),
        "conditions": list(map(str, conditions)),
        "selected_harmonics_hz": [round(float(freq), 4) for freq in selected_harmonics],
        "finite_input_method_version": FREQUENCY_DOMAIN_QC_INTEGRITY_METHOD_VERSION,
        "finite_input_fingerprint": str(finite_input_fingerprint),
        "provisional_harmonic_metadata": _json_safe(
            {key: value for key, value in (provisional_metadata or {}).items()
             if key not in _HARMONIC_CACHE_ANNOTATIONS}
        ),
        "source_fingerprint": str(source_fingerprint),
        "workbooks": workbooks,
        "flags": [
            {
                "participant_id": _normalize_participant_id(flag.get("participant_id")),
                "condition": str(flag.get("condition") or ""),
                "electrode": _normalize_electrode(flag.get("electrode")),
                "roi": str(flag.get("roi") or ""),
                "metric": str(flag.get("metric") or ""),
                "value_uv": _json_safe(flag.get("value_uv")),
                "abs_summed_bca_uv": round(_finding_abs_value(flag), 6),
                "severity": str(flag.get("severity") or ""),
                "finding_fingerprint": str(
                    flag.get("finding_fingerprint") or ""
                ),
                "independent_qc_fingerprint": str(
                    flag.get("independent_qc_fingerprint") or ""
                ),
            }
            for flag in flags if not is_roi_frequency_qc_entry(flag)
        ],
    }
    if recording_assignments is not None:
        payload.update(
            {
                "identity_scope": "recording",
                "recording_assignments": [
                    {
                        "recording_id": str(recording_id),
                        **{
                            key: row.get(key)
                            for key in (
                                "participant_id",
                                "group_id",
                                "session_id",
                                "source_id",
                                "visit_index",
                                "days_from_baseline",
                            )
                        },
                    }
                    for recording_id, row in sorted(
                        recording_assignments.items(),
                        key=lambda item: str(item[0]).casefold(),
                    )
                    if recording_id in subjects
                ],
                "flags": [
                    {
                        "recording_id": _normalize_recording_id(
                            flag.get("recording_id")
                        ),
                        "participant_id": _normalize_participant_id(
                            flag.get("participant_id")
                        ),
                        "session_id": str(flag.get("session_id") or ""),
                        "condition": str(flag.get("condition") or ""),
                        "electrode": _normalize_electrode(flag.get("electrode")),
                        "roi": str(flag.get("roi") or ""),
                        "metric": str(flag.get("metric") or ""),
                        "value_uv": _json_safe(flag.get("value_uv")),
                        "abs_summed_bca_uv": round(
                            _finding_abs_value(flag),
                            6,
                        ),
                        "severity": str(flag.get("severity") or ""),
                        "finding_fingerprint": str(
                            flag.get("finding_fingerprint") or ""
                        ),
                        "independent_qc_fingerprint": str(
                            flag.get("independent_qc_fingerprint") or ""
                        ),
                    }
                    for flag in flags if not is_roi_frequency_qc_entry(flag)
                ],
            }
        )
    fingerprint = _hash_payload(payload)
    if previous_review_evidence is not None:
        previous_fingerprint = str(previous_review_evidence.get("analysis_fingerprint") or "")
        previous_metadata = previous_review_evidence.get("provisional_harmonic_metadata")
        if previous_fingerprint == fingerprint:
            return fingerprint
        if isinstance(previous_metadata, Mapping):
            stable_previous_metadata = _json_safe({
                key: value for key, value in previous_metadata.items()
                if key not in _HARMONIC_CACHE_ANNOTATIONS
            })
            if stable_previous_metadata == payload["provisional_harmonic_metadata"]:
                # Compatibility for validated receipts written before cache
                # bookkeeping was separated from scientific review identity.
                # Every current source, setting and finding must still match;
                # only the three saved cache annotations may be substituted.
                previous_payload = {
                    **payload,
                    "provisional_harmonic_metadata": _json_safe(previous_metadata),
                }
                if previous_fingerprint and _hash_payload(previous_payload) == previous_fingerprint:
                    return previous_fingerprint
    return fingerprint


def _source_workbook_rows(
    *,
    project_root: Path,
    subjects: Sequence[str],
    conditions: Sequence[str],
    subject_data: Mapping[str, Mapping[str, str]],
    recording_assignments: Mapping[str, Mapping[str, object]] | None = None,
) -> list[dict[str, object]]:
    workbooks: list[dict[str, object]] = []
    for subject in subjects:
        for condition in conditions:
            file_path = subject_data.get(subject, {}).get(condition)
            if not file_path:
                continue
            path = Path(file_path)
            try:
                stat = path.stat()
                size = int(stat.st_size)
                mtime = int(stat.st_mtime_ns)
            except OSError:
                size = None
                mtime = None
            workbook: dict[str, object] = {
                    "subject": str(subject),
                    "condition": str(condition),
                    "path": _manifest_safe_path(project_root, path),
                    "size_bytes": size,
                    "mtime_ns": mtime,
                }
            if recording_assignments is not None:
                assignment = recording_assignments.get(str(subject), {})
                workbook.update(
                    {
                        "recording_id": str(subject),
                        "participant_id": str(
                            assignment.get("participant_id") or ""
                        ),
                        "group_id": str(assignment.get("group_id") or ""),
                        "session_id": str(assignment.get("session_id") or ""),
                        "source_id": str(assignment.get("source_id") or ""),
                        "visit_index": assignment.get("visit_index"),
                        "days_from_baseline": assignment.get(
                            "days_from_baseline"
                        ),
                    }
                )
            workbooks.append(workbook)
    return workbooks


def _decision_fingerprint(
    *,
    analysis_fingerprint: str,
    auto_electrodes: Sequence[Mapping[str, object]],
    auto_participants: Sequence[Mapping[str, object]],
    manual_participants: Sequence[Mapping[str, object]],
    auto_recording_electrodes: Sequence[Mapping[str, object]] | None = None,
    auto_recordings: Sequence[Mapping[str, object]] | None = None,
    manual_recordings: Sequence[Mapping[str, object]] | None = None,
    review_decisions: Sequence[Mapping[str, object]] | None = None,
    identity_scope: str,
) -> str:
    payload = {
        "analysis_fingerprint": str(analysis_fingerprint),
        "identity_scope": str(identity_scope).strip().casefold(),
        "auto_electrodes": _json_safe(_normalize_auto_electrode_entries(auto_electrodes)),
        "auto_participants": _json_safe(_normalize_auto_participant_entries(auto_participants)),
        "manual_participants": _json_safe(_normalize_manual_entries(manual_participants)),
        "review_decisions": _json_safe(
            [
                {
                    key: value
                    for key, value in dict(item).items()
                    if key != "reviewed_at"
                }
                for item in sorted(
                    (row for row in (review_decisions or ())
                     if not is_roi_frequency_qc_entry(row)),
                    key=lambda row: str(
                        row.get("decision_fingerprint")
                        or row.get("finding_fingerprint")
                        or ""
                    ),
                )
            ]
        ),
    }
    if (
        auto_recording_electrodes is not None
        or auto_recordings is not None
        or manual_recordings is not None
    ):
        payload.update(
            {
                "auto_recording_electrodes": _json_safe(
                    _normalize_auto_recording_electrode_entries(
                        auto_recording_electrodes
                    )
                ),
                "auto_recordings": _json_safe(
                    _normalize_auto_recording_entries(auto_recordings)
                ),
                "manual_recordings": _json_safe(
                    _normalize_manual_recording_entries(manual_recordings)
                ),
            }
        )
    return _hash_payload(payload)


def _write_frequency_domain_qc_text_report(
    project_root: Path,
    *,
    report: Mapping[str, object],
    manual_participants: Sequence[Mapping[str, object]],
    manual_recordings: Sequence[Mapping[str, object]],
    review_decisions: Sequence[Mapping[str, object]],
    decision_fingerprint: str,
    reviewed_at: str,
) -> Path:
    qc_folder = project_root / QUALITY_CHECK_FOLDER
    qc_folder.mkdir(parents=True, exist_ok=True)
    path = qc_folder / FREQUENCY_DOMAIN_QC_REPORT_NAME
    thresholds = (
        report.get("thresholds")
        if isinstance(report.get("thresholds"), Mapping)
        else {}
    )
    protocol = (
        report.get("frequency_protocol")
        if isinstance(report.get("frequency_protocol"), Mapping)
        else {}
    )
    lines = [
        "Experimental Summed-BCA Review",
        "",
        str(report.get("screening_explanation") or SUMMED_BCA_SCREENING_BRIEF_TEXT),
        "No summed-BCA threshold automatically excludes data.",
        "An outcome-informed exclusion is exploratory and requires sensitivity reporting.",
        "",
        f"Reviewed at: {reviewed_at}",
        f"Decision fingerprint: {decision_fingerprint}",
        f"Project: {project_root}",
        f"Identity scope: {report.get('identity_scope') or 'participant'}",
        f"Screening status: {report.get('screening_status') or ''}",
        "",
        "Project protocol and candidate harmonic state",
        (
            "- Presentation / oddball rate: "
            f"{protocol.get('presentation_rate_hz', '')} / "
            f"{protocol.get('oddball_rate_hz', '')} Hz"
        ),
        (
            "- Expected analyzed oddball cycles / duration: "
            f"{protocol.get('expected_analyzed_oddball_cycles', '')} / "
            f"{protocol.get('analyzed_duration_seconds', '')} s"
        ),
        "- Harmonics: "
        + (
            ", ".join(
                f"{float(freq):g} Hz"
                for freq in report.get("selected_harmonics_hz", []) or []
            )
            or "None"
        ),
        f"- Analysis fingerprint: {report.get('analysis_fingerprint') or ''}",
        "",
        "Experimental review thresholds",
        f"- Warning: > {thresholds.get('warning_summed_bca_uv', 10)} uV",
        f"- Strong: > {thresholds.get('strong_warning_summed_bca_uv', 50)} uV",
        f"- Extreme: > {thresholds.get('extreme_review_summed_bca_uv', 250)} uV",
        (
            "- Concentrated review: at least "
            f"{thresholds.get('concentrated_review_flagged_cells', 5)} flagged cells"
        ),
        (
            "- Broad review: at least "
            f"{thresholds.get('broad_extreme_review_unique_electrodes', 11)} "
            "unique extreme-band electrodes"
        ),
        "",
        "Reviewed findings and decisions",
    ]
    normalized_decisions = [dict(item) for item in review_decisions if not is_roi_frequency_qc_entry(item)]
    if normalized_decisions:
        for entry in normalized_decisions:
            evidence = (
                entry.get("evidence")
                if isinstance(entry.get("evidence"), Mapping)
                else {}
            )
            identity = str(entry.get("recording_id") or entry.get("participant_id") or "")
            target = str(entry.get("electrode") or "")
            signed = evidence.get("summed_bca_uv")
            absolute = evidence.get("abs_summed_bca_uv")
            if absolute is None:
                absolute = evidence.get("value_uv")
            harmonics = evidence.get("selected_harmonics_hz") or []
            lines.append(
                f"- {identity} / {entry.get('condition') or ''} / {target}: "
                f"decision={entry.get('decision') or ''}; "
                f"signed={signed}; absolute={absolute}; "
                f"band={evidence.get('band_crossed') or evidence.get('severity') or ''}; "
                f"harmonics={harmonics}; reason={entry.get('reason') or '(retain)'}"
            )
    else:
        lines.append("- None")

    lines.extend(["", "Preserved whole-participant exclusions"])
    manual_entries = _normalize_manual_entries(manual_participants)
    if manual_entries:
        for entry in manual_entries:
            lines.append(f"- {entry['participant_id']}: {entry['reason']}")
    else:
        lines.append("- None")

    if str(report.get("identity_scope") or "") == "recording":
        lines.extend(["", "Preserved whole-recording exclusions"])
        normalized_manual_recordings = _normalize_manual_recording_entries(
            manual_recordings
        )
        if normalized_manual_recordings:
            for entry in normalized_manual_recordings:
                lines.append(f"- {entry['recording_id']}: {entry['reason']}")
        else:
            lines.append("- None")
    legacy = _iter_mapping_entries(report.get("legacy_machine_suggestions"))
    lines.extend(["", "Preserved legacy machine suggestions (inactive)"])
    if legacy:
        for entry in legacy:
            lines.append(f"- {_json_safe(entry)}")
    else:
        lines.append("- None")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _find_first_bca_columns(
    subjects: Sequence[str],
    conditions: Sequence[str],
    subject_data: Mapping[str, Mapping[str, str]],
) -> list[object]:
    from Main_App.io import read_xlsx_sheet_header

    for subject in subjects:
        for condition in conditions:
            file_path = subject_data.get(subject, {}).get(condition)
            if not file_path:
                continue
            try:
                return [
                    column
                    for column in read_xlsx_sheet_header(file_path, sheet_name="BCA (uV)")
                    if column != "Electrode"
                ]
            except Exception:
                logger.debug("frequency_domain_qc_bca_header_read_failed", exc_info=True)
    return []


def _harmonic_selection_settings(project: Any) -> Any:
    from Tools.Stats.analysis.dv_policy_settings import (
        GROUP_SIGNIFICANT_ELECTRODE_SCOPE_ROI_UNION,
        GROUP_SIGNIFICANT_POLICY_NAME,
        GROUP_SIGNIFICANT_SUMMATION_THROUGH_HIGHEST,
        normalize_dv_policy,
    )

    raw_preprocessing: Mapping[str, object] = (
        getattr(project, "preprocessing", {}) or {}
    )
    project_root = getattr(project, "project_root", None)
    if project_root not in (None, ""):
        manifest_path = Path(project_root).resolve(strict=False) / "project.json"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            manifest = None
        if isinstance(manifest, Mapping) and isinstance(
            manifest.get("preprocessing"), Mapping
        ):
            raw_preprocessing = manifest["preprocessing"]
    try:
        preprocessing = normalize_preprocessing_settings(raw_preprocessing)
    except ValueError:
        preprocessing = normalize_preprocessing_settings({})
    policy: dict[str, object] = {
        "name": preprocessing.get(
            "harmonic_selection_policy",
            GROUP_SIGNIFICANT_POLICY_NAME,
        ),
        "harmonic_selection_profile": raw_preprocessing.get(
            "harmonic_selection_profile"
        ),
        "harmonic_selection_profile_version": raw_preprocessing.get(
            "harmonic_selection_profile_version"
        ),
        "fixed_harmonic_frequencies_hz": preprocessing.get(
            "fixed_harmonic_frequencies_hz",
            "",
        ),
        "fixed_harmonic_input_mode": raw_preprocessing.get(
            "fixed_harmonic_input_mode"
        ),
        "fixed_harmonic_upper_harmonic_index": raw_preprocessing.get(
            "fixed_harmonic_upper_harmonic_index"
        ),
        "fixed_harmonic_upper_frequency_hz": raw_preprocessing.get(
            "fixed_harmonic_upper_frequency_hz"
        ),
        "fixed_harmonic_auto_exclude_base": preprocessing.get(
            "fixed_harmonic_auto_exclude_base",
            True,
        ),
        "group_significant_selection_electrodes": raw_preprocessing.get(
            "group_significant_selection_electrodes"
        ),
        "group_significant_summation_method": preprocessing.get(
            "group_significant_summation_method",
            GROUP_SIGNIFICANT_SUMMATION_THROUGH_HIGHEST,
        ),
    }
    if "group_significant_electrode_scope" in raw_preprocessing:
        policy["group_significant_electrode_scope"] = raw_preprocessing[
            "group_significant_electrode_scope"
        ]
    elif raw_preprocessing.get("harmonic_selection_profile") in (None, ""):
        policy["group_significant_electrode_scope"] = preprocessing.get(
            "group_significant_electrode_scope",
            GROUP_SIGNIFICANT_ELECTRODE_SCOPE_ROI_UNION,
        )
    return normalize_dv_policy(policy)


def _filter_preprocessing_manual_exclusions(project: Any, subjects: list[str]) -> list[str]:
    preprocessing = getattr(project, "preprocessing", {}) or {}
    excluded = set(
        normalize_manual_excluded_participants(
            preprocessing.get("manual_excluded_participants", [])
        )
    )
    return [subject for subject in subjects if str(subject).upper() not in excluded]


def _filter_preprocessing_manual_recording_exclusions(
    project: Any,
    recording_ids: list[str],
    *,
    recording_assignments: Mapping[str, Mapping[str, object]],
) -> list[str]:
    preprocessing = getattr(project, "preprocessing", {}) or {}
    excluded_participants = {
        str(value).casefold()
        for value in normalize_manual_excluded_participants(
            preprocessing.get("manual_excluded_participants", [])
        )
    }
    excluded_recordings = {
        str(value).casefold()
        for value in normalize_manual_excluded_recordings(
            preprocessing.get("manual_excluded_recordings", [])
        )
    }
    return [
        recording_id
        for recording_id in recording_ids
        if recording_id.casefold() not in excluded_recordings
        and str(
            recording_assignments.get(recording_id, {}).get("participant_id") or ""
        ).casefold()
        not in excluded_participants
    ]


def _recording_assignments_from_index(
    dataset_index: ProjectDatasetIndex,
) -> dict[str, dict[str, object]]:
    assignments: dict[str, dict[str, object]] = {}
    group_by_participant = dataset_index.participant_group_id_map()
    for recording_id in dataset_index.recording_ids:
        recording = dataset_index.recordings.get(recording_id)
        if recording is None:
            raise RuntimeError(
                "Repeated-session frequency-domain QC is missing the canonical "
                f"recording assignment for {recording_id}."
            )
        session = dataset_index.sessions.get(recording.session_id)
        assignments[recording_id] = {
            "recording_id": recording_id,
            "participant_id": recording.participant_id,
            "group_id": group_by_participant.get(recording.participant_id, ""),
            "session_id": recording.session_id,
            "session_label": session.label if session is not None else recording.session_id,
            "source_id": recording.source_id,
            "visit_index": recording.visit_index,
            "days_from_baseline": recording.days_from_baseline,
        }
    return assignments


def _filter_to_completed_subjects(
    *,
    project_root: Path,
    subjects: list[str],
    subject_data: dict[str, dict[str, str]],
) -> tuple[list[str], dict[str, dict[str, str]]]:
    from Main_App.processing.processing_ledger import load_ledger

    try:
        ledger = load_ledger(project_root)
    except Exception:
        return subjects, subject_data
    entries = ledger.get("entries") if isinstance(ledger, Mapping) else None
    if not isinstance(entries, Mapping):
        return subjects, subject_data
    completed = {
        str(pid).upper()
        for pid, entry in entries.items()
        if isinstance(entry, Mapping) and str(entry.get("status") or "") == "completed"
    }
    if not completed:
        return subjects, subject_data
    filtered_subjects = [subject for subject in subjects if subject.upper() in completed]
    return filtered_subjects, {
        subject: dict(subject_data.get(subject, {})) for subject in filtered_subjects
    }


def _ordered_conditions(project: Any, scanned_conditions: list[str]) -> list[str]:
    scanned = [str(condition) for condition in scanned_conditions]
    seen: set[str] = set()
    ordered: list[str] = []
    event_map = getattr(project, "event_map", {}) or {}
    if isinstance(event_map, Mapping):
        for condition in event_map.keys():
            text = str(condition)
            if text in scanned and text not in seen:
                ordered.append(text)
                seen.add(text)
    for condition in scanned:
        if condition not in seen:
            ordered.append(condition)
            seen.add(condition)
    return ordered


def _filter_subject_data(
    subject_data: dict[str, dict[str, str]],
    conditions: list[str],
) -> dict[str, dict[str, str]]:
    condition_set = set(conditions)
    return {
        subject: {
            condition: path
            for condition, path in (condition_map or {}).items()
            if condition in condition_set and Path(path).exists()
        }
        for subject, condition_map in subject_data.items()
    }


def _experimental_summed_bca_settings(
    project: Any,
    project_root: Path,
) -> SummedBcaScreeningSettings:
    raw_settings = getattr(project, "experimental_qc_settings", None)
    manifest_path = project_root / "project.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        manifest = None
    if isinstance(manifest, Mapping) and "experimental_qc" in manifest:
        raw_settings = manifest.get("experimental_qc")
    try:
        return normalize_experimental_qc_settings(
            raw_settings
        ).summed_bca_screening
    except (TypeError, ValueError) as exc:
        raise RuntimeError(
            "The project experimental summed-BCA settings are invalid. Correct "
            "them in Settings before frequency-domain review."
        ) from exc


def _project_protocol_review_metadata(
    project: Any,
    project_root: Path,
) -> dict[str, object]:
    raw_protocol = getattr(project, "frequency_protocol", None)
    if raw_protocol is None:
        try:
            manifest = json.loads(
                (project_root / "project.json").read_text(encoding="utf-8")
            )
        except (OSError, json.JSONDecodeError):
            manifest = {}
        raw_protocol = manifest.get("frequency_protocol")
    try:
        protocol = normalize_frequency_protocol(raw_protocol)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(
            "Frequency-domain QC requires a valid project frequency protocol."
        ) from exc
    if (
        not protocol.is_ready
        or protocol.oddball_rate_hz is None
        or protocol.presentation_rate_hz is None
        or protocol.expected_analyzed_oddball_cycles is None
    ):
        raise RuntimeError(
            "Frequency-domain QC requires the project presentation rate, oddball "
            "rate/recurrence, and expected analyzed oddball-cycle count."
        )
    cycles = int(protocol.expected_analyzed_oddball_cycles)
    duration = float(Fraction(cycles, 1) / protocol.oddball_rate_hz)
    return {
        "frequency_protocol_version": protocol.version,
        "frequency_protocol_fingerprint": protocol.fingerprint,
        "presentation_rate_hz": float(protocol.presentation_rate_hz),
        "oddball_rate_hz": float(protocol.oddball_rate_hz),
        "oddball_every_n": int(protocol.oddball_every_n),
        "expected_analyzed_oddball_cycles": cycles,
        "analyzed_duration_seconds": duration,
    }


def _harmonic_selection_fingerprint(
    metadata: Mapping[str, object],
) -> str:
    return str(
        metadata.get("selection_fingerprint")
        or metadata.get("selection_input_fingerprint")
        or ""
    ).strip()


def _strict_embedded_fingerprint(
    value: object,
) -> tuple[str, dict[str, object]]:
    if not isinstance(value, Mapping) or not value:
        return "none_available", {}
    payload = dict(value)
    recorded = str(payload.pop("fingerprint", "") or "")
    try:
        encoded = json.dumps(
            payload,
            sort_keys=True,
            separators=(",", ":"),
            ensure_ascii=False,
            allow_nan=False,
        ).encode("utf-8")
    except (TypeError, ValueError):
        return "invalid", {}
    expected = hashlib.sha256(encoded).hexdigest()
    if not recorded or recorded != expected:
        return "invalid", {}
    return "current", {**payload, "fingerprint": recorded}


def _normalized_processing_qc_entry(
    entry: Mapping[str, object] | None,
) -> dict[str, object]:
    if not isinstance(entry, Mapping):
        payload: dict[str, object] = {"status": "processing_entry_missing"}
        payload["fingerprint"] = _hash_payload(payload)
        return payload

    from Main_App.processing.interpolation_burden import (
        InterpolationBurdenError,
        normalize_interpolation_burden,
    )
    from Main_App.processing.preprocessing_outcome import (
        PreprocessingOutcomeError,
        normalize_preprocessing_outcome,
    )

    errors: list[str] = []
    try:
        outcome = normalize_preprocessing_outcome(entry)
        outcome_payload = outcome.to_payload()
        if not outcome.is_current:
            errors.append("preprocessing_outcome_not_current")
    except (PreprocessingOutcomeError, TypeError, ValueError):
        outcome_payload = {}
        errors.append("preprocessing_outcome_invalid")

    raw_burden = entry.get("interpolation_burden")
    try:
        burden_payload = normalize_interpolation_burden(raw_burden).to_payload()  # type: ignore[arg-type]
    except (InterpolationBurdenError, TypeError, ValueError):
        burden_payload = {}
        errors.append(
            "interpolation_burden_missing"
            if raw_burden in (None, {})
            else "interpolation_burden_invalid"
        )

    kurtosis_status, kurtosis_payload = _strict_embedded_fingerprint(
        entry.get("kurtosis_qc_evidence")
    )
    plan_status, plan_payload = _strict_embedded_fingerprint(
        entry.get("kurtosis_decision_plan")
    )
    if kurtosis_status == "invalid":
        errors.append("kurtosis_evidence_invalid")
    if plan_status == "invalid":
        errors.append("kurtosis_decision_plan_invalid")

    payload = {
        "status": "current" if not errors else "current_source_invalid",
        "errors": errors,
        "preprocessing_outcome": outcome_payload,
        "interpolation_burden": burden_payload,
        "kurtosis_evidence_status": kurtosis_status,
        "kurtosis_qc_evidence": kurtosis_payload,
        "kurtosis_decision_plan_status": plan_status,
        "kurtosis_decision_plan": plan_payload,
    }
    payload["fingerprint"] = _hash_payload(payload)
    return payload


_INDEPENDENT_QC_CACHE_MAX_BYTES = 8 * 1024 * 1024
_INDEPENDENT_QC_LEDGER_MAX_BYTES = 32 * 1024 * 1024


def _independent_qc_context_fits_cache(value: _IndependentQcContext) -> bool:
    """Bound retained Python containers as well as the ledger's serialized size."""
    seen: set[int] = set()
    size = 0

    def fits(item: object) -> bool:
        nonlocal size
        if id(item) in seen:
            return True
        seen.add(id(item))
        size += sys.getsizeof(item)
        if size > _INDEPENDENT_QC_CACHE_MAX_BYTES:
            return False
        if isinstance(item, dict):
            return all(fits(key) and fits(child) for key, child in item.items())
        if isinstance(item, (tuple, list)):
            return all(fits(child) for child in item)
        return item is None or type(item) in (str, bool, int, float)

    return fits((value.source_identity, value.cells, value.processing_entries))


def _load_independent_qc_context(project_root: Path) -> _IndependentQcContext:
    """Reuse only current ledger-derived evidence inside one validation operation.

    Review decisions and workbook validation remain at their original call sites.
    Every hit rehashes the ledger; live path resolution also detects a retargeted
    project/ledger link, which resolved file snapshots alone cannot detect.
    """
    from Main_App.processing import post_processing_context as context
    from Main_App.processing.processing_ledger import ledger_path

    if not context.validation_scope_active():
        return _read_independent_qc_context(project_root)

    def live_key() -> tuple[str, str]:
        return (
            str(project_root.resolve()),
            str(ledger_path(project_root).resolve()),
        )

    key = None
    files = ()
    try:
        candidate_key = live_key()
        if Path(candidate_key[1]).stat().st_size <= _INDEPENDENT_QC_LEDGER_MAX_BYTES:
            files = context.capture_validation_files([candidate_key[1]], hash_contents=True)
            if all(identity is not None for _, identity in files):
                key = candidate_key
                cached = context.cached_validation("independent_qc_ledger", (key, files))
                if cached is not context.CACHE_MISS:
                    # Reject edits during detachment as well as initial validation.
                    if live_key() == key and context.validation_files_unchanged(files):
                        return cached
                    key = None
    except (OSError, RuntimeError, MemoryError):
        # Missing/unreadable paths and cache-only allocation failures must retain
        # the uncached loader's diagnostics and error precedence.
        key = None

    ledger_snapshot = None
    if key is not None:
        try:
            with Path(key[1]).open("rb") as stream:
                encoded_ledger = stream.read(_INDEPENDENT_QC_LEDGER_MAX_BYTES + 1)
            if (
                len(encoded_ledger) <= _INDEPENDENT_QC_LEDGER_MAX_BYTES
                and hashlib.sha256(encoded_ledger).hexdigest() == files[0][1][-1]
            ):
                parsed_ledger = json.loads(encoded_ledger.decode("utf-8"))
                if isinstance(parsed_ledger, dict):
                    if not isinstance(parsed_ledger.get("entries"), dict):
                        parsed_ledger["entries"] = {}
                    parsed_ledger.setdefault("schema_version", 1)
                    ledger_snapshot = parsed_ledger
        except (OSError, ValueError, RuntimeError, MemoryError):
            pass
        if ledger_snapshot is None:
            key = None

    # The cache must bind the context to the exact bytes it parsed, rather than
    # rereading a mutable path that could change and return to its original target.
    # Keep normalizer execution outside cache-only exception handling.
    value = (
        _read_independent_qc_context(project_root)
        if ledger_snapshot is None
        else _read_independent_qc_context(project_root, ledger_snapshot=ledger_snapshot)
    )
    if ledger_snapshot is not None and value.source_identity.get("status") != "current":
        # The original loader owns warnings and precedence for invalid inputs.
        return _read_independent_qc_context(project_root)
    if key is None:
        return value
    try:
        if (
            value.source_identity.get("status") == "current"
            and _independent_qc_context_fits_cache(value)
            and live_key() == key
        ):
            context.remember_validation(
                "independent_qc_ledger", (key, files), value, files=files,
                max_namespace_entries=1,
            )
    except (OSError, RuntimeError, MemoryError):
        pass
    return value


def _read_independent_qc_context(
    project_root: Path,
    *,
    ledger_snapshot: Mapping[str, object] | None = None,
) -> _IndependentQcContext:
    from Main_App.processing.processing_ledger import load_ledger
    from Main_App.processing.roi_coverage import (
        ROI_COVERAGE_STAGE_PRE_REVIEW,
        RoiCoverageGateError,
        _roi_coverage_from_ledger,
        load_roi_coverage,
    )

    try:
        coverage = (
            load_roi_coverage(project_root, stage=ROI_COVERAGE_STAGE_PRE_REVIEW)
            if ledger_snapshot is None
            else _roi_coverage_from_ledger(
                ledger_snapshot, stage=ROI_COVERAGE_STAGE_PRE_REVIEW,
            )
        )
    except (OSError, TypeError, ValueError, RoiCoverageGateError) as exc:
        core = {
            "version": FREQUENCY_DOMAIN_QC_INDEPENDENT_EVIDENCE_VERSION,
            "status": "pre_review_coverage_invalid",
            "reason": str(exc),
            "pre_review_roi_coverage_fingerprint": "",
            "processing_entries": [],
        }
        return _IndependentQcContext(
            source_identity={**core, "fingerprint": _hash_payload(core)},
            cells={},
            processing_entries={},
        )
    if coverage is None:
        core = {
            "version": FREQUENCY_DOMAIN_QC_INDEPENDENT_EVIDENCE_VERSION,
            "status": "pre_review_coverage_missing",
            "reason": "Current QC-21 pre-review coverage was not available.",
            "pre_review_roi_coverage_fingerprint": "",
            "processing_entries": [],
        }
        return _IndependentQcContext(
            source_identity={**core, "fingerprint": _hash_payload(core)},
            cells={},
            processing_entries={},
        )

    ledger = load_ledger(project_root) if ledger_snapshot is None else ledger_snapshot
    raw_entries = ledger.get("entries")
    entries = raw_entries if isinstance(raw_entries, Mapping) else {}
    entry_by_identity = {
        str(key).casefold(): value
        for key, value in entries.items()
        if isinstance(value, Mapping)
    }
    cells: dict[tuple[str, str], dict[str, object]] = {}
    processing_entries: dict[str, dict[str, object]] = {}
    for cell in coverage.cells:
        cell_payload = cell.to_payload()
        cells[(cell.recording_id.casefold(), cell.condition_label.casefold())] = (
            cell_payload
        )
        recording_key = cell.recording_id.casefold()
        if recording_key not in processing_entries:
            processing_entries[recording_key] = _normalized_processing_qc_entry(
                entry_by_identity.get(recording_key)
            )
    status = (
        "current"
        if all(
            str(entry.get("status") or "") == "current"
            for entry in processing_entries.values()
        )
        else "current_source_invalid"
    )
    core = {
        "version": FREQUENCY_DOMAIN_QC_INDEPENDENT_EVIDENCE_VERSION,
        "status": status,
        "reason": "",
        "pre_review_roi_coverage_fingerprint": coverage.fingerprint,
        "processing_entries": [
            {
                "recording_id": recording_id,
                "fingerprint": str(entry.get("fingerprint") or ""),
                "status": str(entry.get("status") or ""),
            }
            for recording_id, entry in sorted(processing_entries.items())
        ],
    }
    return _IndependentQcContext(
        source_identity={**core, "fingerprint": _hash_payload(core)},
        cells=cells,
        processing_entries=processing_entries,
    )


def _expected_scalp_channels_by_subject_condition(
    context: _IndependentQcContext,
    *,
    excluded_conditions: Iterable[tuple[str, str]] = (),
) -> dict[tuple[str, str], tuple[str, ...]] | None:
    # Source coverage precedes review. Only explicit whole-condition decisions
    # remove that source from the subsequent harmonic pool; an unexplained
    # absent file must still fail its existing required-source check.
    excluded_keys = {
        (str(identity).casefold(), str(condition).casefold())
        for identity, condition in excluded_conditions
    }
    expected: dict[tuple[str, str], tuple[str, ...]] = {}
    for key, cell in context.cells.items():
        if key in excluded_keys:
            continue
        source = cell.get("source_evidence")
        if not isinstance(source, Mapping):
            continue
        channels = tuple(
            str(channel).strip()
            for channel in source.get("expected_scalp_channels") or ()
            if str(channel).strip()
        )
        if channels:
            expected[key] = channels
    return expected or None


def _attach_independent_qc_evidence(
    finding: dict[str, object],
    context: _IndependentQcContext,
) -> None:
    identity = _normalize_recording_id(
        finding.get("recording_id") or finding.get("participant_id")
    )
    condition = str(finding.get("condition") or "").strip()
    cell = context.cells.get((identity.casefold(), condition.casefold()))
    if cell is None:
        source_status = str(context.source_identity.get("status") or "")
        status = "current_source_cell_missing"
        evidence = [
            {
                "source": "qc21_pre_review_coverage",
                "status": status,
                "source_status": source_status or "unavailable",
                "authority": "context_only",
                "reason": str(context.source_identity.get("reason") or ""),
                "source_fingerprint": str(
                    context.source_identity.get("fingerprint") or ""
                ),
            }
        ]
        finding["independent_qc"] = evidence
        finding["independent_qc_status"] = status
        finding["independent_qc_source_status"] = source_status or "unavailable"
        finding["independent_qc_authority"] = "context_only"
        finding["independent_qc_fingerprint"] = _hash_payload(
            {
                "status": finding["independent_qc_status"],
                "source_status": finding["independent_qc_source_status"],
                "evidence": evidence,
            }
        )
        return

    evidence: list[dict[str, object]] = []
    source = cell.get("source_evidence")
    source = dict(source) if isinstance(source, Mapping) else {}
    electrode = _normalize_electrode(finding.get("electrode"))
    coverage_row: dict[str, object] = {
        "source": "qc21_pre_review_coverage",
        "status": "current",
        "authority": "context_only",
        "cell_fingerprint": str(cell.get("fingerprint") or ""),
        "source_evidence_fingerprint": str(source.get("fingerprint") or ""),
    }
    target_channels: list[str] = []
    if electrode:
        retained = [str(value) for value in source.get("expected_scalp_channels") or []]
        observed = [str(value) for value in source.get("observed_scalp_channels") or []]
        interpolated = [
            str(value)
            for value in source.get("successfully_interpolated_channels") or []
        ]
        target_channels = [electrode]
        coverage_row.update(
            {
                "electrode": electrode,
                "in_retained_scalp": electrode in retained,
                "observed_in_source": electrode in observed,
                "successfully_interpolated": electrode in interpolated,
            }
        )
    evidence.append(coverage_row)

    processing = context.processing_entries.get(identity.casefold())
    if not isinstance(processing, Mapping):
        evidence.append(
            {
                "source": "processing_ledger",
                "status": "processing_entry_missing",
                "authority": "context_only",
            }
        )
    else:
        burden = processing.get("interpolation_burden")
        burden = dict(burden) if isinstance(burden, Mapping) else {}
        evidence.append(
            {
                "source": "interpolation_burden",
                "authority": "context_only",
                "status": (
                    str(burden.get("status") or "current_source_invalid")
                ),
                "fingerprint": str(burden.get("fingerprint") or ""),
                "numerator": burden.get("numerator"),
                "denominator": burden.get("denominator"),
                "percentage": burden.get("percentage"),
                "requires_review": burden.get("requires_review"),
                "successfully_interpolated_channels": list(
                    burden.get("successfully_interpolated_channels") or []
                ),
                "target_interpolated_channels": [
                    channel
                    for channel in target_channels
                    if channel
                    in set(burden.get("successfully_interpolated_channels") or [])
                ],
            }
        )
        kurtosis = processing.get("kurtosis_qc_evidence")
        plan = processing.get("kurtosis_decision_plan")
        kurtosis_status = str(
            processing.get("kurtosis_evidence_status") or "none_available"
        )
        plan_status = str(
            processing.get("kurtosis_decision_plan_status") or "none_available"
        )
        channel_rows = [
            dict(row)
            for row in _iter_mapping_entries(
                kurtosis.get("channels") if isinstance(kurtosis, Mapping) else None
            )
            if _normalize_electrode(row.get("channel")) in set(target_channels)
        ]
        decision_rows = [
            dict(row)
            for row in _iter_mapping_entries(
                plan.get("channel_decisions") if isinstance(plan, Mapping) else None
            )
            if _normalize_electrode(row.get("channel")) in set(target_channels)
        ]
        evidence.append(
            {
                "source": "kurtosis_qc",
                "authority": "context_only",
                "status": (
                    "none_available"
                    if kurtosis_status == plan_status == "none_available"
                    else "current"
                    if kurtosis_status in {"current", "none_available"}
                    and plan_status in {"current", "none_available"}
                    else "current_source_invalid"
                ),
                "evidence_fingerprint": str(
                    kurtosis.get("fingerprint")
                    if isinstance(kurtosis, Mapping)
                    else ""
                ),
                "decision_plan_fingerprint": str(
                    plan.get("fingerprint") if isinstance(plan, Mapping) else ""
                ),
                "channels": channel_rows,
                "channel_decisions": decision_rows,
            }
        )

    invalid = any(
        str(row.get("status") or "").endswith("invalid")
        or str(row.get("status") or "").endswith("missing")
        for row in evidence
    )
    status = "current_source_invalid" if invalid else "available"
    finding["independent_qc"] = evidence
    finding["independent_qc_status"] = status
    finding["independent_qc_authority"] = "context_only"
    finding["independent_qc_fingerprint"] = _hash_payload(
        {"status": status, "evidence": evidence}
    )


def _technical_status_rows(report: Mapping[str, object]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for category, key in (
        ("technical_integrity_failure", "technical_integrity_failures"),
        ("unavailable_by_method", "unavailable_by_method"),
    ):
        rows.extend(
            {"category": category, **dict(item)}
            for item in _iter_mapping_entries(report.get(key))
            if not is_roi_frequency_qc_entry(item)
        )
    rows.extend(
        {"category": "cohort_context", **dict(item)}
        for item in _iter_mapping_entries(report.get("cohort_relative_rows"))
        if str(item.get("status") or "") != "complete"
        and not is_roi_frequency_qc_entry(item)
    )
    source = report.get("independent_qc_source")
    if isinstance(source, Mapping) and str(source.get("status") or "") != "current":
        rows.append({"category": "independent_qc_source", **dict(source)})
    return rows


def _build_review_evidence_payload(
    report: Mapping[str, object],
) -> dict[str, object]:
    identity_scope = str(report.get("identity_scope") or "participant").casefold()
    core: dict[str, object] = {
        "version": FREQUENCY_DOMAIN_QC_REVIEW_EVIDENCE_VERSION,
        "analysis_fingerprint": str(report.get("analysis_fingerprint") or ""),
        "identity_scope": identity_scope,
        "screening_enabled": bool(report.get("screening_enabled")),
        "condition_specific_interpolation_enabled": bool(
            report.get("condition_specific_interpolation_enabled")
        ),
        "screening_status": str(report.get("screening_status") or ""),
        "screening_policy_version": str(
            report.get("screening_policy_version") or ""
        ),
        "screening_settings": dict(report.get("screening_settings") or {}),
        "thresholds": dict(report.get("thresholds") or {}),
        "selected_harmonics_hz": list(report.get("selected_harmonics_hz") or []),
        "harmonic_policy": str(report.get("harmonic_policy") or ""),
        "harmonic_selection_fingerprint": str(
            report.get("harmonic_selection_fingerprint") or ""
        ),
        "provisional_harmonic_metadata": dict(
            report.get("provisional_harmonic_metadata") or {}
        ),
        "frequency_protocol": dict(report.get("frequency_protocol") or {}),
        "frequency_protocol_fingerprint": str(
            dict(report.get("frequency_protocol") or {}).get(
                "frequency_protocol_fingerprint"
            )
        ),
        "source_workbooks": [
            dict(item)
            for item in _iter_mapping_entries(report.get("source_workbooks"))
        ],
        "source_fingerprint": str(report.get("source_fingerprint") or ""),
        "independent_qc_source": dict(
            report.get("independent_qc_source") or {}
        ),
        "ordinary_findings": [
            dict(item) for item in _iter_mapping_entries(report.get("flags"))
            if not is_roi_frequency_qc_entry(item)
        ],
        "cohort_findings": [
            dict(item)
            for item in _iter_mapping_entries(report.get("cohort_relative_flags"))
            if not is_roi_frequency_qc_entry(item)
        ],
        "cohort_rows": [
            dict(item)
            for item in _iter_mapping_entries(report.get("cohort_relative_rows"))
            if not is_roi_frequency_qc_entry(item)
        ],
        "reconfirmation_findings": [
            dict(item)
            for item in _iter_mapping_entries(report.get("reconfirmation_findings"))
            if not is_roi_frequency_qc_entry(item)
        ],
        "finite_input_status": dict(report.get("finite_input_status") or {}),
        "technical_integrity_failures": [
            dict(item)
            for item in _iter_mapping_entries(
                report.get("technical_integrity_failures")
            )
            if not is_roi_frequency_qc_entry(item)
        ],
        "unavailable_by_method": [
            dict(item)
            for item in _iter_mapping_entries(report.get("unavailable_by_method"))
            if not is_roi_frequency_qc_entry(item)
        ],
        "technical_statuses": _technical_status_rows(report),
    }
    return {**core, "evidence_fingerprint": _hash_payload(core)}


def _source_workbooks_are_current(
    project_root: Path,
    rows: object,
) -> bool:
    raw_rows = rows if isinstance(rows, list) else None
    if raw_rows is None:
        return False
    for row in raw_rows:
        if not isinstance(row, Mapping):
            return False
        raw_path = str(row.get("path") or "")
        if not raw_path:
            return False
        path = Path(raw_path)
        if not path.is_absolute():
            path = project_root / path
        try:
            stat = path.resolve(strict=True).stat()
        except OSError:
            return False
        if row.get("size_bytes") != int(stat.st_size):
            return False
        if row.get("mtime_ns") != int(stat.st_mtime_ns):
            return False
    return True


def _validated_review_evidence_from_state(
    project_root: Path,
    state: Mapping[str, object],
) -> dict[str, object] | None:
    raw = state.get("review_evidence")
    if not isinstance(raw, Mapping):
        return None
    payload = dict(raw)
    recorded = str(payload.pop("evidence_fingerprint", "") or "")
    if (
        payload.get("version") != FREQUENCY_DOMAIN_QC_REVIEW_EVIDENCE_VERSION
        or not recorded
        or _hash_payload(payload) != recorded
    ):
        return None
    if payload.get("roi_definition_fingerprint") or payload.get("cohort_fingerprint"):
        return None
    if any(
        is_roi_frequency_qc_entry(row)
        for key in (
            "ordinary_findings", "cohort_findings", "cohort_rows",
            "reconfirmation_findings", "technical_statuses",
        )
        for row in _iter_mapping_entries(payload.get(key))
    ):
        return None
    last_review = state.get("last_review")
    if not isinstance(last_review, Mapping):
        return None
    if (
        str(payload.get("analysis_fingerprint") or "")
        != str(last_review.get("analysis_fingerprint") or "")
        or recorded != str(last_review.get("evidence_fingerprint") or "")
        or str(payload.get("identity_scope") or "")
        != str(last_review.get("identity_scope") or "")
    ):
        return None
    for finding in [
        *_iter_mapping_entries(payload.get("ordinary_findings")),
        *_iter_mapping_entries(payload.get("cohort_findings")),
        *_iter_mapping_entries(payload.get("reconfirmation_findings")),
    ]:
        if str(finding.get("finding_fingerprint") or "") != (
            _frequency_qc_finding_fingerprint(finding)
        ):
            return None
    if not _source_workbooks_are_current(
        project_root,
        payload.get("source_workbooks"),
    ):
        return None
    saved_independent = payload.get("independent_qc_source")
    current_independent = _load_independent_qc_context(project_root).source_identity
    if not isinstance(saved_independent, Mapping) or dict(saved_independent) != (
        current_independent
    ):
        return None
    return {**payload, "evidence_fingerprint": recorded}


def load_current_frequency_qc_review_evidence(
    project_root: str | Path,
) -> dict[str, object]:
    """Load fingerprint-current saved QC-17 evidence or fail closed."""

    root = Path(project_root).resolve()
    decisions = resolve_frequency_qc_coverage_decisions(root)
    if not decisions.review_complete:
        raise RuntimeError(
            "Saved experimental summed-BCA decisions are missing, stale, or "
            "tampered. Run post-processing and complete QC-17 review again."
        )
    state = load_frequency_domain_qc_state(root)
    evidence = _validated_review_evidence_from_state(root, state)
    if evidence is None:
        raise RuntimeError(
            "Saved experimental summed-BCA evidence is missing, stale, or tampered. "
            "Run post-processing and complete QC-17 review again."
        )
    return evidence


def _frequency_qc_finding_fingerprint(
    finding: Mapping[str, object],
) -> str:
    return _hash_payload(
        {
            key: _json_safe(finding.get(key))
            for key in (
                "finding_type",
                "participant_id",
                "recording_id",
                "session_id",
                "visit_index",
                "condition",
                "electrode",
                "roi",
                "metric",
                "summed_bca_uv",
                "abs_summed_bca_uv",
                "value_uv",
                "robust_center_uv",
                "robust_spread_uv",
                "robust_spread_method",
                "robust_score",
                "threshold_used",
                "absolute_floor_used_uv",
                "severity",
                "selected_harmonics_hz",
                "selection_fingerprint",
                "harmonic_selection_fingerprint",
                "frequency_protocol_fingerprint",
                "expected_analyzed_oddball_cycles",
                "analyzed_duration_seconds",
                "roi_definition_fingerprint",
                "cohort_fingerprint",
                "independent_qc_status",
                "independent_qc_authority",
                "independent_qc_fingerprint",
                "independent_qc",
            )
        }
    )


def _review_decisions_from_state(
    state: Mapping[str, object],
    *, include_retired_roi: bool = False,
) -> list[dict[str, object]]:
    if not bool(state.get("review_complete")):
        return []
    normalized: list[dict[str, object]] = []
    for raw in _iter_mapping_entries(state.get("review_decisions")):
        if not include_retired_roi and is_roi_frequency_qc_entry(raw):
            continue
        row = dict(raw)
        if row.get("version") != FREQUENCY_DOMAIN_QC_DECISION_VERSION:
            continue
        recorded = str(row.pop("decision_fingerprint", "") or "")
        reviewed_at = row.pop("reviewed_at", None)
        if not recorded or _hash_payload(row) != recorded:
            continue
        row["decision_fingerprint"] = recorded
        if reviewed_at not in (None, ""):
            row["reviewed_at"] = reviewed_at
        normalized.append(row)
    return sorted(
        normalized,
        key=lambda item: str(item.get("finding_fingerprint") or ""),
    )


def _current_review_decisions(
    *,
    report_flags: Sequence[Mapping[str, object]],
    state: Mapping[str, object],
) -> list[dict[str, object]]:
    current_fingerprints = {
        str(flag.get("finding_fingerprint") or "")
        for flag in report_flags
        if str(flag.get("finding_fingerprint") or "")
    }
    analysis_fingerprints = {
        str(flag.get("analysis_fingerprint") or "")
        for flag in report_flags
        if str(flag.get("analysis_fingerprint") or "")
    }
    decisions = [
        item
        for item in _review_decisions_from_state(state)
        if str(item.get("finding_fingerprint") or "") in current_fingerprints
        and item.get("decision") in REVIEW_DECISIONS
    ]
    if analysis_fingerprints and any(
        str(item.get("analysis_fingerprint") or "") not in analysis_fingerprints
        for item in decisions
    ):
        return []
    return decisions


def _resolved_reconfirmation_retains(
    *,
    state: Mapping[str, object],
    analysis_fingerprint: str,
    review_evidence: Mapping[str, object] | None,
) -> list[dict[str, object]]:
    """Keep resolved review receipts while their validated evidence is unchanged.

    A prior exclusion changed to Retain no longer produces a reconfirmation
    finding. Its receipt still belongs to that completed review; dropping it
    would change the decision fingerprint and reopen already-decided findings.
    The caller supplies only independently validated saved review evidence.
    """

    if (
        review_evidence is None
        or str(review_evidence.get("analysis_fingerprint") or "")
        != analysis_fingerprint
    ):
        return []
    resolved_findings = {
        str(finding.get("finding_fingerprint") or "")
        for finding in _iter_mapping_entries(
            review_evidence.get("reconfirmation_findings")
        )
        if finding.get("finding_type")
        == "prior_outcome_informed_exclusion_reconfirmation"
    }
    return [
        decision for decision in _review_decisions_from_state(state)
        if decision.get("decision") == DECISION_RETAIN
        and decision.get("analysis_fingerprint") == analysis_fingerprint
        and decision.get("replaces_decision_fingerprint")
        and decision.get("finding_fingerprint") in resolved_findings
    ]


def _merge_review_decision_rows(
    *groups: Sequence[Mapping[str, object]],
) -> list[dict[str, object]]:
    by_fingerprint: dict[str, dict[str, object]] = {}
    for group in groups:
        for item in group:
            fingerprint = str(item.get("decision_fingerprint") or "")
            if not fingerprint:
                continue
            by_fingerprint[fingerprint] = dict(item)
    return [by_fingerprint[key] for key in sorted(by_fingerprint)]


def _superseded_narrow_decisions(
    state: Mapping[str, object], current: Sequence[Mapping[str, object]],
) -> list[dict[str, object]]:
    """Keep retired exclusions and completed repair approvals as audit only."""
    current_ids = {str(row.get("decision_fingerprint") or "") for row in current}
    return _merge_review_decision_rows(
        _iter_mapping_entries(state.get("retired_review_decisions")),
        [
            row for row in _review_decisions_from_state(
                {**state, "review_complete": True}, include_retired_roi=True,
            )
            if (is_roi_frequency_qc_entry(row)
                or row.get("decision") not in {*_BROAD_EXCLUSION_DECISIONS, DECISION_RETAIN})
            and str(row.get("decision_fingerprint") or "") not in current_ids
        ],
    )


def _review_exclusion_reconfirmation_findings(
    state: Mapping[str, object],
    *,
    evidence_context_fingerprint: str,
    selected_harmonics: Sequence[float],
    harmonic_selection_fingerprint: str,
    protocol_metadata: Mapping[str, object],
    independent_qc_context: _IndependentQcContext,
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    """Carry stable exclusions and reopen those bound to older evidence."""

    reconfirmation_findings: list[dict[str, object]] = []
    stable_exclusions: list[dict[str, object]] = []
    for decision in _review_decisions_from_state(state):
        action = str(decision.get("decision") or "")
        if action in {DECISION_RETAIN, DECISION_INTERPOLATE_CONDITION_ELECTRODE}:
            continue
        retired_action = action not in _BROAD_EXCLUSION_DECISIONS
        if not retired_action and str(decision.get("analysis_fingerprint") or "") == str(
            evidence_context_fingerprint
        ):
            stable_exclusions.append(decision)
            continue
        evidence = (
            dict(decision.get("evidence"))
            if isinstance(decision.get("evidence"), Mapping)
            else {}
        )
        signed_value = evidence.get("summed_bca_uv")
        absolute_value = evidence.get("abs_summed_bca_uv")
        if absolute_value is None:
            absolute_value = evidence.get("value_uv")
        finding: dict[str, object] = {
            "finding_type": "prior_outcome_informed_exclusion_reconfirmation",
            "participant_id": _normalize_participant_id(
                decision.get("participant_id")
            ),
            "recording_id": _normalize_recording_id(
                decision.get("recording_id")
            ),
            "session_id": str(decision.get("session_id") or ""),
            "visit_index": decision.get("visit_index"),
            "condition": str(decision.get("condition") or ""),
            "electrode": _normalize_electrode(decision.get("electrode")),
            "roi": str(decision.get("roi") or ""),
            "summed_bca_uv": signed_value,
            "abs_summed_bca_uv": absolute_value,
            "value_uv": evidence.get("value_uv"),
            "metric": evidence.get("metric"),
            "severity": "reconfirmation",
            "band_crossed": (
                "reconfirmation required; prior band "
                f"{evidence.get('band_crossed') or evidence.get('severity') or 'unknown'}"
            ),
            "selected_harmonics_hz": [
                round(float(value), 4) for value in selected_harmonics
            ],
            "selected_harmonic_count": len(selected_harmonics),
            "harmonic_selection_fingerprint": harmonic_selection_fingerprint,
            "prior_selected_harmonics_hz": list(
                evidence.get("selected_harmonics_hz") or []
            ),
            "prior_decision": str(decision.get("decision") or ""),
            "prior_reason": str(decision.get("reason") or ""),
            "replaces_decision_fingerprint": str(
                decision.get("decision_fingerprint") or ""
            ),
            "reconfirmation_reason": (
                "Electrode and ROI exclusions are no longer supported. Choose Retain, "
                "an enabled electrode repair, or a broader exclusion."
                if retired_action else
                "The candidate harmonic/cohort evidence changed after this "
                "outcome-informed exclusion. Confirm or revise its scope."
            ),
            "frequency_protocol_fingerprint": protocol_metadata.get(
                "frequency_protocol_fingerprint"
            ),
            "expected_analyzed_oddball_cycles": protocol_metadata.get(
                "expected_analyzed_oddball_cycles"
            ),
            "analyzed_duration_seconds": protocol_metadata.get(
                "analyzed_duration_seconds"
            ),
            "prior_independent_qc_fingerprint": str(
                evidence.get("independent_qc_fingerprint") or ""
            ),
        }
        _attach_independent_qc_evidence(finding, independent_qc_context)
        finding["finding_fingerprint"] = _frequency_qc_finding_fingerprint(
            {
                **finding,
                "evidence_context_fingerprint": evidence_context_fingerprint,
                "prior_decision_fingerprint": decision.get(
                    "decision_fingerprint"
                ),
            }
        )
        reconfirmation_findings.append(finding)
    return (
        sorted(
            reconfirmation_findings,
            key=lambda item: str(item.get("finding_fingerprint") or ""),
        ),
        stable_exclusions,
    )


def _legacy_machine_suggestions_from_state(
    state: Mapping[str, object],
) -> list[dict[str, object]]:
    return _merge_legacy_machine_suggestions(state, None)


def _merge_legacy_machine_suggestions(
    state: Mapping[str, object],
    report_value: object,
) -> list[dict[str, object]]:
    rows = [
        dict(item)
        for item in _iter_mapping_entries(state.get("legacy_machine_suggestions"))
    ]
    rows.extend(dict(item) for item in _iter_mapping_entries(report_value))
    for state_key, suggestion_type in (
        ("auto_participant_electrode_exclusions", "participant_electrode"),
        ("auto_participant_exclusions", "participant"),
        ("auto_recording_electrode_exclusions", "recording_electrode"),
        ("auto_recording_exclusions", "recording"),
    ):
        for item in _iter_mapping_entries(state.get(state_key)):
            rows.append(
                {
                    **dict(item),
                    "legacy_source_field": state_key,
                    "suggestion_type": suggestion_type,
                    "authority": "review_only_legacy_suggestion",
                    "migration_version": FREQUENCY_DOMAIN_QC_METHOD_VERSION,
                }
            )
    by_fingerprint: dict[str, dict[str, object]] = {}
    for row in rows:
        fingerprint = _hash_payload(row)
        by_fingerprint[fingerprint] = {**row, "suggestion_fingerprint": fingerprint}
    return [by_fingerprint[key] for key in sorted(by_fingerprint)]


def _updated_review_history(
    state: Mapping[str, object],
    *,
    report: Mapping[str, object],
    decision_fingerprint: str,
    reviewed_at: str,
) -> list[dict[str, object]]:
    history = [
        dict(item)
        for item in _iter_mapping_entries(state.get("review_history"))
    ]
    history.append(
        {
            "loop_version": FREQUENCY_DOMAIN_QC_LOOP_VERSION,
            "analysis_fingerprint": str(report.get("analysis_fingerprint") or ""),
            "decision_fingerprint": str(decision_fingerprint),
            "reviewed_at": reviewed_at,
            "selected_harmonics_hz": list(
                report.get("selected_harmonics_hz") or []
            ),
            "subject_count": len(report.get("recordings") or report.get("subjects") or []),
        }
    )
    return history[-FREQUENCY_DOMAIN_QC_MAX_REVIEW_ITERATIONS:]


def _frequency_qc_review_loop_status(
    state: Mapping[str, object],
    *,
    analysis_fingerprint: str,
) -> dict[str, object]:
    """Reject an oscillating or unbounded review/recompute sequence."""

    history = [
        dict(item)
        for item in _iter_mapping_entries(state.get("review_history"))
        if item.get("loop_version") == FREQUENCY_DOMAIN_QC_LOOP_VERSION
        and str(item.get("analysis_fingerprint") or "")
    ]
    fingerprints = [
        str(item.get("analysis_fingerprint") or "") for item in history
    ]
    if (
        len(fingerprints) >= FREQUENCY_DOMAIN_QC_MAX_REVIEW_ITERATIONS
        and analysis_fingerprint != fingerprints[-1]
    ):
        raise RuntimeError(
            "Experimental summed-BCA review reached its bounded recompute limit "
            f"({FREQUENCY_DOMAIN_QC_MAX_REVIEW_ITERATIONS} reviewed states). "
            "Resolve the changing inclusion decisions before rerunning."
        )
    return {
        "method_version": FREQUENCY_DOMAIN_QC_LOOP_VERSION,
        "completed_review_iterations": len(fingerprints),
        "next_review_iteration": len(fingerprints) + 1,
        "maximum_review_iterations": FREQUENCY_DOMAIN_QC_MAX_REVIEW_ITERATIONS,
        "current_analysis_fingerprint": analysis_fingerprint,
        "previous_analysis_fingerprint": fingerprints[-1] if fingerprints else "",
        "status": (
            "stable_review_state"
            if fingerprints and analysis_fingerprint == fingerprints[-1]
            else "revisited_review_state"
            if analysis_fingerprint in fingerprints
            else "new_review_state"
        ),
    }


def _require_non_oscillating_review_decision(
    state: Mapping[str, object],
    *,
    analysis_fingerprint: str,
    decision_fingerprint: str,
) -> None:
    history = [
        dict(item)
        for item in _iter_mapping_entries(state.get("review_history"))
        if item.get("loop_version") == FREQUENCY_DOMAIN_QC_LOOP_VERSION
    ]
    if not history:
        return
    repeated_pair = any(
        str(item.get("analysis_fingerprint") or "") == analysis_fingerprint
        and str(item.get("decision_fingerprint") or "") == decision_fingerprint
        for item in history
    )
    if repeated_pair and str(history[-1].get("analysis_fingerprint") or "") != (
        analysis_fingerprint
    ):
        raise RuntimeError(
            "Experimental summed-BCA review is oscillating: the same review "
            "state and decision were already followed by a different candidate "
            "state. Revise the inclusion choice or project protocol before "
            "continuing."
        )


def _metadata_from_manifest(manifest: Mapping[str, object] | None) -> dict[str, object]:
    current: Mapping[str, object] = manifest if isinstance(manifest, Mapping) else {}
    for key in FREQUENCY_DOMAIN_QC_METADATA_PATH:
        value = current.get(key)
        if not isinstance(value, Mapping):
            return {}
        current = value
    return dict(current)


def _set_metadata_in_manifest(manifest: dict[str, object], state: Mapping[str, object]) -> None:
    current: dict[str, object] = manifest
    for key in FREQUENCY_DOMAIN_QC_METADATA_PATH[:-1]:
        child = current.get(key)
        if not isinstance(child, dict):
            child = {}
            current[key] = child
        current = child
    current[FREQUENCY_DOMAIN_QC_METADATA_PATH[-1]] = _json_safe(dict(state))


def _read_manifest(manifest_path: Path) -> dict[str, object]:
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def _write_manifest_if_changed(manifest_path: Path, manifest: Mapping[str, object]) -> None:
    new_payload = json.dumps(_json_safe(dict(manifest)), sort_keys=True, separators=(",", ":"))
    if manifest_path.exists():
        try:
            current = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            current = {}
        current_payload = json.dumps(current, sort_keys=True, separators=(",", ":"))
        if current_payload == new_payload:
            return
    manifest_path.write_text(
        json.dumps(_json_safe(dict(manifest)), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _auto_electrode_entries_from_state(state: Mapping[str, object]) -> list[dict[str, object]]:
    return _normalize_auto_electrode_entries(state.get("auto_participant_electrode_exclusions"))


def _auto_participant_entries_from_state(state: Mapping[str, object]) -> list[dict[str, object]]:
    return _normalize_auto_participant_entries(state.get("auto_participant_exclusions"))


def _manual_entries_from_state(state: Mapping[str, object]) -> list[dict[str, object]]:
    return _normalize_manual_entries(state.get("manual_participant_exclusions"))


def _auto_recording_electrode_entries_from_state(
    state: Mapping[str, object],
) -> list[dict[str, object]]:
    return _normalize_auto_recording_electrode_entries(
        state.get("auto_recording_electrode_exclusions")
    )


def _auto_recording_entries_from_state(
    state: Mapping[str, object],
) -> list[dict[str, object]]:
    return _normalize_auto_recording_entries(state.get("auto_recording_exclusions"))


def _manual_recording_entries_from_state(
    state: Mapping[str, object],
) -> list[dict[str, object]]:
    return _normalize_manual_recording_entries(state.get("manual_recording_exclusions"))


def _normalize_auto_electrode_entries(value: object) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for item in _iter_mapping_entries(value):
        pid = _normalize_participant_id(item.get("participant_id"))
        electrode = _normalize_electrode(item.get("electrode"))
        if not pid or not electrode:
            continue
        conditions = sorted({str(condition) for condition in item.get("triggering_conditions", []) or []})
        entries.append(
            {
                "participant_id": pid,
                "electrode": electrode,
                "reason": str(item.get("reason") or "abs summed BCA exceeded hard electrode threshold"),
                "threshold_uv": float(item.get("threshold_uv") or DEFAULT_FREQUENCY_DOMAIN_QC_THRESHOLDS.hard_electrode_summed_bca_uv),
                "max_abs_summed_bca_uv": float(item.get("max_abs_summed_bca_uv") or 0.0),
                "triggering_conditions": conditions,
                "source": str(item.get("source") or "automatic_frequency_domain_qc"),
            }
        )
    return sorted(entries, key=lambda entry: (entry["participant_id"], entry["electrode"]))


def _normalize_auto_participant_entries(value: object) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for item in _iter_mapping_entries(value):
        pid = _normalize_participant_id(item.get("participant_id"))
        if not pid:
            continue
        entries.append(
            {
                "participant_id": pid,
                "reason": str(item.get("reason") or "more than 10 unique electrodes exceeded hard electrode threshold"),
                "hard_excluded_electrode_count": int(item.get("hard_excluded_electrode_count") or 0),
                "source": str(item.get("source") or "automatic_frequency_domain_qc"),
            }
        )
    return sorted(entries, key=lambda entry: entry["participant_id"])


def _normalize_manual_entries(value: object) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for item in _iter_mapping_entries(value):
        pid = _normalize_participant_id(item.get("participant_id"))
        if not pid:
            continue
        reason = str(item.get("reason") or WARNING_REASON_UNUSUAL_VALUES)
        if reason not in MANUAL_EXCLUSION_REASONS and reason != "No reason provided":
            reason = WARNING_REASON_UNUSUAL_VALUES
        entry = {
            "participant_id": pid,
            "reason": reason,
            "source": str(item.get("source") or "manual_qc_review"),
        }
        if item.get("added_at"):
            entry["added_at"] = str(item.get("added_at"))
        if item.get("updated_at"):
            entry["updated_at"] = str(item.get("updated_at"))
        entries.append(entry)
    return sorted(entries, key=lambda entry: entry["participant_id"])


def _normalize_auto_recording_electrode_entries(
    value: object,
) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for item in _iter_mapping_entries(value):
        recording_id = _normalize_recording_id(item.get("recording_id"))
        electrode = _normalize_electrode(item.get("electrode"))
        if not recording_id or not electrode:
            continue
        conditions = sorted(
            {
                str(condition)
                for condition in item.get("triggering_conditions", []) or []
            }
        )
        entries.append(
            {
                "recording_id": recording_id,
                "participant_id": _normalize_participant_id(
                    item.get("participant_id")
                ),
                "session_id": str(item.get("session_id") or ""),
                "visit_index": _optional_int(item.get("visit_index")),
                "electrode": electrode,
                "reason": str(
                    item.get("reason")
                    or "abs summed BCA exceeded hard electrode threshold"
                ),
                "threshold_uv": float(
                    item.get("threshold_uv")
                    or DEFAULT_FREQUENCY_DOMAIN_QC_THRESHOLDS.hard_electrode_summed_bca_uv
                ),
                "max_abs_summed_bca_uv": float(
                    item.get("max_abs_summed_bca_uv") or 0.0
                ),
                "triggering_conditions": conditions,
                "source": str(
                    item.get("source") or "automatic_frequency_domain_qc"
                ),
            }
        )
    return sorted(
        entries,
        key=lambda entry: (str(entry["recording_id"]), str(entry["electrode"])),
    )


def _normalize_auto_recording_entries(value: object) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for item in _iter_mapping_entries(value):
        recording_id = _normalize_recording_id(item.get("recording_id"))
        if not recording_id:
            continue
        entries.append(
            {
                "recording_id": recording_id,
                "participant_id": _normalize_participant_id(
                    item.get("participant_id")
                ),
                "session_id": str(item.get("session_id") or ""),
                "visit_index": _optional_int(item.get("visit_index")),
                "reason": str(
                    item.get("reason")
                    or (
                        "more than 10 unique electrodes exceeded hard electrode "
                        "threshold in this recording"
                    )
                ),
                "hard_excluded_electrode_count": int(
                    item.get("hard_excluded_electrode_count") or 0
                ),
                "source": str(
                    item.get("source") or "automatic_frequency_domain_qc"
                ),
            }
        )
    return sorted(entries, key=lambda entry: str(entry["recording_id"]))


def _normalize_manual_recording_entries(value: object) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for item in _iter_mapping_entries(value):
        recording_id = str(item.get("recording_id") or "").strip()
        if not recording_id:
            continue
        reason = str(item.get("reason") or WARNING_REASON_UNUSUAL_VALUES)
        if reason not in MANUAL_EXCLUSION_REASONS and reason != "No reason provided":
            reason = WARNING_REASON_UNUSUAL_VALUES
        entry: dict[str, object] = {
            "recording_id": recording_id,
            "participant_id": str(item.get("participant_id") or "").strip(),
            "session_id": str(item.get("session_id") or ""),
            "reason": reason,
            "source": str(item.get("source") or "manual_qc_review"),
        }
        if item.get("added_at"):
            entry["added_at"] = str(item.get("added_at"))
        if item.get("updated_at"):
            entry["updated_at"] = str(item.get("updated_at"))
        entries.append(entry)
    return sorted(entries, key=lambda entry: str(entry["recording_id"]).casefold())


def _iter_mapping_entries(value: object) -> list[Mapping[str, object]]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes, bytearray)):
        return []
    return [item for item in value if isinstance(item, Mapping)]


def _normalize_participant_id(value: object) -> str:
    text = str(value or "").strip().upper()
    return text


def _normalize_recording_id(value: object) -> str:
    return str(value or "").strip().upper()


def _optional_int(value: object) -> int | None:
    try:
        return int(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def _normalize_electrode(value: object) -> str:
    return str(value or "").strip().upper()


def _hash_payload(payload: Mapping[str, object]) -> str:
    normalized = json.dumps(_json_safe(dict(payload)), sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _json_safe(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set, frozenset)):
        return [_json_safe(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, float):
        if not np.isfinite(value):
            return None
        return float(value)
    return value


def _now_utc_iso() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat()


def _manifest_safe_path(project_root: Path, path: Path) -> str:
    try:
        root = project_root.resolve()
        resolved = Path(path).resolve(strict=False)
        if resolved == root or root in resolved.parents:
            return str(resolved.relative_to(root))
        return str(resolved)
    except OSError:
        return str(path)


__all__ = [
    "DECISION_EXCLUDE_CONDITION",
    "DECISION_EXCLUDE_CONDITION_ELECTRODE",
    "DECISION_INTERPOLATE_CONDITION_ELECTRODE",
    "DECISION_EXCLUDE_PARTICIPANT",
    "DECISION_EXCLUDE_RECORDING",
    "DECISION_RETAIN",
    "DEFAULT_FREQUENCY_DOMAIN_QC_THRESHOLDS",
    "FREQUENCY_DOMAIN_QC_INTEGRITY_METHOD_VERSION",
    "FREQUENCY_DOMAIN_QC_REPORT_NAME",
    "FREQUENCY_DOMAIN_QC_REVIEW_EVIDENCE_VERSION",
    "MANUAL_EXCLUSION_REASONS",
    "REVIEW_DECISIONS",
    "SUMMED_BCA_SCREENING_BRIEF_TEXT",
    "FrequencyDomainCoverageDecisions",
    "FrequencyDomainExclusions",
    "FrequencyDomainQcIntegrityError",
    "FrequencyDomainQcThresholds",
    "REPEATED_FREQUENCY_DOMAIN_QC_METHOD_VERSION",
    "active_frequency_domain_exclusions",
    "frequency_domain_exclusions_from_state",
    "apply_frequency_domain_qc_decision",
    "clear_manual_frequency_domain_participant_exclusions",
    "clear_manual_frequency_domain_recording_exclusions",
    "filter_frequency_domain_recordings",
    "filter_frequency_domain_subjects",
    "frequency_domain_excluded_electrodes_for_subject",
    "frequency_domain_excluded_electrodes_for_recording",
    "is_frequency_domain_output_stale",
    "load_frequency_domain_qc_state",
    "load_current_frequency_qc_review_evidence",
    "mark_frequency_domain_outputs_current",
    "mark_frequency_domain_outputs_stale",
    "require_frequency_domain_qc_complete",
    "resolve_frequency_qc_coverage_decisions",
    "run_frequency_domain_qc_review",
    "sync_frequency_domain_qc_automatic_state",
    "thresholds_summary_lines",
    "validate_frequency_domain_qc_review_decisions",
]
