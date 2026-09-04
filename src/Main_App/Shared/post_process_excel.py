from contextlib import contextmanager
import hashlib
import os
from pathlib import Path
import shutil
import tempfile
from time import perf_counter
from typing import Any, Dict, List, Optional

import numpy as np
from openpyxl import load_workbook
import pandas as pd


NEIGHBOR_OFFSETS = [*range(-11, 0), *range(1, 12)]
_COLUMN_WIDTH_CHUNK_SIZE = 1024
FFT_METADATA_SHEET_NAME = "FFT Metadata"
SPECTRAL_ELIGIBILITY_SHEET_NAME = "Spectral Eligibility"
SPECTRAL_METRIC_QC_SHEET_NAME = "Spectral Metric QC"
FFT_METADATA_COLUMN_MAP = {
    "file_name": "Source File",
    "condition_label": "Condition",
    "repetition_index": "FFT Input Index",
    "fs": "Sampling Frequency (Hz)",
    "N": "FFT Sample Count (N)",
    "T_sec": "FFT Duration (s)",
    "df_hz": "FFT Bin Width (Hz)",
    "crop_mode": "Crop Mode",
}
WORKBOOK_WRITE_RECEIPT_VERSION = "workbook_write_receipt_v1"


def _elapsed_ms(started_at: float) -> int:
    return int((perf_counter() - started_at) * 1000)


def _log_excel_timing(
    stage: str,
    started_at: float,
    *,
    path: str,
    sheet_name: str | None = None,
    rows: int | None = None,
    cols: int | None = None,
    timing_sink: list[dict[str, object]] | None = None,
) -> None:
    elapsed_ms = _elapsed_ms(started_at)
    record = {
        "source": "excel",
        "stage": stage,
        "elapsed_ms": elapsed_ms,
        "path": path,
        "sheet": sheet_name,
        "rows": rows,
        "cols": cols,
    }
    if timing_sink is not None:
        timing_sink.append(record)


def _column_widths(frame: pd.DataFrame) -> tuple[int, ...]:
    """Return the established ``max(str length) + 4`` widths in bounded blocks."""

    widths: list[int] = []
    for start in range(0, len(frame.columns), _COLUMN_WIDTH_CHUNK_SIZE):
        chunk = frame.iloc[:, start : start + _COLUMN_WIDTH_CHUNK_SIZE]
        if chunk.empty:
            value_lengths = [0] * len(chunk.columns)
        else:
            string_lengths = chunk.astype(str).map(len)
            value_lengths = [
                int(value)
                for value in string_lengths.max(axis=0).to_numpy()
            ]
        widths.extend(
            max(len(str(header_name)), value_lengths[index]) + 4
            for index, header_name in enumerate(chunk.columns)
        )
    return tuple(widths)


def _apply_column_widths(
    worksheet: Any,
    frame: pd.DataFrame,
    center_format: Any,
) -> None:
    for column_index, width in enumerate(_column_widths(frame)):
        worksheet.set_column(
            column_index,
            column_index,
            width,
            center_format,
        )


def _can_write_finite_metric_frame_direct(frame: pd.DataFrame) -> bool:
    """Return whether direct XlsxWriter cells reproduce pandas output exactly."""

    if frame.empty or len(frame.columns) < 2:
        return False
    if not all(isinstance(column, str) for column in frame.columns):
        return False
    electrode_values = frame.iloc[:, 0]
    if frame.columns[0] != "Electrode" or not electrode_values.map(
        lambda value: isinstance(value, str)
    ).all():
        return False
    numeric = frame.iloc[:, 1:]
    if not all(dtype == np.dtype(np.float64) for dtype in numeric.dtypes):
        return False
    return bool(np.all(np.isfinite(numeric.to_numpy(dtype=float, copy=False))))


def _write_dataframe_to_excel(
    writer: pd.ExcelWriter,
    *,
    sheet_name: str,
    frame: pd.DataFrame,
) -> Any:
    """Write one frame, batching only the exact finite metric-sheet shape."""

    if not _can_write_finite_metric_frame_direct(frame):
        frame.to_excel(writer, sheet_name=sheet_name, index=False)
        return writer.sheets[sheet_name]

    # Let pandas create the header and its cached style exactly as before.
    frame.iloc[:0].to_excel(writer, sheet_name=sheet_name, index=False)
    worksheet = writer.sheets[sheet_name]
    for column_index in range(len(frame.columns)):
        worksheet.write_column(
            1,
            column_index,
            frame.iloc[:, column_index].tolist(),
        )
    return worksheet


def _path_volume(path: Path) -> str:
    absolute = path.expanduser().absolute()
    return os.path.splitdrive(os.fspath(absolute))[0].casefold()


def _should_stage_workbook_locally(destination: Path) -> bool:
    """Return whether final XLSX assembly would otherwise target another volume."""

    destination_volume = _path_volume(destination)
    temporary_volume = _path_volume(Path(tempfile.gettempdir()))
    return bool(
        destination_volume
        and temporary_volume
        and destination_volume != temporary_volume
    )


def _artifact_identity(path: Path) -> dict[str, object] | None:
    if not path.is_file():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
        "sha256": digest.hexdigest(),
    }


def workbook_artifact_identity(
    path: str | os.PathLike[str],
) -> dict[str, object] | None:
    """Return the immutable identity used by current-run workbook receipts."""

    return _artifact_identity(Path(path))


def _validate_workbook_schema(
    path: Path,
    expected_headers: Dict[str, List[str]],
) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        actual_sheets = list(workbook.sheetnames)
        expected_sheets = list(expected_headers)
        if actual_sheets != expected_sheets:
            raise ValueError(
                "Written workbook sheet order does not match the requested schema: "
                f"expected={expected_sheets!r}, actual={actual_sheets!r}."
            )
        validated_headers: dict[str, list[str]] = {}
        for sheet_name, headers in expected_headers.items():
            worksheet = workbook[sheet_name]
            actual_headers = [
                str(cell.value) if cell.value is not None else ""
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            ]
            if actual_headers != headers:
                raise ValueError(
                    "Written workbook headers do not match the requested schema "
                    f"for {sheet_name!r}: expected={headers!r}, "
                    f"actual={actual_headers!r}."
                )
            validated_headers[sheet_name] = actual_headers
        return {
            "status": "passed",
            "sheet_names": actual_sheets,
            "headers": validated_headers,
        }
    finally:
        workbook.close()


@contextmanager
def _workbook_write_target(destination: Path):
    """Yield a staging path and atomically replace the destination on success."""

    if not _should_stage_workbook_locally(destination):
        publish_descriptor, publish_name = tempfile.mkstemp(
            dir=destination.parent,
            prefix=f".{destination.stem}.",
            suffix=".tmp.xlsx",
        )
        os.close(publish_descriptor)
        publish_path = Path(publish_name)
        try:
            yield os.fspath(publish_path)
            os.replace(publish_path, destination)
        finally:
            publish_path.unlink(missing_ok=True)
        return

    stage_descriptor, stage_name = tempfile.mkstemp(
        prefix="fpvs-workbook-",
        suffix=".xlsx",
    )
    os.close(stage_descriptor)
    stage_path = Path(stage_name)
    publish_path: Path | None = None
    try:
        yield os.fspath(stage_path)
        publish_descriptor, publish_name = tempfile.mkstemp(
            dir=destination.parent,
            prefix=f".{destination.stem}.",
            suffix=".xlsx.tmp",
        )
        os.close(publish_descriptor)
        publish_path = Path(publish_name)
        shutil.copyfile(stage_path, publish_path)
        os.replace(publish_path, destination)
        publish_path = None
    finally:
        stage_path.unlink(missing_ok=True)
        if publish_path is not None:
            publish_path.unlink(missing_ok=True)


def build_fft_neighbors_rows(
    *,
    file_name: str,
    condition_label: str,
    condition_id: str,
    repetition_index: str,
    electrode_names: List[str],
    fft_amplitudes: np.ndarray,
    freqs: np.ndarray,
    fs: float,
    n_samples: int,
    target_freq: float = 1.2,
    crop_mode: str = "55_onbin",
    n55: Optional[int] = None,
    first55_samp: Optional[int] = None,
    last55_samp: Optional[int] = None,
    n_step: Optional[int] = None,
    fallback_reason: str = "",
) -> List[Dict[str, Any]]:
    """Build per-channel FFT neighbor rows (±11 bins, excluding center bin)."""
    rows: List[Dict[str, Any]] = []
    if len(freqs) == 0:
        return rows

    if crop_mode not in {
        "55_onbin",
        "project_marker_plan_target_grid_v2",
    }:
        raise ValueError(
            "Locked FFT crop required for FFT-neighbor export: "
            f"crop_mode={crop_mode}, fallback_reason={fallback_reason or 'unknown'}."
        )
    if not n_step:
        raise ValueError("Locked FFT crop metadata requires N_step for FFT-neighbor export.")
    exact_position = target_freq * n_samples / fs
    k0 = int(round(exact_position))
    if abs(exact_position - k0) >= 1e-9:
        raise ValueError(
            "FFT-neighbor target is not locked to an FFT bin: "
            f"target={target_freq}, N={n_samples}, fs={fs}, k={exact_position:.12g}. "
            "Nearest-bin fallback is disabled."
        )
    if not (0 <= k0 < len(freqs)):
        raise ValueError(
            "FFT-neighbor target bin is outside the FFT frequency grid: "
            f"target={target_freq}, N={n_samples}, fs={fs}, k={k0}."
        )
    f_bin_hz = float(freqs[k0]) if 0 <= k0 < len(freqs) else np.nan
    n_mod_step = int(n_samples % n_step)

    if n_mod_step != 0:
        raise ValueError(
            f"Invalid on-bin metadata for FFT neighbors row: N={n_samples}, N_step={n_step}, N_mod_step={n_mod_step}"
        )

    for chan_idx, channel_name in enumerate(electrode_names):
        row: Dict[str, Any] = {
            "file_name": file_name,
            "condition_label": condition_label,
            "condition_id": condition_id,
            "repetition_index": repetition_index,
            "channel_or_roi": channel_name,
            "target": f"{float(target_freq):g}Hz",
            "fs": float(fs),
            "N": int(n_samples),
            "T_sec": float(n_samples / fs) if fs else np.nan,
            "df_hz": float(fs / n_samples) if n_samples else np.nan,
            "k0": int(k0),
            "f_bin_hz": f_bin_hz,
            "crop_mode": crop_mode,
            "n55": int(n55) if n55 is not None else np.nan,
            "first55_samp": int(first55_samp) if first55_samp is not None else np.nan,
            "last55_samp": int(last55_samp) if last55_samp is not None else np.nan,
            "N_step": int(n_step) if n_step is not None else np.nan,
            "N_mod_step": n_mod_step,
            "fallback_reason": "",
            "warning": "",
        }
        out_of_bounds = []
        for offset in NEIGHBOR_OFFSETS:
            col_name = f"amp_m{abs(offset)}" if offset < 0 else f"amp_p{offset}"
            neighbor_idx = k0 + offset
            if 0 <= neighbor_idx < len(freqs):
                row[col_name] = float(fft_amplitudes[chan_idx, neighbor_idx])
            else:
                row[col_name] = np.nan
                out_of_bounds.append(offset)
        if out_of_bounds:
            row["warning"] = (
                "neighbor bins out of range: "
                + ",".join(str(offset) for offset in out_of_bounds)
            )
        rows.append(row)
    return rows


def build_fft_metadata_frame(fft_neighbors_df: pd.DataFrame) -> pd.DataFrame:
    """Return one user-facing FFT-grid row per exported FFT input."""

    source_columns = tuple(FFT_METADATA_COLUMN_MAP)
    if fft_neighbors_df.empty or any(
        column not in fft_neighbors_df.columns for column in source_columns
    ):
        return pd.DataFrame(columns=tuple(FFT_METADATA_COLUMN_MAP.values()))
    return (
        fft_neighbors_df.loc[:, source_columns]
        .drop_duplicates(ignore_index=True)
        .rename(columns=FFT_METADATA_COLUMN_MAP)
    )


def write_results_workbook(
    full_excel_path: str,
    dataframes_to_save: Dict[str, pd.DataFrame],
    fft_neighbors_df: Optional[pd.DataFrame] = None,
    spectral_eligibility_df: Optional[pd.DataFrame] = None,
    spectral_metric_qc_df: Optional[pd.DataFrame] = None,
    timing_sink: list[dict[str, object]] | None = None,
) -> dict[str, object]:
    """Write results workbook with consistent formatting and optional debug sheet."""
    workbook_started = perf_counter()
    destination = Path(full_excel_path)
    prior_artifact = _artifact_identity(destination)
    expected_headers: Dict[str, List[str]] = {
        str(sheet_name): [str(column) for column in frame.columns]
        for sheet_name, frame in dataframes_to_save.items()
    }
    fft_metadata_df: pd.DataFrame | None = None
    if fft_neighbors_df is not None and not fft_neighbors_df.empty:
        expected_headers["FFT and neighbors"] = [
            str(column) for column in fft_neighbors_df.columns
        ]
        fft_metadata_df = build_fft_metadata_frame(fft_neighbors_df)
        if not fft_metadata_df.empty:
            expected_headers[FFT_METADATA_SHEET_NAME] = [
                str(column) for column in fft_metadata_df.columns
            ]
    for sheet_name, audit_frame in (
        (SPECTRAL_ELIGIBILITY_SHEET_NAME, spectral_eligibility_df),
        (SPECTRAL_METRIC_QC_SHEET_NAME, spectral_metric_qc_df),
    ):
        if audit_frame is not None and not audit_frame.empty:
            expected_headers[sheet_name] = [
                str(column) for column in audit_frame.columns
            ]
    schema_validation: dict[str, object]
    try:
        with _workbook_write_target(destination) as workbook_path:
            with pd.ExcelWriter(workbook_path, engine="xlsxwriter") as writer:
                workbook = writer.book
                center_fmt = workbook.add_format(
                    {"align": "center", "valign": "vcenter"}
                )

                for sheet_name, df_to_write in dataframes_to_save.items():
                    sheet_started = perf_counter()
                    write_started = perf_counter()
                    worksheet = _write_dataframe_to_excel(
                        writer,
                        sheet_name=sheet_name,
                        frame=df_to_write,
                    )
                    _log_excel_timing(
                        "sheet_to_excel",
                        write_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(df_to_write),
                        cols=len(df_to_write.columns),
                        timing_sink=timing_sink,
                    )
                    worksheet.freeze_panes(1, 0)
                    widths_started = perf_counter()
                    _apply_column_widths(
                        worksheet,
                        df_to_write,
                        center_fmt,
                    )
                    _log_excel_timing(
                        "sheet_column_widths",
                        widths_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(df_to_write),
                        cols=len(df_to_write.columns),
                        timing_sink=timing_sink,
                    )
                    _log_excel_timing(
                        "sheet_total",
                        sheet_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(df_to_write),
                        cols=len(df_to_write.columns),
                        timing_sink=timing_sink,
                    )

                if fft_neighbors_df is not None and not fft_neighbors_df.empty:
                    sheet_name = "FFT and neighbors"
                    sheet_started = perf_counter()
                    write_started = perf_counter()
                    fft_neighbors_df.to_excel(
                        writer,
                        sheet_name=sheet_name,
                        index=False,
                    )
                    _log_excel_timing(
                        "sheet_to_excel",
                        write_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(fft_neighbors_df),
                        cols=len(fft_neighbors_df.columns),
                        timing_sink=timing_sink,
                    )
                    worksheet = writer.sheets[sheet_name]
                    worksheet.freeze_panes(1, 0)
                    widths_started = perf_counter()
                    _apply_column_widths(
                        worksheet,
                        fft_neighbors_df,
                        center_fmt,
                    )
                    _log_excel_timing(
                        "sheet_column_widths",
                        widths_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(fft_neighbors_df),
                        cols=len(fft_neighbors_df.columns),
                        timing_sink=timing_sink,
                    )
                    _log_excel_timing(
                        "sheet_total",
                        sheet_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(fft_neighbors_df),
                        cols=len(fft_neighbors_df.columns),
                        timing_sink=timing_sink,
                    )

                    if fft_metadata_df is not None and not fft_metadata_df.empty:
                        sheet_name = FFT_METADATA_SHEET_NAME
                        sheet_started = perf_counter()
                        write_started = perf_counter()
                        fft_metadata_df.to_excel(
                            writer,
                            sheet_name=sheet_name,
                            index=False,
                        )
                        _log_excel_timing(
                            "sheet_to_excel",
                            write_started,
                            path=full_excel_path,
                            sheet_name=sheet_name,
                            rows=len(fft_metadata_df),
                            cols=len(fft_metadata_df.columns),
                            timing_sink=timing_sink,
                        )
                        worksheet = writer.sheets[sheet_name]
                        worksheet.freeze_panes(1, 0)
                        widths_started = perf_counter()
                        _apply_column_widths(
                            worksheet,
                            fft_metadata_df,
                            center_fmt,
                        )
                        _log_excel_timing(
                            "sheet_column_widths",
                            widths_started,
                            path=full_excel_path,
                            sheet_name=sheet_name,
                            rows=len(fft_metadata_df),
                            cols=len(fft_metadata_df.columns),
                            timing_sink=timing_sink,
                        )
                        _log_excel_timing(
                            "sheet_total",
                            sheet_started,
                            path=full_excel_path,
                            sheet_name=sheet_name,
                            rows=len(fft_metadata_df),
                            cols=len(fft_metadata_df.columns),
                            timing_sink=timing_sink,
                        )

                for sheet_name, audit_frame in (
                    (SPECTRAL_ELIGIBILITY_SHEET_NAME, spectral_eligibility_df),
                    (SPECTRAL_METRIC_QC_SHEET_NAME, spectral_metric_qc_df),
                ):
                    if audit_frame is None or audit_frame.empty:
                        continue
                    sheet_started = perf_counter()
                    write_started = perf_counter()
                    audit_frame.to_excel(
                        writer,
                        sheet_name=sheet_name,
                        index=False,
                    )
                    _log_excel_timing(
                        "sheet_to_excel",
                        write_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(audit_frame),
                        cols=len(audit_frame.columns),
                        timing_sink=timing_sink,
                    )
                    worksheet = writer.sheets[sheet_name]
                    worksheet.freeze_panes(1, 0)
                    widths_started = perf_counter()
                    _apply_column_widths(
                        worksheet,
                        audit_frame,
                        center_fmt,
                    )
                    _log_excel_timing(
                        "sheet_column_widths",
                        widths_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(audit_frame),
                        cols=len(audit_frame.columns),
                        timing_sink=timing_sink,
                    )
                    _log_excel_timing(
                        "sheet_total",
                        sheet_started,
                        path=full_excel_path,
                        sheet_name=sheet_name,
                        rows=len(audit_frame),
                        cols=len(audit_frame.columns),
                        timing_sink=timing_sink,
                    )
            schema_validation = _validate_workbook_schema(
                Path(workbook_path),
                expected_headers,
            )
    finally:
        _log_excel_timing(
            "workbook_write_total",
            workbook_started,
            path=full_excel_path,
            timing_sink=timing_sink,
        )
    artifact = _artifact_identity(destination)
    if artifact is None:
        raise RuntimeError(
            f"Workbook publication completed without a readable artifact: {destination}"
        )
    return {
        "version": WORKBOOK_WRITE_RECEIPT_VERSION,
        "status": "written",
        "path": str(destination.resolve()),
        "prior_artifact": prior_artifact,
        "artifact": artifact,
        "schema_validation": schema_validation,
    }
