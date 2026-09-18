"""Atomic project-local exports for Free Harmonic Clustering Analysis."""

from __future__ import annotations

import csv
from dataclasses import asdict
from datetime import UTC, datetime
from enum import Enum
import hashlib
from importlib.metadata import PackageNotFoundError, version as package_version
import json
import os
import platform
from pathlib import Path, PurePosixPath, PureWindowsPath
import re
import shutil
from typing import TYPE_CHECKING, Callable, Iterable, Mapping, Sequence
from uuid import uuid4

import numpy as np
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import FPVS_TOOLBOX_VERSION

from .models import (
    ClusterPermutationResult,
    ClusterRecord,
    CohortWorkbook,
    ExportArtifact,
    ExportReceipt,
    FreeHarmonicCancelledError,
    FreeHarmonicInputError,
    PreparedContrast,
    PreparedRepeatedSessionBatch,
    RepeatedSessionContrastFamily,
    SENSOR_ADJACENCY_VERSION,
)
from .visualization import (
    ClusterMapData,
    build_cluster_map_data,
    build_repeated_cluster_map_data,
)
from .reporting import EXPLORATORY_CRITERION

if TYPE_CHECKING:
    from .api import RepeatedSessionBatchResult
    from .reporting import RepeatedSessionReport


TOOL_TITLE = "Free Harmonic Clustering Analysis"
EXPORT_SCHEMA_VERSION = 2
DEFAULT_RESULTS_SUBFOLDER = Path(
    "3 - Statistical Analysis Results",
    TOOL_TITLE,
)
MANIFEST_FILENAME = "manifest.json"
HUMAN_WORKBOOK_FILENAME = "Free_Harmonic_Clustering_Results.xlsx"
REPEATED_SESSION_WORKBOOK_FILENAME = "Free_Harmonic_Clustering_Repeated_Session_Batch.xlsx"
REPEATED_SESSION_EXPORT_SCHEMA_VERSION = 1
_EXPLORATORY_DESCRIPTION = (
    EXPLORATORY_CRITERION + " These did not survive Holm within their prespecified "
    "family. These are not confirmed findings or "
    "pointwise sensor/harmonic effects."
)
_NO_EXPLORATORY_FINDINGS = "No runs met the exploratory reporting criteria."
HUMAN_WORKBOOK_SHEETS: tuple[str, ...] = (
    "Run Summary",
    "Significant Clusters",
    "All Clusters",
    "Cluster Membership",
    "Harmonic Selection",
    "Participants and Exclusions",
    "Methods and Provenance",
    "Node Statistics",
    "Null Distribution",
)
_RUN_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")


def _validate_run_id(value: object) -> str:
    run_id = str(value or "").strip()
    if run_id in {".", ".."} or _RUN_ID.fullmatch(run_id) is None:
        raise ValueError(
            "run_id must start with an ASCII letter or digit and contain only letters, digits, '.', '_', or '-'."
        )
    return run_id


def _default_run_id() -> str:
    timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    return f"{timestamp}-{uuid4().hex[:8]}"


def _relative_to(path: Path, root: Path, *, label: str) -> Path:
    try:
        return path.relative_to(root)
    except ValueError as exc:
        raise FreeHarmonicInputError(f"{label} must remain beneath the managed project root: {root}") from exc


def _resolve_project_root(prepared: PreparedContrast) -> Path:
    root = Path(prepared.project_root).expanduser().resolve(strict=True)
    if not root.is_dir():
        raise FreeHarmonicInputError(f"Managed project root is not a directory: {root}")
    request_root = Path(prepared.request.project_root).expanduser().resolve(strict=True)
    if request_root != root:
        raise FreeHarmonicInputError("Prepared contrast and request refer to different managed project roots.")
    return root


def resolve_run_destination(
    prepared: PreparedContrast,
    *,
    run_id: str | None = None,
    destination: str | Path | None = None,
) -> tuple[str, Path]:
    """Resolve one never-overwritten run directory beneath the managed project."""

    root = _resolve_project_root(prepared)
    if destination is None:
        resolved_run_id = _validate_run_id(run_id or _default_run_id())
        target = root / DEFAULT_RESULTS_SUBFOLDER / resolved_run_id
    else:
        raw_target = Path(destination).expanduser()
        if not raw_target.is_absolute():
            raw_target = root / raw_target
        target = raw_target.resolve(strict=False)
        resolved_run_id = _validate_run_id(run_id or target.name)
        if target.name != resolved_run_id:
            raise ValueError("destination directory name must equal run_id.")

    target = target.resolve(strict=False)
    relative = _relative_to(target, root, label="Run destination")
    if not relative.parts:
        raise FreeHarmonicInputError("Run destination cannot be the project root itself.")
    return resolved_run_id, target


def _validate_source_workbook(record: CohortWorkbook, project_root: Path) -> str:
    source = Path(record.source_path).expanduser().resolve(strict=False)
    actual = _relative_to(source, project_root, label="Source workbook").as_posix()
    provided_text = str(record.project_relative_path).strip().replace("\\", "/")
    provided_posix = PurePosixPath(provided_text)
    provided_windows = PureWindowsPath(provided_text)
    if (
        not provided_text
        or provided_posix.is_absolute()
        or provided_windows.is_absolute()
        or bool(provided_windows.drive)
        or ".." in provided_posix.parts
    ):
        raise FreeHarmonicInputError("Source workbook provenance must use a project-relative path.")
    normalized = provided_posix.as_posix()
    if normalized.casefold() != actual.casefold():
        raise FreeHarmonicInputError(
            "Source workbook path and project-relative provenance do not match: "
            f"{record.participant_id} / {record.condition}."
        )
    return actual


def _validate_result(prepared: PreparedContrast, result: ClusterPermutationResult) -> None:
    if result.design is not prepared.request.design:
        raise ValueError("Prepared contrast and cluster result designs do not match.")
    if result.sensor_adjacency_version != prepared.method.sensor_adjacency_version:
        raise ValueError("Prepared method and cluster result spatial adjacency versions do not match.")
    expected_shape = (len(prepared.sensor_names), len(prepared.harmonics_hz))
    if result.observed_t.shape != expected_shape:
        raise ValueError(f"Cluster result shape must be {expected_shape}; got {result.observed_t.shape}.")
    labels = np.asarray(result.cluster_labels, dtype=np.int64)
    records = {cluster.cluster_id: cluster for cluster in result.clusters}
    if len(records) != len(result.clusters):
        raise ValueError("Cluster IDs must be unique.")
    observed_ids = {int(value) for value in np.unique(labels) if int(value) != 0}
    if observed_ids != set(records):
        raise ValueError("Cluster labels and cluster records contain different IDs.")
    flat_labels = labels.reshape(-1)
    occupied: set[int] = set()
    sensor_count, harmonic_count = expected_shape
    for cluster in result.clusters:
        _validate_cluster_record(
            cluster,
            flat_labels=flat_labels,
            sensor_count=sensor_count,
            harmonic_count=harmonic_count,
            occupied=occupied,
        )


def _validate_cluster_record(
    cluster: ClusterRecord,
    *,
    flat_labels: np.ndarray,
    sensor_count: int,
    harmonic_count: int,
    occupied: set[int],
) -> None:
    nodes = tuple(cluster.node_indices)
    sensors = tuple(cluster.sensor_indices)
    harmonics = tuple(cluster.harmonic_indices)
    if not (len(nodes) == len(sensors) == len(harmonics)):
        raise ValueError("Cluster node, sensor, and harmonic indices must align.")
    if len(set(nodes)) != len(nodes) or occupied.intersection(nodes):
        raise ValueError("Cluster node membership must be unique across clusters.")
    node_count = sensor_count * harmonic_count
    for node, sensor, harmonic in zip(nodes, sensors, harmonics, strict=True):
        if node >= node_count or sensor >= sensor_count or harmonic >= harmonic_count:
            raise ValueError("Cluster membership index exceeds the result tensor.")
        expected_node = sensor * harmonic_count + harmonic
        if node != expected_node or int(flat_labels[node]) != cluster.cluster_id:
            raise ValueError("Cluster membership does not match sensor-major labels.")
    occupied.update(nodes)


def _write_csv(path: Path, fieldnames: Sequence[str], rows: Iterable[Mapping[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)
        stream.flush()
        os.fsync(stream.fileno())


def _write_npz(path: Path, prepared: PreparedContrast, result: ClusterPermutationResult) -> None:
    plan = prepared.frequency_plan
    target_bin_frequencies = plan.selected_frequencies_hz[plan.target_selected_indices]
    target_frequency_errors = target_bin_frequencies - plan.candidate_harmonics_hz
    with path.open("wb") as stream:
        np.savez_compressed(
            stream,
            observed_t=np.asarray(result.observed_t),
            cluster_labels=np.asarray(result.cluster_labels),
            null_positive_max_mass=np.asarray(result.null_positive_max_mass),
            null_negative_min_mass=np.asarray(result.null_negative_min_mass),
            normalized_values_a=np.asarray(prepared.values_a),
            normalized_values_b=np.asarray(prepared.values_b),
            selected_snr_a=np.asarray(prepared.snr_a),
            selected_snr_b=np.asarray(prepared.snr_b),
            sensor_names=np.asarray(prepared.sensor_names, dtype=np.str_),
            harmonic_orders=np.asarray(prepared.harmonic_orders),
            harmonics_hz=np.asarray(prepared.harmonics_hz),
            participant_ids_a=np.asarray(prepared.participant_ids_a, dtype=np.str_),
            participant_ids_b=np.asarray(prepared.participant_ids_b, dtype=np.str_),
            full_frequency_columns=np.asarray(
                plan.full_frequency_columns,
                dtype=np.str_,
            ),
            full_frequencies_hz=np.asarray(plan.full_frequencies_hz),
            selected_frequency_columns=np.asarray(
                plan.selected_frequency_columns,
                dtype=np.str_,
            ),
            selected_frequencies_hz=np.asarray(plan.selected_frequencies_hz),
            candidate_orders=np.asarray(plan.candidate_orders),
            candidate_harmonics_hz=np.asarray(plan.candidate_harmonics_hz),
            excluded_base_orders=np.asarray(plan.excluded_base_orders),
            excluded_base_harmonics_hz=np.asarray(plan.excluded_base_harmonics_hz),
            harmonic_selection_mode=np.asarray(
                prepared.selection.selection_mode.value,
                dtype=np.str_,
            ),
            fixed_highest_harmonic_order=np.asarray(
                -1
                if prepared.selection.fixed_highest_harmonic_order is None
                else prepared.selection.fixed_highest_harmonic_order,
                dtype=np.int64,
            ),
            highest_detected_harmonic_order=np.asarray(
                -1 if prepared.selection.highest_detected_order is None else prepared.selection.highest_detected_order,
                dtype=np.int64,
            ),
            target_selected_indices=np.asarray(plan.target_selected_indices),
            noise_selected_indices=np.asarray(plan.noise_selected_indices),
            target_bin_frequencies_hz=np.asarray(target_bin_frequencies),
            target_frequency_errors_hz=np.asarray(target_frequency_errors),
        )
        stream.flush()
        os.fsync(stream.fileno())


def _write_manifest(path: Path, payload: Mapping[str, object]) -> None:
    encoded = json.dumps(
        payload,
        indent=2,
        sort_keys=True,
        ensure_ascii=False,
        allow_nan=False,
    ).encode("utf-8")
    with path.open("wb") as stream:
        stream.write(encoded)
        stream.write(b"\n")
        stream.flush()
        os.fsync(stream.fileno())


def _sha256_file(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def _cluster_summary_rows(
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for cluster in sorted(result.clusters, key=lambda item: item.cluster_id):
        effect_audit = _cluster_effect_audit(prepared, cluster)
        sensors = tuple(dict.fromkeys(prepared.sensor_names[index] for index in cluster.sensor_indices))
        orders = tuple(dict.fromkeys(int(prepared.harmonic_orders[index]) for index in cluster.harmonic_indices))
        frequencies = tuple(dict.fromkeys(float(prepared.harmonics_hz[index]) for index in cluster.harmonic_indices))
        rows.append(
            {
                "cluster_id": cluster.cluster_id,
                "sign": cluster.sign,
                "mass": cluster.mass,
                "p_value": cluster.p_value,
                "conservative_p_value": cluster.conservative_p_value,
                "adjusted_two_sided_p_value": cluster.adjusted_two_sided_p_value,
                "tie_count": cluster.tie_count,
                "p_ci_low": cluster.p_ci_low,
                "p_ci_high": cluster.p_ci_high,
                "confidence_interval_straddles_alpha": (cluster.confidence_interval_straddles_alpha),
                "significant_cluster_level": cluster.significant,
                "node_count": len(cluster.node_indices),
                "n_a": len(prepared.participant_ids_a),
                "n_b": len(prepared.participant_ids_b),
                "arm_a_normalized_cluster_node_mean": effect_audit["arm_a_mean"],
                "arm_b_normalized_cluster_node_mean": effect_audit["arm_b_mean"],
                "arm_a_minus_b_raw_difference": effect_audit["difference"],
                "arm_a_minus_b_normalized_difference": effect_audit["difference"],
                "effect_denominator_sd": effect_audit["denominator_sd"],
                "effect_denominator_kind": effect_audit["denominator_kind"],
                "effect_direction": effect_audit["direction"],
                "effect_size_direction": effect_audit["direction"],
                "descriptive_effect_label": ("post-selection/shape-dependent normalized cluster-node mean"),
                "effect_value_scale": ("participant-arm L2-normalized sensor x retained-harmonic tensor"),
                "sensor_count": len(sensors),
                "harmonic_count": len(orders),
                "sensors": "|".join(sensors),
                "harmonic_orders": "|".join(str(value) for value in orders),
                "harmonics_hz": "|".join(f"{value:g}" for value in frequencies),
                "effect_size": "" if cluster.effect_size is None else cluster.effect_size,
                "effect_size_kind": cluster.effect_size_kind or "",
            }
        )
    return rows


def _cluster_effect_audit(
    prepared: PreparedContrast,
    cluster: ClusterRecord,
) -> dict[str, object]:
    nodes = np.asarray(cluster.node_indices, dtype=np.int64)
    arm_a = np.asarray(prepared.values_a).reshape(len(prepared.values_a), -1)[:, nodes]
    arm_b = np.asarray(prepared.values_b).reshape(len(prepared.values_b), -1)[:, nodes]
    participant_means_a = np.mean(arm_a, axis=1)
    participant_means_b = np.mean(arm_b, axis=1)
    mean_a = float(np.mean(participant_means_a))
    mean_b = float(np.mean(participant_means_b))
    difference = mean_a - mean_b
    if prepared.request.design.value == "paired_conditions":
        denominator_kind = "paired_difference_sd"
        denominator = float(np.std(participant_means_a - participant_means_b, ddof=1))
    else:
        denominator_kind = "pooled_within_arm_sd"
        n_a = participant_means_a.size
        n_b = participant_means_b.size
        variance_a = float(np.var(participant_means_a, ddof=1))
        variance_b = float(np.var(participant_means_b, ddof=1))
        denominator = float(np.sqrt(((n_a - 1) * variance_a + (n_b - 1) * variance_b) / (n_a + n_b - 2)))
    if not np.isfinite(denominator):
        raise ValueError("Cluster descriptive effect denominator must be finite.")
    if difference > 0.0:
        direction = "arm_a_greater"
    elif difference < 0.0:
        direction = "arm_b_greater"
    else:
        direction = "no_direction"
    return {
        "arm_a_mean": mean_a,
        "arm_b_mean": mean_b,
        "difference": difference,
        "denominator_sd": denominator,
        "denominator_kind": denominator_kind,
        "direction": direction,
    }


def _cluster_membership_rows(
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for cluster in sorted(result.clusters, key=lambda item: item.cluster_id):
        for node, sensor, harmonic in zip(
            cluster.node_indices,
            cluster.sensor_indices,
            cluster.harmonic_indices,
            strict=True,
        ):
            rows.append(
                {
                    "cluster_id": cluster.cluster_id,
                    "sign": cluster.sign,
                    "node_index": node,
                    "sensor_index": sensor,
                    "sensor": prepared.sensor_names[sensor],
                    "harmonic_index": harmonic,
                    "harmonic_order": int(prepared.harmonic_orders[harmonic]),
                    "harmonic_hz": float(prepared.harmonics_hz[harmonic]),
                    "observed_t": float(result.observed_t[sensor, harmonic]),
                    "cluster_mass": cluster.mass,
                    "cluster_p_value": cluster.p_value,
                    "cluster_significant": cluster.significant,
                }
            )
    return rows


def _node_statistic_rows(
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
) -> list[dict[str, object]]:
    cluster_by_id = {cluster.cluster_id: cluster for cluster in result.clusters}
    rows: list[dict[str, object]] = []
    harmonic_count = len(prepared.harmonics_hz)
    for sensor_index, sensor in enumerate(prepared.sensor_names):
        for harmonic_index, harmonic_hz in enumerate(prepared.harmonics_hz):
            cluster_id = int(result.cluster_labels[sensor_index, harmonic_index])
            cluster = cluster_by_id.get(cluster_id)
            rows.append(
                {
                    "node_index": sensor_index * harmonic_count + harmonic_index,
                    "sensor_index": sensor_index,
                    "sensor": sensor,
                    "harmonic_index": harmonic_index,
                    "harmonic_order": int(prepared.harmonic_orders[harmonic_index]),
                    "harmonic_hz": float(harmonic_hz),
                    "observed_t": float(result.observed_t[sensor_index, harmonic_index]),
                    "cluster_id": cluster_id,
                    "cluster_sign": "" if cluster is None else cluster.sign,
                    "cluster_p_value": "" if cluster is None else cluster.p_value,
                    "cluster_significant": "" if cluster is None else cluster.significant,
                    "pointwise_significance_claimed": False,
                }
            )
    return rows


def _harmonic_selection_rows(prepared: PreparedContrast) -> list[dict[str, object]]:
    selection = prepared.selection
    selected = {int(index) for index in selection.selected_candidate_indices}
    selection_mode = selection.selection_mode.value
    selection_ceiling = int(selection.selected_orders[-1])
    automatic = selection_mode == "automatic"
    rows: list[dict[str, object]] = []
    for index, (order, harmonic_hz) in enumerate(
        zip(selection.candidate_orders, selection.candidate_harmonics_hz, strict=True)
    ):
        rows.append(
            {
                "selection_mode": selection_mode,
                "selection_ceiling_order": selection_ceiling,
                "z_threshold_used_for_selection": automatic,
                "harmonic_order": int(order),
                "harmonic_hz": float(harmonic_hz),
                "eligible_nonbase": True,
                "arm_a_z": float(selection.arm_a_z[index]),
                "arm_b_z": float(selection.arm_b_z[index]),
                "detected_arm_a": bool(selection.detected_arm_a[index]),
                "detected_arm_b": bool(selection.detected_arm_b[index]),
                "retained_fill_through": index in selected,
                "exclusion_reason": (
                    "" if index in selected else ("above_highest_detected" if automatic else "above_fixed_ceiling")
                ),
            }
        )
    for order, harmonic_hz in zip(
        selection.excluded_base_orders,
        selection.excluded_base_harmonics_hz,
        strict=True,
    ):
        rows.append(
            {
                "selection_mode": selection_mode,
                "selection_ceiling_order": selection_ceiling,
                "z_threshold_used_for_selection": automatic,
                "harmonic_order": int(order),
                "harmonic_hz": float(harmonic_hz),
                "eligible_nonbase": False,
                "arm_a_z": "",
                "arm_b_z": "",
                "detected_arm_a": False,
                "detected_arm_b": False,
                "retained_fill_through": False,
                "exclusion_reason": "base_rate_overlap",
            }
        )
    return sorted(rows, key=lambda row: int(row["harmonic_order"]))


def _participant_rows(prepared: PreparedContrast) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for arm, label, condition, participant_ids in (
        (
            "a",
            prepared.arm_a_label,
            prepared.request.condition_a,
            prepared.participant_ids_a,
        ),
        (
            "b",
            prepared.arm_b_label,
            prepared.request.condition_b or prepared.request.condition_a,
            prepared.participant_ids_b,
        ),
    ):
        for index, participant_id in enumerate(participant_ids):
            rows.append(
                {
                    "status": "Included",
                    "exclusion_reason": "",
                    "arm": arm,
                    "arm_label": label,
                    "condition": condition,
                    "participant_index": index,
                    "participant_id": participant_id,
                }
            )
    for exclusion in prepared.provenance.participant_condition_exclusions:
        rows.append(
            {
                "status": "Excluded",
                "exclusion_reason": exclusion.reason,
                "arm": "",
                "arm_label": "",
                "condition": exclusion.condition,
                "participant_index": "",
                "participant_id": exclusion.participant_id,
            }
        )
    return rows


def _participant_and_exclusion_rows(
    prepared: PreparedContrast,
) -> list[dict[str, object]]:
    """Return one human-readable cohort audit including excluded IDs."""

    rows = [dict(row) for row in _participant_rows(prepared)]
    exclusion_sets = (
        (
            "Processing ledger not completed",
            prepared.provenance.ledger_excluded_participants,
        ),
        (
            "Manual project exclusion",
            prepared.provenance.manual_excluded_participants,
        ),
        (
            "Frequency-domain QC exclusion",
            prepared.provenance.frequency_qc_excluded_participants,
        ),
        (
            "Incomplete paired-condition record",
            prepared.provenance.incomplete_pair_participants,
        ),
    )
    for reason, participant_ids in exclusion_sets:
        for participant_id in participant_ids:
            rows.append(
                {
                    "status": "Excluded",
                    "exclusion_reason": reason,
                    "arm": "",
                    "arm_label": "",
                    "condition": "",
                    "participant_index": "",
                    "participant_id": participant_id,
                }
            )
    return rows


def _source_workbook_rows(
    prepared: PreparedContrast,
    project_root: Path,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for record in prepared.source_workbooks:
        project_relative_path = _validate_source_workbook(record, project_root)
        source_stat = Path(record.source_path).stat()
        rows.append(
            {
                "arm": record.arm,
                "arm_label": record.arm_label,
                "participant_id": record.participant_id,
                "condition": record.condition,
                "group_id": record.group_id or "",
                "group_label": record.group_label or "",
                "project_relative_path": project_relative_path,
                "source_size_bytes": int(source_stat.st_size),
                "source_mtime_ns": int(source_stat.st_mtime_ns),
                "header_read_seconds": record.header_read_seconds,
                "amplitude_read_seconds": record.amplitude_read_seconds,
            }
        )
    return rows


def _null_extrema_rows(result: ClusterPermutationResult) -> list[dict[str, object]]:
    return [
        {
            "permutation_index": index,
            "positive_max_mass": float(positive),
            "negative_min_mass": float(negative),
        }
        for index, (positive, negative) in enumerate(
            zip(
                result.null_positive_max_mass,
                result.null_negative_min_mass,
                strict=True,
            )
        )
    ]


def _finite_array_summary(values: np.ndarray) -> dict[str, object]:
    array = np.asarray(values, dtype=np.float64)
    finite = array[np.isfinite(array)]
    return {
        "count": int(array.size),
        "finite_count": int(finite.size),
        "nonfinite_count": int(array.size - finite.size),
        "minimum": None if not finite.size else float(np.min(finite)),
        "maximum": None if not finite.size else float(np.max(finite)),
    }


def _dependency_version(distribution: str) -> str | None:
    try:
        return package_version(distribution)
    except PackageNotFoundError:
        return None


_TITLE_FILL = PatternFill("solid", fgColor="17365D")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_HEADER_FILL = PatternFill("solid", fgColor="2F75B5")
_SIGNIFICANT_FILL = PatternFill("solid", fgColor="E2F0D9")
_NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
_WHITE_FONT = Font(name="Arial", color="FFFFFF", bold=True)
_TITLE_FONT = Font(name="Arial", color="FFFFFF", bold=True, size=16)
_BODY_FONT = Font(name="Arial", size=10)
_HEADER_FONT = Font(name="Arial", color="FFFFFF", bold=True, size=10)
_SECTION_FONT = Font(name="Arial", color="17365D", bold=True, size=11)
_THIN_GRAY = Side(style="thin", color="D9E2F3")
_BOTTOM_BORDER = Border(bottom=_THIN_GRAY)

_HEADER_LABELS = {
    "cluster_id": "Cluster ID",
    "sign": "Sign",
    "mass": "Cluster Mass",
    "p_value": "Raw Tail p",
    "conservative_p_value": "Conservative p",
    "adjusted_two_sided_p_value": "Doubled Two-Sided p",
    "p_ci_low": "Monte Carlo p CI Low",
    "p_ci_high": "Monte Carlo p CI High",
    "confidence_interval_straddles_alpha": "p CI Straddles Alpha",
    "significant_cluster_level": "Significant",
    "cluster_significant": "Cluster Significant",
    "pointwise_significance_claimed": "Pointwise Claim",
    "harmonic_hz": "Harmonic (Hz)",
    "harmonics_hz": "Harmonics (Hz)",
    "arm_a_z": "Arm A z",
    "arm_b_z": "Arm B z",
    "observed_t": "Observed t",
    "n_a": "Arm A n",
    "n_b": "Arm B n",
    "arm_a_normalized_cluster_node_mean": "Arm A Normalized Mean",
    "arm_b_normalized_cluster_node_mean": "Arm B Normalized Mean",
    "arm_a_minus_b_raw_difference": "Arm A - B Difference",
    "effect_denominator_sd": "Effect Denominator SD",
    "effect_size": "Effect Size",
    "effect_size_kind": "Effect Size Type",
    "selection_mode": "Selection Mode",
    "selection_ceiling_order": "Selection Ceiling Order",
    "z_threshold_used_for_selection": "z Threshold Selected Domain",
    "retained_fill_through": "Retained",
    "eligible_nonbase": "Eligible Non-Base",
    "exclusion_reason": "Exclusion Reason",
    "participant_id": "Participant ID",
    "participant_index": "Participant Index",
    "arm_label": "Arm Label",
    "permutation_index": "Permutation Index",
    "positive_max_mass": "Positive Maximum Mass",
    "negative_min_mass": "Negative Minimum Mass",
}


def _header_label(field_name: str) -> str:
    return _HEADER_LABELS.get(
        field_name,
        field_name.replace("_", " ").title(),
    )


def _excel_scalar(value: object) -> object:
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return value.as_posix()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if value is None:
        return ""
    if isinstance(value, (tuple, list, set)):
        return ", ".join(str(item) for item in value)
    return value


def _set_excel_value(cell: Cell, value: object) -> None:
    """Write data without allowing untrusted text to become a formula."""

    scalar = _excel_scalar(value)
    if isinstance(scalar, str):
        cell.value = scalar
        cell.data_type = "s"
        if scalar.startswith(("=", "+", "-", "@")):
            cell.number_format = "@"
    else:
        cell.value = scalar


def _number_format(field_name: str) -> str | None:
    if field_name in {
        "cluster_id",
        "tie_count",
        "node_count",
        "n_a",
        "n_b",
        "sensor_count",
        "harmonic_count",
        "node_index",
        "sensor_index",
        "harmonic_index",
        "harmonic_order",
        "participant_index",
        "permutation_index",
        "selection_ceiling_order",
        "participant_count",
        "observed_cluster_count",
        "raw_significant_cluster_count",
        "degrees_of_freedom",
        "derived_seed",
        "visit_index",
    }:
        return "0"
    if field_name in {
        "p_value",
        "conservative_p_value",
        "adjusted_two_sided_p_value",
        "p_ci_low",
        "p_ci_high",
        "cluster_p_value",
        "global_two_sided_cluster_p_value",
        "holm_within_family_p_value",
        "holm_all_batch_p_value",
        "run_global_two_sided_p_value",
        "run_holm_within_family_p_value",
        "run_holm_all_batch_p_value",
    }:
        return "0.0000"
    if "hz" in field_name:
        return "0.0000"
    if field_name.endswith("seconds"):
        return "0.000"
    if field_name in {
        "mass",
        "cluster_mass",
        "observed_t",
        "effect_size",
        "effect_denominator_sd",
        "arm_a_z",
        "arm_b_z",
        "z_score",
        "z_threshold",
        "arm_a_normalized_cluster_node_mean",
        "arm_b_normalized_cluster_node_mean",
        "arm_a_minus_b_raw_difference",
        "positive_max_mass",
        "negative_min_mass",
    }:
        return "0.0000"
    return None


def _set_sheet_title(
    sheet: Worksheet,
    *,
    title: str,
    description: str,
    column_count: int,
) -> None:
    last_column = get_column_letter(max(1, column_count))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet.merge_cells(f"A2:{last_column}2")
    _set_excel_value(sheet["A1"], title)
    _set_excel_value(sheet["A2"], description)
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].font = _TITLE_FONT
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet["A2"].fill = _NOTE_FILL
    sheet["A2"].font = Font(name="Arial", color="404040", italic=True, size=10)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 32
    sheet.sheet_view.showGridLines = False


def _set_reasonable_widths(
    sheet: Worksheet,
    *,
    fields: Sequence[str],
    rows: Sequence[Mapping[str, object]],
) -> None:
    for column_index, field_name in enumerate(fields, start=1):
        longest = len(_header_label(field_name))
        for row in rows[:500]:
            value = str(_excel_scalar(row.get(field_name, "")))
            longest = max(longest, min(len(value), 80))
        if field_name in {
            "sensors",
            "harmonic_orders",
            "harmonics_hz",
            "exclusion_reason",
            "effect_direction",
            "effect_size_kind",
            "item",
            "value",
            "notes",
        }:
            maximum = 48
        else:
            maximum = 28
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest + 2, 11),
            maximum,
        )


def _write_table_sheet(
    sheet: Worksheet,
    *,
    title: str,
    description: str,
    fields: Sequence[str],
    rows: Sequence[Mapping[str, object]],
    empty_message: str = "No rows were produced.",
    significant_field: str | None = None,
) -> None:
    _set_sheet_title(
        sheet,
        title=title,
        description=description,
        column_count=len(fields),
    )
    header_row = 4
    for column_index, field_name in enumerate(fields, start=1):
        cell = sheet.cell(row=header_row, column=column_index)
        _set_excel_value(cell, _header_label(field_name))
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    sheet.row_dimensions[header_row].height = 32
    if rows:
        for row_index, row in enumerate(rows, start=header_row + 1):
            significant = significant_field is not None and bool(row.get(significant_field, False))
            for column_index, field_name in enumerate(fields, start=1):
                cell = sheet.cell(row=row_index, column=column_index)
                _set_excel_value(cell, row.get(field_name, ""))
                cell.font = _BODY_FONT
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=field_name
                    in {
                        "sensors",
                        "harmonic_orders",
                        "harmonics_hz",
                        "exclusion_reason",
                        "effect_direction",
                        "effect_size_kind",
                        "item",
                        "value",
                        "notes",
                    },
                )
                cell.border = _BOTTOM_BORDER
                number_format = _number_format(field_name)
                if number_format is not None and not isinstance(cell.value, str):
                    cell.number_format = number_format
                if significant:
                    cell.fill = _SIGNIFICANT_FILL
        last_row = header_row + len(rows)
    else:
        last_row = header_row
        last_column = get_column_letter(max(1, len(fields)))
        sheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=len(fields),
        )
        _set_excel_value(sheet.cell(row=header_row + 1, column=1), empty_message)
        sheet.cell(row=header_row + 1, column=1).font = Font(
            name="Arial",
            italic=True,
            color="666666",
        )
        sheet.cell(row=header_row + 1, column=1).fill = _NOTE_FILL
        sheet.cell(row=header_row + 1, column=1).alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[last_column].width = max(
            sheet.column_dimensions[last_column].width or 0,
            12,
        )
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(fields))}{last_row}"
    sheet.freeze_panes = f"A{header_row + 1}"
    _set_reasonable_widths(sheet, fields=fields, rows=rows)


def _cluster_sort_key(row: Mapping[str, object]) -> tuple[bool, float, int]:
    p_value = float(row.get("p_value", 1.0))
    cluster_id = int(row.get("cluster_id", 0))
    return (not bool(row.get("significant_cluster_level", False)), p_value, cluster_id)


def _methods_and_provenance_rows(
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []

    def add(category: str, item: str, value: object, notes: str = "") -> None:
        rows.append(
            {
                "category": category,
                "item": item,
                "value": value,
                "notes": notes,
            }
        )

    for field_name, value in asdict(prepared.method).items():
        add("Method specification", _header_label(field_name), value)
    add(
        "Normalization",
        "Participant normalization",
        "Global L2 over each participant-arm sensor x retained-harmonic matrix",
        "The estimand is the relative sensor/harmonic response distribution.",
    )
    add("Analysis", "Design", prepared.request.design.value)
    add("Analysis", "Degrees of freedom", result.degrees_of_freedom)
    add(
        "Analysis",
        "Cluster-forming t threshold",
        result.cluster_forming_threshold,
    )
    add("Analysis", "Permutations evaluated", result.permutations_evaluated)
    add("Analysis", "RNG algorithm", result.rng_algorithm)
    add("Analysis", "Permutation seed", result.seed)
    add(
        "Analysis",
        "Permutation assignment SHA-256",
        result.permutation_assignment_hash,
    )
    add(
        "Inference",
        "Primary cluster p-value",
        "Raw sign-specific Monte Carlo tail p",
        f"Judged against alpha {result.cluster_alpha_per_tail:g} per direction.",
    )
    add(
        "Inference",
        "Secondary p-value",
        "Doubled two-sided p",
        "Reported as a secondary descriptive conversion, not an extra family correction.",
    )
    add(
        "Inference",
        "Inference scope",
        "Cluster level only",
        "Individual sensor-harmonic nodes are not pointwise significant claims.",
    )
    add("Adjacency", "Spatial version", result.sensor_adjacency_version)
    add(
        "Adjacency",
        "Spatial fingerprint SHA-256",
        result.sensor_adjacency_fingerprint,
    )
    add("Adjacency", "Spatial edge count", len(result.sensor_adjacency_edges))
    add(
        "Adjacency",
        "Spatial derivation",
        (
            "Independent clean-room FieldTrip-style compressed BioSemi64 "
            "reconstruction; fixed 169-edge MNE subset plus 28 audited additions"
            if result.sensor_adjacency_version == SENSOR_ADJACENCY_VERSION and len(result.sensor_adjacency_edges) == 197
            else "Explicit versioned edge table recorded in manifest.json"
        ),
        "This is not represented as the authors' unpublished adjacency matrix.",
    )
    add("Adjacency", "Harmonic adjacency", "Complete within sensor")
    add("Adjacency", "Flattening order", "Sensor-major")
    provenance = prepared.provenance
    add("Input provenance", "Source sheet", provenance.source_sheet)
    add("Input provenance", "Source workbook count", provenance.workbook_count)
    add(
        "Input provenance",
        "Participant-condition exclusion count",
        len(provenance.participant_condition_exclusions),
    )
    add("Input provenance", "Grid fingerprint", provenance.grid_fingerprint)
    if provenance.full_fft_provenance_method_version:
        add(
            "Input provenance",
            "Neutral FullFFT provenance method",
            provenance.full_fft_provenance_method_version,
        )
        add(
            "Input provenance",
            "Neutral FullFFT source fingerprint",
            provenance.full_fft_source_fingerprint,
        )
        add(
            "Input provenance",
            "Neutral FullFFT cohort fingerprint",
            provenance.full_fft_cohort_fingerprint,
        )
        add(
            "Input provenance",
            "Neutral FullFFT QC fingerprint",
            provenance.full_fft_frequency_qc_fingerprint,
        )
        add(
            "Input provenance",
            "Neutral FullFFT processing/export fingerprint",
            provenance.full_fft_processing_export_fingerprint,
        )
    add(
        "Input provenance",
        "Selected-column fingerprint",
        provenance.selected_columns_fingerprint,
    )
    add(
        "Input provenance",
        "Frequency resolution (Hz)",
        provenance.frequency_resolution_hz,
    )
    add("Timing", "Header reads (s)", provenance.header_read_seconds)
    add("Timing", "Amplitude reads (s)", provenance.amplitude_read_seconds)
    add(
        "Timing",
        "Numeric preparation (s)",
        provenance.numeric_preparation_seconds,
    )
    add("Timing", "Preparation total (s)", provenance.total_seconds)
    for warning in result.warnings:
        add("Warnings", "Analysis warning", warning)
    for diagnostic in provenance.dataset_diagnostics:
        add("Warnings", "Dataset diagnostic", diagnostic)
    add(
        "Reference",
        "Hermann et al. article",
        "https://doi.org/10.1111/psyp.70361",
    )
    add(
        "Reference",
        "Public full text",
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC13379596/",
    )
    add(
        "Reference",
        "Declared public code/data project",
        "https://github.com/users/oliver-hermann1/projects/1",
    )
    add(
        "Reference",
        "Implementation status",
        "Paper-faithful clean-room implementation; not author-validated",
    )
    return rows


def _write_run_summary_sheet(
    sheet: Worksheet,
    *,
    run_id: str,
    created_at_utc: str,
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
    cluster_rows: Sequence[Mapping[str, object]],
) -> None:
    _set_sheet_title(
        sheet,
        title=TOOL_TITLE,
        description=(
            "Human-readable run summary. Raw tail p-values are the primary Hermann-compatible cluster-level results."
        ),
        column_count=9,
    )
    sheet.freeze_panes = "A4"
    sheet.merge_cells("A4:I4")
    _set_excel_value(sheet["A4"], "Run Overview")
    sheet["A4"].fill = _SECTION_FILL
    sheet["A4"].font = _SECTION_FONT
    sheet["A4"].alignment = Alignment(vertical="center")
    contrast = f"{prepared.arm_a_label} - {prepared.arm_b_label}"
    overview = (
        ("Run ID", run_id),
        ("Created (UTC)", created_at_utc),
        ("Design", prepared.request.design.value),
        ("Contrast", contrast),
        ("Condition A", prepared.request.condition_a),
        ("Condition B", prepared.request.condition_b or ""),
        ("Arm A n", len(prepared.participant_ids_a)),
        ("Arm B n", len(prepared.participant_ids_b)),
        ("Selection Mode", prepared.selection.selection_mode.value),
        (
            "Retained Harmonics",
            ", ".join(
                f"H{int(order)} ({float(frequency):g} Hz)"
                for order, frequency in zip(
                    prepared.harmonic_orders,
                    prepared.harmonics_hz,
                    strict=True,
                )
            ),
        ),
        ("Permutations", result.permutations_evaluated),
        ("Significant Clusters", sum(cluster.significant for cluster in result.clusters)),
        ("Spatial Adjacency", result.sensor_adjacency_version),
    )
    for index, (label, value) in enumerate(overview, start=5):
        _set_excel_value(sheet.cell(index, 1), label)
        _set_excel_value(sheet.cell(index, 2), value)
        sheet.cell(index, 1).font = Font(name="Arial", bold=True, color="404040")
        sheet.cell(index, 2).font = _BODY_FONT
        sheet.cell(index, 1).border = _BOTTOM_BORDER
        sheet.cell(index, 2).border = _BOTTOM_BORDER
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 60

    significant = sorted(
        (row for row in cluster_rows if bool(row.get("significant_cluster_level", False))),
        key=_cluster_sort_key,
    )
    section_row = 5 + len(overview) + 1
    sheet.merge_cells(
        start_row=section_row,
        start_column=1,
        end_row=section_row,
        end_column=9,
    )
    _set_excel_value(sheet.cell(section_row, 1), "Significant Results")
    sheet.cell(section_row, 1).fill = _SECTION_FILL
    sheet.cell(section_row, 1).font = _SECTION_FONT
    summary_fields = (
        "cluster_id",
        "sign",
        "sensors",
        "harmonics_hz",
        "mass",
        "p_value",
        "adjusted_two_sided_p_value",
        "effect_size",
        "effect_size_kind",
    )
    header_row = section_row + 1
    for column_index, field_name in enumerate(summary_fields, start=1):
        cell = sheet.cell(header_row, column_index)
        _set_excel_value(cell, _header_label(field_name))
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    if significant:
        for row_index, row in enumerate(significant, start=header_row + 1):
            for column_index, field_name in enumerate(summary_fields, start=1):
                cell = sheet.cell(row_index, column_index)
                _set_excel_value(cell, row.get(field_name, ""))
                cell.font = _BODY_FONT
                cell.fill = _SIGNIFICANT_FILL
                cell.border = _BOTTOM_BORDER
                number_format = _number_format(field_name)
                if number_format and not isinstance(cell.value, str):
                    cell.number_format = number_format
    else:
        sheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=9,
        )
        _set_excel_value(
            sheet.cell(header_row + 1, 1),
            "No clusters met the Hermann-compatible per-direction threshold.",
        )
        sheet.cell(header_row + 1, 1).fill = _NOTE_FILL
        sheet.cell(header_row + 1, 1).font = Font(
            name="Arial",
            italic=True,
            color="666666",
        )
    for column_index, width in enumerate(
        (12, 12, 36, 20, 15, 14, 20, 14, 24),
        start=1,
    ):
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _write_human_workbook(
    path: Path,
    *,
    run_id: str,
    created_at_utc: str,
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
    cluster_rows: Sequence[Mapping[str, object]],
    membership_rows: Sequence[Mapping[str, object]],
    harmonic_rows: Sequence[Mapping[str, object]],
    participant_rows: Sequence[Mapping[str, object]],
    node_rows: Sequence[Mapping[str, object]],
    null_rows: Sequence[Mapping[str, object]],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = HUMAN_WORKBOOK_SHEETS[0]
    for sheet_name in HUMAN_WORKBOOK_SHEETS[1:]:
        workbook.create_sheet(sheet_name)
    _write_run_summary_sheet(
        summary,
        run_id=run_id,
        created_at_utc=created_at_utc,
        prepared=prepared,
        result=result,
        cluster_rows=cluster_rows,
    )

    cluster_fields = (
        "cluster_id",
        "sign",
        "effect_direction",
        "mass",
        "p_value",
        "conservative_p_value",
        "adjusted_two_sided_p_value",
        "p_ci_low",
        "p_ci_high",
        "confidence_interval_straddles_alpha",
        "significant_cluster_level",
        "node_count",
        "sensors",
        "harmonic_orders",
        "harmonics_hz",
        "effect_size",
        "effect_size_kind",
        "arm_a_normalized_cluster_node_mean",
        "arm_b_normalized_cluster_node_mean",
        "arm_a_minus_b_raw_difference",
        "n_a",
        "n_b",
    )
    sorted_clusters = sorted(cluster_rows, key=_cluster_sort_key)
    significant_clusters = [row for row in sorted_clusters if bool(row.get("significant_cluster_level", False))]
    _write_table_sheet(
        workbook["Significant Clusters"],
        title="Significant Clusters",
        description=(
            "Clusters meeting the raw sign-specific Monte Carlo p < .025 threshold, sorted by ascending raw tail p."
        ),
        fields=cluster_fields,
        rows=significant_clusters,
        empty_message=("No clusters met the Hermann-compatible per-direction threshold."),
        significant_field="significant_cluster_level",
    )
    _write_table_sheet(
        workbook["All Clusters"],
        title="All Observed Clusters",
        description=(
            "All sign-specific observed clusters; significant clusters are "
            "listed first, followed by ascending raw tail p."
        ),
        fields=cluster_fields,
        rows=sorted_clusters,
        significant_field="significant_cluster_level",
    )
    membership_fields = (
        "cluster_id",
        "sign",
        "node_index",
        "sensor_index",
        "sensor",
        "harmonic_index",
        "harmonic_order",
        "harmonic_hz",
        "observed_t",
        "cluster_mass",
        "cluster_p_value",
        "cluster_significant",
    )
    _write_table_sheet(
        workbook["Cluster Membership"],
        title="Cluster Membership",
        description=(
            "Sensor-harmonic nodes belonging to each cluster. Inference is at "
            "the cluster level, not the individual-node level."
        ),
        fields=membership_fields,
        rows=membership_rows,
        significant_field="cluster_significant",
    )
    harmonic_fields = (
        "selection_mode",
        "selection_ceiling_order",
        "z_threshold_used_for_selection",
        "harmonic_order",
        "harmonic_hz",
        "eligible_nonbase",
        "arm_a_z",
        "arm_b_z",
        "detected_arm_a",
        "detected_arm_b",
        "retained_fill_through",
        "exclusion_reason",
    )
    _write_table_sheet(
        workbook["Harmonic Selection"],
        title="Harmonic Selection Audit",
        description=(
            "Eligible oddball harmonics, base-rate overlaps, grand-spectrum z "
            "audit, and the retained complete fill-through domain."
        ),
        fields=harmonic_fields,
        rows=harmonic_rows,
    )
    participant_fields = (
        "status",
        "exclusion_reason",
        "arm",
        "arm_label",
        "condition",
        "participant_index",
        "participant_id",
    )
    _write_table_sheet(
        workbook["Participants and Exclusions"],
        title="Participants and Exclusions",
        description=("Included analysis-arm membership plus project/QC exclusions recorded during preparation."),
        fields=participant_fields,
        rows=participant_rows,
    )
    _write_table_sheet(
        workbook["Methods and Provenance"],
        title="Methods and Provenance",
        description=(
            "Versioned scientific settings, inference scope, adjacency "
            "identity, input fingerprints, timings, warnings, and references."
        ),
        fields=("category", "item", "value", "notes"),
        rows=_methods_and_provenance_rows(prepared, result),
    )
    node_fields = (
        "node_index",
        "sensor_index",
        "sensor",
        "harmonic_index",
        "harmonic_order",
        "harmonic_hz",
        "observed_t",
        "cluster_id",
        "cluster_sign",
        "cluster_p_value",
        "cluster_significant",
        "pointwise_significance_claimed",
    )
    _write_table_sheet(
        workbook["Node Statistics"],
        title="Node Statistics",
        description=(
            "Observed sensor-harmonic t statistics and cluster assignments. "
            "These rows do not make pointwise significance claims."
        ),
        fields=node_fields,
        rows=node_rows,
        significant_field="cluster_significant",
    )
    _write_table_sheet(
        workbook["Null Distribution"],
        title="Permutation Null Distribution",
        description=(
            "Separate positive maximum and negative minimum cluster-mass "
            "extrema for every evaluated whole-participant permutation."
        ),
        fields=(
            "permutation_index",
            "positive_max_mass",
            "negative_min_mass",
        ),
        rows=null_rows,
    )
    workbook.save(path)
    with path.open("rb+") as stream:
        os.fsync(stream.fileno())


def _frequency_plan_manifest(prepared: PreparedContrast) -> dict[str, object]:
    plan = prepared.frequency_plan
    target_frequencies = plan.selected_frequencies_hz[plan.target_selected_indices]
    target_errors = target_frequencies - plan.candidate_harmonics_hz
    return {
        "electrode_column": plan.electrode_column,
        "selected_index_basis": "zero-based selected_frequency_columns",
        "noise_window_rule": ("physical +/- half-width inclusive; target and immediate adjacent FFT bins excluded"),
        "frequency_resolution_hz": plan.frequency_resolution_hz,
        "grid_fingerprint": plan.grid_fingerprint,
        "selected_columns_fingerprint": plan.selected_columns_fingerprint,
        "full_frequency_column_count": len(plan.full_frequency_columns),
        "selected_frequency_column_count": len(plan.selected_frequency_columns),
        "selected_frequency_columns": list(plan.selected_frequency_columns),
        "selected_frequencies_hz": [float(value) for value in plan.selected_frequencies_hz],
        "candidate_orders": [int(value) for value in plan.candidate_orders],
        "candidate_harmonics_hz": [float(value) for value in plan.candidate_harmonics_hz],
        "excluded_base_orders": [int(value) for value in plan.excluded_base_orders],
        "excluded_base_harmonics_hz": [float(value) for value in plan.excluded_base_harmonics_hz],
        "target_selected_indices": [int(value) for value in plan.target_selected_indices],
        "noise_selected_indices": [[int(value) for value in row] for row in plan.noise_selected_indices],
        "target_bin_frequencies_hz": [float(value) for value in target_frequencies],
        "target_frequency_errors_hz": [float(value) for value in target_errors],
    }


def _manifest_payload(
    *,
    run_id: str,
    project_root: Path,
    destination: Path,
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
    source_rows: Sequence[Mapping[str, object]],
    artifacts: Sequence[Mapping[str, object]],
    created_at_utc: str,
) -> dict[str, object]:
    selection = prepared.selection
    provenance = prepared.provenance
    return {
        "schema_version": EXPORT_SCHEMA_VERSION,
        "status": "complete",
        "run_id": run_id,
        "created_at_utc": created_at_utc,
        "tool": {
            "title": TOOL_TITLE,
            "method_version": prepared.method.method_version,
            "export_schema_version": EXPORT_SCHEMA_VERSION,
            "validation_label": "paper-faithful-not-author-validated",
        },
        "software": {
            "fpvs_toolbox_version": FPVS_TOOLBOX_VERSION,
            "python_version": platform.python_version(),
            "numpy_version": np.__version__,
            "scipy_version": _dependency_version("scipy"),
        },
        "references": {
            "article_doi_url": "https://doi.org/10.1111/psyp.70361",
            "article_full_text_url": ("https://pmc.ncbi.nlm.nih.gov/articles/PMC13379596/"),
            "article_pubmed_url": "https://pubmed.ncbi.nlm.nih.gov/42469955/",
            "article_declared_code_data_url": ("https://github.com/users/oliver-hermann1/projects/1"),
            "clean_room_license_boundary": (
                "The referenced authors' code uses restrictive academic/"
                "non-commercial/share-alike terms; this MIT-licensed Toolbox "
                "implementation is independent and copies no author code."
            ),
        },
        "claims": {
            "familywise_error_control": "weak-FWER",
            "familywise_error_scope": "one-declared-contrast-complete-null",
            "cluster_level_inference_only": True,
            "pointwise_sensor_harmonic_significance": False,
            "cross_contrast_family_corrected": False,
            "adaptive_harmonic_selection_uses_analysis_data": (selection.selection_mode.value == "automatic"),
            "adaptive_selection_confirmatory_caveat": (
                (
                    "The paper-faithful harmonic domain is selected from the "
                    "observed analysis arms before permutation and remains fixed; "
                    "confirmatory work should prefer a preregistered or independent "
                    "domain."
                )
                if selection.selection_mode.value == "automatic"
                else (
                    "The retained fill-through domain used the declared fixed "
                    "highest eligible oddball harmonic and did not depend on "
                    "crossing the observed-arm z threshold."
                )
            ),
            "exchangeability_assumption": (
                "Within-participant whole-tensor swaps are exchangeable under the paired null."
                if prepared.request.design.value == "paired_conditions"
                else "Unrestricted whole-participant group labels are "
                "exchangeable under the independent-groups global null; "
                "unmodeled strata, matching, or family structure invalidate "
                "this permutation scheme."
            ),
            "method_fidelity": "paper-faithful",
            "author_validation": "not-author-validated",
        },
        "project": {
            "project_root": ".",
            "output_directory": _relative_to(
                destination,
                project_root,
                label="Output directory",
            ).as_posix(),
        },
        "contrast": {
            "design": prepared.request.design.value,
            "condition_a": prepared.request.condition_a,
            "condition_b": prepared.request.condition_b,
            "group_ids": list(prepared.request.group_ids),
            "arm_a_label": prepared.arm_a_label,
            "arm_b_label": prepared.arm_b_label,
            "participant_ids_a": list(prepared.participant_ids_a),
            "participant_ids_b": list(prepared.participant_ids_b),
            "participant_count_a": len(prepared.participant_ids_a),
            "participant_count_b": len(prepared.participant_ids_b),
        },
        "method": asdict(prepared.method),
        "normalization": {
            "method": ("participant-arm global L2 across sensor x retained harmonic tensor"),
            "estimand": "relative sensor/harmonic response distribution",
            "descriptive_cluster_effects": ("post-selection/shape-dependent normalized cluster-node means"),
        },
        "frequency_plan": _frequency_plan_manifest(prepared),
        "preparation": {
            "source_sheet": provenance.source_sheet,
            "grid_fingerprint": provenance.grid_fingerprint,
            "selected_columns_fingerprint": provenance.selected_columns_fingerprint,
            "frequency_resolution_hz": provenance.frequency_resolution_hz,
            "full_frequency_column_count": provenance.full_frequency_column_count,
            "selected_frequency_column_count": (provenance.selected_frequency_column_count),
            "workbook_count": provenance.workbook_count,
            "neutral_full_fft_provenance": {
                "method_version": (provenance.full_fft_provenance_method_version),
                "source_fingerprint": provenance.full_fft_source_fingerprint,
                "cohort_fingerprint": provenance.full_fft_cohort_fingerprint,
                "frequency_qc_fingerprint": (provenance.full_fft_frequency_qc_fingerprint),
                "processing_export_fingerprint": (provenance.full_fft_processing_export_fingerprint),
            },
            "timing_seconds": {
                "header_read": provenance.header_read_seconds,
                "amplitude_read": provenance.amplitude_read_seconds,
                "numeric_preparation": provenance.numeric_preparation_seconds,
                "total": provenance.total_seconds,
                "reader_phases": dict(provenance.reader_phase_seconds),
            },
            "cohort_filters": {
                "ledger_filter_applied": provenance.ledger_filter_applied,
                "completed_participants": list(provenance.completed_participants),
                "ledger_excluded_participants": list(provenance.ledger_excluded_participants),
                "manual_excluded_participants": list(provenance.manual_excluded_participants),
                "frequency_qc_excluded_participants": list(provenance.frequency_qc_excluded_participants),
                "incomplete_pair_participants": list(provenance.incomplete_pair_participants),
                "participant_condition_exclusions": [
                    asdict(row) for row in provenance.participant_condition_exclusions
                ],
                "dataset_diagnostics": list(provenance.dataset_diagnostics),
            },
        },
        "harmonic_selection": {
            "mode": selection.selection_mode.value,
            "fixed_highest_harmonic_order": (selection.fixed_highest_harmonic_order),
            "z_threshold": selection.z_threshold,
            "z_ddof": selection.z_ddof,
            "highest_detected_order": selection.highest_detected_order,
            "selected_ceiling_order": int(selection.selected_orders[-1]),
            "selected_orders": [int(value) for value in selection.selected_orders],
            "selected_harmonics_hz": [float(value) for value in selection.selected_harmonics_hz],
            "candidate_count": int(selection.candidate_orders.size),
            "selected_count": int(selection.selected_orders.size),
        },
        "analysis": {
            "permutations_evaluated": result.permutations_evaluated,
            "degrees_of_freedom": result.degrees_of_freedom,
            "cluster_forming_threshold": result.cluster_forming_threshold,
            "cluster_entry_alpha": result.cluster_entry_alpha,
            "cluster_alpha_per_tail": result.cluster_alpha_per_tail,
            "cluster_count": len(result.clusters),
            "significant_cluster_count": sum(cluster.significant for cluster in result.clusters),
            "rng_algorithm": result.rng_algorithm,
            "seed": result.seed,
            "permutation_assignment_hash": result.permutation_assignment_hash,
            "warnings": list(result.warnings),
            "timing_seconds": dict(result.timing_seconds),
            "null_positive_max_mass": _finite_array_summary(result.null_positive_max_mass),
            "null_negative_min_mass": _finite_array_summary(result.null_negative_min_mass),
        },
        "adjacency": {
            "version": result.sensor_adjacency_version,
            "fingerprint_sha256": result.sensor_adjacency_fingerprint,
            "edge_count": len(result.sensor_adjacency_edges),
            "derivation": (
                "Independent clean-room FieldTrip-style compressed BioSemi64 "
                "reconstruction: fixed 169-edge MNE Delaunay subset plus 28 "
                "audited neighbour additions; not the authors' unpublished matrix"
                if result.sensor_adjacency_version == SENSOR_ADJACENCY_VERSION
                and len(result.sensor_adjacency_edges) == 197
                else "Explicit caller-provided versioned spatial edge table"
            ),
            "base_edge_count": (
                169
                if result.sensor_adjacency_version == SENSOR_ADJACENCY_VERSION
                and len(result.sensor_adjacency_edges) == 197
                else None
            ),
            "added_edge_count": (
                28
                if result.sensor_adjacency_version == SENSOR_ADJACENCY_VERSION
                and len(result.sensor_adjacency_edges) == 197
                else None
            ),
            "sensor_edges": [list(edge) for edge in result.sensor_adjacency_edges],
            "harmonic_adjacency": "complete-within-sensor",
            "flattening_order": "sensor-major",
        },
        "source_workbooks": [dict(row) for row in source_rows],
        "artifacts": [dict(artifact) for artifact in artifacts],
    }


def _artifact_manifest_row(
    *,
    role: str,
    staging_path: Path,
    destination: Path,
    project_root: Path,
) -> dict[str, object]:
    digest, size = _sha256_file(staging_path)
    final_path = destination / staging_path.name
    return {
        "role": role,
        "path": _relative_to(final_path, project_root, label="Artifact").as_posix(),
        "sha256": digest,
        "size_bytes": size,
    }


def _check_export_cancelled(cancel_check: Callable[[], bool] | None) -> None:
    if cancel_check is not None and cancel_check():
        raise FreeHarmonicCancelledError("Free Harmonic Clustering export was cancelled.")


def _export_cluster_maps(
    data: ClusterMapData,
    output_directory: Path,
    *,
    cancel_check: Callable[[], bool] | None,
) -> tuple[Path, ...]:
    """Write descriptive maps inside the caller's unpublished staging tree."""

    from .render_cluster_maps import cluster_map_caption, export_cluster_map_figures

    _check_export_cancelled(cancel_check)
    output_directory.mkdir(parents=True, exist_ok=False)
    figure_paths = export_cluster_map_figures(data, output_directory, cancel_check=cancel_check)
    _check_export_cancelled(cancel_check)
    metadata_path = output_directory / "cluster_maps.json"
    _write_manifest(
        metadata_path,
        {
            "schema_version": 1,
            "purpose": "descriptive harmonic slices of existing FHC clusters",
            "run_label": data.run_label,
            "arm_a_label": data.arm_a_label,
            "arm_b_label": data.arm_b_label,
            "value_label": data.value_label,
            "background": (
                "Analyzed arm-mean A-minus-B difference; paired subtraction "
                "precedes averaging for paired runs."
            ),
            "sensor_names": list(data.sensor_names),
            "harmonic_orders": list(data.harmonic_orders),
            "harmonics_hz": list(data.harmonics_hz),
            "mean_difference": data.mean_difference.tolist(),
            "cluster_labels": data.cluster_labels.tolist(),
            "clusters": [asdict(cluster) for cluster in data.clusters],
            "symmetric_color_limit": data.color_limit,
            "multiplicity_note": data.multiplicity_note,
            "caption": cluster_map_caption(data),
            "figures": [path.relative_to(output_directory).as_posix() for path in figure_paths],
            "pointwise_significance_claimed": False,
            "reselected_or_reclustered": False,
        },
    )
    return (*figure_paths, metadata_path)


def export_free_harmonic_run(
    prepared: PreparedContrast,
    result: ClusterPermutationResult,
    *,
    run_id: str | None = None,
    destination: str | Path | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> ExportReceipt:
    """Publish one complete result bundle using a manifest-last directory commit."""

    if not isinstance(prepared, PreparedContrast):
        raise TypeError("prepared must be a PreparedContrast.")
    if not isinstance(result, ClusterPermutationResult):
        raise TypeError("result must be a ClusterPermutationResult.")
    _check_export_cancelled(cancel_check)
    project_root = _resolve_project_root(prepared)
    resolved_run_id, final_directory = resolve_run_destination(
        prepared,
        run_id=run_id,
        destination=destination,
    )
    _validate_result(prepared, result)
    source_rows = _source_workbook_rows(prepared, project_root)
    cluster_rows = _cluster_summary_rows(prepared, result)
    membership_rows = _cluster_membership_rows(prepared, result)
    node_rows = _node_statistic_rows(prepared, result)
    harmonic_rows = _harmonic_selection_rows(prepared)
    machine_participant_rows = _participant_rows(prepared)
    human_participant_rows = _participant_and_exclusion_rows(prepared)
    null_rows = _null_extrema_rows(result)
    created_at_utc = datetime.now(UTC).isoformat()
    if final_directory.exists():
        raise FileExistsError(f"Free-harmonic run already exists: {final_directory}")

    parent = final_directory.parent
    _relative_to(parent.resolve(strict=False), project_root, label="Output parent")
    parent.mkdir(parents=True, exist_ok=True)
    staging = parent / f".{resolved_run_id}.staging-{uuid4().hex}"
    _relative_to(staging.resolve(strict=False), project_root, label="Staging directory")
    staging.mkdir(exist_ok=False)

    csv_specs: tuple[tuple[str, str, Sequence[str], Sequence[Mapping[str, object]]], ...] = (
        (
            "cluster_summary",
            "cluster_summary.csv",
            (
                "cluster_id",
                "sign",
                "mass",
                "p_value",
                "conservative_p_value",
                "adjusted_two_sided_p_value",
                "tie_count",
                "p_ci_low",
                "p_ci_high",
                "confidence_interval_straddles_alpha",
                "significant_cluster_level",
                "node_count",
                "n_a",
                "n_b",
                "arm_a_normalized_cluster_node_mean",
                "arm_b_normalized_cluster_node_mean",
                "arm_a_minus_b_raw_difference",
                "arm_a_minus_b_normalized_difference",
                "effect_denominator_sd",
                "effect_denominator_kind",
                "effect_direction",
                "effect_size_direction",
                "descriptive_effect_label",
                "effect_value_scale",
                "sensor_count",
                "harmonic_count",
                "sensors",
                "harmonic_orders",
                "harmonics_hz",
                "effect_size",
                "effect_size_kind",
            ),
            cluster_rows,
        ),
        (
            "cluster_membership",
            "cluster_membership.csv",
            (
                "cluster_id",
                "sign",
                "node_index",
                "sensor_index",
                "sensor",
                "harmonic_index",
                "harmonic_order",
                "harmonic_hz",
                "observed_t",
                "cluster_mass",
                "cluster_p_value",
                "cluster_significant",
            ),
            membership_rows,
        ),
        (
            "node_statistics",
            "node_statistics.csv",
            (
                "node_index",
                "sensor_index",
                "sensor",
                "harmonic_index",
                "harmonic_order",
                "harmonic_hz",
                "observed_t",
                "cluster_id",
                "cluster_sign",
                "cluster_p_value",
                "cluster_significant",
                "pointwise_significance_claimed",
            ),
            node_rows,
        ),
        (
            "harmonic_selection",
            "harmonic_selection.csv",
            (
                "selection_mode",
                "selection_ceiling_order",
                "z_threshold_used_for_selection",
                "harmonic_order",
                "harmonic_hz",
                "eligible_nonbase",
                "arm_a_z",
                "arm_b_z",
                "detected_arm_a",
                "detected_arm_b",
                "retained_fill_through",
                "exclusion_reason",
            ),
            harmonic_rows,
        ),
        (
            "participants",
            "participants.csv",
            (
                "status",
                "exclusion_reason",
                "arm",
                "arm_label",
                "condition",
                "participant_index",
                "participant_id",
            ),
            machine_participant_rows,
        ),
        (
            "source_workbooks",
            "source_workbooks.csv",
            (
                "arm",
                "arm_label",
                "participant_id",
                "condition",
                "group_id",
                "group_label",
                "project_relative_path",
                "source_size_bytes",
                "source_mtime_ns",
                "header_read_seconds",
                "amplitude_read_seconds",
            ),
            source_rows,
        ),
        (
            "null_extrema",
            "null_extrema.csv",
            ("permutation_index", "positive_max_mass", "negative_min_mass"),
            null_rows,
        ),
    )

    artifact_rows: list[dict[str, object]] = []
    try:
        for role, filename, fieldnames, rows in csv_specs:
            _check_export_cancelled(cancel_check)
            path = staging / filename
            _write_csv(path, fieldnames, rows)
            artifact_rows.append(
                _artifact_manifest_row(
                    role=role,
                    staging_path=path,
                    destination=final_directory,
                    project_root=project_root,
                )
            )

        arrays_path = staging / "arrays.npz"
        _write_npz(arrays_path, prepared, result)
        artifact_rows.append(
            _artifact_manifest_row(
                role="compressed_arrays",
                staging_path=arrays_path,
                destination=final_directory,
                project_root=project_root,
            )
        )

        workbook_path = staging / HUMAN_WORKBOOK_FILENAME
        _write_human_workbook(
            workbook_path,
            run_id=resolved_run_id,
            created_at_utc=created_at_utc,
            prepared=prepared,
            result=result,
            cluster_rows=cluster_rows,
            membership_rows=membership_rows,
            harmonic_rows=harmonic_rows,
            participant_rows=human_participant_rows,
            node_rows=node_rows,
            null_rows=null_rows,
        )
        artifact_rows.append(
            _artifact_manifest_row(
                role="human_readable_workbook",
                staging_path=workbook_path,
                destination=final_directory,
                project_root=project_root,
            )
        )

        map_data = build_cluster_map_data(prepared, result)
        for map_path in _export_cluster_maps(map_data, staging / "cluster_maps", cancel_check=cancel_check):
            artifact_rows.append(
                _batch_artifact_manifest_row(
                    role="cluster_map_data" if map_path.suffix == ".json" else "cluster_map_figure",
                    staging_path=map_path,
                    staging_root=staging,
                    destination=final_directory,
                    project_root=project_root,
                )
            )
        _check_export_cancelled(cancel_check)
        manifest_payload = _manifest_payload(
            run_id=resolved_run_id,
            project_root=project_root,
            destination=final_directory,
            prepared=prepared,
            result=result,
            source_rows=source_rows,
            artifacts=artifact_rows,
            created_at_utc=created_at_utc,
        )
        manifest_staging_path = staging / MANIFEST_FILENAME
        _write_manifest(manifest_staging_path, manifest_payload)

        _check_export_cancelled(cancel_check)
        if final_directory.exists():
            raise FileExistsError(f"Free-harmonic run appeared during export: {final_directory}")
        os.replace(staging, final_directory)
    except BaseException:
        if staging.exists():
            shutil.rmtree(staging)
        raise

    receipt_artifacts: list[ExportArtifact] = []
    for artifact in artifact_rows:
        receipt_artifacts.append(
            ExportArtifact(
                role=str(artifact["role"]),
                path=project_root / str(artifact["path"]),
                sha256=str(artifact["sha256"]),
                size_bytes=int(artifact["size_bytes"]),
            )
        )
    manifest_path = final_directory / MANIFEST_FILENAME
    manifest_sha256, manifest_size = _sha256_file(manifest_path)
    receipt_artifacts.append(
        ExportArtifact(
            role="manifest",
            path=manifest_path,
            sha256=manifest_sha256,
            size_bytes=manifest_size,
        )
    )
    return ExportReceipt(
        output_directory=final_directory,
        manifest_path=manifest_path,
        artifacts=tuple(receipt_artifacts),
    )


def _batch_family_tensor_semantics(
    family: RepeatedSessionContrastFamily,
    *,
    session_a_label: str,
    session_b_label: str,
) -> str:
    if family is RepeatedSessionContrastFamily.SESSION_AVERAGED_GROUPS:
        return (
            "Complete-pair participant mean SNR across both sessions, followed "
            "by one global L2 normalization across the retained sensor x "
            "harmonic tensor; independent comparison of canonical groups."
        )
    if family is RepeatedSessionContrastFamily.PAIRED_SESSIONS_WITHIN_GROUP:
        return (
            f"{session_a_label} and {session_b_label} SNR tensors were separately "
            "L2-normalized within participant; inference uses the paired "
            f"{session_a_label} minus {session_b_label} normalized-pattern change."
        )
    return (
        f"Each participant's separately L2-normalized {session_a_label} minus "
        f"{session_b_label} sensor x harmonic tensor was formed first; those "
        "normalized-pattern change tensors were compared between canonical "
        "groups and were not renormalized after subtraction."
    )


def _batch_run_slug(family_id: str, condition: str) -> str:
    readable = re.sub(
        r"[^A-Za-z0-9._-]+",
        "-",
        f"{family_id}-{condition}",
    ).strip("-._")
    digest = hashlib.sha256(f"{family_id.casefold()}\0{condition.casefold()}".encode("utf-8")).hexdigest()[:10]
    return f"{(readable or 'contrast')[:80]}-{digest}"


def _batch_summary_rows(
    prepared: PreparedRepeatedSessionBatch,
    result: "RepeatedSessionBatchResult",
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    session_a, session_b = prepared.sessions
    for outcome in result.outcomes:
        run = outcome.prepared_run
        contrast = run.prepared
        cluster_result = outcome.result
        rows.append(
            {
                "family_id": run.family_id,
                "family": run.family.value,
                "family_label": run.family_label,
                "condition": run.condition,
                "group_id": run.group_id or "",
                "design": contrast.request.design.value,
                "arm_a_label": contrast.arm_a_label,
                "arm_b_label": contrast.arm_b_label,
                "n_a": len(contrast.participant_ids_a),
                "n_b": len(contrast.participant_ids_b),
                "global_two_sided_cluster_p_value": (outcome.global_two_sided_p_value),
                "holm_within_family_p_value": (outcome.holm_within_family_p_value),
                "holm_all_batch_p_value": outcome.holm_all_batch_p_value,
                "significant_raw_global": (outcome.global_two_sided_p_value <= 0.05),
                "significant_holm_within_family": (outcome.holm_within_family_p_value <= 0.05),
                "significant_holm_all_batch": (outcome.holm_all_batch_p_value <= 0.05),
                "observed_cluster_count": len(cluster_result.clusters),
                "raw_significant_cluster_count": sum(cluster.significant for cluster in cluster_result.clusters),
                "degrees_of_freedom": cluster_result.degrees_of_freedom,
                "derived_seed": outcome.derived_seed,
                "permutation_assignment_hash": (cluster_result.permutation_assignment_hash),
                "tensor_semantics_id": run.tensor_semantics.value,
                "tensor_semantics": _batch_family_tensor_semantics(
                    run.family,
                    session_a_label=session_a.label,
                    session_b_label=session_b.label,
                ),
            }
        )
    return rows


def _batch_cluster_rows(
    result: "RepeatedSessionBatchResult",
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for outcome in result.outcomes:
        run = outcome.prepared_run
        for cluster_row in _cluster_summary_rows(run.prepared, outcome.result):
            arm_a_mean = cluster_row.pop("arm_a_normalized_cluster_node_mean")
            arm_b_mean = cluster_row.pop("arm_b_normalized_cluster_node_mean")
            tensor_difference = cluster_row.pop("arm_a_minus_b_raw_difference")
            cluster_row.pop("arm_a_minus_b_normalized_difference")
            cluster_row["arm_a_tensor_cluster_node_mean"] = arm_a_mean
            cluster_row["arm_b_tensor_cluster_node_mean"] = arm_b_mean
            cluster_row["arm_a_minus_b_tensor_difference"] = tensor_difference
            cluster_row["descriptive_effect_label"] = (
                "post-selection/shape-dependent cluster-node mean on the "
                "declared repeated-session tensor"
            )
            cluster_row["effect_value_scale"] = run.tensor_semantics.value
            rows.append(
                {
                    "family_id": run.family_id,
                    "family": run.family.value,
                    "family_label": run.family_label,
                    "condition": run.condition,
                    "group_id": run.group_id or "",
                    "run_global_two_sided_p_value": (outcome.global_two_sided_p_value),
                    "run_holm_within_family_p_value": (outcome.holm_within_family_p_value),
                    "run_holm_all_batch_p_value": (outcome.holm_all_batch_p_value),
                    **cluster_row,
                }
            )
    return rows


def _exploratory_batch_rows(
    report: "RepeatedSessionReport",
    cluster_rows: Sequence[Mapping[str, object]],
    membership_rows: Sequence[Mapping[str, object]],
) -> tuple[list[dict[str, object]], list[Mapping[str, object]], list[Mapping[str, object]]]:
    """Select existing audit rows using the shared reporting classification."""

    findings = []
    keys = set()
    for row in report.rows:
        if not row.is_exploratory:
            continue
        keys.update(
            (row.family_id, row.condition, cluster_id)
            for cluster_id in row.exploratory_cluster_ids
        )
        finding = {
            "family_id": row.family_id,
            "family_label": row.family_label,
            "condition": row.condition,
            "global_two_sided_cluster_p_value": row.global_p,
            "holm_within_family_p_value": row.holm_family_p,
            "holm_all_batch_p_value": row.holm_batch_p,
            "cluster_ids": "|".join(str(value) for value in row.exploratory_cluster_ids),
        }
        # Keep detailed prose in bounded, wrapped worksheet rows. The complete
        # narrative is also exported separately without Excel's cell limit.
        for paragraph in row.detail_text.splitlines():
            if paragraph.strip():
                for offset in range(0, len(paragraph), 4000):
                    findings.append({**finding, "notes": paragraph[offset:offset + 4000]})

    def selected(row: Mapping[str, object]) -> bool:
        return (row["family_id"], row["condition"], row["cluster_id"]) in keys

    return (
        findings,
        [row for row in cluster_rows if selected(row)],
        [row for row in membership_rows if selected(row)],
    )


def _write_exploratory_report(path: Path, report: "RepeatedSessionReport") -> None:
    """Write the shared detailed narrative inside the atomic run staging area."""

    lines = [
        "# Exploratory FHC findings", "", _EXPLORATORY_DESCRIPTION, "",
        f"{report.exploratory_count} of {len(report.rows)} condition-by-family runs met these criteria.",
        "",
    ]
    for row in report.rows:
        if row.is_exploratory:
            detail = row.detail_text.replace("\n", "  \n")
            lines.extend([f"## {row.condition} | {row.family_label}", "", detail, ""])
    if not report.exploratory_count:
        lines.extend([_NO_EXPLORATORY_FINDINGS, ""])
    lines.extend([
        "Exact candidate cluster rows: exploratory_clusters.csv. Exact sensor-harmonic "
        "membership: exploratory_cluster_membership.csv. The full batch, all clusters, "
        "and their original p-values remain in the primary workbook and CSV tables.",
        "",
    ])
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        stream.write("\n".join(lines))
        stream.flush()
        os.fsync(stream.fileno())


def _batch_membership_rows(
    result: "RepeatedSessionBatchResult",
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for outcome in result.outcomes:
        run = outcome.prepared_run
        for membership in _cluster_membership_rows(run.prepared, outcome.result):
            rows.append(
                {
                    "family_id": run.family_id,
                    "condition": run.condition,
                    "group_id": run.group_id or "",
                    **membership,
                }
            )
    return rows


def _batch_node_rows(
    result: "RepeatedSessionBatchResult",
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for outcome in result.outcomes:
        run = outcome.prepared_run
        for node in _node_statistic_rows(run.prepared, outcome.result):
            rows.append(
                {
                    "family_id": run.family_id,
                    "condition": run.condition,
                    "group_id": run.group_id or "",
                    **node,
                }
            )
    return rows


def _batch_cohort_audit_rows(
    prepared: PreparedRepeatedSessionBatch,
) -> list[dict[str, object]]:
    return [
        {
            "participant_id": row.participant_id,
            "group_id": row.group_id,
            "condition": row.condition,
            "available_session_ids": "|".join(row.available_session_ids),
            "missing_session_ids": "|".join(row.missing_session_ids),
            "recording_ids_by_session": "|".join(
                f"{session_id}:{recording_id}" for session_id, recording_id in row.recording_ids_by_session
            ),
            "excluded_recording_ids": "|".join(row.excluded_recording_ids),
            "excluded_session_ids": "|".join(row.excluded_session_ids),
            "exclusion_reasons": "|".join(row.exclusion_reasons),
            "included_complete_pair": row.included_complete_pair,
        }
        for row in prepared.cohort_audit
    ]


def _batch_shared_selection_rows(
    prepared: PreparedRepeatedSessionBatch,
) -> list[dict[str, object]]:
    selection = prepared.shared_selection
    audit = prepared.shared_selection_audit
    retained = {int(index) for index in selection.selected_candidate_indices}
    rows: list[dict[str, object]] = []
    for cell_index, cell_label in enumerate(audit.cell_labels):
        for harmonic_index, (order, harmonic_hz) in enumerate(
            zip(
                selection.candidate_orders,
                selection.candidate_harmonics_hz,
                strict=True,
            )
        ):
            rows.append(
                {
                    "cell_label": cell_label,
                    "group_id": audit.cell_group_ids[cell_index],
                    "session_id": audit.cell_session_ids[cell_index],
                    "condition": audit.cell_conditions[cell_index],
                    "participant_count": audit.cell_participant_counts[cell_index],
                    "harmonic_order": int(order),
                    "harmonic_hz": float(harmonic_hz),
                    "z_score": float(audit.z_scores[cell_index, harmonic_index]),
                    "strict_z_detected": bool(audit.detected[cell_index, harmonic_index]),
                    "retained_shared_fill_through": harmonic_index in retained,
                    "z_threshold": selection.z_threshold,
                }
            )
    return rows


def _batch_source_rows(
    prepared: PreparedRepeatedSessionBatch,
    project_root: Path,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    seen: set[str] = set()
    for record in prepared.source_workbooks:
        relative = _validate_source_workbook(record, project_root)
        key = relative.casefold()
        if key in seen:
            continue
        seen.add(key)
        stat = Path(record.source_path).stat()
        rows.append(
            {
                "recording_id": record.recording_id or "",
                "participant_id": record.participant_id,
                "group_id": record.group_id or "",
                "group_label": record.group_label or "",
                "session_id": record.session_id or "",
                "session_label": record.session_label or "",
                "visit_index": record.visit_index or "",
                "condition": record.condition,
                "project_relative_path": relative,
                "source_size_bytes": int(stat.st_size),
                "source_mtime_ns": int(stat.st_mtime_ns),
                "header_read_seconds": record.header_read_seconds,
                "amplitude_read_seconds": record.amplitude_read_seconds,
            }
        )
    return rows


def _batch_methods_rows(
    prepared: PreparedRepeatedSessionBatch,
    result: "RepeatedSessionBatchResult",
) -> list[dict[str, object]]:
    session_a, session_b = prepared.sessions
    provenance = prepared.provenance
    plan = prepared.frequency_plan
    first_result = result.outcomes[0].result
    rows: list[dict[str, object]] = [
        {
            "category": "Batch contract",
            "item": "batch_version",
            "value": prepared.request.batch_version,
            "notes": "Versioned repeated-session extension; legacy one-contrast behavior is unchanged.",
        },
        {
            "category": "Calibration",
            "item": "calibration_status",
            "value": "not covered by legacy powered receipt",
            "notes": (
                "The legacy v2 numerical permutation core is reused unchanged, "
                "but its powered automatic-domain receipt does not validate the "
                "multi-cell shared selector or this repeated-session batch of "
                f"{len(prepared.contrast_runs)} runs."
            ),
        },
        {
            "category": "Harmonic domain",
            "item": "shared_domain_fingerprint",
            "value": prepared.shared_domain_fingerprint,
            "notes": "One frozen candidate domain was supplied to every batch run.",
        },
        {
            "category": "Input provenance",
            "item": "source_sheet",
            "value": provenance.source_sheet,
            "notes": "Original processing-owned FullFFT amplitudes; standard selected-harmonic outputs were not consumed.",
        },
        {
            "category": "Input provenance",
            "item": "grid_fingerprint",
            "value": provenance.grid_fingerprint,
            "notes": "Exact FullFFT frequency-grid identity.",
        },
        {
            "category": "Input provenance",
            "item": "selected_columns_fingerprint",
            "value": provenance.selected_columns_fingerprint,
            "notes": "Exact target/noise-column selection identity.",
        },
        {
            "category": "Input provenance",
            "item": "full_fft_provenance_method_version",
            "value": provenance.full_fft_provenance_method_version,
            "notes": "Neutral processing-owned FullFFT provenance contract.",
        },
        {
            "category": "Input provenance",
            "item": "full_fft_source_fingerprint",
            "value": provenance.full_fft_source_fingerprint,
            "notes": "Canonical active source-workbook identity.",
        },
        {
            "category": "Input provenance",
            "item": "full_fft_cohort_fingerprint",
            "value": provenance.full_fft_cohort_fingerprint,
            "notes": "Canonical processing cohort identity.",
        },
        {
            "category": "Input provenance",
            "item": "full_fft_frequency_qc_fingerprint",
            "value": provenance.full_fft_frequency_qc_fingerprint,
            "notes": "Frequency-domain QC identity used by neutral FullFFT provenance.",
        },
        {
            "category": "Input provenance",
            "item": "full_fft_processing_export_fingerprint",
            "value": provenance.full_fft_processing_export_fingerprint,
            "notes": "Processing/export identity used by neutral FullFFT provenance.",
        },
        {
            "category": "Frequency plan",
            "item": "frequency_resolution_hz",
            "value": plan.frequency_resolution_hz,
            "notes": "Resolved from the canonical FullFFT grid.",
        },
        {
            "category": "Frequency plan",
            "item": "noise_window_rule",
            "value": "physical +/- half-width inclusive; target and immediate adjacent FFT bins excluded",
            "notes": (
                "noise_half_width_hz="
                f"{prepared.method.noise_half_width_hz:g}; exact resolved bins "
                "are listed below"
            ),
        },
        {
            "category": "Adjacency",
            "item": "sensor_adjacency_version",
            "value": first_result.sensor_adjacency_version,
            "notes": "Shared by every run in the repeated-session batch.",
        },
        {
            "category": "Adjacency",
            "item": "sensor_adjacency_fingerprint_sha256",
            "value": first_result.sensor_adjacency_fingerprint,
            "notes": "Exact ordered sensor-edge graph identity.",
        },
        {
            "category": "Adjacency",
            "item": "sensor_adjacency_edges",
            "value": json.dumps(
                [list(edge) for edge in first_result.sensor_adjacency_edges],
                separators=(",", ":"),
            ),
            "notes": f"{len(first_result.sensor_adjacency_edges)} undirected sensor edges; harmonic adjacency is complete within sensor.",
        },
        {
            "category": "Multiplicity",
            "item": "within_family",
            "value": result.within_family_method,
            "notes": (
                "Applied to run-global two-sided cluster p values across all "
                "declared conditions separately for each scientific family."
            ),
        },
        {
            "category": "Multiplicity",
            "item": "all_batch",
            "value": result.all_batch_method,
            "notes": "Conservative sensitivity correction across every batch run.",
        },
        {
            "category": "Inference",
            "item": "run_global_p_definition",
            "value": "minimum cluster adjusted_two_sided_p_value; 1 when no observed clusters",
            "notes": (
                "The signed max-cluster null already controls electrode x "
                "harmonic multiplicity within a run. Cross-run Holm values are "
                "run-level annotations and are not cluster-specific p values."
            ),
        },
        {
            "category": "Session direction",
            "item": "arm_a_minus_arm_b",
            "value": f"{session_a.label} minus {session_b.label}",
            "notes": (
                f"Visit {session_a.visit_index} minus Visit {session_b.visit_index}; "
                "session/phase is confounded with fixed visit order, elapsed "
                "time, and retest effects."
            ),
        },
    ]
    for family in RepeatedSessionContrastFamily:
        rows.append(
            {
                "category": "Tensor semantics",
                "item": family.value,
                "value": _batch_family_tensor_semantics(
                    family,
                    session_a_label=session_a.label,
                    session_b_label=session_b.label,
                ),
                "notes": "FHC tests normalized sensor x harmonic response distributions, not total Raw BCA magnitude.",
            }
        )
    for field_name, value in asdict(prepared.method).items():
        rows.append(
            {
                "category": "Method setting",
                "item": field_name,
                "value": value,
                "notes": ("Base batch seed; each run receives a stable derived seed." if field_name == "seed" else ""),
            }
        )
    for harmonic_index, (order, harmonic_hz) in enumerate(
        zip(plan.candidate_orders, plan.candidate_harmonics_hz, strict=True)
    ):
        target_index = int(plan.target_selected_indices[harmonic_index])
        noise_indices = tuple(
            int(value) for value in plan.noise_selected_indices[harmonic_index]
        )
        rows.append(
            {
                "category": "Frequency-bin audit",
                "item": f"H{int(order)} ({float(harmonic_hz):g} Hz)",
                "value": json.dumps(
                    {
                        "target_selected_index": target_index,
                        "target_column": plan.selected_frequency_columns[target_index],
                        "target_bin_hz": float(plan.selected_frequencies_hz[target_index]),
                        "target_error_hz": float(
                            plan.selected_frequencies_hz[target_index]
                            - harmonic_hz
                        ),
                        "noise_selected_indices": list(noise_indices),
                        "noise_columns": [
                            plan.selected_frequency_columns[index]
                            for index in noise_indices
                        ],
                        "noise_frequencies_hz": [
                            float(plan.selected_frequencies_hz[index])
                            for index in noise_indices
                        ],
                    },
                    separators=(",", ":"),
                ),
                "notes": "Exact target and physical neighboring-noise bins used for participant SNR.",
            }
        )
    for outcome in result.outcomes:
        rows.append(
            {
                "category": "Derived seed mapping",
                "item": f"{outcome.family_id} | {outcome.condition}",
                "value": outcome.derived_seed,
                "notes": outcome.result.permutation_assignment_hash,
            }
        )
    for warning in dict.fromkeys(warning for outcome in result.outcomes for warning in outcome.result.warnings):
        rows.append(
            {
                "category": "Inference warning",
                "item": "warning",
                "value": warning,
                "notes": "",
            }
        )
    return rows


def _write_batch_arrays(
    path: Path,
    *,
    outcome: object,
    prepared: PreparedRepeatedSessionBatch,
) -> None:
    run = outcome.prepared_run
    contrast = run.prepared
    result = outcome.result
    session_a, session_b = prepared.sessions
    semantics = _batch_family_tensor_semantics(
        run.family,
        session_a_label=session_a.label,
        session_b_label=session_b.label,
    )
    if run.family is RepeatedSessionContrastFamily.SESSION_AVERAGED_GROUPS:
        tensor_a_key = "l2_normalized_participant_mean_snr_group_a"
        tensor_b_key = "l2_normalized_participant_mean_snr_group_b"
    elif run.family is RepeatedSessionContrastFamily.PAIRED_SESSIONS_WITHIN_GROUP:
        tensor_a_key = "l2_normalized_session_a_snr"
        tensor_b_key = "l2_normalized_session_b_snr"
    else:
        tensor_a_key = "normalized_snr_session_change_group_a"
        tensor_b_key = "normalized_snr_session_change_group_b"
    metadata = {
        "family_id": run.family_id,
        "family": run.family.value,
        "condition": run.condition,
        "group_id": run.group_id,
        "arm_a_label": contrast.arm_a_label,
        "arm_b_label": contrast.arm_b_label,
        "tensor_a_key": tensor_a_key,
        "tensor_b_key": tensor_b_key,
        "tensor_semantics_id": run.tensor_semantics.value,
        "tensor_semantics": semantics,
        "shared_domain_fingerprint": prepared.shared_domain_fingerprint,
        "global_two_sided_cluster_p_value": (outcome.global_two_sided_p_value),
        "holm_within_family_p_value": outcome.holm_within_family_p_value,
        "holm_all_batch_p_value": outcome.holm_all_batch_p_value,
        "derived_seed": outcome.derived_seed,
    }
    payload = {
        tensor_a_key: np.asarray(contrast.values_a),
        tensor_b_key: np.asarray(contrast.values_b),
        "participant_ids_a": np.asarray(contrast.participant_ids_a, dtype=np.str_),
        "participant_ids_b": np.asarray(contrast.participant_ids_b, dtype=np.str_),
        "sensor_names": np.asarray(contrast.sensor_names, dtype=np.str_),
        "harmonic_orders": np.asarray(contrast.harmonic_orders),
        "harmonics_hz": np.asarray(contrast.harmonics_hz),
        "observed_t": np.asarray(result.observed_t),
        "cluster_labels": np.asarray(result.cluster_labels),
        "null_positive_max_mass": np.asarray(result.null_positive_max_mass),
        "null_negative_min_mass": np.asarray(result.null_negative_min_mass),
        "metadata_json": np.asarray(
            json.dumps(metadata, sort_keys=True, ensure_ascii=False),
            dtype=np.str_,
        ),
    }
    with path.open("wb") as stream:
        np.savez_compressed(stream, **payload)
        stream.flush()
        os.fsync(stream.fileno())


def _write_repeated_session_workbook(
    path: Path,
    *,
    summary_rows: Sequence[Mapping[str, object]],
    cluster_rows: Sequence[Mapping[str, object]],
    membership_rows: Sequence[Mapping[str, object]],
    node_rows: Sequence[Mapping[str, object]],
    cohort_rows: Sequence[Mapping[str, object]],
    selection_rows: Sequence[Mapping[str, object]],
    source_rows: Sequence[Mapping[str, object]],
    methods_rows: Sequence[Mapping[str, object]],
    exploratory_rows: Sequence[Mapping[str, object]],
) -> None:
    sheets = (
        "Batch Summary",
        "Family-Significant Runs",
        "All Clusters",
        "Cluster Membership",
        "Node Statistics",
        "Cohort Audit",
        "Shared Harmonic Selection",
        "Source Workbooks",
        "Methods and Provenance",
        "Exploratory Findings",
    )
    workbook = Workbook()
    workbook.active.title = sheets[0]
    for name in sheets[1:]:
        workbook.create_sheet(name)
    summary_fields = (
        "family_id",
        "family_label",
        "condition",
        "group_id",
        "design",
        "arm_a_label",
        "arm_b_label",
        "n_a",
        "n_b",
        "global_two_sided_cluster_p_value",
        "holm_within_family_p_value",
        "holm_all_batch_p_value",
        "significant_raw_global",
        "significant_holm_within_family",
        "significant_holm_all_batch",
        "observed_cluster_count",
        "raw_significant_cluster_count",
        "degrees_of_freedom",
        "derived_seed",
        "tensor_semantics_id",
        "tensor_semantics",
    )
    _write_table_sheet(
        workbook["Batch Summary"],
        title="Repeated-Session FHC Batch Summary",
        description=(
            "Run-global cluster p values and Holm corrections. Cross-condition "
            "adjustments are run-level annotations, not cluster-specific p values."
        ),
        fields=summary_fields,
        rows=summary_rows,
        significant_field="significant_holm_within_family",
    )
    _write_table_sheet(
        workbook["Family-Significant Runs"],
        title="Family-Corrected Significant Runs",
        description=(
            "Runs with Holm-adjusted p < .05 across the declared conditions in their prespecified scientific family."
        ),
        fields=summary_fields,
        rows=[row for row in summary_rows if bool(row["significant_holm_within_family"])],
        empty_message="No run survived Holm correction within its scientific family.",
        significant_field="significant_holm_within_family",
    )
    cluster_fields = (
        "family_id",
        "family_label",
        "condition",
        "group_id",
        "run_global_two_sided_p_value",
        "run_holm_within_family_p_value",
        "run_holm_all_batch_p_value",
        "cluster_id",
        "sign",
        "mass",
        "p_value",
        "conservative_p_value",
        "adjusted_two_sided_p_value",
        "significant_cluster_level",
        "node_count",
        "sensors",
        "harmonic_orders",
        "harmonics_hz",
        "effect_size",
        "effect_size_kind",
        "effect_direction",
        "n_a",
        "n_b",
    )
    _write_table_sheet(
        workbook["All Clusters"],
        title="All Observed Clusters",
        description=(
            "Raw conditional max-cluster inference. Individual nodes are not "
            "pointwise significant and Holm values apply only to the run."
        ),
        fields=cluster_fields,
        rows=cluster_rows,
        significant_field="significant_cluster_level",
    )
    membership_fields = (
        "family_id",
        "condition",
        "group_id",
        "cluster_id",
        "sign",
        "node_index",
        "sensor",
        "harmonic_order",
        "harmonic_hz",
        "observed_t",
        "cluster_mass",
        "cluster_p_value",
        "cluster_significant",
    )
    _write_table_sheet(
        workbook["Cluster Membership"],
        title="Cluster Membership",
        description="Sensor-harmonic membership only; no pointwise inference is claimed.",
        fields=membership_fields,
        rows=membership_rows,
        significant_field="cluster_significant",
    )
    node_fields = (
        "family_id",
        "condition",
        "group_id",
        "node_index",
        "sensor_index",
        "sensor",
        "harmonic_index",
        "harmonic_order",
        "harmonic_hz",
        "observed_t",
        "cluster_id",
        "cluster_sign",
        "cluster_p_value",
        "cluster_significant",
        "pointwise_significance_claimed",
    )
    _write_table_sheet(
        workbook["Node Statistics"],
        title="Node Statistics",
        description=(
            "Observed sensor-harmonic t values and cluster membership for every "
            "batch run; no pointwise significance is claimed."
        ),
        fields=node_fields,
        rows=node_rows,
        significant_field="cluster_significant",
    )
    cohort_fields = (
        "participant_id",
        "group_id",
        "condition",
        "available_session_ids",
        "missing_session_ids",
        "recording_ids_by_session",
        "excluded_recording_ids",
        "excluded_session_ids",
        "exclusion_reasons",
        "included_complete_pair",
    )
    _write_table_sheet(
        workbook["Cohort Audit"],
        title="Complete-Pair Cohort Audit",
        description=(
            "Availability and explicit analysis-only recording exclusions for "
            "every participant x condition cell. Missing sessions reduce coverage."
        ),
        fields=cohort_fields,
        rows=cohort_rows,
    )
    selection_fields = (
        "cell_label",
        "group_id",
        "session_id",
        "condition",
        "participant_count",
        "harmonic_order",
        "harmonic_hz",
        "z_score",
        "strict_z_detected",
        "retained_shared_fill_through",
        "z_threshold",
    )
    _write_table_sheet(
        workbook["Shared Harmonic Selection"],
        title="Shared Multi-Cell Harmonic Selection",
        description=(
            "One shared domain selected before inference and reused unchanged for every condition and contrast family."
        ),
        fields=selection_fields,
        rows=selection_rows,
    )
    source_fields = (
        "recording_id",
        "participant_id",
        "group_id",
        "session_id",
        "session_label",
        "visit_index",
        "condition",
        "project_relative_path",
        "source_size_bytes",
        "source_mtime_ns",
        "header_read_seconds",
        "amplitude_read_seconds",
    )
    _write_table_sheet(
        workbook["Source Workbooks"],
        title="Recording-Aware Source Workbooks",
        description="Original FullFFT workbooks consumed through the canonical project index.",
        fields=source_fields,
        rows=source_rows,
    )
    _write_table_sheet(
        workbook["Methods and Provenance"],
        title="Methods and Provenance",
        description=(
            "Tensor estimands, fixed-order caveat, multiplicity, calibration "
            "scope, seeds, fingerprints, and numerical warnings."
        ),
        fields=("category", "item", "value", "notes"),
        rows=methods_rows,
    )
    exploratory_fields = (
        "family_id", "family_label", "condition", "global_two_sided_cluster_p_value",
        "holm_within_family_p_value", "holm_all_batch_p_value", "cluster_ids", "notes",
    )
    sheet = workbook["Exploratory Findings"]
    _write_table_sheet(
        sheet,
        title="Exploratory Findings — Not Significant After Family Holm",
        description=_EXPLORATORY_DESCRIPTION,
        fields=exploratory_fields,
        rows=exploratory_rows,
        empty_message=_NO_EXPLORATORY_FINDINGS,
    )
    sheet.row_dimensions[2].height = 48
    sheet.column_dimensions["H"].width = 110
    for row_index, row in enumerate(exploratory_rows, start=5):
        sheet.row_dimensions[row_index].height = min(
            409, max(30, 15 * (2 + len(str(row["notes"])) // 100)),
        )
        for column_index in (4, 5, 6):
            sheet.cell(row_index, column_index).number_format = "0.########"
    workbook.save(path)
    with path.open("rb+") as stream:
        os.fsync(stream.fileno())


def _batch_artifact_manifest_row(
    *,
    role: str,
    staging_path: Path,
    staging_root: Path,
    destination: Path,
    project_root: Path,
) -> dict[str, object]:
    digest, size = _sha256_file(staging_path)
    relative_in_bundle = staging_path.relative_to(staging_root)
    final_path = destination / relative_in_bundle
    return {
        "role": role,
        "path": _relative_to(
            final_path,
            project_root,
            label="Batch artifact",
        ).as_posix(),
        "sha256": digest,
        "size_bytes": size,
    }


def export_repeated_session_batch(
    prepared: PreparedRepeatedSessionBatch,
    result: "RepeatedSessionBatchResult",
    *,
    run_id: str | None = None,
    destination: str | Path | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> ExportReceipt:
    """Atomically publish one corrected repeated-session FHC batch bundle."""

    if not isinstance(prepared, PreparedRepeatedSessionBatch):
        raise TypeError("prepared must be a PreparedRepeatedSessionBatch.")
    _check_export_cancelled(cancel_check)
    outcomes = tuple(getattr(result, "outcomes", ()))
    if len(outcomes) != len(prepared.contrast_runs):
        raise ValueError("Batch result does not match the prepared contrast count.")
    if int(getattr(result, "base_seed", -1)) != prepared.method.seed:
        raise ValueError("Batch result base seed does not match the prepared method.")
    from .analysis import (
        adjust_batch_cluster_p_values,
        derive_repeated_session_run_seed,
    )

    expected_multiplicity = adjust_batch_cluster_p_values(
        tuple(
            (outcome.family_id, outcome.condition, outcome.result)
            for outcome in outcomes
        )
    )
    first_adjacency: tuple[str, str, tuple[tuple[int, int], ...]] | None = None
    for expected, outcome, expected_adjustment in zip(
        prepared.contrast_runs,
        outcomes,
        expected_multiplicity,
        strict=True,
    ):
        _check_export_cancelled(cancel_check)
        actual = outcome.prepared_run
        if actual is not expected:
            raise ValueError(
                "Batch result must reference the exact prepared contrast objects."
            )
        if (
            actual.family_id.casefold() != expected.family_id.casefold()
            or actual.condition.casefold() != expected.condition.casefold()
        ):
            raise ValueError("Batch result order does not match preparation order.")
        _validate_result(actual.prepared, outcome.result)
        expected_seed = derive_repeated_session_run_seed(
            prepared.method.seed,
            family_id=expected.family_id,
            condition=expected.condition,
        )
        if outcome.result.seed != outcome.derived_seed or outcome.derived_seed != expected_seed:
            raise ValueError("Batch result seed mapping is inconsistent.")
        expected_values = (
            expected_adjustment.global_two_sided_p_value,
            expected_adjustment.holm_within_family_p_value,
            expected_adjustment.holm_all_batch_p_value,
        )
        actual_values = (
            outcome.global_two_sided_p_value,
            outcome.holm_within_family_p_value,
            outcome.holm_all_batch_p_value,
        )
        if any(
            not np.isclose(actual_value, expected_value, rtol=0.0, atol=1e-15)
            for actual_value, expected_value in zip(
                actual_values,
                expected_values,
                strict=True,
            )
        ):
            raise ValueError(
                "Batch result run-global or Holm p values are inconsistent "
                "with the cluster results."
            )
        adjacency = (
            outcome.result.sensor_adjacency_version,
            outcome.result.sensor_adjacency_fingerprint,
            outcome.result.sensor_adjacency_edges,
        )
        if first_adjacency is None:
            first_adjacency = adjacency
        elif adjacency != first_adjacency:
            raise ValueError("Every batch run must use the same sensor adjacency graph.")

    first = prepared.contrast_runs[0].prepared
    project_root = _resolve_project_root(first)
    if Path(prepared.project_root).resolve(strict=False) != project_root:
        raise FreeHarmonicInputError("Prepared batch and contrast refer to different project roots.")
    resolved_run_id, final_directory = resolve_run_destination(
        first,
        run_id=run_id,
        destination=destination,
    )
    if final_directory.exists():
        raise FileExistsError(f"Free-harmonic batch already exists: {final_directory}")

    summary_rows = _batch_summary_rows(prepared, result)
    cluster_rows = _batch_cluster_rows(result)
    membership_rows = _batch_membership_rows(result)
    node_rows = _batch_node_rows(result)
    cohort_rows = _batch_cohort_audit_rows(prepared)
    selection_rows = _batch_shared_selection_rows(prepared)
    source_rows = _batch_source_rows(prepared, project_root)
    methods_rows = _batch_methods_rows(prepared, result)
    from .reporting import build_repeated_session_report

    report = build_repeated_session_report(result)
    exploratory_rows, exploratory_clusters, exploratory_membership = _exploratory_batch_rows(
        report, cluster_rows, membership_rows,
    )
    created_at_utc = datetime.now(UTC).isoformat()

    parent = final_directory.parent
    _relative_to(parent.resolve(strict=False), project_root, label="Output parent")
    parent.mkdir(parents=True, exist_ok=True)
    staging = parent / f".{resolved_run_id}.staging-{uuid4().hex}"
    _relative_to(staging.resolve(strict=False), project_root, label="Staging directory")
    staging.mkdir(exist_ok=False)
    artifacts: list[dict[str, object]] = []

    csv_specs = (
        ("batch_summary", "batch_summary.csv", tuple(summary_rows[0]) if summary_rows else (), summary_rows),
        (
            "all_clusters",
            "all_clusters.csv",
            tuple(cluster_rows[0])
            if cluster_rows
            else ("family_id", "family_label", "condition", "group_id", "cluster_id"),
            cluster_rows,
        ),
        (
            "cluster_membership",
            "cluster_membership.csv",
            tuple(membership_rows[0]) if membership_rows else ("family_id", "condition", "group_id", "cluster_id"),
            membership_rows,
        ),
        (
            "node_statistics",
            "node_statistics.csv",
            tuple(node_rows[0]) if node_rows else ("family_id", "condition", "group_id", "node_index"),
            node_rows,
        ),
        ("cohort_audit", "cohort_audit.csv", tuple(cohort_rows[0]) if cohort_rows else (), cohort_rows),
        (
            "shared_harmonic_selection",
            "shared_harmonic_selection.csv",
            tuple(selection_rows[0]) if selection_rows else (),
            selection_rows,
        ),
        ("source_workbooks", "source_workbooks.csv", tuple(source_rows[0]) if source_rows else (), source_rows),
        ("methods_and_provenance", "methods_and_provenance.csv", ("category", "item", "value", "notes"), methods_rows),
        (
            "exploratory_clusters", "exploratory_clusters.csv",
            tuple(cluster_rows[0]) if cluster_rows else ("family_id", "condition", "cluster_id"),
            exploratory_clusters,
        ),
        (
            "exploratory_cluster_membership", "exploratory_cluster_membership.csv",
            tuple(membership_rows[0]) if membership_rows else ("family_id", "condition", "cluster_id"),
            exploratory_membership,
        ),
    )
    try:
        _check_export_cancelled(cancel_check)
        exploratory_path = staging / "exploratory_findings.md"
        _write_exploratory_report(exploratory_path, report)
        artifacts.append(
            _batch_artifact_manifest_row(
                role="exploratory_findings", staging_path=exploratory_path,
                staging_root=staging, destination=final_directory, project_root=project_root,
            )
        )
        for role, filename, fields, rows in csv_specs:
            _check_export_cancelled(cancel_check)
            if not fields:
                raise ValueError(f"Batch export has no columns for {role}.")
            path = staging / filename
            _write_csv(path, fields, rows)
            artifacts.append(
                _batch_artifact_manifest_row(
                    role=role,
                    staging_path=path,
                    staging_root=staging,
                    destination=final_directory,
                    project_root=project_root,
                )
            )

        arrays_directory = staging / "contrast_arrays"
        arrays_directory.mkdir()
        array_mapping: list[dict[str, object]] = []
        map_mapping: list[dict[str, object]] = []
        for outcome in outcomes:
            _check_export_cancelled(cancel_check)
            slug = _batch_run_slug(outcome.family_id, outcome.condition)
            path = arrays_directory / f"{slug}.npz"
            _write_batch_arrays(path, outcome=outcome, prepared=prepared)
            artifacts.append(
                _batch_artifact_manifest_row(
                    role="contrast_arrays",
                    staging_path=path,
                    staging_root=staging,
                    destination=final_directory,
                    project_root=project_root,
                )
            )
            array_mapping.append(
                {
                    "family_id": outcome.family_id,
                    "condition": outcome.condition,
                    "path": f"contrast_arrays/{path.name}",
                }
            )
            map_directory = staging / "cluster_maps" / slug
            for map_path in _export_cluster_maps(
                build_repeated_cluster_map_data(outcome),
                map_directory,
                cancel_check=cancel_check,
            ):
                artifacts.append(
                    _batch_artifact_manifest_row(
                        role="cluster_map_data" if map_path.suffix == ".json" else "cluster_map_figure",
                        staging_path=map_path,
                        staging_root=staging,
                        destination=final_directory,
                        project_root=project_root,
                    )
                )
            map_mapping.append(
                {
                    "family_id": outcome.family_id,
                    "condition": outcome.condition,
                    "data_path": (map_directory / "cluster_maps.json").relative_to(staging).as_posix(),
                }
            )

        workbook_path = staging / REPEATED_SESSION_WORKBOOK_FILENAME
        _write_repeated_session_workbook(
            workbook_path,
            summary_rows=summary_rows,
            cluster_rows=cluster_rows,
            membership_rows=membership_rows,
            node_rows=node_rows,
            cohort_rows=cohort_rows,
            selection_rows=selection_rows,
            source_rows=source_rows,
            methods_rows=methods_rows,
            exploratory_rows=exploratory_rows,
        )
        artifacts.append(
            _batch_artifact_manifest_row(
                role="human_readable_workbook",
                staging_path=workbook_path,
                staging_root=staging,
                destination=final_directory,
                project_root=project_root,
            )
        )

        session_a, session_b = prepared.sessions
        provenance = prepared.provenance
        first_result = outcomes[0].result
        manifest = {
            "schema_version": REPEATED_SESSION_EXPORT_SCHEMA_VERSION,
            "analysis_kind": "repeated_session_free_harmonic_clustering_batch",
            "run_id": resolved_run_id,
            "created_at_utc": created_at_utc,
            "project": {
                "project_root": ".",
                "output_directory": _relative_to(
                    final_directory,
                    project_root,
                    label="Output directory",
                ).as_posix(),
            },
            "batch": {
                "version": prepared.request.batch_version,
                "conditions": list(prepared.conditions),
                "group_ids_a_minus_b": list(prepared.request.group_ids),
                "session_ids_a_minus_b": list(prepared.request.session_ids),
                "session_labels_a_minus_b": [session_a.label, session_b.label],
                "visit_indices_a_minus_b": [
                    session_a.visit_index,
                    session_b.visit_index,
                ],
                "recording_exclusions": [asdict(row) for row in prepared.request.recording_exclusions],
                "base_seed": result.base_seed,
                "shared_domain_fingerprint": prepared.shared_domain_fingerprint,
            },
            "method": asdict(prepared.method),
            "frequency_plan": _frequency_plan_manifest(first),
            "preparation": {
                "source_sheet": provenance.source_sheet,
                "grid_fingerprint": provenance.grid_fingerprint,
                "selected_columns_fingerprint": (
                    provenance.selected_columns_fingerprint
                ),
                "frequency_resolution_hz": provenance.frequency_resolution_hz,
                "full_frequency_column_count": (
                    provenance.full_frequency_column_count
                ),
                "selected_frequency_column_count": (
                    provenance.selected_frequency_column_count
                ),
                "workbook_count": provenance.workbook_count,
                "neutral_full_fft_provenance": {
                    "method_version": (
                        provenance.full_fft_provenance_method_version
                    ),
                    "source_fingerprint": (
                        provenance.full_fft_source_fingerprint
                    ),
                    "cohort_fingerprint": (
                        provenance.full_fft_cohort_fingerprint
                    ),
                    "frequency_qc_fingerprint": (
                        provenance.full_fft_frequency_qc_fingerprint
                    ),
                    "processing_export_fingerprint": (
                        provenance.full_fft_processing_export_fingerprint
                    ),
                },
                "timing_seconds": {
                    "header_read": provenance.header_read_seconds,
                    "amplitude_read": provenance.amplitude_read_seconds,
                    "numeric_preparation": (
                        provenance.numeric_preparation_seconds
                    ),
                    "total": provenance.total_seconds,
                    "reader_phases": dict(provenance.reader_phase_seconds),
                },
                "cohort_filters": {
                    "ledger_filter_applied": provenance.ledger_filter_applied,
                    "completed_recordings": list(
                        provenance.completed_recordings
                    ),
                    "ledger_excluded_recordings": list(
                        provenance.ledger_excluded_recordings
                    ),
                    "manual_excluded_participants": list(
                        provenance.manual_excluded_participants
                    ),
                    "frequency_qc_excluded_participants": list(
                        provenance.frequency_qc_excluded_participants
                    ),
                    "frequency_qc_excluded_recordings": list(
                        provenance.frequency_qc_excluded_recordings
                    ),
                    "incomplete_pair_participants": list(
                        provenance.incomplete_pair_participants
                    ),
                    "participant_condition_exclusions": [
                        asdict(row)
                        for row in provenance.participant_condition_exclusions
                    ],
                    "analysis_recording_exclusions": [
                        asdict(row)
                        for row in provenance.request_recording_exclusions
                    ],
                    "dataset_diagnostics": list(
                        provenance.dataset_diagnostics
                    ),
                },
            },
            "harmonic_selection": {
                "mode": prepared.shared_selection.selection_mode.value,
                "z_threshold": prepared.shared_selection.z_threshold,
                "z_ddof": prepared.shared_selection.z_ddof,
                "highest_detected_order": (
                    prepared.shared_selection.highest_detected_order
                ),
                "selected_ceiling_order": int(
                    prepared.shared_selection.selected_orders[-1]
                ),
                "selected_orders": [
                    int(value)
                    for value in prepared.shared_selection.selected_orders
                ],
                "selected_harmonics_hz": [
                    float(value)
                    for value in prepared.shared_selection.selected_harmonics_hz
                ],
                "candidate_count": int(
                    prepared.shared_selection.candidate_orders.size
                ),
                "selected_count": int(
                    prepared.shared_selection.selected_orders.size
                ),
                "shared_cell_audit_artifact": "shared_harmonic_selection.csv",
            },
            "adjacency": {
                "version": first_result.sensor_adjacency_version,
                "fingerprint_sha256": first_result.sensor_adjacency_fingerprint,
                "edge_count": len(first_result.sensor_adjacency_edges),
                "sensor_edges": [
                    list(edge) for edge in first_result.sensor_adjacency_edges
                ],
                "harmonic_adjacency": "complete-within-sensor",
                "flattening_order": "sensor-major",
            },
            "multiplicity": {
                "run_global_p": ("minimum cluster adjusted_two_sided_p_value; 1 when no observed clusters"),
                "within_family": result.within_family_method,
                "all_batch": result.all_batch_method,
                "cluster_specific_cross_condition_adjustment": False,
            },
            "calibration": {
                "legacy_numerical_core_reused_unchanged": True,
                "legacy_powered_receipt_validates_repeated_batch": False,
                "status": (
                    "The powered legacy receipt does not validate the multi-cell "
                    "shared selector or the repeated-session multi-run batch."
                ),
            },
            "claims": {
                "whole_participant_exchangeability_only": True,
                "cluster_level_inference_only": True,
                "pointwise_sensor_harmonic_significance": False,
                "session_phase_confound": (
                    "Session/phase-at-visit is confounded with fixed visit order, elapsed time, and retest effects."
                ),
                "estimand": ("normalized sensor x harmonic response distribution; not total Raw BCA magnitude"),
            },
            "tensor_semantics": {
                family.value: _batch_family_tensor_semantics(
                    family,
                    session_a_label=session_a.label,
                    session_b_label=session_b.label,
                )
                for family in RepeatedSessionContrastFamily
            },
            "results": summary_rows,
            "exploratory_reporting": {
                "criterion": EXPLORATORY_CRITERION,
                "qualifying_run_count": report.exploratory_count,
                "qualifying_cluster_count": len(exploratory_clusters),
                "report": "exploratory_findings.md",
                "cluster_rows": "exploratory_clusters.csv",
                "membership_rows": "exploratory_cluster_membership.csv",
                "changes_inference": False,
                "pointwise_significance_claimed": False,
            },
            "contrast_arrays": array_mapping,
            "cluster_maps": map_mapping,
            "source_workbooks": source_rows,
            "artifacts": artifacts,
        }
        manifest_path = staging / MANIFEST_FILENAME
        _check_export_cancelled(cancel_check)
        _write_manifest(manifest_path, manifest)
        _check_export_cancelled(cancel_check)
        if final_directory.exists():
            raise FileExistsError(f"Free-harmonic batch appeared during export: {final_directory}")
        os.replace(staging, final_directory)
    except BaseException:
        if staging.exists():
            shutil.rmtree(staging)
        raise

    receipt_artifacts = [
        ExportArtifact(
            role=str(row["role"]),
            path=project_root / str(row["path"]),
            sha256=str(row["sha256"]),
            size_bytes=int(row["size_bytes"]),
        )
        for row in artifacts
    ]
    final_manifest = final_directory / MANIFEST_FILENAME
    manifest_sha256, manifest_size = _sha256_file(final_manifest)
    receipt_artifacts.append(
        ExportArtifact(
            role="manifest",
            path=final_manifest,
            sha256=manifest_sha256,
            size_bytes=manifest_size,
        )
    )
    return ExportReceipt(
        output_directory=final_directory,
        manifest_path=final_manifest,
        artifacts=tuple(receipt_artifacts),
    )


__all__ = [
    "DEFAULT_RESULTS_SUBFOLDER",
    "EXPORT_SCHEMA_VERSION",
    "HUMAN_WORKBOOK_FILENAME",
    "HUMAN_WORKBOOK_SHEETS",
    "MANIFEST_FILENAME",
    "REPEATED_SESSION_EXPORT_SCHEMA_VERSION",
    "REPEATED_SESSION_WORKBOOK_FILENAME",
    "TOOL_TITLE",
    "export_free_harmonic_run",
    "export_repeated_session_batch",
    "resolve_run_destination",
]
