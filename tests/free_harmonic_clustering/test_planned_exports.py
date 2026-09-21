"""Frozen-family export and read-only report contracts."""

from dataclasses import replace
import hashlib
import json

import numpy as np
from openpyxl import load_workbook
import pytest

from Tools.Free_Harmonic_Clustering import planned_exports
from Tools.Free_Harmonic_Clustering.analysis import holm_adjust_p_values
from Tools.Free_Harmonic_Clustering.analysis_plan import AnalysisFamily, AnalysisPlan, FAMILY_LABELS, PlannedComparison
from Tools.Free_Harmonic_Clustering.models import FreeHarmonicCancelledError, FreeHarmonicInputError
from Tools.Free_Harmonic_Clustering.planned_analysis import (
    PlannedAnalysisResult,
    PlannedComparisonOutcome,
    derive_planned_comparison_seed,
)
from Tools.Free_Harmonic_Clustering.planned_inputs import PreparedAnalysisPlan
from Tools.Free_Harmonic_Clustering.planned_reporting import build_analysis_plan_report, build_planned_cluster_map_data
from tests.free_harmonic_clustering.test_exports import _stub_map_figures
from tests.free_harmonic_clustering.test_repeated_session_batch_inference import _batch_fixture


@pytest.fixture
def planned_result(tmp_path, monkeypatch):
    _stub_map_figures(monkeypatch)
    prepared, legacy = _batch_fixture(tmp_path)
    kinds = (
        AnalysisFamily.BETWEEN_GROUPS,
        AnalysisFamily.WITHIN_GROUP_VISITS,
        AnalysisFamily.WITHIN_GROUP_VISITS,
        AnalysisFamily.GROUP_VISIT_CHANGE,
    )
    comparisons = tuple(
        PlannedComparison(
            comparison_id=f"comparison_{index}",
            family_id=kind.value,
            family_label=FAMILY_LABELS[kind],
            label=run.family_label,
            kind=kind,
            design=run.prepared.request.design,
            condition_a=run.condition,
            condition_b=run.prepared.request.condition_b,
            group_ids=run.prepared.request.group_ids,
            session_ids=prepared.request.session_ids,
        )
        for index, (kind, run) in enumerate(zip(kinds, prepared.contrast_runs, strict=True))
    )
    plan = AnalysisPlan(
        project_root=prepared.project_root,
        group_ids=prepared.request.group_ids,
        conditions=prepared.conditions,
        session_ids=prepared.request.session_ids,
        families=tuple(dict.fromkeys(kinds)),
        condition_pairs=(),
        recording_exclusions=(),
        comparisons=comparisons,
        group_labels=tuple(row.label for row in prepared.groups),
        session_labels=tuple(row.label for row in prepared.sessions),
    )
    contrasts = tuple(
        replace(
            run.prepared,
            request=replace(
                run.prepared.request,
                contrast_family_id=comparison.comparison_id,
            ),
        )
        for run, comparison in zip(prepared.contrast_runs, comparisons, strict=True)
    )
    snapshot = PreparedAnalysisPlan(
        plan,
        prepared.method,
        contrasts,
        prepared.shared_selection,
        prepared.shared_selection_audit,
        prepared.frequency_plan,
        prepared.provenance,
        prepared.source_workbooks,
        prepared.cohort_audit,
    )
    # Same stored raw p-values; one formerly separate visit test now fails the
    # shared visit family. Export must not reinterpret clusters as Holm tests.
    globals_ = (0.01, 0.04, 0.03, 0.5)
    family = (0.01, 0.06, 0.06, 0.5)
    full = holm_adjust_p_values(globals_)
    outcomes = []
    for index, (comparison, prior) in enumerate(zip(comparisons, legacy.outcomes, strict=True)):
        seed = derive_planned_comparison_seed(prepared.method.seed, comparison.comparison_id)
        contrast = replace(contrasts[index], method=replace(prepared.method, seed=seed))
        cluster = replace(
            prior.result.clusters[0],
            p_value=globals_[index] / 2,
            adjusted_two_sided_p_value=globals_[index],
            significant=globals_[index] < 0.05,
        )
        raw = replace(prior.result, seed=seed, clusters=(cluster,))
        outcomes.append(
            PlannedComparisonOutcome(comparison, contrast, raw, seed, globals_[index], family[index], full[index])
        )
    return PlannedAnalysisResult(plan, snapshot, tuple(outcomes))


def test_plan_export_keeps_family_identity_and_unmodified_sources(planned_result):
    root = planned_result.plan.project_root
    project_before = (root / "project.json").read_bytes()
    receipt = planned_exports.export_analysis_plan_result(planned_result, run_id="families")
    manifest = json.loads(receipt.manifest_path.read_text(encoding="utf-8"))
    assert (root / "project.json").read_bytes() == project_before
    assert manifest["plan_fingerprint"] == planned_result.plan.fingerprint
    assert len(manifest["comparisons"]) == 4
    assert manifest["results"][1]["family_comparison_count"] == 2
    assert manifest["results"][1]["holm_within_family_p_value"] == 0.06
    assert len({row["comparison_id"] for row in manifest["comparisons"]}) == 4
    assert "legacy single-contrast calibration does not validate" in manifest["validation"]
    saved_plan = json.loads((receipt.output_directory / "analysis_plan.json").read_text(encoding="utf-8"))
    assert saved_plan["plan_fingerprint"] == planned_result.plan.fingerprint
    assert saved_plan["version"] == "fhc_analysis_families_v2"
    workbook = load_workbook(receipt.output_directory / planned_exports.PLANNED_WORKBOOK_FILENAME, read_only=True)
    try:
        assert {"Families", "Comparisons", "Participants", "Cohort Audit", "Methods and Provenance"} <= set(
            workbook.sheetnames
        )
        assert workbook["Comparisons"]["J6"].value == 0.06
    finally:
        workbook.close()
    for artifact in receipt.artifacts:
        assert hashlib.sha256(artifact.path.read_bytes()).hexdigest() == artifact.sha256
    array_path = receipt.output_directory / "comparisons" / "0004" / "arrays.npz"
    with np.load(array_path, allow_pickle=False) as arrays:
        assert "normalized_values_a" not in arrays
        np.testing.assert_array_equal(arrays["analyzed_values_a"], planned_result.outcomes[3].contrast.values_a)
        metadata = json.loads(str(arrays["metadata_json"]))
        assert "no renormalization" in metadata["tensor_semantics"]
    with pytest.raises(FileExistsError):
        planned_exports.export_analysis_plan_result(planned_result, run_id="families")


def test_report_preserves_comparison_indices_and_exploratory_boundary(planned_result):
    report = build_analysis_plan_report(planned_result)
    assert [row.run_index for row in report.rows if row.is_exploratory] == [1, 2]
    assert report.rows[1].family_label == report.rows[2].family_label
    assert report.rows[1].comparison_label != report.rows[2].comparison_label
    assert "Averaged visits:" in report.rows[0].detail_text
    assert "Visit order:" not in report.rows[0].detail_text
    data = build_planned_cluster_map_data(planned_result.outcomes[3])
    assert "visit-change difference" in data.value_label
    regrouped = replace(
        planned_result.outcomes[3],
        comparison=replace(planned_result.outcomes[3].comparison, family_id="another_correction_family"),
    )
    assert build_planned_cluster_map_data(regrouped).value_label == data.value_label
    assert "does not pass" in report.rows[1].detail_text.casefold()
    assert "family Holm p = 0.06" in build_planned_cluster_map_data(planned_result.outcomes[1]).multiplicity_note


def test_export_rejects_recomputed_correction_or_seed_tampering(planned_result):
    bad = replace(planned_result.outcomes[1], family_adjusted_p=0.04)
    changed = replace(planned_result, outcomes=(planned_result.outcomes[0], bad, *planned_result.outcomes[2:]))
    with pytest.raises(ValueError, match="complete analysis plan"):
        planned_exports.export_analysis_plan_result(changed, run_id="bad")
    first = planned_result.outcomes[0]
    wrong_seed = first.derived_seed + 1
    bad = replace(
        first,
        derived_seed=wrong_seed,
        result=replace(first.result, seed=wrong_seed),
        contrast=replace(first.contrast, method=replace(first.contrast.method, seed=wrong_seed)),
    )
    changed = replace(planned_result, outcomes=(bad, *planned_result.outcomes[1:]))
    with pytest.raises(ValueError, match="seed"):
        planned_exports.export_analysis_plan_result(changed, run_id="bad-seed")


def test_plan_export_is_cancel_safe_and_contained(planned_result, monkeypatch, tmp_path):
    with pytest.raises(FreeHarmonicInputError, match="beneath"):
        planned_exports.export_analysis_plan_result(planned_result, destination=tmp_path / "outside")
    parent = planned_result.plan.project_root / "3 - Statistical Analysis Results" / "Free Harmonic Clustering Analysis"
    monkeypatch.setattr(
        planned_exports,
        "_write_plan_arrays",
        lambda *_args: (_ for _ in ()).throw(FreeHarmonicCancelledError("cancelled during arrays")),
    )
    with pytest.raises(FreeHarmonicCancelledError):
        planned_exports.export_analysis_plan_result(planned_result, run_id="cancelled")
    assert not (parent / "cancelled").exists()
    assert not list(parent.glob(".*.staging-*"))
