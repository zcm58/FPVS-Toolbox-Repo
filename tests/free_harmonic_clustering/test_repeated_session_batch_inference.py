from __future__ import annotations

import csv
from dataclasses import replace
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import pytest

from Tools.Free_Harmonic_Clustering import analysis, api, exports
from Tools.Free_Harmonic_Clustering.models import (
    AnalysisDesign,
    FreeHarmonicCancelledError,
    PreparedRepeatedSessionBatch,
    PreparedRepeatedSessionContrast,
    ProjectContrastRequest,
    ProjectGroupOption,
    ProjectSessionOption,
    RepeatedSessionBatchRequest,
    RepeatedSessionCohortAuditRow,
    RepeatedSessionContrastFamily,
    RepeatedSessionTensorSemantics,
    SharedHarmonicSelectionAudit,
)

from tests.free_harmonic_clustering.test_exports import _prepared_and_result, _stub_map_figures


@pytest.fixture(autouse=True)
def _map_figures(monkeypatch: pytest.MonkeyPatch) -> None:
    _stub_map_figures(monkeypatch)


def _batch_fixture(
    tmp_path: Path,
) -> tuple[PreparedRepeatedSessionBatch, api.RepeatedSessionBatchResult]:
    base_prepared, base_result = _prepared_and_result(tmp_path)
    project_root = base_prepared.project_root
    groups = (
        ProjectGroupOption("anxious", "Anxious"),
        ProjectGroupOption("non_anxious", "Non-Anxious"),
    )
    sessions = (
        ProjectSessionOption("visit_2", "Follicular Phase", 2),
        ProjectSessionOption("visit_1", "Luteal Phase", 1),
    )
    request = RepeatedSessionBatchRequest(
        project_root=project_root,
        conditions=("Neutral Angry",),
        group_ids=(groups[0].group_id, groups[1].group_id),
        session_ids=(sessions[0].session_id, sessions[1].session_id),
    )
    run_specs = (
        (
            RepeatedSessionContrastFamily.SESSION_AVERAGED_GROUPS,
            "session_averaged_groups",
            "Session-averaged groups",
            None,
            AnalysisDesign.INDEPENDENT_GROUPS,
            groups[0].group_id,
        ),
        (
            RepeatedSessionContrastFamily.PAIRED_SESSIONS_WITHIN_GROUP,
            "paired_sessions_within_group:anxious",
            "Follicular minus luteal within Anxious",
            groups[0].group_id,
            AnalysisDesign.PAIRED_CONDITIONS,
            groups[0].group_id,
        ),
        (
            RepeatedSessionContrastFamily.PAIRED_SESSIONS_WITHIN_GROUP,
            "paired_sessions_within_group:non_anxious",
            "Follicular minus luteal within Non-Anxious",
            groups[1].group_id,
            AnalysisDesign.PAIRED_CONDITIONS,
            groups[1].group_id,
        ),
        (
            RepeatedSessionContrastFamily.GROUP_SESSION_CHANGE,
            "group_session_change",
            "Group difference in follicular-minus-luteal change",
            None,
            AnalysisDesign.INDEPENDENT_GROUPS,
            groups[0].group_id,
        ),
    )
    prepared_runs: list[PreparedRepeatedSessionContrast] = []
    p_values = (0.01, 0.04, 0.02, 0.50)
    raw_results = []
    for family, family_id, family_label, group_id, design, paired_group in run_specs:
        group_ids = (
            (groups[0].group_id, groups[1].group_id) if design is AnalysisDesign.INDEPENDENT_GROUPS else (paired_group,)
        )
        run_request = ProjectContrastRequest(
            project_root=project_root,
            design=design,
            condition_a="Neutral Angry",
            condition_b=("Neutral Angry" if design is AnalysisDesign.PAIRED_CONDITIONS else None),
            group_ids=group_ids,
            session_ids=request.session_ids,
            contrast_family_id=family_id,
        )
        if design is AnalysisDesign.PAIRED_CONDITIONS:
            if paired_group == groups[0].group_id:
                participant_ids = base_prepared.participant_ids_a
                values = base_prepared.values_a
                snr = base_prepared.snr_a
            else:
                participant_ids = base_prepared.participant_ids_b
                values = base_prepared.values_b
                snr = base_prepared.snr_b
            arm_a_label = sessions[0].label
            arm_b_label = sessions[1].label
            run_prepared = replace(
                base_prepared,
                request=run_request,
                arm_a_label=arm_a_label,
                arm_b_label=arm_b_label,
                participant_ids_a=participant_ids,
                participant_ids_b=participant_ids,
                values_a=values,
                values_b=np.flip(values, axis=2),
                snr_a=snr,
                snr_b=np.flip(snr, axis=2),
            )
        else:
            run_prepared = replace(base_prepared, request=run_request)
        wrapper = PreparedRepeatedSessionContrast(
            family=family,
            family_id=family_id,
            family_label=family_label,
            condition="Neutral Angry",
            group_id=group_id,
            tensor_semantics=(
                RepeatedSessionTensorSemantics.SESSION_NORMALIZED_PAIRED_PROFILE
                if family is RepeatedSessionContrastFamily.PAIRED_SESSIONS_WITHIN_GROUP
                else (
                    RepeatedSessionTensorSemantics.NORMALIZED_SESSION_CHANGE
                    if family is RepeatedSessionContrastFamily.GROUP_SESSION_CHANGE
                    else RepeatedSessionTensorSemantics.SESSION_AVERAGED_NORMALIZED_PROFILE
                )
            ),
            prepared=run_prepared,
        )
        prepared_runs.append(wrapper)
        derived_seed = analysis.derive_repeated_session_run_seed(
            base_prepared.method.seed,
            family_id=family_id,
            condition="Neutral Angry",
        )
        cluster = replace(
            base_result.clusters[0],
            adjusted_two_sided_p_value=p_values[len(raw_results)],
            p_value=min(1.0, p_values[len(raw_results)] / 2.0),
            significant=p_values[len(raw_results)] < 0.05,
            effect_size_kind=("paired_dz" if design is AnalysisDesign.PAIRED_CONDITIONS else "pooled_cohen_d"),
        )
        run_result = replace(
            base_result,
            design=design,
            clusters=(cluster,),
            cluster_labels=np.array([[1, 0], [0, 0]]),
            seed=derived_seed,
        )
        raw_results.append(run_result)

    fingerprint = "shared-domain-sha"
    provenance = replace(
        base_prepared.provenance,
        repeated_session_batch_version=request.batch_version,
        shared_domain_fingerprint=fingerprint,
    )
    prepared_runs = [replace(row, prepared=replace(row.prepared, provenance=provenance)) for row in prepared_runs]
    source_workbooks = tuple(
        replace(
            row,
            recording_id=f"{row.participant_id}_visit_2",
            session_id=sessions[0].session_id,
            session_label=sessions[0].label,
            visit_index=sessions[0].visit_index,
        )
        for row in base_prepared.source_workbooks
    )
    cohort_audit = tuple(
        RepeatedSessionCohortAuditRow(
            participant_id=participant_id,
            group_id=group.group_id,
            condition="Neutral Angry",
            available_session_ids=request.session_ids,
            missing_session_ids=(),
            recording_ids_by_session=(
                (sessions[0].session_id, f"{participant_id}_visit_2"),
                (sessions[1].session_id, f"{participant_id}_visit_1"),
            ),
            included_complete_pair=True,
        )
        for group, participant_ids in zip(
            groups,
            (base_prepared.participant_ids_a, base_prepared.participant_ids_b),
            strict=True,
        )
        for participant_id in participant_ids
    )
    shared_audit = SharedHarmonicSelectionAudit(
        cell_labels=(
            "Anxious | Follicular | Neutral Angry",
            "Anxious | Luteal | Neutral Angry",
            "Non-Anxious | Follicular | Neutral Angry",
            "Non-Anxious | Luteal | Neutral Angry",
        ),
        cell_group_ids=("anxious", "anxious", "non_anxious", "non_anxious"),
        cell_session_ids=("visit_2", "visit_1", "visit_2", "visit_1"),
        cell_conditions=("Neutral Angry",) * 4,
        cell_participant_counts=(2, 2, 2, 2),
        z_scores=np.array([[4.0, 2.0], [3.8, 2.1], [3.7, 3.5], [3.6, 3.4]]),
        detected=np.array([[True, False], [True, False], [True, True], [True, True]]),
    )
    prepared_batch = PreparedRepeatedSessionBatch(
        request=request,
        method=base_prepared.method,
        project_root=project_root,
        conditions=request.conditions,
        groups=groups,
        sessions=sessions,
        contrast_runs=tuple(prepared_runs),
        shared_selection=base_prepared.selection,
        shared_selection_audit=shared_audit,
        frequency_plan=base_prepared.frequency_plan,
        sensor_names=base_prepared.sensor_names,
        shared_domain_fingerprint=fingerprint,
        source_workbooks=source_workbooks,
        cohort_audit=cohort_audit,
        provenance=provenance,
    )
    multiplicity = analysis.adjust_batch_cluster_p_values(
        tuple((run.family_id, run.condition, result) for run, result in zip(prepared_runs, raw_results, strict=True))
    )
    outcomes = tuple(
        api.RepeatedSessionContrastOutcome(
            prepared_run=run,
            result=run_result,
            derived_seed=run_result.seed,
            global_two_sided_p_value=adjusted.global_two_sided_p_value,
            holm_within_family_p_value=adjusted.holm_within_family_p_value,
            holm_all_batch_p_value=adjusted.holm_all_batch_p_value,
        )
        for run, run_result, adjusted in zip(
            prepared_runs,
            raw_results,
            multiplicity,
            strict=True,
        )
    )
    return prepared_batch, api.RepeatedSessionBatchResult(
        outcomes=outcomes,
        base_seed=base_prepared.method.seed,
    )


def test_batch_export_is_atomic_corrected_and_tensor_semantic(tmp_path: Path) -> None:
    prepared, result = _batch_fixture(tmp_path)

    receipt = exports.export_repeated_session_batch(
        prepared,
        result,
        run_id="batch-001",
    )

    manifest = json.loads(receipt.manifest_path.read_text(encoding="utf-8"))
    assert manifest["calibration"]["legacy_powered_receipt_validates_repeated_batch"] is False
    assert manifest["multiplicity"]["cluster_specific_cross_condition_adjustment"] is False
    assert len(manifest["cluster_maps"]) == len(prepared.contrast_runs)
    for mapping, outcome in zip(manifest["cluster_maps"], result.outcomes, strict=True):
        assert mapping["family_id"] == outcome.family_id
        assert mapping["condition"] == outcome.condition
        map_path = receipt.output_directory / mapping["data_path"]
        maps = json.loads(map_path.read_text(encoding="utf-8"))
        assert outcome.prepared_run.family_label in maps["run_label"]
        assert "Raw cluster p values are within-run" in maps["multiplicity_note"]
        assert f"Holm within family p = {outcome.holm_within_family_p_value:.4g}" in maps["multiplicity_note"]
        assert "not individual clusters" in maps["multiplicity_note"]
        np.testing.assert_allclose(
            maps["mean_difference"],
            outcome.prepared_run.prepared.values_a.mean(axis=0)
            - outcome.prepared_run.prepared.values_b.mean(axis=0),
        )
    assert manifest["frequency_plan"]["noise_selected_indices"]
    assert manifest["preparation"]["grid_fingerprint"] == prepared.provenance.grid_fingerprint
    assert (
        manifest["preparation"]["neutral_full_fft_provenance"][
            "source_fingerprint"
        ]
        == prepared.provenance.full_fft_source_fingerprint
    )
    assert manifest["adjacency"]["fingerprint_sha256"] == result.outcomes[0].result.sensor_adjacency_fingerprint
    assert manifest["adjacency"]["sensor_edges"]
    assert len(manifest["results"]) == 4
    assert {row["holm_within_family_p_value"] for row in manifest["results"]} == {
        0.01,
        0.04,
        0.02,
        0.5,
    }
    interaction_path = next(
        receipt.output_directory / row["path"]
        for row in manifest["contrast_arrays"]
        if row["family_id"] == "group_session_change"
    )
    with np.load(interaction_path) as arrays:
        assert "normalized_snr_session_change_group_a" in arrays.files
        assert "normalized_snr_session_change_group_b" in arrays.files
        assert "snr_a" not in arrays.files
        metadata = json.loads(str(arrays["metadata_json"]))
        assert "not renormalized after subtraction" in metadata["tensor_semantics"]

    with (receipt.output_directory / "all_clusters.csv").open(
        encoding="utf-8",
        newline="",
    ) as stream:
        cluster_rows = list(csv.DictReader(stream))
    interaction_cluster = next(
        row for row in cluster_rows if row["family_id"] == "group_session_change"
    )
    assert "arm_a_tensor_cluster_node_mean" in interaction_cluster
    assert "arm_a_normalized_cluster_node_mean" not in interaction_cluster
    assert "arm_a_minus_b_raw_difference" not in interaction_cluster
    assert (
        interaction_cluster["effect_value_scale"]
        == RepeatedSessionTensorSemantics.NORMALIZED_SESSION_CHANGE.value
    )

    workbook = load_workbook(
        receipt.output_directory / exports.REPEATED_SESSION_WORKBOOK_FILENAME,
        read_only=True,
    )
    assert "Batch Summary" in workbook.sheetnames
    assert "Node Statistics" in workbook.sheetnames
    workbook.close()
    assert not list(receipt.output_directory.parent.glob(".batch-001.staging-*"))


def test_batch_map_cancellation_discards_all_condition_family_outputs(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from Tools.Free_Harmonic_Clustering import render_cluster_maps

    prepared, result = _batch_fixture(tmp_path)
    figure_exporter = render_cluster_maps.export_cluster_map_figures
    calls = 0
    cancel_requested = False

    def cancel_second_run(data: object, output_dir: Path, *, cancel_check: object = None) -> tuple[Path, ...]:
        nonlocal calls, cancel_requested
        calls += 1
        paths = figure_exporter(data, output_dir, cancel_check=cancel_check)
        if calls == 2:
            cancel_requested = True
        return paths

    monkeypatch.setattr(render_cluster_maps, "export_cluster_map_figures", cancel_second_run)
    with pytest.raises(FreeHarmonicCancelledError):
        exports.export_repeated_session_batch(
            prepared,
            result,
            run_id="cancelled-batch-maps",
            cancel_check=lambda: cancel_requested,
        )
    assert calls == 2
    parent = prepared.project_root / exports.DEFAULT_RESULTS_SUBFOLDER
    assert not (parent / "cancelled-batch-maps").exists()
    assert not list(parent.glob(".cancelled-batch-maps.staging-*"))


def test_batch_export_recomputes_run_global_and_holm_annotations(
    tmp_path: Path,
) -> None:
    prepared, result = _batch_fixture(tmp_path)
    bad_outcome = replace(
        result.outcomes[0],
        holm_all_batch_p_value=min(
            1.0,
            result.outcomes[0].holm_all_batch_p_value + 0.1,
        ),
    )
    bad_result = replace(
        result,
        outcomes=(bad_outcome, *result.outcomes[1:]),
    )

    with pytest.raises(ValueError, match="Holm p values are inconsistent"):
        exports.export_repeated_session_batch(
            prepared,
            bad_result,
            run_id="bad-p-values",
        )

    assert not (
        prepared.project_root
        / exports.DEFAULT_RESULTS_SUBFOLDER
        / "bad-p-values"
    ).exists()


def test_batch_analysis_uses_stable_per_cell_seeds_and_total_progress(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    prepared, existing = _batch_fixture(tmp_path)
    calls: list[tuple[AnalysisDesign, int]] = []
    progress_rows: list[tuple[int, int]] = []

    def fake_run(
        _arm_a: np.ndarray,
        _arm_b: np.ndarray,
        *,
        design: AnalysisDesign,
        method: object,
        progress: object,
        **_kwargs: object,
    ) -> object:
        calls.append((design, method.seed))
        progress(0, method.n_permutations)
        progress(method.n_permutations, method.n_permutations)
        template = existing.outcomes[len(calls) - 1].result
        return replace(template, seed=method.seed)

    monkeypatch.setattr(analysis, "run_cluster_permutation", fake_run)

    result = api.analyze_prepared_repeated_session_batch(
        prepared,
        progress=lambda done, total: progress_rows.append((done, total)),
    )

    assert [design for design, _seed in calls] == [
        AnalysisDesign.INDEPENDENT_GROUPS,
        AnalysisDesign.PAIRED_CONDITIONS,
        AnalysisDesign.PAIRED_CONDITIONS,
        AnalysisDesign.INDEPENDENT_GROUPS,
    ]
    assert [seed for _design, seed in calls] == [
        analysis.derive_repeated_session_run_seed(
            prepared.method.seed,
            family_id=run.family_id,
            condition=run.condition,
        )
        for run in prepared.contrast_runs
    ]
    assert progress_rows[-1] == (
        len(prepared.contrast_runs) * prepared.method.n_permutations,
        len(prepared.contrast_runs) * prepared.method.n_permutations,
    )
    assert len(result.outcomes) == 4
