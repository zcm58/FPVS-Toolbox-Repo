from __future__ import annotations

import csv
from dataclasses import replace
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import pytest

from Tools.Free_Harmonic_Clustering import exports
from Tools.Free_Harmonic_Clustering.models import (
    AnalysisDesign,
    ClusterPermutationResult,
    ClusterRecord,
    CohortWorkbook,
    FreeHarmonicInputError,
    FreeHarmonicMethodSpec,
    FrequencyWindowPlan,
    HarmonicSelection,
    ParticipantConditionExclusion,
    PreparationProvenance,
    PreparedContrast,
    ProjectContrastRequest,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _prepared_and_result(tmp_path: Path) -> tuple[PreparedContrast, ClusterPermutationResult]:
    project_root = tmp_path / "project"
    project_root.mkdir()
    (project_root / "project.json").write_text("{}", encoding="utf-8")
    participant_ids_a = ("A1", "A2")
    participant_ids_b = ("B1", "B2")
    source_workbooks: list[CohortWorkbook] = []
    for arm, label, group_id, group_label, participant_ids in (
        ("a", "Anxious", "anxious", "Anxious", participant_ids_a),
        ("b", "Non-Anxious", "non_anxious", "Non-Anxious", participant_ids_b),
    ):
        for participant_id in participant_ids:
            source = (
                project_root
                / "1 - Excel Data Files"
                / "Neutral Angry"
                / group_label
                / f"{participant_id}_Neutral Angry_Results.xlsx"
            )
            source.parent.mkdir(parents=True, exist_ok=True)
            source.write_bytes(b"synthetic workbook")
            source_workbooks.append(
                CohortWorkbook(
                    arm=arm,
                    arm_label=label,
                    participant_id=participant_id,
                    condition="Neutral Angry",
                    group_id=group_id,
                    group_label=group_label,
                    source_path=source,
                    project_relative_path=source.relative_to(project_root).as_posix(),
                    header_read_seconds=0.01,
                    amplitude_read_seconds=0.02,
                )
            )

    method = FreeHarmonicMethodSpec(
        n_permutations=3,
        seed=7,
        sensor_adjacency_version="test-adjacency-v1",
    )
    request = ProjectContrastRequest(
        project_root=project_root,
        design=AnalysisDesign.INDEPENDENT_GROUPS,
        condition_a="Neutral Angry",
        group_ids=("anxious", "non_anxious"),
    )
    frequency_plan = FrequencyWindowPlan(
        full_frequency_columns=(
            "0.0000_Hz",
            "1.2000_Hz",
            "2.4000_Hz",
            "3.6000_Hz",
            "4.8000_Hz",
        ),
        full_frequencies_hz=np.array([0.0, 1.2, 2.4, 3.6, 4.8]),
        selected_frequency_columns=(
            "0.0000_Hz",
            "1.2000_Hz",
            "2.4000_Hz",
            "3.6000_Hz",
            "4.8000_Hz",
        ),
        selected_frequencies_hz=np.array([0.0, 1.2, 2.4, 3.6, 4.8]),
        candidate_orders=np.array([1, 2]),
        candidate_harmonics_hz=np.array([1.2, 2.4]),
        excluded_base_orders=np.array([], dtype=np.int64),
        excluded_base_harmonics_hz=np.array([], dtype=np.float64),
        target_selected_indices=np.array([1, 2]),
        noise_selected_indices=np.array([[0, 3], [0, 4]]),
        frequency_resolution_hz=1.2,
        grid_fingerprint="grid-sha",
        selected_columns_fingerprint="selected-sha",
    )
    selection = HarmonicSelection(
        candidate_orders=np.array([1, 2]),
        candidate_harmonics_hz=np.array([1.2, 2.4]),
        arm_a_z=np.array([4.0, 2.0]),
        arm_b_z=np.array([2.5, 3.5]),
        detected_arm_a=np.array([True, False]),
        detected_arm_b=np.array([False, True]),
        selected_candidate_indices=np.array([0, 1]),
        selected_orders=np.array([1, 2]),
        selected_harmonics_hz=np.array([1.2, 2.4]),
        excluded_base_orders=np.array([], dtype=np.int64),
        excluded_base_harmonics_hz=np.array([], dtype=np.float64),
        z_threshold=3.29,
        z_ddof=1,
        highest_detected_order=2,
    )
    snr_a = np.arange(1.0, 9.0).reshape(2, 2, 2)
    snr_b = np.arange(2.0, 10.0).reshape(2, 2, 2)
    values_a = snr_a / np.linalg.norm(snr_a, axis=(1, 2), keepdims=True)
    values_b = snr_b / np.linalg.norm(snr_b, axis=(1, 2), keepdims=True)
    prepared = PreparedContrast(
        request=request,
        method=method,
        project_root=project_root,
        arm_a_label="Anxious",
        arm_b_label="Non-Anxious",
        participant_ids_a=participant_ids_a,
        participant_ids_b=participant_ids_b,
        sensor_names=("Fp1", "Fp2"),
        harmonic_orders=np.array([1, 2]),
        harmonics_hz=np.array([1.2, 2.4]),
        snr_a=snr_a,
        snr_b=snr_b,
        values_a=values_a,
        values_b=values_b,
        selection=selection,
        frequency_plan=frequency_plan,
        source_workbooks=tuple(source_workbooks),
        provenance=PreparationProvenance(
            source_sheet="FullFFT Amplitude (uV)",
            grid_fingerprint="grid-sha",
            selected_columns_fingerprint="selected-sha",
            frequency_resolution_hz=1.2,
            full_frequency_column_count=5,
            selected_frequency_column_count=5,
            workbook_count=4,
            header_read_seconds=0.04,
            amplitude_read_seconds=0.08,
            numeric_preparation_seconds=0.01,
            total_seconds=0.13,
            reader_phase_seconds=(("worksheet_xml", 0.08),),
            ledger_filter_applied=True,
            completed_participants=("A1", "A2", "B1", "B2"),
            participant_condition_exclusions=(
                ParticipantConditionExclusion(
                    participant_id="P4",
                    condition="Neutral Angry",
                ),
            ),
            full_fft_provenance_method_version="neutral-full-fft-v1",
            full_fft_source_fingerprint="full-fft-source-sha",
            full_fft_cohort_fingerprint="full-fft-cohort-sha",
            full_fft_frequency_qc_fingerprint="full-fft-qc-sha",
            full_fft_processing_export_fingerprint="processing-export-sha",
        ),
    )
    clusters = (
        ClusterRecord(
            cluster_id=1,
            sign="positive",
            mass=2.0,
            p_value=0.02,
            conservative_p_value=0.03,
            adjusted_two_sided_p_value=0.04,
            tie_count=0,
            p_ci_low=0.01,
            p_ci_high=0.05,
            confidence_interval_straddles_alpha=True,
            significant=True,
            node_indices=(0,),
            sensor_indices=(0,),
            harmonic_indices=(0,),
            effect_size=0.8,
            effect_size_kind="pooled_cohen_d",
        ),
        ClusterRecord(
            cluster_id=-1,
            sign="negative",
            mass=-3.0,
            p_value=0.2,
            conservative_p_value=0.21,
            adjusted_two_sided_p_value=0.4,
            tie_count=1,
            p_ci_low=0.1,
            p_ci_high=0.3,
            confidence_interval_straddles_alpha=False,
            significant=False,
            node_indices=(1,),
            sensor_indices=(0,),
            harmonic_indices=(1,),
            effect_size=-0.5,
            effect_size_kind="pooled_cohen_d",
        ),
    )
    result = ClusterPermutationResult(
        design=AnalysisDesign.INDEPENDENT_GROUPS,
        observed_t=np.array([[2.0, -3.0], [0.5, -0.2]]),
        cluster_labels=np.array([[1, -1], [0, 0]]),
        clusters=clusters,
        null_positive_max_mass=np.array([0.0, 1.0, 2.0]),
        null_negative_min_mass=np.array([0.0, -1.0, -2.0]),
        permutations_evaluated=3,
        degrees_of_freedom=2,
        cluster_forming_threshold=2.0,
        cluster_entry_alpha=0.01,
        cluster_alpha_per_tail=0.025,
        sensor_adjacency_version="test-adjacency-v1",
        sensor_adjacency_fingerprint="a" * 64,
        sensor_adjacency_edges=(("Fp1", "Fp2"),),
        rng_algorithm="PCG64",
        seed=7,
        permutation_assignment_hash="b" * 64,
        warnings=("synthetic validation",),
        timing_seconds=(("total", 0.5),),
    )
    return prepared, result


def test_export_publishes_complete_hashed_project_relative_bundle(tmp_path: Path) -> None:
    prepared, result = _prepared_and_result(tmp_path)

    receipt = exports.export_free_harmonic_run(prepared, result, run_id="run-001")

    expected = {
        "arrays.npz",
        "cluster_membership.csv",
        "cluster_summary.csv",
        "harmonic_selection.csv",
        "Free_Harmonic_Clustering_Results.xlsx",
        "manifest.json",
        "node_statistics.csv",
        "null_extrema.csv",
        "participants.csv",
        "source_workbooks.csv",
    }
    assert {path.name for path in receipt.output_directory.iterdir()} == expected
    assert receipt.output_directory == (
        prepared.project_root / "3 - Statistical Analysis Results" / "Free Harmonic Clustering Analysis" / "run-001"
    )
    manifest_text = receipt.manifest_path.read_text(encoding="utf-8")
    manifest = json.loads(manifest_text)
    assert "NaN" not in manifest_text
    assert str(prepared.project_root) not in manifest_text
    assert manifest["claims"]["familywise_error_control"] == "weak-FWER"
    assert manifest["tool"]["validation_label"] == ("paper-faithful-not-author-validated")
    assert manifest["claims"]["pointwise_sensor_harmonic_significance"] is False
    assert manifest["software"]["fpvs_toolbox_version"]
    assert manifest["references"]["article_doi_url"] == ("https://doi.org/10.1111/psyp.70361")
    assert manifest["normalization"]["estimand"] == ("relative sensor/harmonic response distribution")
    assert manifest["frequency_plan"]["target_selected_indices"] == [1, 2]
    assert manifest["frequency_plan"]["noise_selected_indices"] == [[0, 3], [0, 4]]
    assert manifest["frequency_plan"]["target_frequency_errors_hz"] == [0.0, 0.0]
    assert manifest["schema_version"] == 2
    assert len(manifest["artifacts"]) == 9
    assert manifest["harmonic_selection"]["mode"] == "automatic"
    assert manifest["preparation"]["neutral_full_fft_provenance"] == {
        "method_version": "neutral-full-fft-v1",
        "source_fingerprint": "full-fft-source-sha",
        "cohort_fingerprint": "full-fft-cohort-sha",
        "frequency_qc_fingerprint": "full-fft-qc-sha",
        "processing_export_fingerprint": "processing-export-sha",
    }
    assert manifest["preparation"]["cohort_filters"][
        "participant_condition_exclusions"
    ] == [
        {
            "participant_id": "P4",
            "condition": "Neutral Angry",
            "reason": "Project participant-condition exclusion",
        }
    ]
    for artifact in manifest["artifacts"]:
        path = prepared.project_root / artifact["path"]
        assert artifact["sha256"] == _sha256(path)
        assert artifact["size_bytes"] == path.stat().st_size
    assert {artifact.role for artifact in receipt.artifacts} == {
        *(artifact["role"] for artifact in manifest["artifacts"]),
        "manifest",
    }

    with (receipt.output_directory / "source_workbooks.csv").open(encoding="utf-8", newline="") as stream:
        source_rows = list(csv.DictReader(stream))
    assert len(source_rows) == 4
    assert all(not Path(row["project_relative_path"]).is_absolute() for row in source_rows)
    assert all(".." not in Path(row["project_relative_path"]).parts for row in source_rows)
    assert all(int(row["source_size_bytes"]) > 0 for row in source_rows)
    assert all(int(row["source_mtime_ns"]) > 0 for row in source_rows)
    with (receipt.output_directory / "cluster_summary.csv").open(encoding="utf-8", newline="") as stream:
        cluster_rows = list(csv.DictReader(stream))
    assert cluster_rows[0]["descriptive_effect_label"] == (
        "post-selection/shape-dependent normalized cluster-node mean"
    )
    assert cluster_rows[0]["effect_denominator_kind"] == "pooled_within_arm_sd"
    assert cluster_rows[0]["arm_a_minus_b_raw_difference"]
    assert cluster_rows[0]["effect_value_scale"].startswith("participant-arm L2")
    assert cluster_rows[0]["n_a"] == "2"
    assert cluster_rows[0]["n_b"] == "2"
    with (receipt.output_directory / "participants.csv").open(
        encoding="utf-8",
        newline="",
    ) as stream:
        participant_rows = list(csv.DictReader(stream))
    assert any(
        row["status"] == "Excluded"
        and row["participant_id"] == "P4"
        and row["condition"] == "Neutral Angry"
        and row["exclusion_reason"]
        == "Project participant-condition exclusion"
        for row in participant_rows
    )
    with np.load(receipt.output_directory / "arrays.npz", allow_pickle=False) as arrays:
        assert arrays["observed_t"].shape == (2, 2)
        assert arrays["normalized_values_a"].shape == (2, 2, 2)
        assert arrays["sensor_names"].tolist() == ["Fp1", "Fp2"]
        assert arrays["selected_frequency_columns"].tolist() == [
            "0.0000_Hz",
            "1.2000_Hz",
            "2.4000_Hz",
            "3.6000_Hz",
            "4.8000_Hz",
        ]
        assert arrays["target_selected_indices"].tolist() == [1, 2]
        assert arrays["noise_selected_indices"].tolist() == [[0, 3], [0, 4]]
        assert arrays["target_frequency_errors_hz"].tolist() == [0.0, 0.0]
        assert arrays["harmonic_selection_mode"].item() == "automatic"

    workbook_path = receipt.output_directory / exports.HUMAN_WORKBOOK_FILENAME
    workbook = load_workbook(workbook_path, data_only=False)
    assert tuple(workbook.sheetnames) == exports.HUMAN_WORKBOOK_SHEETS
    for sheet_name in exports.HUMAN_WORKBOOK_SHEETS[1:]:
        sheet = workbook[sheet_name]
        assert sheet.freeze_panes == "A5"
        assert sheet.auto_filter.ref is not None
        assert sheet.sheet_view.showGridLines is False
    significant_sheet = workbook["Significant Clusters"]
    all_sheet = workbook["All Clusters"]
    assert significant_sheet["A5"].value == 1
    assert all_sheet["A5"].value == 1
    assert all_sheet["A6"].value == -1
    assert significant_sheet["E5"].number_format == "0.0000"
    assert significant_sheet["A5"].fill.fgColor.rgb.endswith("E2F0D9")
    assert all(
        cell.data_type != "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    participants_sheet = workbook["Participants and Exclusions"]
    assert any(
        participants_sheet.cell(row=row, column=1).value == "Excluded"
        and participants_sheet.cell(row=row, column=5).value == "Neutral Angry"
        and participants_sheet.cell(row=row, column=7).value == "P4"
        for row in range(5, participants_sheet.max_row + 1)
    )


def test_export_rejects_destination_outside_project(tmp_path: Path) -> None:
    prepared, result = _prepared_and_result(tmp_path)
    outside = tmp_path / "outside" / "run-001"

    with pytest.raises(FreeHarmonicInputError, match="managed project root"):
        exports.export_free_harmonic_run(
            prepared,
            result,
            destination=outside,
        )

    assert not outside.exists()


def test_explicit_destination_is_the_exact_final_run_directory(tmp_path: Path) -> None:
    prepared, result = _prepared_and_result(tmp_path)

    receipt = exports.export_free_harmonic_run(
        prepared,
        result,
        run_id="exact-run",
        destination=Path("custom-results") / "exact-run",
    )

    assert receipt.output_directory == prepared.project_root / "custom-results" / "exact-run"
    assert receipt.manifest_path == receipt.output_directory / "manifest.json"
    assert not (receipt.output_directory / "exact-run").exists()


def test_destination_name_must_match_explicit_run_id(tmp_path: Path) -> None:
    prepared, result = _prepared_and_result(tmp_path)

    with pytest.raises(ValueError, match="must equal run_id"):
        exports.export_free_harmonic_run(
            prepared,
            result,
            run_id="expected-name",
            destination=Path("custom-results") / "different-name",
        )


def test_export_never_overwrites_existing_run(tmp_path: Path) -> None:
    prepared, result = _prepared_and_result(tmp_path)
    destination = (
        prepared.project_root / "3 - Statistical Analysis Results" / "Free Harmonic Clustering Analysis" / "same-run"
    )
    destination.mkdir(parents=True)
    marker = destination / "keep.txt"
    marker.write_text("keep", encoding="utf-8")

    with pytest.raises(FileExistsError, match="already exists"):
        exports.export_free_harmonic_run(prepared, result, run_id="same-run")

    assert marker.read_text(encoding="utf-8") == "keep"
    assert list(destination.iterdir()) == [marker]


def test_export_cleans_staging_when_manifest_write_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    prepared, result = _prepared_and_result(tmp_path)

    def fail_manifest(path: Path, payload: object) -> None:
        raise OSError("synthetic manifest failure")

    monkeypatch.setattr(exports, "_write_manifest", fail_manifest)
    with pytest.raises(OSError, match="synthetic manifest failure"):
        exports.export_free_harmonic_run(prepared, result, run_id="failed-run")

    parent = prepared.project_root / exports.DEFAULT_RESULTS_SUBFOLDER
    assert not (parent / "failed-run").exists()
    assert not list(parent.glob(".failed-run.staging-*"))


def test_export_rejects_non_relative_source_provenance(tmp_path: Path) -> None:
    prepared, result = _prepared_and_result(tmp_path)
    first = prepared.source_workbooks[0]
    bad = replace(
        first,
        source_path=tmp_path / "outside.xlsx",
        project_relative_path="../outside.xlsx",
    )
    prepared = replace(
        prepared,
        source_workbooks=(bad, *prepared.source_workbooks[1:]),
    )

    with pytest.raises(FreeHarmonicInputError, match="Source workbook"):
        exports.export_free_harmonic_run(prepared, result, run_id="bad-source")


def test_export_rejects_cluster_membership_that_disagrees_with_labels(
    tmp_path: Path,
) -> None:
    prepared, result = _prepared_and_result(tmp_path)
    result = replace(result, cluster_labels=np.array([[1, 0], [0, 0]]))

    with pytest.raises(ValueError, match="different IDs"):
        exports.export_free_harmonic_run(prepared, result, run_id="bad-labels")


def test_human_workbook_forces_project_controlled_formula_like_text_to_string(
    tmp_path: Path,
) -> None:
    prepared, result = _prepared_and_result(tmp_path)
    prepared = replace(
        prepared,
        arm_a_label="=2+3",
        arm_b_label="@SUM(A1:A2)",
    )

    receipt = exports.export_free_harmonic_run(
        prepared,
        result,
        run_id="safe-text",
    )

    workbook = load_workbook(
        receipt.output_directory / exports.HUMAN_WORKBOOK_FILENAME,
        data_only=False,
    )
    dangerous = [
        cell
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.startswith(("=", "+", "-", "@"))
    ]
    assert dangerous
    assert all(cell.data_type == "s" for cell in dangerous)
