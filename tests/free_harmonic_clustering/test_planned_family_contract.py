"""Independent scientific-boundary regressions for the planned FHC workflow."""

from __future__ import annotations

from collections import Counter
import csv
from dataclasses import replace
import json
from pathlib import Path

import numpy as np
import pytest

from Tools.Free_Harmonic_Clustering.analysis_plan import (
    AnalysisFamily,
    build_analysis_plan,
)
from Tools.Free_Harmonic_Clustering.models import FreeHarmonicInputError, FreeHarmonicMethodSpec
from Tools.Free_Harmonic_Clustering.planned_inputs import prepare_analysis_plan
from Tools.Free_Harmonic_Clustering.preparation import l2_normalize_snr

from tests.free_harmonic_clustering.test_inputs import (
    _write_project,
    _write_repeated_project,
    _install_frequency_exclusions,
    _install_reader_doubles,
)


def _repeated_root(tmp_path: Path, **kwargs: object) -> Path:
    return _write_repeated_project(
        tmp_path / "Repeated",
        participants_by_group={
            "bc_group": ("BC1", "BC2", "BC3"),
            "control_group": ("C1", "C2", "C3"),
        },
        conditions=("Angry", "Happy", "Sad", "Neutral"),
        **kwargs,
    )


def test_default_repeated_plan_pools_both_groups_visit_tests(tmp_path: Path) -> None:
    root = _repeated_root(tmp_path)
    plan = build_analysis_plan(
        root,
        group_ids=("bc_group", "control_group"),
        conditions=("Angry", "Happy", "Sad", "Neutral"),
        session_ids=("luteal_phase", "follicular_phase"),
    )

    assert Counter(row.family_id for row in plan.comparisons) == {
        "between_groups": 4,
        "within_group_visits": 8,
        "group_visit_change": 4,
    }
    visits = [row for row in plan.comparisons if row.kind is AnalysisFamily.WITHIN_GROUP_VISITS]
    assert len({row.comparison_id for row in visits}) == 8
    assert all(row.session_ids == ("luteal_phase", "follicular_phase") for row in visits)
    assert {row.group_ids for row in visits if row.condition == "Angry"} == {
        ("bc_group",), ("control_group",),
    }


def test_single_group_repeated_plan_does_not_require_a_second_group(tmp_path: Path) -> None:
    root = _repeated_root(tmp_path)
    plan = build_analysis_plan(
        root,
        group_ids=("bc_group",),
        conditions=("Happy", "Sad"),
        session_ids=("luteal_phase", "follicular_phase"),
    )

    assert len(plan.comparisons) == 2
    assert {row.family_id for row in plan.comparisons} == {"within_group_visits"}
    assert all(row.group_ids == ("bc_group",) for row in plan.comparisons)


def test_three_group_plan_counts_pairwise_comparisons_in_one_family(tmp_path: Path) -> None:
    groups = {key: (f"Group {key.upper()}", key) for key in ("a", "b", "c")}
    participants = {f"{group}{number}": group for group in groups for number in (1, 2)}
    root, _paths = _write_project(
        tmp_path / "Flat",
        groups=groups,
        participant_groups=participants,
        participant_conditions={participant: ("Faces", "Objects") for participant in participants},
    )
    plan = build_analysis_plan(root, group_ids=tuple(groups), conditions=("Faces", "Objects"))

    assert len(plan.comparisons) == 6
    assert {row.family_id for row in plan.comparisons} == {"between_groups"}
    assert Counter(row.group_ids for row in plan.comparisons) == {
        ("a", "b"): 2, ("a", "c"): 2, ("b", "c"): 2,
    }
    assert all("Group " in row.label for row in plan.comparisons)


def test_reversed_condition_pair_cannot_double_enter_holm_family(tmp_path: Path) -> None:
    root = _repeated_root(tmp_path)
    with pytest.raises(FreeHarmonicInputError, match="reversed duplicates"):
        build_analysis_plan(
            root,
            group_ids=("bc_group",),
            conditions=("Happy", "Sad"),
            session_ids=("luteal_phase", "follicular_phase"),
            families=(AnalysisFamily.BETWEEN_CONDITIONS,),
            condition_pairs=(("Happy", "Sad"), ("Sad", "Happy")),
        )


def test_unrelated_selected_condition_does_not_enter_reference_pair_plan(tmp_path: Path) -> None:
    root = _repeated_root(tmp_path)
    plan = build_analysis_plan(
        root,
        group_ids=("bc_group",),
        conditions=("Angry", "Happy", "Sad"),
        session_ids=("luteal_phase", "follicular_phase"),
        families=(AnalysisFamily.BETWEEN_CONDITIONS,),
        condition_pairs=(("Happy", "Sad"),),
    )

    assert plan.conditions == ("Happy", "Sad")
    assert len(plan.comparisons) == 1
    assert plan.comparisons[0].condition_a == "Happy"
    assert plan.comparisons[0].condition_b == "Sad"


def test_condition_exclusion_stays_local_with_one_shared_source_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = _repeated_root(
        tmp_path,
        recording_condition_exclusions={"BC3__follicular_phase": ("Happy",)},
    )
    _install_frequency_exclusions(monkeypatch)
    headers, amplitudes = _install_reader_doubles(monkeypatch)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"),
        conditions=("Angry", "Happy"),
        session_ids=("luteal_phase", "follicular_phase"),
        families=(AnalysisFamily.BETWEEN_GROUPS, AnalysisFamily.BETWEEN_CONDITIONS),
    )

    prepared = prepare_analysis_plan(plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6))
    rows = dict(zip(plan.comparisons, prepared.contrasts, strict=True))
    angry = next(value for comparison, value in rows.items()
                 if comparison.kind is AnalysisFamily.BETWEEN_GROUPS and comparison.condition == "Angry")
    happy = next(value for comparison, value in rows.items()
                 if comparison.kind is AnalysisFamily.BETWEEN_GROUPS and comparison.condition == "Happy")
    paired = next(value for comparison, value in rows.items()
                  if comparison.kind is AnalysisFamily.BETWEEN_CONDITIONS and comparison.group_ids == ("bc_group",))
    assert angry.participant_ids_a == ("BC1", "BC2", "BC3")
    assert happy.participant_ids_a == ("BC1", "BC2")
    assert paired.participant_ids_a == paired.participant_ids_b == ("BC1", "BC2")
    assert len(amplitudes) == len(set(amplitudes)) == len(headers) == 22
    assert all("BC3__follicular_phase_Happy" not in path.name for path in amplitudes)
    cell_counts = dict(zip(
        zip(prepared.shared_selection_audit.cell_group_ids,
            prepared.shared_selection_audit.cell_conditions, strict=True),
        prepared.shared_selection_audit.cell_participant_counts, strict=True,
    ))
    assert cell_counts[("bc_group", "Angry")] == 3
    assert cell_counts[("bc_group", "Happy")] == 2


def test_average_and_delta_keep_distinct_normalization_recipes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = _repeated_root(tmp_path)
    _install_frequency_exclusions(monkeypatch)

    def varied_profile(path: Path, frame: object) -> None:
        sensors = np.arange(64, dtype=np.float64)
        is_first = "luteal_phase" in path.name
        frame["1.200000_Hz"] = 20.0 + sensors * (0.2 if is_first else 0.8)
        frame["2.400000_Hz"] = 12.0 + sensors[::-1] * (0.4 if is_first else 0.1)

    _install_reader_doubles(monkeypatch, mutate_frame=varied_profile)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Happy",),
        session_ids=("luteal_phase", "follicular_phase"),
    )
    prepared = prepare_analysis_plan(plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6))
    rows = dict(zip(plan.comparisons, prepared.contrasts, strict=True))
    pooled = next(value for comparison, value in rows.items() if comparison.kind is AnalysisFamily.BETWEEN_GROUPS)
    paired = next(value for comparison, value in rows.items()
                  if comparison.kind is AnalysisFamily.WITHIN_GROUP_VISITS and comparison.group_ids == ("bc_group",))
    delta = next(value for comparison, value in rows.items() if comparison.kind is AnalysisFamily.GROUP_VISIT_CHANGE)

    np.testing.assert_allclose(pooled.snr_a, (paired.snr_a + paired.snr_b) / 2.0)
    np.testing.assert_allclose(pooled.values_a, l2_normalize_snr(pooled.snr_a))
    np.testing.assert_allclose(delta.values_a, paired.values_a - paired.values_b)
    assert not np.allclose(pooled.values_a, (paired.values_a + paired.values_b) / 2.0)
    assert not np.allclose(np.linalg.norm(delta.values_a, axis=(1, 2)), 1.0)
    assert paired.arm_a_label == "Luteal Phase"
    assert paired.arm_b_label == "Follicular Phase"


def test_bad_electrode_on_cell_outside_complete_condition_pairs_does_not_block(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = _repeated_root(
        tmp_path,
        recording_condition_exclusions={"BC3__follicular_phase": ("Happy",)},
    )
    _install_frequency_exclusions(
        monkeypatch,
        recording_condition_electrodes={("BC3__luteal_phase", "Angry"): frozenset({"P10"})},
    )
    _headers, amplitudes = _install_reader_doubles(monkeypatch)
    plan = build_analysis_plan(
        root, group_ids=("bc_group",), conditions=("Angry", "Happy"),
        session_ids=("luteal_phase", "follicular_phase"),
        families=(AnalysisFamily.BETWEEN_CONDITIONS,),
    )

    prepared = prepare_analysis_plan(plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6))
    assert prepared.contrasts[0].participant_ids_a == ("BC1", "BC2")
    assert len(amplitudes) == 8
    assert all("BC3" not in path.name for path in amplitudes)


def test_insufficient_planned_cell_blocks_all_spectral_reads(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = _repeated_root(
        tmp_path,
        recording_condition_exclusions={
            "BC2__follicular_phase": ("Happy",),
            "BC3__follicular_phase": ("Happy",),
        },
    )
    _install_frequency_exclusions(monkeypatch)
    headers, amplitudes = _install_reader_doubles(monkeypatch)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Angry", "Happy"),
        session_ids=("luteal_phase", "follicular_phase"),
    )

    with pytest.raises(FreeHarmonicInputError, match="at least two"):
        prepare_analysis_plan(plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6))
    assert headers == amplitudes == []


def test_tampered_plan_cannot_drop_a_test_from_its_declared_family(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = _repeated_root(tmp_path)
    _install_frequency_exclusions(monkeypatch)
    headers, amplitudes = _install_reader_doubles(monkeypatch)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Happy",),
        session_ids=("luteal_phase", "follicular_phase"),
    )

    with pytest.raises((FreeHarmonicInputError, ValueError)):
        malformed = replace(plan, comparisons=plan.comparisons[:-1])
        prepare_analysis_plan(malformed, FreeHarmonicMethodSpec(max_harmonic_hz=3.6))
    assert headers == amplitudes == []


def test_prepared_tensor_order_cannot_be_reassigned_to_other_comparison_labels(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = _repeated_root(tmp_path)
    _install_frequency_exclusions(monkeypatch)
    _install_reader_doubles(monkeypatch)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Happy",),
        session_ids=("luteal_phase", "follicular_phase"),
    )
    prepared = prepare_analysis_plan(plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6))

    with pytest.raises(ValueError, match="comparison|request|order"):
        replace(prepared, contrasts=tuple(reversed(prepared.contrasts)))


def test_comparison_seed_is_unchanged_when_other_families_or_conditions_are_added(tmp_path: Path) -> None:
    from Tools.Free_Harmonic_Clustering.planned_analysis import derive_planned_comparison_seed

    root = _repeated_root(tmp_path)
    small = build_analysis_plan(
        root, group_ids=("bc_group",), conditions=("Happy",),
        session_ids=("luteal_phase", "follicular_phase"),
        families=(AnalysisFamily.WITHIN_GROUP_VISITS,),
    )
    full = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Angry", "Happy", "Sad"),
        session_ids=("luteal_phase", "follicular_phase"),
    )
    comparison = next(row for row in full.comparisons if row.comparison_id == small.comparisons[0].comparison_id)
    assert derive_planned_comparison_seed(2026, comparison.comparison_id) == derive_planned_comparison_seed(
        2026, small.comparisons[0].comparison_id,
    )
    regrouped = replace(comparison, family_id="review_only_label")
    assert derive_planned_comparison_seed(2026, regrouped.comparison_id) == derive_planned_comparison_seed(
        2026, comparison.comparison_id,
    )


def test_pooled_visit_holm_uses_both_group_tests_and_original_global_p(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from Tools.Free_Harmonic_Clustering import planned_analysis
    from tests.free_harmonic_clustering.test_exports import _prepared_and_result

    root = _repeated_root(tmp_path)
    _install_frequency_exclusions(monkeypatch)
    _install_reader_doubles(monkeypatch)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Happy",),
        session_ids=("luteal_phase", "follicular_phase"),
    )
    prepared = prepare_analysis_plan(plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6, n_permutations=3))
    _template_prepared, template = _prepared_and_result(tmp_path)
    raw_p = iter((0.01, 0.04, 0.02, 0.50))

    def completed_comparison(contrast: object, **_kwargs: object) -> object:
        p = next(raw_p)
        cluster = replace(template.clusters[0], p_value=p / 2,
                          adjusted_two_sided_p_value=p, significant=p < 0.05)
        labels = np.zeros(contrast.values_a.shape[1:], dtype=np.int64)
        labels[0, 0] = 1
        return replace(template, design=contrast.request.design, seed=contrast.method.seed,
                       observed_t=np.where(labels, 4.0, 0.0), cluster_labels=labels,
                       clusters=(cluster,))

    monkeypatch.setattr(planned_analysis, "analyze_prepared_contrast", completed_comparison)
    result = planned_analysis.analyze_prepared_analysis_plan(prepared)

    assert [row.global_p for row in result.outcomes] == pytest.approx([0.01, 0.04, 0.02, 0.50])
    assert [row.family_adjusted_p for row in result.outcomes] == pytest.approx([0.01, 0.04, 0.04, 0.50])
    assert [row.batch_adjusted_p for row in result.outcomes] == pytest.approx([0.04, 0.08, 0.06, 0.50])
    assert [row.result.clusters[0].p_value for row in result.outcomes] == pytest.approx([0.005, 0.02, 0.01, 0.25])
    with pytest.raises(ValueError, match="exactly one outcome"):
        replace(result, outcomes=result.outcomes[:-1])


def test_real_planned_pipeline_exports_frozen_sources_tensors_and_families(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from openpyxl import load_workbook

    from Tools.Free_Harmonic_Clustering.planned_analysis import run_analysis_plan
    from Tools.Free_Harmonic_Clustering.planned_exports import (
        PLANNED_WORKBOOK_FILENAME, export_analysis_plan_result,
    )
    from tests.free_harmonic_clustering.test_exports import _stub_map_figures

    root = _repeated_root(
        tmp_path,
        recording_condition_exclusions={"BC3__follicular_phase": ("Happy",)},
    )
    _install_frequency_exclusions(monkeypatch)
    _stub_map_figures(monkeypatch)

    def varied_profile(path: Path, frame: object) -> None:
        sensors = np.arange(64, dtype=np.float64)
        participant = path.name.split("__", 1)[0]
        number = int("".join(value for value in participant if value.isdigit()))
        visit = 1.0 if "luteal_phase" in path.name else -0.3
        condition = 0.4 if "Happy" in path.name else -0.2
        group = 0.3 if participant.startswith("BC") else -0.2
        frame["1.200000_Hz"] = 20.0 + sensors * (0.3 + number * 0.02 + (visit + condition + group) * 0.04)
        frame["2.400000_Hz"] = 15.0 + np.sin(sensors / (number + 2)) + sensors[::-1] * (0.15 + visit * 0.03)

    headers, amplitudes = _install_reader_doubles(monkeypatch, mutate_frame=varied_profile)
    plan = build_analysis_plan(
        root, group_ids=("bc_group", "control_group"), conditions=("Angry", "Happy"),
        session_ids=("luteal_phase", "follicular_phase"),
        families=tuple(AnalysisFamily),
    )
    manifest_before = (root / "project.json").read_bytes()
    progress = []
    result = run_analysis_plan(
        plan, FreeHarmonicMethodSpec(max_harmonic_hz=3.6, n_permutations=7, seed=83),
        progress_callback=lambda stage, done, total: progress.append((stage, done, total)),
    )
    receipt = export_analysis_plan_result(result, run_id="planned-contract-integration")

    assert len(result.outcomes) == 10
    assert len(headers) == len(amplitudes) == len(set(amplitudes)) == 22
    assert ("inference", 70, 70) in progress
    assert (root / "project.json").read_bytes() == manifest_before
    assert all(path.read_bytes() == b"selected-reader-test-double" for path in amplitudes)
    manifest = json.loads(receipt.manifest_path.read_text(encoding="utf-8"))
    assert manifest["plan_fingerprint"] == plan.fingerprint
    assert manifest["plan"]["version"] == plan.version
    assert manifest["preparation"]["full_fft_source_fingerprint"] == "source-fingerprint"
    assert manifest["preparation"]["shared_domain_fingerprint"] == result.prepared.shared_domain_fingerprint
    assert manifest["adjacency"]["fingerprint"] == result.outcomes[0].result.sensor_adjacency_fingerprint
    assert manifest["selected_harmonic_orders"] == result.prepared.shared_selection.selected_orders.tolist()
    assert Counter(row["family_id"] for row in manifest["results"]) == {
        "between_groups": 2, "between_conditions": 2,
        "within_group_visits": 4, "group_visit_change": 2,
    }

    for record, outcome in zip(manifest["comparisons"], result.outcomes, strict=True):
        assert record["comparison_id"] == outcome.comparison.comparison_id
        assert record["seed"] == outcome.derived_seed
        assert record["permutations_evaluated"] == 7
        with np.load(receipt.output_directory / record["directory"] / "arrays.npz", allow_pickle=False) as arrays:
            np.testing.assert_array_equal(arrays["analyzed_values_a"], outcome.contrast.values_a)
            np.testing.assert_array_equal(arrays["analyzed_values_b"], outcome.contrast.values_b)
            np.testing.assert_array_equal(arrays["null_positive_max_mass"], outcome.result.null_positive_max_mass)
            metadata = json.loads(str(arrays["metadata_json"]))
            assert metadata["comparison_id"] == record["comparison_id"]
            if record["family_id"] == "group_visit_change":
                assert "no renormalization" in metadata["tensor_semantics"]
            elif record["family_id"] in {"between_groups", "between_conditions"}:
                assert "averaged across both visits" in metadata["tensor_semantics"]
    with (receipt.output_directory / "harmonic_selection.csv").open(encoding="utf-8", newline="") as stream:
        selector = list(csv.DictReader(stream))
    assert {int(row["participants"]) for row in selector
            if row["group_id"] == "bc_group" and row["condition"] == "Happy"} == {2}
    workbook = load_workbook(receipt.output_directory / PLANNED_WORKBOOK_FILENAME, read_only=True)
    try:
        assert {"Comparisons", "Families", "Participants", "Cohort Audit", "Harmonic Selection"}.issubset(workbook.sheetnames)
    finally:
        workbook.close()
    assert not list(receipt.output_directory.parent.glob(".planned-contract-integration.staging-*"))
