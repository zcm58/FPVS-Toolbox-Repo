"""Reporting-only repeated FHC bundles from stored synthetic inference."""

from __future__ import annotations

import csv
from contextlib import closing
from dataclasses import replace
import hashlib
import json

import numpy as np
from openpyxl import load_workbook
import pytest

from Tools.Free_Harmonic_Clustering import analysis, api, exports
from Tools.Free_Harmonic_Clustering.models import AnalysisDesign
from Tools.Free_Harmonic_Clustering.reporting import (
    EXPLORATORY_CRITERION,
    build_repeated_session_report,
)
from tests.free_harmonic_clustering.test_exports import _stub_map_figures
from tests.free_harmonic_clustering.test_repeated_session_batch_inference import _batch_fixture


@pytest.fixture(autouse=True)
def _map_figures(monkeypatch):
    _stub_map_figures(monkeypatch)


def _two_condition_fixture(tmp_path):
    prepared, original = _batch_fixture(tmp_path)
    conditions = ("Neutral Angry", "Neutral Happy")
    # The first family qualifies at both conditions. The second family has
    # family Holm exactly .05; the third has unadjusted global exactly .05.
    p_values = ((.03, .025, .05, .6), (.04, .8, .05, .9))
    runs = []
    results = []
    for condition_index, condition in enumerate(conditions):
        for family_index, outcome in enumerate(original.outcomes):
            prior = outcome.prepared_run
            request = replace(
                prior.prepared.request,
                condition_a=condition,
                condition_b=condition if prior.prepared.request.design is AnalysisDesign.PAIRED_CONDITIONS else None,
            )
            run = replace(prior, condition=condition, prepared=replace(prior.prepared, request=request))
            runs.append(run)
            raw = p_values[condition_index][family_index]
            positive = replace(
                outcome.result.clusters[0], p_value=raw / 2,
                adjusted_two_sided_p_value=raw, significant=raw <= .05,
            )
            # Signed-tail p < .05 alone is insufficient: two-sided p = .08.
            negative = replace(
                positive, cluster_id=-1, sign="negative", mass=-3.,
                p_value=.04, adjusted_two_sided_p_value=.08, significant=False,
                node_indices=(1,), sensor_indices=(0,), harmonic_indices=(1,),
                effect_size=-.5,
            )
            seed = analysis.derive_repeated_session_run_seed(
                original.base_seed, family_id=run.family_id, condition=condition,
            )
            results.append(replace(
                outcome.result, clusters=(positive, negative),
                cluster_labels=np.array([[1, -1], [0, 0]]), seed=seed,
            ))
    prepared = replace(
        prepared, request=replace(prepared.request, conditions=conditions),
        conditions=conditions, contrast_runs=tuple(runs),
    )
    adjustments = analysis.adjust_batch_cluster_p_values(tuple(
        (run.family_id, run.condition, result)
        for run, result in zip(runs, results, strict=True)
    ))
    outcomes = tuple(api.RepeatedSessionContrastOutcome(
        prepared_run=run, result=result, derived_seed=result.seed,
        global_two_sided_p_value=adjusted.global_two_sided_p_value,
        holm_within_family_p_value=adjusted.holm_within_family_p_value,
        holm_all_batch_p_value=adjusted.holm_all_batch_p_value,
    ) for run, result, adjusted in zip(runs, results, adjustments, strict=True))
    return prepared, api.RepeatedSessionBatchResult(outcomes=outcomes, base_seed=original.base_seed)


def _csv_rows(path):
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def _serialized_rows(rows):
    return [{key: "" if value is None else str(value) for key, value in row.items()} for row in rows]


def test_exploratory_bundle_preserves_primary_results_and_exact_membership(tmp_path):
    prepared, result = _two_condition_fixture(tmp_path)
    summaries = exports._batch_summary_rows(prepared, result)
    clusters = exports._batch_cluster_rows(result)
    membership = exports._batch_membership_rows(result)
    before_manifest = (prepared.project_root / "project.json").read_bytes()
    tensors_before = tuple(
        (outcome.prepared_run.prepared.values_a.tobytes(), outcome.result.observed_t.tobytes(),
         outcome.result.null_positive_max_mass.tobytes(), outcome.result.clusters)
        for outcome in result.outcomes
    )

    receipt = exports.export_repeated_session_batch(prepared, result, run_id="exploratory")

    assert (prepared.project_root / "project.json").read_bytes() == before_manifest
    assert tensors_before == tuple(
        (outcome.prepared_run.prepared.values_a.tobytes(), outcome.result.observed_t.tobytes(),
         outcome.result.null_positive_max_mass.tobytes(), outcome.result.clusters)
        for outcome in result.outcomes
    )
    directory = receipt.output_directory
    manifest = json.loads(receipt.manifest_path.read_text(encoding="utf-8"))
    assert manifest["results"] == summaries
    assert _csv_rows(directory / "batch_summary.csv") == _serialized_rows(summaries)
    assert _csv_rows(directory / "all_clusters.csv") == _serialized_rows(clusters)
    assert _csv_rows(directory / "cluster_membership.csv") == _serialized_rows(membership)
    expected_keys = {
        ("session_averaged_groups", "Neutral Angry", "1"),
        ("session_averaged_groups", "Neutral Happy", "1"),
    }
    candidates = _csv_rows(directory / "exploratory_clusters.csv")
    exact_members = _csv_rows(directory / "exploratory_cluster_membership.csv")
    def key(row):
        return row["family_id"], row["condition"], row["cluster_id"]
    assert {key(row) for row in candidates} == expected_keys
    assert candidates == [row for row in _serialized_rows(clusters) if key(row) in expected_keys]
    assert exact_members == [row for row in _serialized_rows(membership) if key(row) in expected_keys]
    assert all(row["harmonic_order"] == "1" for row in exact_members)
    assert not any(row["cluster_id"] == "-1" for row in candidates)
    assert result.outcomes[1].holm_within_family_p_value == .05
    assert result.outcomes[2].global_two_sided_p_value == .05
    assert manifest["exploratory_reporting"]["criterion"] == EXPLORATORY_CRITERION
    assert manifest["exploratory_reporting"]["qualifying_run_count"] == 2
    assert manifest["exploratory_reporting"]["qualifying_cluster_count"] == 2
    assert manifest["exploratory_reporting"]["changes_inference"] is False
    assert manifest["exploratory_reporting"]["pointwise_significance_claimed"] is False
    report = build_repeated_session_report(result)
    narrative = (directory / "exploratory_findings.md").read_text(encoding="utf-8")
    assert "2 of 8" in narrative
    assert EXPLORATORY_CRITERION in narrative
    assert "not confirmed findings" in narrative
    assert "not established as pointwise significant" in narrative
    for row in report.rows:
        if row.is_exploratory:
            assert row.detail_text in narrative.replace("  \n", "\n")
            assert row.detail_text.replace("\n", "  \n") in narrative
    with closing(load_workbook(directory / exports.REPEATED_SESSION_WORKBOOK_FILENAME)) as workbook:
        sheet = workbook["Exploratory Findings"]
        notes = "\n".join(str(row[7]) for row in sheet.iter_rows(min_row=5, values_only=True))
        for row in report.rows:
            if row.is_exploratory:
                for line in row.detail_text.splitlines():
                    if line.strip():
                        assert line in notes
        assert sheet["D5"].value == .03
        assert sheet["E5"].value == .06
        assert sheet["D5"].number_format == "0.########"
        assert "All Clusters" in workbook.sheetnames
    new_roles = {"exploratory_findings", "exploratory_clusters", "exploratory_cluster_membership"}
    assert new_roles.issubset({item.role for item in receipt.artifacts})
    for artifact in manifest["artifacts"]:
        artifact_path = prepared.project_root / artifact["path"]
        assert artifact_path.is_relative_to(prepared.project_root)
        assert hashlib.sha256(artifact_path.read_bytes()).hexdigest() == artifact["sha256"]
        assert artifact_path.stat().st_size == artifact["size_bytes"]


def test_zero_exploratory_findings_still_exports_explicit_report(tmp_path):
    prepared, result = _batch_fixture(tmp_path)
    # One condition per family: some runs fail the secondary all-batch layer,
    # but no run fails its primary family layer while having global p < .05.
    assert any(row.holm_all_batch_p_value > .05 and row.holm_within_family_p_value <= .05 for row in result.outcomes)

    receipt = exports.export_repeated_session_batch(prepared, result, run_id="no-exploratory")

    directory = receipt.output_directory
    assert _csv_rows(directory / "exploratory_clusters.csv") == []
    assert _csv_rows(directory / "exploratory_cluster_membership.csv") == []
    assert "No runs met" in (directory / "exploratory_findings.md").read_text(encoding="utf-8")
    with closing(load_workbook(directory / exports.REPEATED_SESSION_WORKBOOK_FILENAME)) as workbook:
        assert "No runs met" in workbook["Exploratory Findings"]["A5"].value


def test_exploratory_report_failure_discards_atomic_bundle(tmp_path, monkeypatch):
    prepared, result = _batch_fixture(tmp_path)
    before = (prepared.project_root / "project.json").read_bytes()

    def fail_report(*_args):
        raise OSError("Exploratory report write failed")

    monkeypatch.setattr(exports, "_write_exploratory_report", fail_report)
    with pytest.raises(OSError, match="Exploratory report write failed"):
        exports.export_repeated_session_batch(prepared, result, run_id="report-failed")

    parent = prepared.project_root / exports.DEFAULT_RESULTS_SUBFOLDER
    assert not (parent / "report-failed").exists()
    assert list(parent.glob(".report-failed.staging-*")) == []
    assert (prepared.project_root / "project.json").read_bytes() == before
