from __future__ import annotations

import json
import math
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

import Main_App.exports.analysis_ready_workbook as export_module
from Main_App.exports.analysis_ready_workbook import (
    ANALYSIS_READY_WORKBOOK_NAME,
    write_analysis_ready_workbook,
)
from Main_App.projects import load_project_dataset_index


def _write_project(root: Path) -> None:
    anxious_raw = root / "Raw" / "Anxious"
    non_anxious_raw = root / "Raw" / "NonAnxious"
    anxious_raw.mkdir(parents=True)
    non_anxious_raw.mkdir(parents=True)
    manifest = {
        "results_folder": ".",
        "subfolders": {"excel": "1 - Excel Data Files"},
        "groups": {
            "anxious": {
                "label": "anxious",
                "folder_name": "Anxious",
                "raw_input_folder": str(anxious_raw),
            },
            "non_anxious": {
                "label": "non-anxious",
                "folder_name": "NonAnxious",
                "raw_input_folder": str(non_anxious_raw),
            },
        },
        "participants": {
            "P1": {"group_id": "anxious"},
            "P2": {"group_id": "non_anxious"},
        },
        "preprocessing": {
            "manual_excluded_participant_conditions": {
                "P2": ["Condition A"],
            }
        },
        "tools": {
            "frequency_domain_qc": {
                "auto_participant_electrode_exclusions": [
                    {
                        "participant_id": "P1",
                        "electrode": "O2",
                        "reason": "test electrode flag",
                        "triggering_conditions": ["Condition A"],
                    }
                ],
                "auto_participant_exclusions": [
                    {
                        "participant_id": "P2",
                        "reason": "test participant flag",
                    }
                ],
            }
        },
    }
    (root / "project.json").write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )


def _write_bca_workbook(
    root: Path,
    *,
    pid: str,
    condition: str,
    group_folder: str,
    values: list[tuple[str, float, float]],
) -> Path:
    folder = root / "1 - Excel Data Files" / condition / group_folder
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{pid}_{condition}_Results.xlsx"
    frame = pd.DataFrame(
        values,
        columns=["Electrode", "1.2000_Hz", "2.4000_Hz"],
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="BCA (uV)", index=False)
    return path


def _selection_metadata() -> dict[str, object]:
    return {
        "harmonic_policy": "group_significant_harmonics",
        "harmonic_selection_profile": "significant_only_exploratory",
        "harmonic_selection_profile_version": "1.0",
        "selection_fingerprint": "c" * 64,
        "final_release_receipt_fingerprint": "test-final-release",
        "selected_harmonics_hz": [1.2, 2.4],
        "included_harmonics_hz": [1.2, 2.4],
        "detected_significant_harmonics_hz": [1.2],
        "selection_electrode_mask": ["O1", "O2"],
        "pooling_method": "balanced_group_condition_then_equal_condition_z",
        "pooling_cells": [
            {
                "group_id": "anxious",
                "condition": "Condition A",
                "participant_count": 1,
                "participant_weight_within_cell": 1.0,
                "group_weight_within_condition": 0.5,
                "condition_weight": 0.5,
                "effective_participant_weight": 0.25,
            }
        ],
        "source_workbook_fingerprints": [
            {
                "subject": "P1",
                "condition": "Condition A",
                "path": "1 - Excel Data Files/Condition A/Anxious/P1_Condition A_Results.xlsx",
                "size_bytes": 123,
                "mtime_ns": 456,
            }
        ],
        "selection_rows": [
            {
                "target_frequency_hz": 1.2,
                "z_score": 3.0,
                "selected": True,
                "included_in_summation": True,
                "excluded_base_rate": False,
            },
            {
                "target_frequency_hz": 2.4,
                "z_score": 2.5,
                "selected": True,
                "included_in_summation": True,
                "excluded_base_rate": False,
            },
        ],
    }


@pytest.fixture(autouse=True)
def _accepted_harmonic_selection(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        export_module,
        "_load_current_processing_selection_metadata",
        lambda _root: _selection_metadata(),
    )


def _released_coverage(
    index,
    *,
    rois: dict[str, list[str]],
    excluded_by_cell: dict[tuple[str, str], set[str]] | None = None,
    excluded_cells: set[tuple[str, str]] | None = None,
):
    excluded_by_cell = excluded_by_cell or {}
    excluded_cells = excluded_cells or set()
    roi_definitions = tuple(
        SimpleNamespace(name=name, electrodes=tuple(electrodes))
        for name, electrodes in rois.items()
    )
    cells = []
    by_key = {}
    for record in [*index.workbooks, *index.excluded_workbooks]:
        identity = str(record.recording_id or record.participant_id)
        key = (identity.casefold(), str(record.condition).casefold())
        frame = pd.read_excel(record.path, sheet_name="BCA (uV)")
        expected = tuple(
            dict.fromkeys(
                frame["Electrode"].astype(str).str.strip().str.upper().tolist()
            )
        )
        excluded = tuple(
            channel
            for channel in expected
            if channel in excluded_by_cell.get(key, set())
        )
        cell_excluded = key in excluded_cells
        memberships = []
        for roi_name, roi_channels in rois.items():
            expected_roi = tuple(roi_channels)
            excluded_roi = tuple(
                channel for channel in expected_roi if channel in excluded
            )
            unavailable = cell_excluded or bool(excluded_roi)
            memberships.append(
                SimpleNamespace(
                    roi_name=roi_name,
                    expected_channels=expected_roi,
                    observed_channels=expected_roi,
                    excluded_channels=excluded_roi,
                    interpolated_channels=(),
                    used_channels=() if unavailable else expected_roi,
                    configured_count=len(expected_roi),
                    interpolated_count=0,
                    all_members_interpolated_warning=False,
                    status="unavailable" if unavailable else "available",
                    reason_codes=(
                        ("reviewed_required_roi_member_exclusion",)
                        if unavailable
                        else ()
                    ),
                    fingerprint=f"coverage-{identity}-{record.condition}-{roi_name}",
                )
            )
        normalization = SimpleNamespace(
            expected_channels=expected,
            observed_channels=expected,
            excluded_channels=excluded,
            interpolated_channels=(),
            used_channels=() if cell_excluded or excluded else expected,
            status="unavailable" if cell_excluded or excluded else "available",
            reason_codes=(
                ("reviewed_whole_scalp_member_exclusion",)
                if cell_excluded or excluded
                else ()
            ),
            fingerprint=f"normalization-{identity}-{record.condition}",
        )
        cell = SimpleNamespace(
            recording_id=identity,
            participant_id=str(record.participant_id),
            condition_label=str(record.condition),
            outcome_status="ready",
            workbook_path=str(Path(record.path).resolve()),
            source_evidence=object(),
            roi_memberships=tuple(memberships),
            whole_scalp_normalization=normalization,
            downstream_cell_excluded=cell_excluded,
            decision_reason_codes=(
                ("reviewed_recording_condition_exclusion",)
                if cell_excluded
                else ()
            ),
            fingerprint=f"cell-{identity}-{record.condition}",
        )
        cells.append(cell)
        by_key[key] = cell
    return SimpleNamespace(
        roi_snapshot=SimpleNamespace(rois=roi_definitions),
        cells=tuple(cells),
        cell_for=lambda identity, condition: by_key.get(
            (str(identity).casefold(), str(condition).casefold())
        ),
    )


def _build_fixture_project(root: Path) -> None:
    _write_project(root)
    _write_bca_workbook(
        root,
        pid="P1",
        condition="Condition A",
        group_folder="Anxious",
        values=[("O1", 1.0, 2.0), ("O2", 3.0, 4.0), ("Cz", 5.0, 6.0)],
    )
    _write_bca_workbook(
        root,
        pid="P1",
        condition="Condition B",
        group_folder="Anxious",
        values=[("O1", 1.0, 1.0), ("O2", 2.0, 2.0), ("Cz", 3.0, 3.0)],
    )
    _write_bca_workbook(
        root,
        pid="P2",
        condition="Condition A",
        group_folder="NonAnxious",
        values=[("O1", 2.0, 2.0), ("O2", 4.0, 4.0), ("Cz", 6.0, 6.0)],
    )


def _rewrite_sources_as_compact_only(root: Path) -> None:
    from Main_App.Shared.post_process_excel import write_results_workbook

    index = load_project_dataset_index(root)
    for record in [*index.workbooks, *index.excluded_workbooks]:
        values = pd.read_excel(record.path, sheet_name="BCA (uV)")
        receipt = write_results_workbook(
            str(record.path), {"BCA (uV)": values},
        )
        companion = record.path.with_name(receipt["condition_companion"]["path"])
        assert companion.is_file()
        # The default writer publishes only a notice in the physical worksheet.
        # All downstream numerical work must therefore read the companion.
        notice = pd.read_excel(record.path, sheet_name="BCA (uV)")
        assert "Electrode" not in notice.columns
        assert "Companion file" in notice.columns


def test_compact_only_sources_preserve_analysis_ready_long_and_wide_exports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "Project"
    _build_fixture_project(root)
    coverage = _released_coverage(
        load_project_dataset_index(root),
        rois={"LOT": ["O1"], "ROT": ["O2"], "Central": ["CZ"]},
        excluded_by_cell={("p1", "condition a"): {"O2"}},
        excluded_cells={("p2", "condition a")},
    )
    # Fixture-authorized QC coverage is unchanged by the storage migration.
    # All source reads, BCA sums, normalization, pivots and Excel writes are real.
    monkeypatch.setattr(
        export_module, "_require_analysis_ready_release",
        lambda _root: (coverage, "test-final-release"),
    )
    sheets = [
        "ROI Long", "Raw BCA Wide", "RMS Normalized Wide",
        "Signed Mean Normalized Wide", "Electrode Long",
        "Whole Scalp Values", "RMS Harmonic Scales", "QC Flags", "ROI Coverage",
    ]
    legacy = write_analysis_ready_workbook(root)
    expected = pd.read_excel(legacy.path, sheet_name=sheets)

    _rewrite_sources_as_compact_only(root)
    compact = write_analysis_ready_workbook(root)
    actual = pd.read_excel(compact.path, sheet_name=sheets)

    assert compact.path == legacy.path
    assert compact.row_count == legacy.row_count == 9
    for sheet in sheets:
        pd.testing.assert_frame_equal(actual[sheet], expected[sheet], check_exact=True)
    long = actual["ROI Long"]
    assert long.loc[long["PID"].eq("P2"), "Raw Summed BCA"].isna().all()
    retained = long.loc[
        long["PID"].eq("P1") & long["Condition"].eq("Condition B")
        & long["ROI"].eq("LOT")
    ].iloc[0]
    assert retained["Raw Summed BCA"] == 2.0
    assert math.isfinite(retained["RMS Normalized BCA"])


def test_compact_only_sources_preserve_stats_ready_long_and_wide_exports(
    tmp_path: Path,
) -> None:
    from Tools.Stats.analysis.dv_policy_settings import FIXED_PREDEFINED_POLICY_NAME
    from Tools.Stats.io.stats_ready_export import (
        LONG_FORMAT_SHEET,
        WIDE_FORMAT_SHEET,
        prepare_stats_ready_export,
    )

    root = tmp_path / "Project"
    _build_fixture_project(root)
    # Include signed BCA and an unused channel to catch value or row-selection
    # changes; the same fixed selected frequencies are used in both runs.
    _write_bca_workbook(
        root, pid="P1", condition="Condition A", group_folder="Anxious",
        values=[("O1", -0.5, 1.25), ("O2", 0.75, -0.25), ("Cz", 0.0, 0.0)],
    )
    index = load_project_dataset_index(root)
    subject_data = {}
    for record in index.workbooks:
        subject_data.setdefault(record.participant_id, {})[record.condition] = str(record.path)
    target = root / "Stats_Ready_Summed_BCA.xlsx"
    inputs = {
        "subjects": ["P1"], "conditions": ["Condition A", "Condition B"],
        "subject_data": subject_data, "base_freq": 6.0,
        "rois": {"Occipital": ["O1", "O2"]},
        "dv_policy": {
            "name": FIXED_PREDEFINED_POLICY_NAME,
            "fixed_harmonic_frequencies_hz": "1.2, 2.4",
        },
        "group_map": {"P1": "anxious"},
        "log_func": lambda _message: None,
        "save_path": target,
    }
    legacy = prepare_stats_ready_export(**inputs)
    sheets = [LONG_FORMAT_SHEET, WIDE_FORMAT_SHEET]
    expected = pd.read_excel(target, sheet_name=sheets)

    _rewrite_sources_as_compact_only(root)
    compact = prepare_stats_ready_export(**inputs)
    actual = pd.read_excel(target, sheet_name=sheets)

    assert compact.workbook_path == legacy.workbook_path == target
    assert compact.row_count == legacy.row_count == 2
    for sheet in sheets:
        pd.testing.assert_frame_equal(actual[sheet], expected[sheet], check_exact=True)
        pd.testing.assert_frame_equal(
            compact.frames[sheet], legacy.frames[sheet], check_exact=True,
        )
    assert actual[LONG_FORMAT_SHEET]["summed_bca_uv"].tolist() == [0.625, 3.0]
    assert actual[WIDE_FORMAT_SHEET].iloc[0, 2:].tolist() == [0.625, 3.0]


def test_full_audit_workbook_keeps_excluded_values_and_flags(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "Project"
    _build_fixture_project(root)
    index = load_project_dataset_index(root)
    assert {record.participant_id for record in index.workbooks} == {"P1"}
    assert {record.participant_id for record in index.excluded_workbooks} == {"P2"}
    coverage = _released_coverage(
        index,
        rois={"LOT": ["O1"], "ROT": ["O2"], "Central": ["CZ"]},
        excluded_by_cell={("p1", "condition a"): {"O2"}},
        excluded_cells={("p2", "condition a")},
    )
    monkeypatch.setattr(
        export_module,
        "_require_analysis_ready_release",
        lambda _root: (coverage, "test-final-release"),
    )

    result = write_analysis_ready_workbook(
        root,
        dataset_index=index,
        selection_metadata=_selection_metadata(),
    )

    assert result.workbook_path.name == ANALYSIS_READY_WORKBOOK_NAME
    assert result.roi_row_count == 9
    assert result.participant_count == 2
    assert result.condition_count == 2
    assert result.flag_count >= 3

    roi_long = pd.read_excel(result.workbook_path, sheet_name="ROI Long")
    assert list(roi_long.columns) == [
        "PID",
        "Group",
        "Condition",
        "ROI",
        "Raw Summed BCA",
        "RMS Normalized BCA",
        "Signed Mean Normalized BCA",
        "Current Toolbox Exclusion",
        "QC Flag",
        "QC Notes",
    ]
    excluded = roi_long.loc[roi_long["PID"].eq("P2")]
    assert len(excluded) == 3
    assert set(excluded["Group"]) == {"non-anxious"}
    assert set(excluded["Current Toolbox Exclusion"]) == {"Yes"}
    assert excluded["Raw Summed BCA"].isna().all()

    p1_lot = roi_long.loc[
        roi_long["PID"].eq("P1")
        & roi_long["Condition"].eq("Condition A")
        & roi_long["ROI"].eq("LOT")
    ].iloc[0]
    assert p1_lot["Raw Summed BCA"] == pytest.approx(3.0)
    assert pd.isna(p1_lot["RMS Normalized BCA"])
    assert pd.isna(p1_lot["Signed Mean Normalized BCA"])

    p1_rot = roi_long.loc[
        roi_long["PID"].eq("P1") & roi_long["Condition"].eq("Condition A") & roi_long["ROI"].eq("ROT")
    ].iloc[0]
    assert p1_rot["Group"] == "anxious"
    assert pd.isna(p1_rot["Raw Summed BCA"])
    assert p1_rot["Current Toolbox Exclusion"] == "Yes"
    assert "O2" in p1_rot["QC Notes"]
    assert pd.isna(p1_rot["RMS Normalized BCA"])
    assert pd.isna(p1_rot["Signed Mean Normalized BCA"])

    electrodes = pd.read_excel(result.workbook_path, sheet_name="Electrode Long")
    flagged_o2 = electrodes.loc[
        electrodes["PID"].eq("P1") & electrodes["Condition"].eq("Condition A") & electrodes["Electrode"].eq("O2")
    ].iloc[0]
    assert flagged_o2["Raw Summed BCA"] == pytest.approx(7.0)
    assert flagged_o2["Current Toolbox Exclusion"] == "Yes"
    p2_electrodes = electrodes.loc[
        electrodes["PID"].eq("P2")
        & electrodes["Condition"].eq("Condition A")
    ]
    assert p2_electrodes["Raw Summed BCA"].notna().all()
    assert p2_electrodes["RMS Normalized BCA"].isna().all()
    assert p2_electrodes["Signed Mean Normalized BCA"].isna().all()
    assert set(p2_electrodes["Current Toolbox Exclusion"]) == {"Yes"}
    assert p2_electrodes["QC Notes"].str.contains("Audit-only raw electrode").all()

    harmonic_scales = pd.read_excel(
        result.workbook_path,
        sheet_name="RMS Harmonic Scales",
    )
    p1_condition_a = harmonic_scales.loc[
        harmonic_scales["PID"].eq("P1") & harmonic_scales["Condition"].eq("Condition A")
    ].sort_values("Harmonic (Hz)")
    assert list(p1_condition_a["Scalp Vector Length"]) == pytest.approx(
        [math.sqrt(1.0**2 + 3.0**2 + 5.0**2), math.sqrt(2.0**2 + 4.0**2 + 6.0**2)]
    )
    assert set(p1_condition_a["Used for RMS Normalization"]) == {"No"}
    p2_condition_a = harmonic_scales.loc[
        harmonic_scales["PID"].eq("P2")
        & harmonic_scales["Condition"].eq("Condition A")
    ]
    assert p2_condition_a["Scalp Vector Length"].isna().all()
    assert set(p2_condition_a["Current Toolbox Exclusion"]) == {"Yes"}

    whole_scalp = pd.read_excel(result.workbook_path, sheet_name="Whole Scalp Values")
    p1_condition_a_whole_scalp = whole_scalp.loc[
        whole_scalp["PID"].eq("P1") & whole_scalp["Condition"].eq("Condition A")
    ].iloc[0]
    assert p1_condition_a_whole_scalp["Descriptive Post-Sum RMS (Not Used for Normalization)"] == pytest.approx(
        math.sqrt((3.0**2 + 7.0**2 + 11.0**2) / 3.0)
    )
    p2_condition_a_whole_scalp = whole_scalp.loc[
        whole_scalp["PID"].eq("P2")
        & whole_scalp["Condition"].eq("Condition A")
    ].iloc[0]
    assert pd.isna(
        p2_condition_a_whole_scalp[
            "Descriptive Post-Sum RMS (Not Used for Normalization)"
        ]
    )
    assert p2_condition_a_whole_scalp["Current Toolbox Exclusion"] == "Yes"

    wide = pd.read_excel(result.workbook_path, sheet_name="Raw BCA Wide")
    p2_wide = wide.loc[wide["PID"].eq("P2")].iloc[0]
    assert pd.isna(p2_wide["Condition B | LOT"])
    assert pd.isna(p2_wide["Condition A | ROT"])

    workbook = load_workbook(result.workbook_path)
    assert workbook.sheetnames == [
        "ROI Long",
        "Raw BCA Wide",
        "RMS Normalized Wide",
        "Signed Mean Normalized Wide",
        "Electrode Long",
        "Whole Scalp Values",
        "RMS Harmonic Scales",
        "QC Flags",
        "ROI Definitions",
        "ROI Coverage",
        "Selection Summary",
        "Harmonic Selection",
        "Analysis Notes",
    ]
    sheet = workbook["ROI Long"]
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].fill.fgColor.rgb in {"00595959", "FF595959"}
    assert sheet["A2"].fill.fgColor.rgb in {"00F2F2F2", "FFF2F2F2"}
    assert sheet["A2"].alignment.horizontal == "center"
    assert all("path" not in str(cell.value or "").casefold() for cell in sheet[1])

    selection_summary = pd.read_excel(
        result.workbook_path,
        sheet_name="Selection Summary",
    )
    summary = dict(
        zip(selection_summary["Summary Item"], selection_summary["Value"])
    )
    assert summary["Harmonic selection profile ID"] == (
        "significant_only_exploratory"
    )
    assert summary["Harmonic selection profile version"] == "1.0"
    assert summary["Selection fingerprint"] == "c" * 64
    assert summary["Frozen/effective selection electrode mask"] == "O1; O2"
    assert summary["Pooling cell participant counts"] == "anxious::Condition A=1"
    assert "P1::Condition A" in summary["Selection source workbook identities"]

    roi_coverage = pd.read_excel(result.workbook_path, sheet_name="ROI Coverage")
    normalization = roi_coverage.loc[
        roi_coverage["PID"].eq("P1")
        & roi_coverage["Condition"].eq("Condition A")
        & roi_coverage["Coverage Type"].eq("Whole-scalp normalization")
    ].iloc[0]
    assert normalization["Status"] == "unavailable"
    assert normalization["Excluded Electrodes"] == "O2"
    assert normalization["Expected Count"] == 3
    assert normalization["Observed Count"] == 3
    assert normalization["Excluded Count"] == 1
    assert normalization["Interpolated Count"] == 0
    assert normalization["Used Count"] == 0
    assert normalization["Coverage Fingerprint"] == (
        "normalization-P1-Condition A"
    )
    assert not roi_long["QC Notes"].fillna("").str.contains(
        r"required electrode\(s\) were unavailable:\s*$",
        regex=True,
    ).any()


def test_full_audit_export_rejects_index_for_another_project(tmp_path: Path) -> None:
    wrong_index = SimpleNamespace(project_root=tmp_path / "Other Project")

    with pytest.raises(ValueError, match="different project root"):
        write_analysis_ready_workbook(tmp_path, dataset_index=wrong_index)


def test_full_audit_export_rejects_stale_same_root_dataset_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "Project"
    _build_fixture_project(root)
    index = load_project_dataset_index(root)
    coverage = _released_coverage(index, rois={"Posterior": ["O1"]})
    monkeypatch.setattr(
        export_module,
        "_require_analysis_ready_release",
        lambda _root: (coverage, "test-final-release"),
    )
    stale_first = replace(index.workbooks[0], group_id="stale-group")
    stale_index = replace(
        index,
        workbooks=(stale_first, *index.workbooks[1:]),
    )

    with pytest.raises(RuntimeError, match="stale participant, group, session"):
        write_analysis_ready_workbook(
            root,
            dataset_index=stale_index,
            selection_metadata=_selection_metadata(),
        )


def test_full_audit_rejects_semantically_modified_harmonic_metadata(
    tmp_path: Path,
) -> None:
    accepted = _selection_metadata()
    tampered = _selection_metadata()
    tampered["included_harmonics_hz"] = [1.2]

    frames, source = export_module._load_selection_frames(
        tmp_path,
        selection_metadata=accepted,
        expected_release_fingerprint="test-final-release",
    )
    assert frames
    assert source == "processing-time metadata"

    with pytest.raises(RuntimeError, match="differs from the current persisted"):
        export_module._load_selection_frames(
            tmp_path,
            selection_metadata=tampered,
            expected_release_fingerprint="test-final-release",
        )


def test_full_audit_repeated_session_export_keeps_both_recordings(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "Repeated Project"
    raw_root = root / "Raw" / "Control"
    (raw_root / "Luteal").mkdir(parents=True)
    (raw_root / "Follicular").mkdir(parents=True)
    manifest = {
        "schema_version": "2.2.0",
        "results_folder": ".",
        "subfolders": {"excel": "1 - Excel Data Files"},
        "groups": {
            "control": {
                "label": "Control",
                "folder_name": "Control",
                "raw_input_folder": str(raw_root),
            }
        },
        "sessions": {
            "luteal": {"label": "Luteal (Visit 1)", "visit_index": 1},
            "follicular": {
                "label": "Follicular (Visit 2)",
                "visit_index": 2,
            },
        },
        "recording_sources": {
            "control_luteal": {
                "group_id": "control",
                "session_id": "luteal",
                "raw_input_folder": str(raw_root / "Luteal"),
            },
            "control_follicular": {
                "group_id": "control",
                "session_id": "follicular",
                "raw_input_folder": str(raw_root / "Follicular"),
            },
        },
        "participants": {"P1": {"group_id": "control"}},
        "preprocessing": {
            "manual_excluded_recordings": ["P1__luteal"],
        },
        "tools": {
            "frequency_domain_qc": {
                "auto_recording_exclusions": [
                    {
                        "recording_id": "P1__luteal",
                        "reason": "test automatic recording flag",
                    }
                ],
                "manual_recording_exclusions": [
                    {
                        "recording_id": "P1__luteal",
                        "reason": "test manual recording flag",
                    }
                ],
                "auto_recording_electrode_exclusions": [
                    {
                        "recording_id": "P1__luteal",
                        "electrode": "O1",
                        "reason": "test recording electrode flag",
                        "triggering_conditions": ["Condition A"],
                    }
                ],
            }
        },
        "recordings": {
            "P1__luteal": {
                "participant_id": "P1",
                "session_id": "luteal",
                "source_id": "control_luteal",
                "raw_file": str(raw_root / "Luteal" / "P1_L.bdf"),
                "visit_index": 1,
                "days_from_baseline": 0,
            },
            "P1__follicular": {
                "participant_id": "P1",
                "session_id": "follicular",
                "source_id": "control_follicular",
                "raw_file": str(raw_root / "Follicular" / "P1_F.bdf"),
                "visit_index": 2,
                "days_from_baseline": 14,
            },
        },
    }
    (root / "project.json").write_text(json.dumps(manifest), encoding="utf-8")
    _write_bca_workbook(
        root,
        pid="P1__luteal",
        condition="Condition A",
        group_folder="Control",
        values=[("O1", 1.0, 2.0)],
    )
    _write_bca_workbook(
        root,
        pid="P1__follicular",
        condition="Condition A",
        group_folder="Control",
        values=[("O1", 3.0, 4.0)],
    )
    index = load_project_dataset_index(root)
    coverage = _released_coverage(
        index,
        rois={"Occipital": ["O1"]},
        excluded_by_cell={("p1__luteal", "condition a"): {"O1"}},
    )
    monkeypatch.setattr(
        export_module,
        "_require_analysis_ready_release",
        lambda _root: (coverage, "test-final-release"),
    )

    result = write_analysis_ready_workbook(
        root,
        selection_metadata=_selection_metadata(),
    )

    roi_long = pd.read_excel(result.path, sheet_name="ROI Long")
    assert list(roi_long["Recording ID"]) == ["P1__luteal", "P1__follicular"]
    assert list(roi_long["Session ID"]) == ["luteal", "follicular"]
    assert list(roi_long["Visit Index"]) == [1, 2]
    assert list(roi_long["Days From Baseline"]) == [0, 14]
    assert set(roi_long["Group ID"]) == {"control"}
    assert pd.isna(roi_long.iloc[0]["Raw Summed BCA"])
    assert roi_long.iloc[1]["Raw Summed BCA"] == pytest.approx(7.0)
    assert list(roi_long["Current Toolbox Exclusion"]) == ["Yes", "No"]
    assert "test automatic recording flag" in str(roi_long.iloc[0]["QC Notes"])
    assert "test manual recording flag" in str(roi_long.iloc[0]["QC Notes"])
    assert "frequency-domain recording" not in str(
        roi_long.iloc[1]["QC Notes"]
    ).casefold()
    electrode_long = pd.read_excel(result.path, sheet_name="Electrode Long")
    luteal_electrode = electrode_long.loc[
        electrode_long["Recording ID"].eq("P1__luteal")
        & electrode_long["Electrode"].eq("O1")
    ].iloc[0]
    follicular_electrode = electrode_long.loc[
        electrode_long["Recording ID"].eq("P1__follicular")
        & electrode_long["Electrode"].eq("O1")
    ].iloc[0]
    assert "test recording electrode flag" in luteal_electrode["QC Notes"]
    assert "test recording electrode flag" not in str(
        follicular_electrode["QC Notes"]
    )
    wide = pd.read_excel(result.path, sheet_name="Raw BCA Wide")
    assert len(wide) == 2
    assert set(wide["Recording ID"]) == {"P1__luteal", "P1__follicular"}
    flags = pd.read_excel(result.path, sheet_name="QC Flags")
    recording_flag = flags.loc[
        flags["Flag Type"].eq("Manual preprocessing recording exclusion")
    ].iloc[0]
    assert recording_flag["Recording ID"] == "P1__luteal"
    assert recording_flag["Session ID"] == "luteal"
    frequency_recording_flags = flags.loc[
        flags["Flag Type"].isin(
            {
                "Automatic frequency-domain recording exclusion",
                "Manual frequency-domain recording exclusion",
                "Automatic frequency-domain recording-electrode exclusion",
            }
        )
    ]
    assert set(frequency_recording_flags["Flag Type"]) == {
        "Automatic frequency-domain recording exclusion",
        "Manual frequency-domain recording exclusion",
        "Automatic frequency-domain recording-electrode exclusion",
    }
    assert set(frequency_recording_flags["Recording ID"]) == {"P1__luteal"}
    assert set(frequency_recording_flags["Session ID"]) == {"luteal"}


def test_full_audit_rejects_partial_selected_harmonic_sums() -> None:
    frame = pd.DataFrame(
        {
            "Electrode": ["O1", "O2"],
            "1.2000_Hz": [0.0, 0.0],
            "2.4000_Hz": [1.0, math.nan],
        }
    )

    with pytest.raises(RuntimeError, match="partial harmonic sums are forbidden"):
        export_module._prepare_electrode_values(
            frame,
            selected_columns=["1.2000_Hz", "2.4000_Hz"],
            expected_scalp_channels=["O1", "O2"],
            normalization_available=True,
        )


def test_atomic_writer_preserves_previous_workbook_when_rebuild_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    target = tmp_path / "analysis.xlsx"
    target.write_bytes(b"previous complete workbook")

    class _FailingWriter:
        def __init__(self, *_args, **_kwargs) -> None:
            pass

        def __enter__(self):
            raise RuntimeError("simulated write failure")

        def __exit__(self, *_args) -> None:
            return None

    monkeypatch.setattr(export_module.pd, "ExcelWriter", _FailingWriter)

    with pytest.raises(RuntimeError, match="simulated write failure"):
        export_module._write_frames_atomically(
            target,
            {"ROI Long": pd.DataFrame({"PID": ["P1"]})},
        )

    assert target.read_bytes() == b"previous complete workbook"
    assert not list(tmp_path.glob("*.tmp.xlsx"))
