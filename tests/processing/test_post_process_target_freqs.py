from __future__ import annotations

import logging
from types import SimpleNamespace

import mne
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from Main_App.Shared.post_process import (
    _can_batch_target_noise,
    _create_output_subfolder,
    _eeg_pick_indices,
    _mean_epochs_float64,
    _resolve_frequency_protocol,
    post_process,
)
from Main_App.processing.spectral_eligibility import SpectralEligibilityError
from Main_App.projects.frequency_protocol import (
    EXPECTED_CYCLES_SOURCE_MANUAL,
    FrequencyProtocol,
)


def _ready_protocol() -> FrequencyProtocol:
    return FrequencyProtocol.from_recurrence(
        3,
        10,
        expected_analyzed_oddball_cycles=36,
        expected_analyzed_oddball_cycles_source=EXPECTED_CYCLES_SOURCE_MANUAL,
    )


@pytest.mark.parametrize(
    "array_kind",
    ["float64", "float32", "noncontiguous"],
)
def test_mean_epochs_float64_matches_copying_expression_byte_exact(
    array_kind: str,
) -> None:
    epoch_data = np.random.default_rng(24).normal(size=(7, 5, 18))
    if array_kind == "float32":
        epoch_data = epoch_data.astype(np.float32)
    elif array_kind == "noncontiguous":
        epoch_data = epoch_data[:, :, ::2]
        assert not epoch_data.flags.c_contiguous

    expected = np.mean(epoch_data.astype(np.float64), axis=0)
    actual = _mean_epochs_float64(epoch_data)

    assert actual.dtype == expected.dtype
    assert actual.shape == expected.shape
    assert actual.tobytes() == expected.tobytes()


def test_mean_epochs_float64_preserves_broadcast_layout_result_bytes() -> None:
    base = np.array(
        [
            [-1.4238250364546312, 1.2637284581291104, -0.8706617379590857],
            [-0.2591732349343976, -0.07534330701052097, -0.740884652085609],
        ],
        dtype=np.float64,
    )
    epoch_data = np.broadcast_to(base, (9, *base.shape))
    assert epoch_data.strides[0] == 0
    assert not epoch_data.flags.c_contiguous
    assert not epoch_data.flags.f_contiguous

    expected = np.mean(epoch_data.astype(np.float64), axis=0)
    actual = _mean_epochs_float64(epoch_data)

    assert actual.tobytes() == expected.tobytes()


@pytest.mark.parametrize("object_kind", ["epochs", "evoked"])
def test_direct_mne_eeg_pick_matches_copy_pick_byte_exact(object_kind: str) -> None:
    rng = np.random.default_rng(1729)
    info = mne.create_info(
        ["Fp1", "EOG", "Cz", "Pz", "STI"],
        sfreq=256.0,
        ch_types=["eeg", "eog", "eeg", "eeg", "stim"],
    )
    info["bads"] = ["Cz"]
    if object_kind == "epochs":
        original = mne.EpochsArray(
            rng.normal(size=(4, 5, 96)),
            info,
            tmin=0.0,
            baseline=None,
            verbose=False,
        )
        is_evoked = False
    else:
        original = mne.EvokedArray(
            rng.normal(size=(5, 96)),
            info,
            tmin=0.0,
            verbose=False,
        )
        is_evoked = True

    original_bytes = (
        original.data.tobytes()
        if is_evoked
        else original.get_data(copy=True).tobytes()
    )
    expected = original.copy().pick(
        "eeg",
        exclude=[] if is_evoked else "bads",
    )
    picks = _eeg_pick_indices(original, is_evoked=is_evoked)
    actual_names = [original.ch_names[int(index)] for index in picks]
    if is_evoked:
        expected_data = expected.data
        actual_data = original.data[picks, :]
    else:
        expected_data = expected.get_data()
        actual_data = original.get_data(picks=picks, copy=True)

    assert actual_names == expected.ch_names
    assert actual_data.dtype == expected_data.dtype
    assert actual_data.shape == expected_data.shape
    assert actual_data.tobytes() == expected_data.tobytes()
    assert (
        original.data.tobytes()
        if is_evoked
        else original.get_data(copy=True).tobytes()
    ) == original_bytes
    assert original.info["bads"] == ["Cz"]


def test_direct_mne_eeg_pick_defers_empty_selection_to_copy_pick() -> None:
    info = mne.create_info(["EOG"], sfreq=128.0, ch_types=["eog"])
    evoked = mne.EvokedArray(np.zeros((1, 24)), info, verbose=False)

    picks = _eeg_pick_indices(evoked, is_evoked=True)

    assert picks.size == 0
    with pytest.raises(ValueError) as expected_error:
        evoked.copy().pick("eeg", exclude=[])
    assert "could not be interpreted" in str(expected_error.value)


def test_target_noise_batch_guard_is_conservative() -> None:
    amplitudes = np.abs(np.random.default_rng(11).normal(size=(4, 96)))
    target_indices = np.array([20, 40, -1], dtype=np.intp)

    assert _can_batch_target_noise(amplitudes, target_indices)

    with_zero = amplitudes.copy()
    with_zero[:, 30] = 0.0
    assert not _can_batch_target_noise(with_zero, target_indices)

    tied = amplitudes.copy()
    tied[:, 30:51] = 1.0
    assert not _can_batch_target_noise(tied, np.array([40], dtype=np.intp))


@pytest.mark.parametrize("nested", [False, True])
def test_resolve_frequency_protocol_from_run_snapshot(nested: bool) -> None:
    protocol = _ready_protocol()
    settings = (
        {"analysis": {"frequency_protocol": protocol.to_manifest()}}
        if nested
        else {"frequency_protocol": protocol}
    )

    resolved = _resolve_frequency_protocol(SimpleNamespace(settings=settings))

    assert resolved == protocol
    assert float(resolved.presentation_rate_hz) == pytest.approx(3.0)
    assert float(resolved.oddball_rate_hz) == pytest.approx(0.3)


def test_resolve_frequency_protocol_rejects_legacy_rate_and_ceiling_settings() -> None:
    app = SimpleNamespace(
        settings={"analysis": {"oddball_freq": 1.2, "bca_upper_limit": 16.8}}
    )

    with pytest.raises(SpectralEligibilityError, match="fallbacks are retired"):
        _resolve_frequency_protocol(app)


def test_create_output_subfolder_routes_condition_then_group(tmp_path) -> None:
    app = SimpleNamespace(log=lambda _message: None)

    output = _create_output_subfolder(
        app,
        tmp_path,
        "Condition A",
        "Control",
    )

    assert output == str((tmp_path / "Condition A" / "Control").resolve())
    assert (tmp_path / "Condition A" / "Control").is_dir()


def test_create_output_subfolder_raises_instead_of_falling_back(
    tmp_path,
    monkeypatch,
) -> None:
    logs: list[str] = []
    app = SimpleNamespace(log=logs.append)

    def fail_mkdir(*_args, **_kwargs):
        raise OSError("disk unavailable")

    monkeypatch.setattr("Main_App.Shared.post_process.os.makedirs", fail_mkdir)

    with pytest.raises(OSError, match="disk unavailable"):
        _create_output_subfolder(app, tmp_path, "Condition A", "Control")

    assert logs and "Processing cannot continue" in logs[-1]
    assert "Saving to parent" not in logs[-1]


def test_post_process_logs_export_timing_when_no_data(tmp_path, caplog) -> None:
    class _PathBox:
        def get(self):
            return str(tmp_path)

    logs: list[str] = []
    app = SimpleNamespace(
        save_folder_path=_PathBox(),
        settings={"frequency_protocol": _ready_protocol()},
        preprocessed_data={},
        data_paths=[],
        log=logs.append,
        export_timing_records=[],
    )

    caplog.set_level(logging.DEBUG, logger="Main_App.Shared.post_process")
    post_process(app, ["CondA"])

    assert "[EXPORT TIMING]" in caplog.text
    assert "stage=condition_skip_no_data" in caplog.text
    assert "stage=post_process_total" in caplog.text
    assert {record["stage"] for record in app.export_timing_records} >= {
        "condition_skip_no_data",
        "post_process_total",
    }


def test_post_process_exports_filter_domain_and_structured_notch_hole(tmp_path) -> None:
    sampling_rate = 128.0
    sample_count = 7_680  # 120 cycles at a 2-Hz oddball rate.
    info = mne.create_info(["Oz"], sampling_rate, ["eeg"])
    with info._unlock():
        info["highpass"] = 0.1
        info["lowpass"] = 60.0
    metadata = pd.DataFrame(
        {
            "crop_mode": ["55_onbin"],
            "n55": [120],
            "first55_samp": [0],
            "last55_samp": [sample_count - 64],
            "N_step": [64],
            "fallback_reason": [""],
        }
    )
    epochs = mne.EpochsArray(
        np.random.default_rng(12).normal(size=(1, 1, sample_count)) * 1e-6,
        info,
        tmin=0.0,
        metadata=metadata,
        verbose=False,
    )
    protocol = FrequencyProtocol.from_recurrence(
        10,
        5,
        expected_analyzed_oddball_cycles=120,
        expected_analyzed_oddball_cycles_source=EXPECTED_CYCLES_SOURCE_MANUAL,
    )
    app = SimpleNamespace(
        save_folder_path=SimpleNamespace(get=lambda: str(tmp_path)),
        settings={
            "frequency_protocol": protocol,
            "high_pass": 0.1,
            "low_pass": 60.0,
            "line_noise_filter_enabled": True,
            "_fpvs_fft_multinotch_applied_centers_hz": [50.0],
        },
        preprocessed_data={"Condition A": [epochs]},
        data_paths=[],
        log=lambda _message: None,
    )

    post_process(app, ["Condition A"])

    workbook_path = next(tmp_path.rglob("*.xlsx"))
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        assert "48.0000_Hz" in [
            cell.value for cell in workbook["BCA (uV)"][1]
        ]
        bca_headers = [cell.value for cell in workbook["BCA (uV)"][1]]
        bca_values = [cell.value for cell in workbook["BCA (uV)"][2]]
        assert bca_values[bca_headers.index("48.0000_Hz")] is not None
        assert bca_values[bca_headers.index("50.0000_Hz")] is None

        eligibility_rows = workbook["Spectral Eligibility"].iter_rows(values_only=True)
        eligibility_headers = list(next(eligibility_rows))
        frequency_index = eligibility_headers.index("Target Frequency (Hz)")
        available_index = eligibility_headers.index("BCA Available")
        reason_index = eligibility_headers.index("Unavailable Reasons")
        fifty_hz = next(
            row for row in eligibility_rows if row[frequency_index] == 50.0
        )
        assert fifty_hz[available_index] is False
        assert "target_inside_applied_notch" in fifty_hz[reason_index]
    finally:
        workbook.close()
