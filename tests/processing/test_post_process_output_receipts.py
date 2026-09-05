from __future__ import annotations

from types import SimpleNamespace

import mne
import numpy as np
import pandas as pd
import pytest

from Main_App.Shared.post_process import post_process
from Main_App.processing.output_integrity import OutputIntegrityError
from Main_App.projects.frequency_protocol import (
    EXPECTED_CYCLES_SOURCE_MANUAL,
    FrequencyProtocol,
)


def _protocol() -> FrequencyProtocol:
    return FrequencyProtocol.from_recurrence(
        3,
        10,
        expected_analyzed_oddball_cycles=36,
        expected_analyzed_oddball_cycles_source=EXPECTED_CYCLES_SOURCE_MANUAL,
    )


def _epochs(*, invalid: float | None = None, sampling_rate: float = 20.0) -> mne.EpochsArray:
    sample_count = int(120 * sampling_rate)
    info = mne.create_info(["Oz"], sampling_rate, ["eeg"])
    with info._unlock():
        info["highpass"] = 0.1
        info["lowpass"] = 9.0
    values = np.random.default_rng(42).normal(
        size=(1, 1, sample_count)
    ) * 1e-6
    if invalid is not None:
        values[0, 0, 100] = invalid
    metadata = pd.DataFrame(
        {
            "crop_mode": ["project_marker_plan_target_grid_v2"],
            "n55": [36],
            "first55_samp": [0],
            "last55_samp": [sample_count - 1],
            "N_step": [int(sampling_rate * 10)],
            "fallback_reason": [""],
            "approved_span_fingerprint": ["a" * 64],
            "marker_plan_fingerprint": ["b" * 64],
            "source_start_sample": [100],
            "source_stop_sample": [100 + sample_count],
            "target_start_sample": [100],
            "target_stop_sample": [100 + sample_count],
            "marker_disposition": ["complete"],
        }
    )
    return mne.EpochsArray(
        values,
        info,
        tmin=0.0,
        metadata=metadata,
        verbose=False,
    )


def _app(tmp_path, epochs: mne.EpochsArray):
    return SimpleNamespace(
        save_folder_path=SimpleNamespace(get=lambda: str(tmp_path)),
        settings={
            "frequency_protocol": _protocol(),
            "high_pass": 0.1,
            "low_pass": 9.0,
            "line_noise_filter_enabled": False,
            "_fpvs_recording_id": "P01__visit_1",
            "output_recording_stem": "P01__visit_1",
            "_fpvs_participant_id": "P01",
            "_fpvs_session_id": "visit_1",
            "_fpvs_expected_plan_run_id": "run-123",
            "_fpvs_processing_fingerprint": "c" * 64,
            "_fpvs_processing_fingerprint_version": "processing-v1",
        },
        preprocessed_data={"Faces": [epochs]},
        data_paths=[],
        log=lambda _message: None,
        export_receipts=[],
    )


def test_current_run_receipt_links_atomic_workbook_and_exact_span(tmp_path):
    app = _app(tmp_path, _epochs())

    post_process(app, ["Faces"])

    assert len(app.export_receipts) == 1
    receipt = app.export_receipts[0]
    assert receipt["version"] == "recording_condition_export_receipt_v1"
    assert receipt["status"] == "written"
    assert receipt["run_id"] == "run-123"
    assert receipt["recording_id"] == "P01__visit_1"
    assert receipt["condition_label"] == "Faces"
    assert receipt["retained_occurrence_count"] == 1
    assert receipt["retained_occurrences"][0]["span_status"] == "exact"
    assert receipt["workbook_write"]["schema_validation"]["status"] == "passed"
    assert receipt["workbook_write"]["artifact"]["sha256"]
    assert {item["value_category"] for item in receipt["finite_integrity"]} == {
        "retained_eeg",
        "bca",
    }


def test_512_hz_export_preserves_full_spectra_and_calculated_snr(tmp_path):
    from openpyxl import load_workbook

    from Main_App.io.spectral_data import read_spectral_sheet, spectral_companion_identity
    from Tools.Stats.analysis.full_snr import compute_full_snr_from_amplitudes

    epochs = _epochs(sampling_rate=512.0)
    app = _app(tmp_path, epochs)
    post_process(app, ["Faces"])
    receipt = app.export_receipts[0]
    assert receipt["status"] == "written", receipt
    path = receipt["path"]
    descriptor = spectral_companion_identity(path)
    assert descriptor == receipt["workbook_write"]["spectral_companion"]

    # Established production expression, including its Nyquist scaling.
    averaged_uv = np.mean(epochs.get_data(copy=True).astype(np.float64), axis=0) * 1e6
    samples = averaged_uv.shape[1]
    frequencies = np.fft.rfftfreq(samples, d=1 / 512.0)
    expected_fft = np.abs(np.fft.fft(averaged_uv, axis=1)[:, :samples // 2 + 1]) / samples * 2
    full_fft = read_spectral_sheet(path, sheet_name="FullFFT Amplitude (uV)")
    assert full_fft.shape == (1, 30_722)
    assert full_fft["Electrode"].tolist() == ["Oz"]
    assert full_fft.iloc[:, 1:].to_numpy().tobytes() == expected_fft.tobytes()
    metadata = full_fft.attrs["spectral_metadata"]
    np.testing.assert_array_equal(metadata["frequencies_hz"], frequencies)
    assert metadata["fft_sample_count"] == 61_440
    assert metadata["sampling_frequency_hz"] == 512.0

    full_snr = read_spectral_sheet(path, sheet_name="FullSNR")
    snr_grid = np.arange(0.5, 9.0 + 0.01, 0.01)
    calculated_snr = compute_full_snr_from_amplitudes(expected_fft)
    expected_snr = np.interp(snr_grid, frequencies, calculated_snr[0])
    np.testing.assert_array_equal(full_snr.iloc[0, 1:].to_numpy(dtype=float), expected_snr)
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook["FullFFT Amplitude (uV)"].max_column == 2
        assert workbook["FullSNR"].max_column == 2
        assert "Spectral Data" in workbook.sheetnames
    finally:
        workbook.close()


def test_compact_export_preserves_calculated_tables_without_excel_copies(
    tmp_path, monkeypatch
):
    import importlib
    from openpyxl import load_workbook
    from Main_App.io.condition_data import CONDITION_DATA_SHEET_NAMES, read_condition_sheet

    module = importlib.import_module("Main_App.Shared.post_process")
    original_write = module.write_results_workbook
    calculated = {}

    def capture_write(*args, **kwargs):
        calculated.update({
            name: frame.copy(deep=True)
            for name, frame in kwargs["dataframes_to_save"].items()
            if name in CONDITION_DATA_SHEET_NAMES
        })
        for argument, name in (
            ("spectral_eligibility_df", "Spectral Eligibility"),
            ("spectral_metric_qc_df", "Spectral Metric QC"),
            ("fft_neighbors_df", "FFT and neighbors"),
        ):
            calculated[name] = kwargs[argument].copy(deep=True)
        return original_write(*args, **kwargs)

    monkeypatch.setattr(module, "write_results_workbook", capture_write)
    app = _app(tmp_path, _epochs())
    post_process(app, ["Faces"])
    receipt = app.export_receipts[0]
    assert receipt["status"] == "written"
    assert receipt["workbook_write"]["condition_companion"]["path"].endswith(".npz")
    path = receipt["path"]
    workbook = load_workbook(path, read_only=True)
    try:
        assert "Condition Data" in workbook.sheetnames
        for name in calculated:
            assert workbook[name].max_row == 2
            assert workbook[name].max_column == 2
    finally:
        workbook.close()

    def no_excel_metrics(*args, **kwargs):
        raise AssertionError("Compact numerical reads must not parse Excel tables")

    monkeypatch.setattr(pd, "read_excel", no_excel_metrics)
    for name, expected in calculated.items():
        actual = read_condition_sheet(path, sheet_name=name)
        pd.testing.assert_frame_equal(actual, expected.reset_index(drop=True), check_exact=True)


@pytest.mark.parametrize("invalid", [np.nan, np.inf, -np.inf])
def test_nonfinite_source_blocks_current_export_and_preserves_prior_workbook(
    tmp_path,
    invalid,
):
    destination = tmp_path / "Faces" / "P01__visit_1_Faces_Results.xlsx"
    destination.parent.mkdir(parents=True)
    prior_bytes = b"prior complete workbook placeholder"
    destination.write_bytes(prior_bytes)
    app = _app(tmp_path, _epochs(invalid=invalid))

    with pytest.raises(OutputIntegrityError, match="NaN or infinite"):
        post_process(app, ["Faces"])

    assert destination.read_bytes() == prior_bytes
    assert len(app.export_receipts) == 1
    receipt = app.export_receipts[0]
    assert receipt["status"] == "blocked"
    assert receipt["failure_stage"] == "retained_signal"
    assert receipt["current_run_artifact"] is None
    assert receipt["prior_artifact"]["sha256"]
    assert receipt["integrity_failure"]["electrode"] == "Oz"
